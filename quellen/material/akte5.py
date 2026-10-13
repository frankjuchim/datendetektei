#!/usr/bin/env python3
"""Erstellungshilfe: Material Akte 5 „Der Bauplan“ (DS 9–11, nur eA) im Akten-Design (Phase 7, Task 6).

Erzeugt unter materialien/Akte5_Bauplan/ (Spec §12.1) und materialien/Warm-up/:
  1_Praesentation/DD_A5_Bauplan.pptx                              Fallbesprechung DS 9–11 (Notizen, Lösungsfolien)
  2_Arbeitsblaetter/DD_A5_AB1_Bauplaene_lesen.docx                 AB 5.1 (Schülerfassung)
  2_Arbeitsblaetter/DD_A5_AB2_Der_Architekturwettbewerb.docx       AB 5.2 (Schülerfassung)
  2_Arbeitsblaetter/DD_A5_AB3_Vom_Plan_zur_Datenbank.docx          AB 5.3 (Schülerfassung)
  5_Loesungen/…_LSG.docx                                           Lösungen mit typischen Fehlern (P-Nummern)
  3_Unplugged/DD_A5_Unplugged_Faeden-Kardinalitaeten.docx          Rollenkarten + Etiketten (DS 9)
  3_Unplugged/DD_A5_Unplugged_Architekturwettbewerb.docx           Anforderungstext, Teamrollen, Bewertungsbogen (DS 10)
  3_Unplugged/DD_A5_Unplugged_Regelkarten.docx                     Regelkarten R1–R6 + Relationenkarten (DS 11)
  4_Lehrerhinweise/DD_A5_Lehrerhinweise.docx                       Überblick, Ziele, Verlauf, Unplugged, Schwierigkeiten
  Warm-up/DD_Warm-up-Karten_A5.docx                                aus seed/warmup/warmup_a5.json

Grundlage: docs/DS_PLANUNG.md (DS 9–11), docs/STORY.md 4.5 und 5.9–5.11, seed/akten/akte5.json (Spur im Lab),
seed/datenbanken/detektei/LIESMICH.md, Spec §11.2 (Korrekturen der alten Präsentation „ER-Modell“). ER-Diagramme in
Chen-Notation nach EH §7 über ER(graph, layout) bzw. er_png; jedes Layout wird auf Überlappungen geprüft
(pruefe_layout). Alle Ergebnistabellen kommen direkt aus der Detektei-DB (abfrage_ergebnis).
Die erzeugten docx/pptx sind danach maßgeblich (Spec §12); dieses Skript ist nur die Erstellungshilfe.

Aufruf: python3 quellen/material/akte5.py   (überschreibt die oben genannten Dateien)
"""
import json
import math
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "werkzeuge"))

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import akten_design as A  # noqa: E402
from akten_design import (  # noqa: E402
    ER, AktenDokument, AktenPraesentation, Auftrag, Bild, Ergebnis, Liste, Loesung, Notiz, Schema, SQL, Tabelle,
    abfrage_ergebnis, warmup_karten_dokument)

MAT = REPO / "materialien"
ORDNER = MAT / "Akte5_Bauplan"
DB = "detektei"
TMP = Path(tempfile.mkdtemp(prefix="dd-akte5-"))
AKTE5 = json.loads((REPO / "seed" / "akten" / "akte5.json").read_text(encoding="utf-8"))
AUFGABEN = {t["id"]: t for t in AKTE5["aufgaben"]}


def ergebnis(sql, max_zeilen=None, db=DB, breiten=None):
    """Ergebnistabelle direkt aus der DB (optional gekürzt mit „…“)."""
    kopf, zeilen = abfrage_ergebnis(sql, db)
    if max_zeilen and len(zeilen) > max_zeilen:
        return Ergebnis(kopf, zeilen[:max_zeilen], abgeschnitten=True, breiten=breiten)
    return Ergebnis(kopf, zeilen, breiten=breiten)


def auftraege(d, liste):
    """Ermittlungsaufträge; in der Schülerfassung bleibt jede Zeile auf einer Seite."""
    tab = d.ermittlungsauftraege(liste)
    if d.fassung != "schueler":
        return tab
    for zeile in tab.rows[1:]:
        A._zeile_nicht_teilen(zeile)
    return tab


def mit_be(text, be):
    """Abi-Text mit „[x BE]“, das nicht umbricht (geschützte Leerzeichen)."""
    return f"{text}\u00a0[{be}\u00a0BE]"


def loesung_sql(sql, *rest, max_zeilen=None, fehler=(), varianten=()):
    return Loesung(SQL(sql), ergebnis(sql, max_zeilen), *rest, fehler=fehler, varianten=varianten)


# ================================================================== ER-Diagramme (Chen-Notation nach EH §7)
def ent(i, name, *attribute):
    """Entitätstyp; Schlüsselattribute mit führendem „_“."""
    return {"id": i, "name": name,
            "attribute": [{"name": a.lstrip("_"), "schluessel": a.startswith("_")} for a in attribute]}


def bez(i, name, enden, *attribute):
    """Beziehungstyp; enden = [(entitaet, kard) | (entitaet, kard, rolle)]."""
    e = []
    for x in enden:
        d = {"entitaet": x[0], "kard": x[1]}
        if len(x) > 2:
            d["rolle"] = x[2]
        e.append(d)
    return {"id": i, "name": name, "enden": e, "attribute": [{"name": a} for a in attribute]}


#: Versatz der Attribut-Ellipsen zum Mittelpunkt des Besitzers (Bauplan-Einheiten), je Anzahl; „oben“, unten = gespiegelt
_RING = {
    1: [(0, -92)],
    2: [(-70, -82), (70, -82)],
    3: [(-82, -82), (0, -128), (82, -82)],
    4: [(-150, -74), (-62, -128), (62, -128), (150, -74)],
    5: [(-178, -70), (-128, -136), (0, -142), (128, -136), (178, -70)],
}
_SEITE = {  # links/rechts: Attribute untereinander neben dem Rechteck
    2: [(-1, -26), (-1, 26)],
    3: [(-1, -50), (-1, 0), (-1, 50)],
    4: [(-1, -75), (-1, -25), (-1, 25), (-1, 75)],
    5: [(-1, -100), (-1, -50), (-1, 0), (-1, 50), (-1, 100)],
}


def platz(layout, graph, besitzer, x, y, seite="oben", versatz=None, abstand=None):
    """Setzt einen Entitäts-/Beziehungstyp nach (x, y) und seine Attribute auf die gewünschte Seite."""
    layout[besitzer] = {"x": x, "y": y}
    el = next(e for e in graph["entitaeten"] + graph["beziehungen"] if e["id"] == besitzer)
    n = len(el["attribute"])
    if not n:
        return
    if versatz is None:
        if seite in ("oben", "unten"):
            versatz = [(dx, dy if seite == "oben" else -dy) for dx, dy in _RING[n]]
        else:
            ab = abstand or 205
            vz = -1 if seite == "links" else 1
            versatz = [(vz * ab, dy) for _, dy in _SEITE[n]] if n > 1 else [(vz * ab, 0)]
    for i, (dx, dy) in enumerate(versatz):
        layout[f"{besitzer}.{i}"] = {"x": x + dx, "y": y + dy}


def _rand_punkte(e, p, n=64):
    """Punkte auf dem Umriss eines Elements (Ellipse, Raute, Rechteck)."""
    hb, hh = A._ausdehnung(e)
    aus = []
    for i in range(n):
        w = 2 * math.pi * i / n
        c, s = math.cos(w), math.sin(w)
        if e["art"] == "attribut":
            aus.append((p["x"] + hb * c, p["y"] + hh * s))
        elif e["art"] == "beziehung":
            t = 1 / (abs(c) / hb + abs(s) / hh)
            aus.append((p["x"] + t * c, p["y"] + t * s))
        else:
            t = min(hb / abs(c) if c else math.inf, hh / abs(s) if s else math.inf)
            aus.append((p["x"] + t * c, p["y"] + t * s))
    return aus


def _innen(e, p, q, schrumpf=0.9):
    """Liegt Punkt q im Element e (Mittelpunkt p)? Ellipse/Raute/Rechteck, leicht verkleinert."""
    hb, hh = A._ausdehnung(e)
    hb, hh = hb * schrumpf, hh * schrumpf
    dx, dy = abs(q[0] - p["x"]), abs(q[1] - p["y"])
    if e["art"] == "attribut":
        return (dx / hb) ** 2 + (dy / hh) ** 2 < 1
    if e["art"] == "beziehung":
        return dx / hb + dy / hh < 1
    return dx < hb and dy < hh


def pruefe_layout(graph, layout, name="ER"):
    """Keine Überlappung von Elementen, keine Linie durch fremde Elemente, Kardinalitäten frei."""
    el = A._er_elemente(graph)
    fehler = []
    for i, a in enumerate(el):
        for b in el[i + 1:]:
            if any(_innen(b, layout[b["id"]], q, schrumpf=1.08) for q in _rand_punkte(a, layout[a["id"]])) or \
                    any(_innen(a, layout[a["id"]], q, schrumpf=1.08) for q in _rand_punkte(b, layout[b["id"]])):
                fehler.append(f"{a['id']} überlappt {b['id']}")
    ent_ids = {e["id"] for e in el if e["art"] == "entitaet"}
    strecken = []
    for e in el:
        if e["art"] == "attribut":
            strecken.append((layout[e["id"]], layout[e["besitzer"]["id"]], {e["id"], e["besitzer"]["id"]}))
    plaettchen = []
    for e in el:
        if e["art"] == "beziehung":
            for k in A.er_kanten(e["el"], layout, ent_ids):
                strecken.append(({"x": k["start"][0], "y": k["start"][1]}, {"x": k["bis"][0], "y": k["bis"][1]},
                                 {e["id"], k["ende"]["entitaet"]}))
                plaettchen.append((k["mitte"], e["id"]))
    for p, q, eigene in strecken:
        for e in el:
            if e["id"] in eigene:
                continue
            for t in range(1, 40):
                pt = (p["x"] + (q["x"] - p["x"]) * t / 40, p["y"] + (q["y"] - p["y"]) * t / 40)
                if _innen(e, layout[e["id"]], pt):
                    fehler.append(f"Linie {sorted(eigene)} durch {e['id']}")
                    break
    for m, bid in plaettchen:
        for e in el:
            if e["art"] == "attribut" and _innen(e, layout[e["id"]], m, schrumpf=1.25):
                fehler.append(f"Kardinalität an {bid} verdeckt {e['id']}")
    if fehler:
        raise AssertionError(f"Layout {name}: " + "; ".join(sorted(set(fehler))))
    return True


def gross(er, breite_cm):
    """Setzt die Breite eines ER-Diagramms im Dokument (cm)."""
    er.breite_cm = breite_cm
    return er


def diagramm(graph, layout, name, breite_cm=None, markiert=()):
    pruefe_layout(graph, layout, name)
    return ER(graph, layout, breite_cm=breite_cm, markiert=markiert)


def _aus_akte(aufgabe, pfad):
    """ER-Block (graph, layout) aus seed/akten/akte5.json – identisch mit dem Lab."""
    o = AUFGABEN[aufgabe] if aufgabe in AUFGABEN else None
    for schritt in pfad:
        o = o[schritt]
    return json.loads(json.dumps(o["graph"])), json.loads(json.dumps(o["layout"]))


# -- D1 Legende: Ermittler –leitet– Fall (wie Intro A5-E1)
def er_leitet():
    b = AKTE5["levels"][0]["intro"][0]
    return diagramm(json.loads(json.dumps(b["graph"])), json.loads(json.dumps(b["layout"])), "leitet")


# -- D2 Person –verdaechtig– Fall mit Grad
def er_verdacht(mit_grad=True, kard=True):
    g = {"entitaeten": [ent("e1", "Person", "_PersonNr", "Vorname", "Nachname"),
                        ent("e2", "Fall", "_FallNr", "Titel", "Status")],
         "beziehungen": [bez("b1", "verdaechtig", [("e1", "n" if kard else ""), ("e2", "m" if kard else "")],
                             *(["Grad"] if mit_grad else []))]}
    L = {}
    platz(L, g, "e1", 156, 174)
    platz(L, g, "e2", 656, 174)
    platz(L, g, "b1", 406, 174, versatz=[(0, 92)])
    return diagramm(g, L, "verdacht")


# -- D3 Fall –gehoertZu– Beweisstueck –gefundenAn– Ort (wie A5-E1-05)
def er_fbo():
    g, L = _aus_akte("A5-E1-05", ["inhalt", 1])
    return diagramm(g, L, "fbo")


# -- D4 Abi-Ecke 5.1: Ermittler –leitet– Fall –verdaechtig– Person
def er_fallverwaltung():
    g = {"entitaeten": [ent("e1", "Ermittler", "_ErmittlerNr", "Name", "Dienstgrad"),
                        ent("e2", "Fall", "_FallNr", "Titel", "Eroeffnet", "Status"),
                        ent("e3", "Person", "_PersonNr", "Vorname", "Nachname", "Jahrgang")],
         "beziehungen": [bez("b1", "leitet", [("e1", "1"), ("e2", "n")]),
                         bez("b2", "verdaechtig", [("e3", "n"), ("e2", "m")], "Grad")]}
    L = {}
    platz(L, g, "e1", 170, 190)
    platz(L, g, "b1", 420, 190)
    platz(L, g, "e2", 670, 190)
    platz(L, g, "b2", 920, 190, versatz=[(0, 92)])
    platz(L, g, "e3", 1170, 190)
    return diagramm(g, L, "fallverwaltung")


# -- D5 Lösung A5-E1-07 (Startlayout des Lab + Kardinalitäten + Fundzeitpunkt)
def er_e107(loesung=True):
    it = AUFGABEN["A5-E1-07"]["interaktion"]
    g, L = json.loads(json.dumps(it["start"]["graph"])), json.loads(json.dumps(it["start"]["layout"]))
    kard = {"b1": ("1", "n"), "b2": ("1", "n"), "b3": ("n", "1")}
    for b in g["beziehungen"]:
        for e, k in zip(b["enden"], kard[b["id"]]):
            e["kard"] = k if loesung else ""
        if loesung and b["id"] == "b3":
            b["attribute"] = [{"name": "Fundzeitpunkt"}]
            L["b3.0"] = {"x": 1078 - 190, "y": 328}
    return diagramm(g, L, "e107")


# -- D6 rekursiv 1:n: Ermittler betreut Ermittler (wie A5-E2-01)
def er_betreut():
    g, L = _aus_akte("A5-E2-01", ["inhalt", 1])
    return diagramm(g, L, "betreut")


# -- D7 rekursiv n:m: Person kennt Person (Seit)
def er_kennt():
    g = {"entitaeten": [ent("e1", "Person", "_PersonNr", "Vorname", "Nachname")],
         "beziehungen": [bez("b1", "kennt", [("e1", "n"), ("e1", "m")], "Seit")]}
    L = {}
    platz(L, g, "e1", 169, 245, seite="unten")
    platz(L, g, "b1", 169, 65, versatz=[(150, 0)])
    return diagramm(g, L, "kennt")


# -- D8 zwei Beziehungstypen: Team spieltHeim/spieltGast Spiel (wie Intro A5-E2)
def er_heim_gast():
    b = AKTE5["levels"][1]["intro"][0]
    return diagramm(json.loads(json.dumps(b["graph"])), json.loads(json.dumps(b["layout"])), "heimgast")


# -- D9 ternär: Person, Fall, Ort –sagtAus– (wie A5-E3-07)
def er_sagt_aus():
    g, L = _aus_akte("A5-E3-07", ["inhalt", 1])
    L["b1.0"], L["b1.1"] = {"x": 632, "y": 332}, {"x": 632, "y": 398}  # Lab: Ellipsen berühren sich knapp
    for k, y in (("e3", 520), ("e3.0", 600), ("e3.1", 600)):  # Ort näher heran: kompakter, auf Folien größer
        L[k]["y"] = y
    return diagramm(g, L, "sagtaus")


# -- D10 Referenzlösung des Architekturwettbewerbs (= Fall-Datenbank der Detektei)
def er_referenz():
    g = {"entitaeten": [ent("e1", "Ermittler", "_ErmittlerNr", "Name", "Dienstgrad"),
                        ent("e2", "Fall", "_FallNr", "Titel", "Eroeffnet", "Status"),
                        ent("e3", "Person", "_PersonNr", "Vorname", "Nachname", "Jahrgang", "Stadtteil"),
                        ent("e4", "Ort", "_OrtNr", "Bezeichnung", "Stadtteil"),
                        ent("e5", "Beweisstueck", "_BeweisNr", "Bezeichnung")],
         "beziehungen": [bez("b1", "betreut", [("e1", "1", "Mentor"), ("e1", "n", "Neuling")]),
                         bez("b2", "leitet", [("e1", "1"), ("e2", "n")]),
                         bez("b3", "verdaechtig", [("e3", "n"), ("e2", "m")], "Grad"),
                         bez("b4", "kennt", [("e3", "n"), ("e3", "m")], "Seit"),
                         bez("b5", "sagtAus", [("e3", "m"), ("e2", "n"), ("e4", "p")], "Zeitpunkt", "Inhalt"),
                         bez("b6", "gehoertZu", [("e2", "1"), ("e5", "n")]),
                         bez("b7", "gefundenAn", [("e4", "1"), ("e5", "n")])]}
    L = {}
    platz(L, g, "e1", 190, 330, seite="unten")
    platz(L, g, "b1", 190, 150)
    platz(L, g, "b2", 470, 330)
    platz(L, g, "e2", 760, 330)
    platz(L, g, "b6", 1060, 330)
    platz(L, g, "e5", 1340, 330)
    platz(L, g, "b3", 560, 560, versatz=[(-150, 0)])
    platz(L, g, "e3", 760, 790, seite="links", abstand=230)
    platz(L, g, "b4", 760, 1010, versatz=[(170, 0)])
    platz(L, g, "b5", 1060, 600, versatz=[(112, -95), (0, 100)])
    platz(L, g, "e4", 1340, 790, seite="unten")
    platz(L, g, "b7", 1340, 560)
    # Attribute von Fall über dem Rechteck, die von Beweisstück oben
    for i, (dx, dy) in enumerate(_RING[4]):
        L[f"e2.{i}"] = {"x": 760 + dx, "y": 330 + dy}
    for i, (dx, dy) in enumerate(_RING[2]):
        L[f"e5.{i}"] = {"x": 1340 + dx, "y": 330 + dy}
    return diagramm(g, L, "referenz")


# -- D11 Erweiterung Vorgängerfall (rekursiv 1:n an Fall)
def er_vorgaenger():
    g = {"entitaeten": [ent("e1", "Fall", "_FallNr", "Titel", "Eroeffnet", "Status")],
         "beziehungen": [bez("b1", "nimmtWiederAuf", [("e1", "1", "Vorgänger"), ("e1", "n", "Wiederaufnahme")])]}
    L = {}
    platz(L, g, "e1", 230, 245, seite="unten")
    platz(L, g, "b1", 230, 65)
    return diagramm(g, L, "vorgaenger")


# -- D12 Abi-Ecke 5.2 (Lösung) bzw. gegebenes Modell der Abi-Ecke 5.3: vertritt, arbeitetMitAn, leitet
def er_erweiterung():
    g = {"entitaeten": [ent("e1", "Ermittler", "_ErmittlerNr", "Name", "Dienstgrad"),
                        ent("e2", "Fall", "_FallNr", "Titel", "Eroeffnet")],
         "beziehungen": [bez("b1", "vertritt", [("e1", "1", "Vertretung"), ("e1", "n", "Vertretene")]),
                         bez("b2", "leitet", [("e1", "1"), ("e2", "n")]),
                         bez("b3", "arbeitetMitAn", [("e1", "n"), ("e2", "m")], "Rolle")]}
    L = {}
    platz(L, g, "e1", 220, 300, seite="unten")
    platz(L, g, "b1", 220, 110)
    platz(L, g, "b2", 520, 190)
    platz(L, g, "b3", 520, 410, versatz=[(0, 90)])
    platz(L, g, "e2", 820, 300, seite="rechts", abstand=200)
    return diagramm(g, L, "erweiterung")


# -- D13 Cold Case 5.2: ternär mit zwei Kanten zu Person
def er_komplize():
    g = {"entitaeten": [ent("e1", "Person", "_PersonNr", "Vorname", "Nachname"),
                        ent("e2", "Fall", "_FallNr", "Titel")],
         "beziehungen": [bez("b1", "istKomplizeVon", [("e1", "m", "Komplize"), ("e1", "n", "Hauptperson"), ("e2", "p")])]}
    L = {}
    platz(L, g, "e1", 300, 110)
    L["e1.0"], L["e1.1"], L["e1.2"] = {"x": 70, "y": 70}, {"x": 70, "y": 150}, {"x": 530, "y": 110}
    platz(L, g, "b1", 300, 330)
    platz(L, g, "e2", 300, 540, seite="unten")
    return diagramm(g, L, "komplize")


# -- D14 Lösung AB 5.3/3: Bibliothek rückwärts
def er_bibliothek():
    g = {"entitaeten": [ent("e1", "Leser", "_LeserNr", "Name", "Jahrgang"),
                        ent("e2", "Buch", "_BuchNr", "Titel", "Erscheinungsjahr"),
                        ent("e3", "Autor", "_AutorNr", "Name")],
         "beziehungen": [bez("b1", "leiht", [("e1", "n"), ("e2", "m")], "Rueckgabe"),
                         bez("b2", "verfasst", [("e3", "1"), ("e2", "n")])]}
    L = {}
    platz(L, g, "e1", 170, 190)
    platz(L, g, "b1", 430, 190, versatz=[(0, 92)])
    platz(L, g, "e2", 700, 190, versatz=[(-112, -82), (0, -140), (122, -82)])
    platz(L, g, "b2", 960, 190)
    platz(L, g, "e3", 1220, 190)
    return diagramm(g, L, "bibliothek")


# -- D15 Folien-Beispiel DS 9: Schueler –besucht– Klasse (eigene Daten, nicht Detektei)
def er_schule(kard=True):
    g = {"entitaeten": [ent("e1", "Schueler", "_SchuelerNr", "Name"), ent("e2", "Klasse", "_KlassenNr", "Raum")],
         "beziehungen": [bez("b1", "besucht", [("e1", "n" if kard else ""), ("e2", "1" if kard else "")])]}
    L = {}
    platz(L, g, "e1", 160, 150)
    platz(L, g, "b1", 410, 150)
    platz(L, g, "e2", 660, 150)
    return diagramm(g, L, "schule")


# -- D16 Folien-Beispiel DS 10: Tierheim (eigener Anforderungstext)
def er_tierheim():
    g = {"entitaeten": [ent("e1", "Pflegekraft", "_PflegeNr", "Name"),
                        ent("e2", "Tier", "_TierNr", "Name", "Art"),
                        ent("e3", "Helfer", "_AusweisNr", "Name")],
         "beziehungen": [bez("b1", "arbeitetEin", [("e1", "1", "Erfahrene"), ("e1", "n", "Neue")]),
                         bez("b2", "versorgt", [("e1", "1"), ("e2", "n")]),
                         bez("b3", "fuehrtAus", [("e3", "n"), ("e2", "m")])]}
    L = {}
    platz(L, g, "e1", 190, 330, seite="unten")
    platz(L, g, "b1", 190, 150)
    platz(L, g, "b2", 470, 330)
    platz(L, g, "e2", 750, 330, seite="unten")
    platz(L, g, "b3", 1020, 330)
    platz(L, g, "e3", 1290, 330, seite="unten")
    return diagramm(g, L, "tierheim")


# -- D17 AB 5.1/1: Ermittler –befragt– Person (Datum am Beziehungstyp)
def er_befragt():
    g = {"entitaeten": [ent("e1", "Ermittler", "_ErmittlerNr", "Name", "Dienstgrad"),
                        ent("e2", "Person", "_PersonNr", "Name", "Stadtteil")],
         "beziehungen": [bez("b1", "befragt", [("e1", "n"), ("e2", "m")], "Datum")]}
    L = {}
    platz(L, g, "e1", 170, 180)
    platz(L, g, "b1", 430, 180, versatz=[(0, 92)])
    platz(L, g, "e2", 690, 180)
    return diagramm(g, L, "befragt")


# -- D18 AB 5.2/4: Wettbewerbsentwurf von Team „Hafenkante“ (absichtlich mit Schwächen)
def er_hafenkante():
    g = {"entitaeten": [ent("e1", "Mentor", "_MentorNr", "Name"),
                        ent("e2", "Ermittler", "_ErmittlerNr", "Name", "Dienstgrad"),
                        ent("e3", "Fall", "_FallNr", "Titel", "Status"),
                        ent("e4", "Person", "_PersonNr", "Name", "Grad"),
                        ent("e5", "Ort", "Bezeichnung", "Stadtteil")],
         "beziehungen": [bez("b1", "betreut", [("e1", "1"), ("e2", "n")]),
                         bez("b2", "leitet", [("e2", "1"), ("e3", "n")]),
                         bez("b3", "verdaechtig", [("e4", "n"), ("e3", "m")]),
                         bez("b4", "sagtAus", [("e4", "n"), ("e3", "m")], "Zeitpunkt"),
                         bez("b5", "betrifft", [("e3", "n"), ("e5", "m")])]}
    L = {}
    platz(L, g, "e1", 170, 150)
    platz(L, g, "b1", 170, 330)
    platz(L, g, "e2", 170, 510, seite="unten")
    platz(L, g, "b2", 450, 510)
    platz(L, g, "e3", 730, 510, seite="unten")
    platz(L, g, "b3", 730, 330, versatz=[])
    platz(L, g, "b4", 1010, 330, versatz=[(170, 0)])
    platz(L, g, "e4", 870, 150)
    platz(L, g, "b5", 1010, 510)
    platz(L, g, "e5", 1290, 510, seite="unten")
    return diagramm(g, L, "hafenkante")


# -- D19 AB 5.3/2: Person –verdaechtig– Fall und Person –kennt– Person
def er_verdacht_kennt():
    g = {"entitaeten": [ent("e1", "Person", "_PersonNr", "Vorname", "Nachname", "Stadtteil"),
                        ent("e2", "Fall", "_FallNr", "Titel", "Status")],
         "beziehungen": [bez("b1", "verdaechtig", [("e1", "n"), ("e2", "m")], "Grad"),
                         bez("b2", "kennt", [("e1", "n", "Person"), ("e1", "m", "Bekannte")], "Seit")]}
    L = {}
    platz(L, g, "e1", 260, 300, seite="unten")
    platz(L, g, "b2", 260, 110, versatz=[(170, 0)])
    platz(L, g, "b1", 560, 300, versatz=[(0, 92)])
    platz(L, g, "e2", 860, 300, seite="unten")
    return diagramm(g, L, "verdacht_kennt")


ALLE_DIAGRAMME = [er_befragt, er_hafenkante, er_verdacht_kennt, er_leitet, er_verdacht, er_fbo, er_fallverwaltung, er_e107, er_betreut, er_kennt, er_heim_gast,
                  er_sagt_aus, er_referenz, er_vorgaenger, er_erweiterung, er_komplize, er_bibliothek, er_schule,
                  er_tierheim]


# ================================================================== Grafiken: Mengendiagramme, Bilder nebeneinander
def mengen_png(panels, pfad, massstab=3.0):
    """Mengendiagramme (Pillow): panels = [(titel, (name_l, [elemente]), (name_r, [elemente]), [(i, j), …])].
    Elemente als Punkte in zwei Ovalen, Linien = Beziehungen zwischen einzelnen Entitäten."""
    from PIL import Image, ImageDraw
    s = massstab
    pw, ph = 330, 300
    bild = Image.new("RGB", (int(pw * len(panels) * s), int(ph * s)), "#FFFFFF")
    d = ImageDraw.Draw(bild)
    tinte, rot, grau = "#" + A.TINTE, "#" + A.SIEGELROT, "#" + A.RAND_STARK
    f_titel, f_name, f_el = A._font("kopf_fett", 19 * s), A._font("kopf", 16 * s), A._font("text", 14 * s)
    for pi, (titel, (nl, el_l), (nr, el_r), linien) in enumerate(panels):
        ox = pi * pw * s
        if titel:
            b = f_titel.getlength(titel)
            d.text((ox + (pw * s - b) / 2, 6 * s), titel, font=f_titel, fill=rot)
        punkte = []
        for seite, (name, el) in enumerate(((nl, el_l), (nr, el_r))):
            cx = ox + (82 if seite == 0 else 248) * s
            oben, unten = 66 * s, (ph - 16) * s
            d.ellipse([cx - 54 * s, oben, cx + 54 * s, unten], outline=tinte, width=int(2 * s))
            b = f_name.getlength(name)
            d.text((cx - b / 2, 36 * s), name, font=f_name, fill=tinte)
            pos = []
            n = len(el)
            for i, e in enumerate(el):
                y = oben + (unten - oben) * (i + 1) / (n + 1)
                px = cx + (22 if seite == 0 else -22) * s
                d.ellipse([px - 4 * s, y - 4 * s, px + 4 * s, y + 4 * s], fill=tinte)
                bb = f_el.getlength(e)
                tx = px - 9 * s - bb if seite == 0 else px + 9 * s
                d.text((tx, y - 9 * s), e, font=f_el, fill=tinte)
                pos.append((px, y))
            punkte.append(pos)
        for i, j in linien:
            d.line([punkte[0][i], punkte[1][j]], fill=rot, width=int(2.2 * s))
        if pi:
            d.line([(ox, 20 * s), (ox, (ph - 10) * s)], fill=grau, width=1)
    bild.save(pfad, dpi=(int(96 * s), int(96 * s)))
    return pfad


def er_bild(er, breite_cm, name):
    """ER-Diagramm als Bild-Datei (für nebeneinander gesetzte Diagramme)."""
    pfad = TMP / f"{name}.png"
    A.er_png(er.graph, er.layout, pfad)
    return pfad


def er_folie(er, breite_cm, name=None):
    """ER-Diagramm als Bild mit fester Breite auf Papiergrund (Folien: ER nutzt sonst die ganze Breite)."""
    pfad = TMP / f"{name or 'f' + str(id(er))}.png"
    A.er_png(er.graph, er.layout, pfad, hintergrund=A.PAPIER)
    return Bild(str(pfad), breite_cm=breite_cm)


def nebeneinander(pfade, name, abstand=40):
    """Mehrere PNGs nebeneinander (gleiche Höhe, unten bündig) zu einem Bild."""
    from PIL import Image
    bilder = [Image.open(p) for p in pfade]
    h = max(b.height for b in bilder)
    w = sum(b.width for b in bilder) + abstand * (len(bilder) - 1)
    neu = Image.new("RGB", (w, h), "#FFFFFF")
    x = 0
    for b in bilder:
        neu.paste(b, (x, (h - b.height) // 2))
        x += b.width + abstand
    pfad = TMP / f"{name}.png"
    neu.save(pfad, dpi=(288, 288))
    return pfad


#: Mengendiagramme der Beweisstück-Kästen (AB 5.1 BS 2, Folie DS 9) – eigene Beispieldaten
MENGEN_BEISPIEL = [
    ("1 : 1  hat", ("Ermittler", ["E1", "E2", "E3"]), ("Dienstausweis", ["D1", "D2", "D3"]), [(0, 1), (1, 0), (2, 2)]),
    ("1 : n  leitet", ("Ermittler", ["E1", "E2", "E3"]), ("Fall", ["F1", "F2", "F3", "F4", "F5"]),
     [(0, 0), (0, 2), (1, 1), (1, 3), (1, 4)]),
    ("n : m  verdaechtig", ("Person", ["P1", "P2", "P3", "P4"]), ("Fall", ["F1", "F2", "F3"]),
     [(0, 0), (1, 0), (1, 1), (2, 1), (2, 2), (3, 2)]),
]
#: Auftrag 2 von AB 5.1 – drei neue Mengendiagramme (ohne Kardinalität)
MENGEN_AUFTRAG = [
    ("a)  gehoertZu", ("Fall", ["F1", "F2", "F3"]), ("Beweisstück", ["B1", "B2", "B3", "B4", "B5"]),
     [(0, 0), (0, 1), (1, 2), (2, 3), (2, 4)]),
    ("b)  nutzt", ("Ermittler", ["E1", "E2", "E3", "E4"]), ("Schließfach", ["S1", "S2", "S3", "S4"]),
     [(0, 1), (1, 0), (2, 3), (3, 2)]),
    ("c)  besucht", ("Ermittler", ["E1", "E2", "E3"]), ("Lehrgang", ["L1", "L2", "L3"]),
     [(0, 0), (0, 1), (1, 0), (2, 1), (2, 2)]),
]


def mengen_bild(panels, name, breite_cm=16.5):
    return Bild(str(mengen_png(panels, TMP / f"{name}.png")), breite_cm=breite_cm)


# ================================================================== AB 5.1 „Baupläne lesen“ (DS 9)
def ab1(fassung):
    lsg = fassung == "loesung"
    d = AktenDokument("ab", 5, "Baupläne lesen", fassung=fassung, ermittlung="5.1", ds=9, material="AB 5.1",
                      datenbank=DB, untertitel="Chen-Notation nach EH §7, Kardinalitäten, Attribute an Beziehungstypen · nur eA")
    d.kopfband()
    d.fall_einstieg("Die Pinnwand der Detektei ist voller Fäden, und Juri hat die rote Wolle aufgebraucht. Edda bringt einen "
                    "alten Bauplan aus dem Polizeiarchiv mit – ihr müsst ihn lesen können, bevor ihr selbst einen entwerft.")
    d.beweisstueck(1, "Chen-Notation nach EH §7",
                   "Ein **ER-Diagramm** (Entity-Relationship-Diagramm) ist der Bauplan einer Datenbank. Der alte Plan aus "
                   "dem Archiv:",
                   gross(er_leitet(), 10.5),
                   Tabelle(["Zeichen", "Fachbegriff", "im Bauplan"],
                           [["Rechteck", "**Entitätstyp** (Singular)", "Ermittler, Fall"],
                            ["Raute", "**Beziehungstyp** (Verb, lowerCamelCase)", "*leitet*"],
                            ["Ellipse", "**Attribut**", "Name, Titel, Status"],
                            ["unterstrichene Ellipse", "**Schlüsselattribut**", "<u>ErmittlerNr</u>, <u>FallNr</u>"],
                            ["1, n, m an den Kanten", "**Kardinalität**", "1 an der Kante zu Ermittler, n an der Kante zu Fall"]],
                           breiten=[4.2, 5.8, 6.6]),
                   "Ein Entitätstyp steht für alle gleichartigen Dinge (alle Fälle der Detektei), eine **Entität** ist ein "
                   "einzelnes davon (der Fall „Das Fundbüro“). Wir sagen Entitätstyp – nicht „Entitätsmenge“.")
    d.beweisstueck(2, "Kardinalitäten 1:1, 1:n, n:m",
                   "Ein **Mengendiagramm** zeigt einzelne Entitäten als Punkte und jede Beziehung als Linie (eigene "
                   "Beispieldaten):",
                   mengen_bild(MENGEN_BEISPIEL, "mengen_bs2", breite_cm=13),
                   Tabelle(["Kardinalität", "Leseschablone – immer in **beide** Richtungen"],
                           [["Ermittler **1** – *hat* – **1** Dienstausweis",
                             "Ein Ermittler hat genau einen Dienstausweis; jeder Dienstausweis gehört genau einem Ermittler."],
                            ["Ermittler **1** – *leitet* – **n** Fall",
                             "Ein Ermittler kann mehrere (n) Fälle leiten; jeder Fall wird von genau einem (1) Ermittler "
                             "geleitet. (E3 leitet gerade keinen Fall – das ist erlaubt.)"],
                            ["Person **n** – *verdaechtig* – **m** Fall",
                             "Eine Person kann in mehreren Fällen verdächtig sein; in einem Fall können mehrere Personen "
                             "verdächtig sein."]], breiten=[5.6, 11.0]),
                   Notiz("Wo steht die Zahl?",
                         "Die Kardinalität steht an der Kante des Entitätstyps, **über den du sprichst**: „Ein Ermittler "
                         "leitet *n Fälle*“ → n an der Kante zu Fall. Wer sie an die eigene Seite schreibt, dreht das Modell um."))
    d.beweisstueck(3, "Kardinalität aus Tabellendaten ableiten",
                   Liste(["Suche die Spalte, die zwei Entitätstypen verbindet (Fremdschlüssel bzw. Beziehungstabelle).",
                          "Kommt **derselbe Wert der einen Seite** mehrfach vor, gehören zu ihm mehrere Entitäten der "
                          "anderen Seite → dort steht **n**.",
                          "Prüfe **beide** Richtungen – erst dann steht die Kardinalität fest."], nummeriert=True),
                   "Beispiel (eigene Daten): In `Dienstausweis` kommt jede ErmittlerNr genau einmal vor.",
                   Tabelle(["AusweisNr", "Ausgestellt", "ErmittlerNr"],
                           [["D-17", "2025-01-02", "1"], ["D-18", "2025-03-14", "3"], ["D-19", "2025-03-14", "4"]],
                           pk=("AusweisNr",), fk=("ErmittlerNr",), breiten=[2.8, 3.2, 3.0]),
                   Notiz("Daten stützen – begründet wird im Sachzusammenhang",
                         "Wie beim Primärschlüssel gilt: Ein Ausschnitt beweist nichts. Dass jede ErmittlerNr *hier* nur "
                         "einmal vorkommt, **stützt** 1:1. Begründet wird die Kardinalität mit der Organisation der "
                         "Detektei: Jede Ermittlerin bekommt genau einen Ausweis."))
    d.beweisstueck(4, "Attribute an Beziehungstypen – der Verdachtsgrad",
                   "Person P3 ist in Fall F2 *hoch* verdächtig, in Fall F5 nur *gering*. Der Grad gehört also weder zur "
                   "Person (sie hat zwei verschiedene Grade) noch zum Fall (jede verdächtige Person hat einen eigenen). "
                   "Er hängt vom **Paar** Person–Fall ab und wird deshalb als Ellipse an die **Raute** gezeichnet:",
                   gross(er_verdacht(), 9.5),
                   "Merkfrage: *Ändert sich der Wert, wenn ich eine andere Person **oder** einen anderen Fall nehme?* "
                   "Wenn ja → Attribut des Beziehungstyps.")

    a1_inhalt = [gross(er_befragt(), 9.0)]
    if not lsg:
        a1_inhalt.append(Tabelle(["Element im Bauplan", "Fachbegriff"],
                                 [["a) Rechteck *Person*", ""], ["b) Raute *befragt*", ""],
                                  ["c) Ellipse *Dienstgrad*", ""], ["d) unterstrichene Ellipse *PersonNr*", ""],
                                  ["e) Ellipse *Datum* an der Raute", ""], ["f) *m* an der Kante zu Person", ""]],
                                 breiten=[6.6, 5.6]))
    a4_inhalt = [Tabelle(*abfrage_ergebnis("SELECT ErmittlerNr, Name, Dienstgrad\nFROM Ermittler\n"
                                           "WHERE ErmittlerNr IN (1, 4, 5, 7, 8, 12);", DB), pk=("ErmittlerNr",),
                         breiten=[2.6, 3.6, 2.6]),
                 Tabelle(*abfrage_ergebnis("SELECT FallNr, Titel, Status, LeiterNr\nFROM Fall\nWHERE FallNr >= 5;", DB),
                         pk=("FallNr",), fk=("LeiterNr",), breiten=[1.6, 5.8, 2.8, 2.2])]
    beleg = "SELECT LeiterNr, COUNT(*) AS Faelle\nFROM Fall\nGROUP BY LeiterNr\nHAVING COUNT(*) > 1;"
    auftraege(d, [
        Auftrag(1, 1, "PAPIER", "Ordne den Elementen des Bauplans die Fachbegriffe nach den Ergänzenden Hinweisen zu.",
                kompetenzen=["K33", "K47"], inhalt=a1_inhalt,
                loesung=Loesung("a) Entitätstyp · b) Beziehungstyp · c) Attribut · d) Schlüsselattribut · e) Attribut des "
                                "Beziehungstyps *befragt* (das Datum hängt vom Paar Ermittler–Person ab) · f) Kardinalität: "
                                "Ein Ermittler kann mehrere Personen befragen, eine Person kann von mehreren Ermittlern "
                                "befragt werden (n:m).",
                                fehler=[("P7", "„Entitätsmenge“ statt Entitätstyp oder „Beziehung“ statt Beziehungstyp – "
                                               "im Abitur werden die Begriffe der EH erwartet."),
                                        ("P9", "Datum als Attribut von Person gedeutet – dann hätte jede Person nur ein "
                                               "Befragungsdatum.")])),
        Auftrag(2, 1, "PAPIER", "Bestimme in den drei Mengendiagrammen die Kardinalität und formuliere sie in **beiden** "
                "Richtungen.", kompetenzen=["K34"],
                inhalt=[mengen_bild(MENGEN_AUFTRAG, "mengen_a2", breite_cm=11.5)] + ([] if lsg else [
                    Tabelle(["", "Kardinalität", "Satz in beiden Richtungen"],
                            [["a)", "", ""], ["b)", "", ""], ["c)", "", ""]], breiten=[0.8, 2.6, 8.8])]),
                loesung=Loesung("a) **Fall 1 : n Beweisstück** – Ein Fall kann mehrere Beweisstücke haben; jedes "
                                "Beweisstück gehört zu genau einem Fall.",
                                "b) **Ermittler 1 : 1 Schließfach** – Jeder Ermittler nutzt genau ein Schließfach; jedes "
                                "Schließfach wird von genau einem Ermittler genutzt.",
                                "c) **Ermittler n : m Lehrgang** – Ein Ermittler kann mehrere Lehrgänge besuchen (E1, E3); "
                                "ein Lehrgang wird von mehreren Ermittlern besucht (L1, L2).",
                                fehler=[("P9", "Nur eine Richtung geprüft: Bei c) sieht E2 mit einer Linie wie 1:n aus – "
                                               "entscheidend ist, dass E1 **und** L1 mehrere Linien haben."),
                                        ("P9", "a) als n:1 notiert, aber das n an die Kante zu Fall geschrieben "
                                               "(Kardinalität an der falschen Kante).")])),
        Auftrag(3, 2, "PAPIER", "**Erläutere** das ER-Diagramm im Sachzusammenhang der Detektei – jede Kardinalität in "
                "beiden Richtungen. Beweisstück 5, die Nachtfalter-Skizze, hat keinen bekannten Fundort: Was sagt das "
                "Diagramm dazu?", kompetenzen=["K33"], inhalt=[er_fbo()], linien=7,
                loesung=Loesung("Entitätstypen **Fall** (Schlüsselattribut FallNr), **Beweisstueck** (BeweisNr) und "
                                "**Ort** (OrtNr).",
                                "*gehoertZu* (1:n): Zu einem Fall können mehrere Beweisstücke gehören; jedes Beweisstück "
                                "gehört zu genau einem Fall.",
                                "*gefundenAn* (n:1): An einem Ort können mehrere Beweisstücke gefunden werden; jedes "
                                "Beweisstück wurde an genau einem Ort gefunden.",
                                "Nachtfalter-Skizze: Die 1 an der Kante zu Ort sagt nur „höchstens ein Fundort“ je "
                                "Beweisstück; dass der Fundort auch **fehlen** darf, zeigt die Notation 1/n/m nicht – das "
                                "gehört als Anmerkung zum Modell (in der Tabelle: OrtNr NULL).",
                                fehler=[("P10", "Kardinalitäten nur genannt („gehoertZu ist 1:n“) – Erläutern verlangt "
                                                "den Sachzusammenhang in ganzen Sätzen."),
                                        ("P9", "Leserichtung vertauscht: „Ein Beweisstück gehört zu mehreren Fällen.“")])),
        Auftrag(4, 2, "PAPIER", "Leite aus den Tabellenausschnitten Ermittler und Fall die Kardinalität von *leitet* ab. "
                "Begründe in beiden Richtungen und erkläre, was die Daten **nicht** beweisen können.",
                kompetenzen=["K34"], inhalt=a4_inhalt, linien=5,
                loesung=Loesung("**Ermittler 1 : n Fall.** Dieselbe LeiterNr kommt mehrfach vor (Imke Tiemann, Nr. 5, "
                                "leitet die Fälle 7, 8 und 11) → ein Ermittler kann mehrere Fälle leiten (n an der Kante "
                                "zu Fall). In jeder Fall-Zeile steht genau eine LeiterNr → jeder Fall hat genau eine "
                                "Leitung (1 an der Kante zu Ermittler). Dass ein Ermittler gar keinen Fall leitet (wie "
                                "Juri Brandes), ist bei 1:n ebenfalls erlaubt.",
                                "Die Daten **stützen** das nur: Dass ein Fall nie zwei Leitungen bekommt, folgt aus der "
                                "Organisation der Detektei („genau eine Leitung je Fall“), nicht aus acht Zeilen. Beleg "
                                "auf der ganzen Fall-Datenbank:", SQL(beleg), ergebnis(beleg),
                                fehler=[("P9", "„Ermittler 5 kommt dreimal vor, also n:m“ – die zweite Richtung (ein Fall "
                                               "→ ein Leiter) wurde nicht geprüft."),
                                        ("P8", "„Im Ausschnitt leitet Ermittler 12 nur einen Fall, also 1:1.“ – ein "
                                               "Ausschnitt beweist keine Kardinalität.")])),
        Auftrag(5, 3, "LAB", "Bauplan-Tisch, Ermittlung 5.1, Aufgabe A5-E1-07 „Fehlende Kardinalitäten“: Setze in Juris "
                "Bauplan alle Kardinalitäten und ergänze das Attribut Fundzeitpunkt. Notiere danach hier, an welches "
                "Element du Fundzeitpunkt gehängt hast, und begründe in einem Satz.",
                kompetenzen=["K34", "K36"], linien=4,
                loesung=Loesung(er_e107(),
                                "Fundzeitpunkt hängt am Beziehungstyp *gefundenAn*: Er beschreibt das Finden eines "
                                "bestimmten Beweisstücks an einem bestimmten Ort. Weil *gefundenAn* 1:n ist, akzeptiert "
                                "das Lab auch Fundzeitpunkt an Beweisstueck (n-Seite) – bei n:m ginge das nicht.",
                                fehler=[("P9", "Fundzeitpunkt an Ort – an einem Ort werden viele Beweisstücke zu "
                                               "verschiedenen Zeiten gefunden."),
                                        ("P9", "*leitet* als n:m gesetzt, weil Ermittler mehrere Fälle haben (zweite "
                                               "Richtung vergessen).")])),
    ])
    d.abi_ecke(mit_be("Gegeben ist das folgende ER-Diagramm zur Fallverwaltung einer Detektei. **Begründen** Sie die "
               "Kardinalitäten der Beziehungstypen *leitet* und *verdaechtig* im Sachzusammenhang.", 4), be=4, operator="begründen", afb="II", kompetenzen=["K34", "K33"], inhalt=[er_fallverwaltung()], linien=8,
               loesung=Loesung("*leitet* (Ermittler 1 : n Fall): Jeder Fall wird von genau einer Ermittlerin bzw. einem "
                               "Ermittler verantwortlich geleitet, damit klar ist, wer zuständig ist **(1 BE)**; eine "
                               "Ermittlerin kann gleichzeitig mehrere Fälle leiten **(1 BE)**.",
                               "*verdaechtig* (Person n : m Fall): Eine Person kann in mehreren Fällen verdächtig sein "
                               "(z. B. bei einer Serie) **(1 BE)**; in einem Fall können mehrere Personen verdächtig sein "
                               "**(1 BE)**. Der Grad hängt vom Paar ab und steht deshalb an der Raute.",
                               "**Bewertung:** je Richtung 1 BE, nur mit Bezug auf die Detektei – Begründen heißt „weil“.",
                               fehler=[("P10", "„1:n, weil an der Kante 1 und n stehen“ – das ist Ablesen, keine "
                                               "Begründung."),
                                       ("P9", "Leserichtung vertauscht („Ein Fall leitet mehrere Ermittler“).")]))
    d.spur_im_lab("5.1 „Baupläne lesen“", "A5-E1-01 … A5-E1-08 (★ → ★★ → ★★★)",
                  werkzeug="Bauplan-Tisch (Modi lesen und ergänzen; Listenmodus auf dem iPad)",
                  text="A5-E1-08 belegt mit SQL auf der Fall-Datenbank, dass Personen in mehreren Fällen verdächtig sind. "
                       "Wer fertig ist: Cold Case A5-E1-CC „1:1 oder 1:n?“ im Lab.")
    d.cold_case("Dienstwagen", "Die Detektei schafft Dienstwagen an. Entscheide begründet, welche Kardinalität der "
                "Beziehungstyp *faehrt* zwischen Ermittler und Dienstwagen hat – für zwei Organisationsformen: (A) Jede "
                "Ermittlerin und jeder Ermittler bekommt einen persönlichen Wagen. (B) Die Senior-Ermittler verwalten den "
                "Fuhrpark: Jeder Wagen ist genau einer Senior-Person zugeordnet, die mehrere Wagen (Kombi, Lastenrad …) "
                "betreuen kann.", kompetenzen=["K34", "K40"], linien=5,
                loesung=Loesung("(A) **1:1** – ein Wagen je Person, jeder Wagen gehört genau einer Person. (B) **Ermittler "
                                "1 : n Dienstwagen** – eine Senior-Person verwaltet mehrere Wagen, jeder Wagen hat genau "
                                "eine verantwortliche Person (n an der Kante zu Dienstwagen).",
                                "Dieselbe Fuhrpark-Liste kann zu beiden Modellen passen; entscheidend ist, was die "
                                "Detektei festlegt. Teilen sich mehrere Personen im Schichtbetrieb die Wagen, wird daraus "
                                "n:m (Lab-Cold-Case A5-E1-CC)."))
    name = "5_Loesungen/DD_A5_AB1_Bauplaene_lesen_LSG.docx" if lsg else "2_Arbeitsblaetter/DD_A5_AB1_Bauplaene_lesen.docx"
    return d.speichern(ORDNER / name)


# ================================================================== AB 5.2 „Der Architekturwettbewerb“ (DS 10)
ANFORDERUNGSTEXT = ("„Die Detektei verwaltet ihre **Ermittler** (Nummer, Name, Dienstgrad). Erfahrene Ermittler **betreuen** "
                    "als Mentor mehrere Neulinge; jeder Neuling hat höchstens einen Mentor. Jeder **Fall** (Nummer, Titel, "
                    "Eröffnungsdatum, Status) wird von genau einem Ermittler **geleitet**. **Personen** (Nummer, Vorname, "
                    "Nachname, Jahrgang, Stadtteil) können in mehreren Fällen **verdächtig** sein, jeweils mit einem "
                    "Verdachtsgrad. Personen **kennen** einander seit einem bestimmten Datum. Personen **sagen** zu einem "
                    "Fall an einem **Ort** (Nummer, Bezeichnung, Stadtteil) zu einem Zeitpunkt **aus**; der Inhalt wird "
                    "festgehalten. **Beweisstücke** (Nummer, Bezeichnung) gehören zu genau einem Fall und wurden an genau "
                    "einem Ort gefunden.“")
#: Kriterien des Bewertungsbogens (DS_PLANUNG DS 10, Unplugged) – zugleich Kernpunkte der Referenzlösung
KRITERIEN = ["alle 5 Entitätstypen (Ermittler, Fall, Person, Ort, Beweisstück) mit Attributen",
             "Schlüsselattribute markiert (unterstrichen)",
             "*leitet* 1:n richtig (n an der Kante zu Fall)",
             "*betreut* rekursiv 1:n mit Rollen Mentor/Neuling",
             "*kennt* rekursiv n:m mit Attribut Seit",
             "*verdaechtig* n:m mit Attribut Grad an der Raute",
             "*sagtAus* ternär (Person, Fall, Ort) mit Zeitpunkt und Inhalt",
             "Beweisstück 1:n zu Fall und 1:n zu Ort",
             "Fachsprache und Notation nach EH §7 (Verben, 1/n/m, Singular)",
             "Lesbarkeit (wenige Kreuzungen, klare Beschriftung)"]


def ab2(fassung):
    lsg = fassung == "loesung"
    d = AktenDokument("ab", 5, "Der Architekturwettbewerb", fassung=fassung, ermittlung="5.2", ds=10, material="AB 5.2",
                      datenbank=DB, untertitel="ER-Modell entwickeln und erweitern: rekursiv, mehrere Beziehungstypen, "
                                               "ternär, Grenzen der Notation · nur eA")
    d.kopfband()
    d.fall_einstieg("Edda Lorenz schreibt einen Wettbewerb aus: Das Team mit dem besten Bauplan für die Fall-Datenbank "
                    "bekommt den Stempel ARCHITEKTUR – und sein Entwurf wird gebaut.")
    d.beweisstueck(1, "Vom Text zum Modell",
                   Liste(["**Substantive** → Entitätstypen (Singular) mit ihren Attributen; je Entitätstyp ein "
                          "**Schlüsselattribut** festlegen und unterstreichen.",
                          "**Verben** → Beziehungstypen in lowerCamelCase, z. B. *versorgt* oder *istPateVon*.",
                          "**Mengenangaben** („genau ein“, „höchstens ein“, „mehrere“) → Kardinalitäten, in **beiden** "
                          "Richtungen lesen.",
                          "Angaben, die vom **Paar** (oder Trio) abhängen → Attribute an der Raute.",
                          "Kontrolle: Jede Information des Textes ist im Modell abgehakt."], nummeriert=True),
                   Tabelle(["Textstelle (Tierheim)", "im Modell"],
                           [["„Jedes **Tier** (Nummer, Name, Art) …“", "Entitätstyp Tier, Schlüsselattribut TierNr"],
                            ["„… wird von genau einer Pflegekraft **versorgt** …“", "Beziehungstyp *versorgt*, 1 an der Kante zu Pflegekraft"],
                            ["„… eine Pflegekraft versorgt mehrere Tiere.“", "n an der Kante zu Tier"]],
                           breiten=[8.4, 8.2]))
    rek = Bild(str(nebeneinander([er_bild(er_betreut(), 6, "bs2a"), er_bild(er_kennt(), 6, "bs2b")], "bs2")),
               breite_cm=9.5)
    d.beweisstueck(2, "Rekursive Beziehungen",
                   "Ein **rekursiver** Beziehungstyp verbindet einen Entitätstyp **mit sich selbst**: zwei Kanten zum "
                   "selben Rechteck, die **Rollen** stehen an den Kanten. Für „Mentor“ entsteht **kein** neuer Entitätstyp – "
                   "Mentor und Neuling sind beide Ermittler.", rek,
                   Tabelle(["Beziehungstyp", "Leseschablone"],
                           [["*betreut* (1:n)", "Ein Ermittler kann als Mentor mehrere Neulinge betreuen; jeder Neuling hat "
                             "höchstens einen Mentor."],
                            ["*kennt* (n:m)", "Eine Person kann mehrere Personen kennen und von mehreren gekannt werden. "
                             "Seit hängt vom Paar ab → Attribut an der Raute."]], breiten=[3.6, 13.0]))
    d.beweisstueck(3, "Mehrere Beziehungstypen zwischen denselben Entitätstypen",
                   "Spielen zwei Entitäten in **verschiedenen Rollen** zusammen, braucht jede Rolle einen eigenen "
                   "Beziehungstyp: Ein Spiel hat ein Heim- **und** ein Gastteam.", gross(er_heim_gast(), 8.0),
                   "Genauso in der Detektei: Eine Person kann einen Fall *melden* und in einem (anderen) Fall *verdaechtig* "
                   "sein – zwei Rauten zwischen Person und Fall.")
    d.beweisstueck(4, "Ternäre Beziehung m:n:p",
                   "Eine Raute, **drei** Kanten: *sagtAus* hält fest, **wer** in **welchem Fall** an **welchem Ort** "
                   "ausgesagt hat. Zeitpunkt und Inhalt hängen vom Trio ab.", gross(er_sagt_aus(), 8.0),
                   "Drei binäre Beziehungen (Person–Fall, Fall–Ort, Person–Ort) reichen nicht: Sagt eine Person in zwei "
                   "Fällen an zwei Orten aus, lässt sich daraus nicht mehr ablesen, wo sie zu welchem Fall ausgesagt hat.")
    d.beweisstueck(5, "Grenzen der Notation und Beurteilen",
                   Tabelle(["Anforderung", "mit 1/n/m?"],
                           [["„Jeder Fall hat **genau eine** Leitung.“", "ja – die 1 an der Kante"],
                            ["„Eine Person kann in **beliebig vielen** Fällen verdächtig sein.“", "ja – n bzw. m"],
                            ["„Jedes Team hat **genau drei** Ermittler.“", "nein – Anmerkung zum Modell"],
                            ["„Ein Mentor muss einen **höheren Dienstgrad** haben.“", "nein – Anmerkung zum Modell"]],
                           breiten=[11.0, 5.6]),
                   "**Beurteilen** heißt: Alternativen benennen (z. B. Ort als Attribut oder als eigener Entitätstyp), "
                   "ihre Folgen abwägen (Redundanz, Auswertbarkeit, Vollständigkeit) und ein begründetes Urteil fällen.")

    a1 = [Tabelle(["ANFORDERUNGSTEXT DES WETTBEWERBS"], [[ANFORDERUNGSTEXT]], breiten=[12.0])]
    if not lsg:
        a1.append("Farben: Entitätstyp ______ · Beziehungstyp ______ · Attribut ______ · Mengenangabe ______")
    auftraege(d, [
        Auftrag(1, 1, "PAPIER", "Markiere im Anforderungstext **Entitätstypen**, **Beziehungstypen**, **Attribute** und "
                "**Mengenangaben** in vier verschiedenen Farben.", kompetenzen=["K35"], inhalt=a1,
                loesung=Loesung("**Entitätstypen:** Ermittler, Fall, Person, Ort, Beweisstück · **Beziehungstypen:** "
                                "betreuen, leiten, verdächtig sein, kennen, aussagen, gehören zu, gefunden an · "
                                "**Attribute:** Nummer, Name, Dienstgrad, Titel, Eröffnungsdatum, Status, Vorname, Nachname, "
                                "Jahrgang, Stadtteil, Bezeichnung, Verdachtsgrad, Datum (Seit), Zeitpunkt, Inhalt · "
                                "**Mengenangaben:** mehrere Neulinge, höchstens einen Mentor, genau einem Ermittler, in "
                                "mehreren Fällen, genau einem Fall, genau einem Ort.",
                                fehler=[("P9", "Verdachtsgrad, Seit, Zeitpunkt, Inhalt als Attribute eines Entitätstyps "
                                               "markiert – sie gehören zu Beziehungen.")])),
        Auftrag(2, 2, "PAPIER", "**Entwickelt** im Team das ER-Modell zum Anforderungstext auf dem Plakat (A2). Skizziere "
                "hier die Beziehungstypen, über die ihr im Team gestritten habt, und eure Entscheidung.",
                kompetenzen=["K35", "K37", "K38"], linien=6,
                loesung=Loesung("Referenzlösung (= Fall-Datenbank der Detektei, im Lab nach A5-E2-07 sichtbar):",
                                gross(er_referenz(), 15.0),
                                "Kernpunkte = Kriterien des Bewertungsbogens: " + " · ".join(KRITERIEN[:8]) + ".",
                                fehler=[("P9", "„Mentor“ als eigener Entitätstyp statt rekursiver Beziehung."),
                                        ("P9", "*sagtAus* in drei binäre Beziehungen zerlegt."),
                                        ("P9", "Grad an Person, Seit an Person, Schlüsselattribute nicht unterstrichen.")])),
        Auftrag(3, 2, "PAPIER", "Neue Anforderung: „Ein Fall kann als **Wiederaufnahme** eines früheren Falls eröffnet "
                "werden; derselbe frühere Fall kann mehrmals wieder aufgenommen werden.“ **Erweitere** das Modell (nur "
                "der betroffene Ausschnitt).", kompetenzen=["K39", "K37"], linien=6,
                loesung=Loesung(gross(er_vorgaenger(), 6.5),
                                "Rekursiver Beziehungstyp an Fall mit den Rollen Vorgänger (1) und Wiederaufnahme (n): Ein "
                                "Fall hat höchstens einen Vorgängerfall, ein Vorgängerfall kann mehrere Wiederaufnahmen "
                                "haben. (Wer begründet „jeder Fall wird höchstens einmal wieder aufgenommen“ annimmt, "
                                "modelliert 1:1 – mit dem Text oben ist 1:n richtig.)",
                                fehler=[("P9", "Neuer Entitätstyp „Vorgängerfall“ – die Rolle gehört an die Kante."),
                                        ("P9", "Attribut VorgaengerNr an Fall – Fremdschlüssel gibt es erst im "
                                               "Relationenschema, im ER-Diagramm ist es eine Beziehung.")])),
        Auftrag(4, 3, "PAPIER", "Im Galerie-Rundgang hängt der Entwurf von Team „Hafenkante“. **Beurteile** ihn mit den "
                "Kriterien des Bewertungsbogens: Was ist gelungen, was fehlt oder ist falsch? Schlage eine Verbesserung vor "
                "und begründe dein Urteil.", kompetenzen=["K40"], inhalt=[gross(er_hafenkante(), 12.0)], linien=7,
                loesung=Loesung("**Gelungen:** *leitet* 1:n und *verdaechtig* n:m richtig; Schlüssel bei Ermittler, Fall, "
                                "Person markiert.",
                                "**Schwächen:** (1) „Mentor“ als eigener Entitätstyp – ein Mentor ist selbst Ermittler, "
                                "seine Daten stünden doppelt → rekursiver Beziehungstyp *betreut* an Ermittler. (2) Grad an "
                                "Person – der Grad hängt vom Paar Person–Fall ab → an *verdaechtig*. (3) *sagtAus* nur "
                                "binär, Ort über *betrifft* – wer wo zu welchem Fall aussagte, geht verloren → ternär. "
                                "(4) Ort ohne Schlüsselattribut. (5) *kennt* und Beweisstück fehlen.",
                                "**Urteil:** Grundgerüst brauchbar, aber für die Ermittlung zu ungenau (etwa 4 von 10 "
                                "Kriterien); Verbesserung zuerst bei sagtAus und Grad, weil sonst Ermittlungsdaten fehlen.",
                                fehler=[("P10", "Nur Fehler aufgezählt, ohne Folgen und ohne begründetes Gesamturteil – "
                                                "Beurteilen verlangt alle drei Schritte.")])),
        Auftrag(5, 3, "LAB", "Bauplan-Tisch, Ermittlung 5.2, Aufgabe A5-E2-07 „Die Ermittlerschule“: Entwickle das "
                "ER-Modell aus dem Anforderungstext. Notiere, welche Kriterien dir der Bauplan-Tisch zuerst als offen "
                "gemeldet hat.", kompetenzen=["K35"], linien=3,
                loesung=Loesung("Entitätstypen Ermittler, Lehrgang, Raum; *leitet* (Ermittler 1 : n Lehrgang), *besucht* "
                                "(Ermittler n : m Lehrgang, Note an der Raute), *betreut* (rekursiv 1:n, Rollen "
                                "Mentor/Neuling), *findetStatt* (Raum 1 : n Lehrgang). Häufig offen: Note an Ermittler, "
                                "fehlende Rollen, zweiter Beziehungstyp Ermittler–Lehrgang.")),
    ])
    d.abi_ecke(mit_be("Gegeben ist das ER-Modell der Detektei (Ausschnitt). Die Detektei möchte zusätzlich erfassen, welche "
               "Ermittler gemeinsam an einem Fall arbeiten, wobei jeweils die Rolle (z. B. Spurensicherung) festgehalten "
               "wird, und welcher Ermittler welchen anderen vertritt. Jeder Ermittler hat höchstens eine feste Vertretung, "
               "kann selbst aber mehrere Kolleginnen und Kollegen vertreten. **Erweitern** Sie das gegebene ER-Modell "
               "entsprechend.", 5), be=5, operator="erweitern", afb="II/III", kompetenzen=["K39", "K37", "K36"],
               inhalt=[gross(er_leitet(), 10.5)], linien=9,
               loesung=Loesung(gross(er_erweiterung(), 12.0),
                               "**Bewertung:** *arbeitetMitAn* als n:m mit Kardinalitäten **(1 BE)**, Attribut Rolle an der "
                               "Raute **(1 BE)**; *vertritt* als rekursiver Beziehungstyp an Ermittler **(1 BE)** mit 1:n "
                               "**(1 BE)** und Rollen an den Kanten; Notation nach EH §7 (Verben, Kanten, keine neuen "
                               "Entitätstypen für Rollen) **(1 BE)**.",
                               fehler=[("P9", "Rolle als Attribut von Ermittler – jeder Ermittler hätte in allen Fällen "
                                              "dieselbe Rolle."),
                                       ("P9", "Entitätstyp „Vertretung“ statt rekursiver Beziehung; *arbeitetMitAn* mit "
                                              "*leitet* zusammengelegt (zwei Rollen = zwei Beziehungstypen).")]))
    d.spur_im_lab("5.2 „Der Architekturwettbewerb“", "A5-E2-01 … A5-E2-06 in der Stunde, A5-E2-07/08 als Hausaufgabe",
                  werkzeug="Bauplan-Tisch (Modi ergänzen und entwickeln)",
                  text="Die Referenzlösung des Wettbewerbs erscheint im Lab erst, wenn du A5-E2-07 gelöst hast. Cold Case "
                       "im Lab: A5-E2-CC „Das Phantom-Netz“.")
    d.cold_case("Das Phantom-Netz", "Edda vermutet, dass das Phantom nicht allein arbeitet: „Eine Person ist **in einem "
                "bestimmten Fall** Komplizin bzw. Komplize einer anderen Person.“ Modelliere diese Anforderung. Welche Art "
                "von Beziehung entsteht – und warum reicht ein zweiter rekursiver Beziehungstyp an Person ohne Fall nicht?",
                kompetenzen=["K37", "K38"], linien=6,
                loesung=Loesung(gross(er_komplize(), 6.5),
                                "Ein **ternärer** Beziehungstyp mit zwei Kanten zu Person (Rollen Komplize/Hauptperson) und "
                                "einer zu Fall – rekursiv und ternär zugleich. Ohne Fall ginge verloren, **in welchem** Fall "
                                "jemand geholfen hat (dieselben zwei Personen können in einem Fall Komplizen sein, in einem "
                                "anderen nicht)."))
    name = ("5_Loesungen/DD_A5_AB2_Der_Architekturwettbewerb_LSG.docx" if lsg
            else "2_Arbeitsblaetter/DD_A5_AB2_Der_Architekturwettbewerb.docx")
    return d.speichern(ORDNER / name)


# ================================================================== AB 5.3 „Vom Plan zur Datenbank“ (DS 11)
#: Überführungsregeln R1–R6 am Schulbeispiel (wie die Intro-Notiz im Lab – nicht an der Detektei, sonst stünden die
#: Lösungen der Aufträge 1, 2 und der Abi-Ecke schon im Beweisstück)
REGELN_SCHULE = [
    ["**R1** Entitätstyp", "Rechteck Kurs", "Kurs(<u>KursNr</u>, Titel) – Schlüsselattribut → Primärschlüssel"],
    ["**R2** 1:n", "Lehrkraft 1 – *leitet* – n Kurs", "Kurs(<u>KursNr</u>, Titel, ↑LehrkraftNr) – Fremdschlüssel auf der "
     "**n-Seite**; Attribute des Beziehungstyps ebenfalls dorthin"],
    ["**R3** n:m", "Schueler n – *belegt* – m Kurs (Note)", "belegt(<u>↑SchuelerNr</u>, <u>↑KursNr</u>, Note) – eigene "
     "Relation, Primärschlüssel aus beiden Fremdschlüsseln"],
    ["**R4** 1:1", "Lehrkraft 1 – *nutzt* – 1 Schliessfach", "Schliessfach(<u>FachNr</u>, Etage, ↑LehrkraftNr) – "
     "Fremdschlüssel auf **einer** Seite (oder beide Relationen zusammenlegen)"],
    ["**R5** rekursiv", "Schueler 1 – *istPateVon* – n Schueler", "Schueler(<u>SchuelerNr</u>, Name, ↑PateNr) – "
     "Fremdschlüssel auf die eigene Relation, benannt nach der Rolle; rekursiv n:m → eigene Relation mit zwei "
     "Fremdschlüsseln auf dieselbe Relation"],
    ["**R6** ternär", "Schueler, Kurs, Raum – *praesentiert*", "praesentiert(<u>↑SchuelerNr</u>, <u>↑KursNr</u>, "
     "<u>↑RaumNr</u>, Thema) – eigene Relation mit drei Fremdschlüsseln im Primärschlüssel"],
]
SELFJOIN = ("SELECT e.Name AS Ermittler, m.Name AS Mentor\nFROM Ermittler e, Ermittler m\nWHERE e.MentorNr = m.ErmittlerNr\n"
            "ORDER BY m.Name, e.Name;")


def ab3(fassung):
    lsg = fassung == "loesung"
    d = AktenDokument("ab", 5, "Vom Plan zur Datenbank", fassung=fassung, ermittlung="5.3", ds=11, material="AB 5.3",
                      datenbank=DB, untertitel="Überführungsregeln R1–R6, Begründungen, Rückrichtung Schema → ER, "
                                               "Self-Join · nur eA")
    d.kopfband()
    d.fall_einstieg("Der Bauplan ist genehmigt. Jetzt müsst ihr ihn in Tabellen gießen – und zwar so, dass das Phantom "
                    "keine Anomalie mehr ausnutzen kann.")
    d.beweisstueck(1, "Überführungsregeln R1–R6 (Beispiel Schule)",
                   Tabelle(["Regel", "im ER-Modell", "im Relationenschema (Kurzschreibweise)"], REGELN_SCHULE,
                           breiten=[2.6, 4.6, 9.4]),
                   "Kurzschreibweise nach EH §7: Primärschlüssel unterstrichen, Fremdschlüssel mit ↑ direkt vor dem Namen.")
    d.beweisstueck(2, "Begründen: Warum n:m eine eigene Relation braucht – 1:n nicht",
                   Liste(["**1:n:** Auf der n-Seite gehört zu jeder Zeile **genau ein** Partner (jeder Kurs hat genau eine "
                          "Lehrkraft) – ein Fremdschlüssel mit **einem** Wert passt in die Zeile.",
                          "**n:m:** Auf **keiner** Seite passt ein einzelner Wert: Eine KursNr-Spalte in Schueler bräuchte "
                          "mehrere Werte in einer Zelle (nicht atomar) – oder die Schülerdaten würden je Kurs wiederholt "
                          "(**Redundanz** → Änderungs-, Einfüge- und Löschanomalien).",
                          "**Attribute des Beziehungstyps** wandern mit: bei 1:n in die Relation der n-Seite (z. B. "
                          "Kurs(…, ↑LehrkraftNr, LeitetSeit)), bei n:m in die neue Relation (Note in belegt)."]))
    buchung = Tabelle(["SpielerNr", "PlatzNr", "Beginn"],
                      [["7", "2", "2026-10-05 17:00"], ["7", "2", "2026-10-12 17:00"], ["9", "1", "2026-10-05 18:00"]],
                      pk=("SpielerNr", "PlatzNr", "Beginn"), fk=("SpielerNr", "PlatzNr"), breiten=[2.6, 2.4, 4.2])
    d.beweisstueck(3, "Zusätzliches Schlüsselattribut bei mehrfach vorkommenden Paaren",
                   "Tennishalle (eigenes Beispiel): Spieler n – *bucht* – m Platz mit dem Attribut Beginn. Spieler 7 bucht "
                   "Platz 2 **jede Woche** – das Paar (7, 2) kommt mehrfach vor und kann allein kein Primärschlüssel sein.",
                   buchung,
                   "Lösung: Ein weiteres Attribut kommt in den Schlüssel – bucht(<u>↑SpielerNr</u>, <u>↑PlatzNr</u>, "
                   "<u>Beginn</u>) – oder die Relation bekommt eine eigene Nummer: Buchung(<u>BuchungNr</u>, ↑SpielerNr, "
                   "↑PlatzNr, Beginn). Prüffrage: *Kann dieselbe Kombination zweimal vorkommen?*")
    d.beweisstueck(4, "Rückrichtung: vom Schema zum ER-Diagramm",
                   Liste(["Relation mit **eigenem** Schlüssel → **Entitätstyp**, ihre Nicht-Fremdschlüssel-Attribute → "
                          "Attribute, der Primärschlüssel → Schlüsselattribut.",
                          "**Fremdschlüssel** in einer Entitätsrelation → Beziehungstyp **1:n** (die Relation mit dem "
                          "Fremdschlüssel ist die n-Seite); zeigt er auf die eigene Relation → **rekursiv**.",
                          "Relation, deren Schlüssel aus **zwei** (drei) Fremdschlüsseln besteht → **n:m** (ternär); ihre "
                          "übrigen Attribute → Attribute an der Raute.",
                          "Fremdschlüssel werden im ER-Diagramm **keine** Attribute. Den Namen des Beziehungstyps (ein "
                          "Verb) wählst du aus dem Sachzusammenhang."]))
    d.beweisstueck(5, "Self-Join über einen rekursiven Fremdschlüssel",
                   "Wer betreut wen? Die Tabelle Ermittler wird **zweimal** mit verschiedenen Aliasen verwendet – `e` für "
                   "die betreute Person, `m` für den Mentor – und über `e.MentorNr = m.ErmittlerNr` verbunden:",
                   SQL(SELFJOIN), ergebnis(SELFJOIN, max_zeilen=5),
                   "(10 Zeilen.) Edda Lorenz und Ulf Witte fehlen als „Ermittler“: Ihre MentorNr ist NULL, der Vergleich "
                   "ist nie wahr. Ohne Aliase wäre `Name` mehrdeutig (P7).")

    sql4 = "SELECT Fall.FallNr, Fall.Titel, Ermittler.Name\nFROM Fall, Ermittler\nWHERE Fall.LeiterNr = Ermittler.ErmittlerNr\nORDER BY Fall.FallNr;"
    a3 = [Schema("Leser(<u>LeserNr</u>, Name, Jahrgang)"), Schema("Autor(<u>AutorNr</u>, Name)"),
          Schema("Buch(<u>BuchNr</u>, Titel, Erscheinungsjahr, ↑AutorNr)"),
          Schema("leiht(<u>↑LeserNr</u>, <u>↑BuchNr</u>, Rueckgabe)")]
    a5 = [Schema("Person(<u>PersonNr</u>, Vorname, Nachname, Stadtteil, ↑FallNr, Grad)")]
    if not lsg:
        a5.append(Tabelle(["PersonNr", "Vorname", "Nachname", "Stadtteil", "FallNr", "Grad"],
                          [["", "", "", "", "", ""]] * 3, breiten=[2.0, 2.0, 2.4, 2.6, 1.6, 1.6]))
    auftraege(d, [
        Auftrag(1, 1, "PAPIER", "**Überführe** das ER-Modell in ein Relationenschema in Kurzschreibweise.",
                kompetenzen=["K41", "K04"], inhalt=[gross(er_fbo(), 11.5)], linien=4,
                loesung=Loesung(Schema("Fall(<u>FallNr</u>, Titel, Eroeffnet, Status)"),
                                Schema("Ort(<u>OrtNr</u>, Bezeichnung, Stadtteil)"),
                                Schema("Beweisstueck(<u>BeweisNr</u>, Bezeichnung, ↑FallNr, ↑OrtNr)"),
                                "Beide Beziehungstypen sind 1:n mit Beweisstueck auf der n-Seite → beide Fremdschlüssel "
                                "in Beweisstueck (R2).",
                                fehler=[("P9", "BeweisNr als Fremdschlüssel in Fall bzw. Ort – dann bräuchte ein Fall "
                                               "mehrere Werte in einer Zelle."),
                                        ("P9", "Eigene Relationen gehoertZu/gefundenAn – bei 1:n überflüssig (kein "
                                               "Fehler im Sinne der Daten, aber nicht die Regel).")])),
        Auftrag(2, 2, "PAPIER", "**Überführe** *verdaechtig* und *kennt* (Person und Fall gibt es schon) und **begründe** "
                "deine Wahl der Primärschlüssel.", kompetenzen=["K41", "K42"],
                inhalt=[gross(er_verdacht_kennt(), 11.5)], linien=6,
                loesung=Loesung(Schema("verdaechtig(<u>↑PersonNr</u>, <u>↑FallNr</u>, Grad)"),
                                Schema("kennt(<u>↑PersonNr</u>, <u>↑BekannterNr</u>, Seit)"),
                                "Beide sind n:m → eigene Relationen (R3, bei *kennt* rekursiv: R5). Eine Person ist in "
                                "einem Fall nur **einmal** verdächtig und zwei Personen kennen sich nur **einmal** → das "
                                "Paar der Fremdschlüssel ist eindeutig und minimal; einzeln wäre keiner eindeutig. "
                                "BekannterNr ist nach der Rolle benannt, weil PersonNr nicht zweimal vorkommen darf; "
                                "beide verweisen auf Person.",
                                fehler=[("P9", "kennt(<u>↑PersonNr</u>, ↑PersonNr, Seit) – gleicher Name zweimal, "
                                               "Schlüssel nur halb."),
                                        ("P9", "Grad in Person oder Fall statt in verdaechtig.")])),
        Auftrag(3, 2, "PAPIER", "**Rekonstruiere** aus dem Relationenschema einer Stadtbücherei das ER-Diagramm "
                "(Entitätstypen, Beziehungstypen mit Kardinalitäten, Attribute).", kompetenzen=["K43"], inhalt=a3,
                linien=7,
                loesung=Loesung(gross(er_bibliothek(), 13.0),
                                "Leser, Buch, Autor haben eigene Schlüssel → Entitätstypen. AutorNr in Buch → *verfasst* "
                                "1:n (Buch auf der n-Seite). leiht hat einen Schlüssel aus zwei Fremdschlüsseln → n:m, "
                                "Rueckgabe an der Raute. (Der Name des Beziehungstyps ist frei wählbar, z. B. *schreibt*.)",
                                fehler=[("P9", "AutorNr als Attribut von Buch eingezeichnet – Fremdschlüssel werden "
                                               "Beziehungen."),
                                        ("P9", "leiht als Entitätstyp mit Kanten 1:n – nur als „aufgelöste "
                                               "Zwischenentität“ gleichwertig, in der Abiturnotation erwartet: n:m.")])),
        Auftrag(4, 2, "LAB", "Gib in der Konsole alle Fälle mit FallNr, Titel und dem Namen der leitenden Ermittlerin bzw. "
                "des leitenden Ermittlers aus (Fremdschlüssel LeiterNr), sortiert nach FallNr.",
                kompetenzen=["K23", "K03"], linien=5,
                loesung=loesung_sql(sql4, "(12 Zeilen, gekürzt.)", max_zeilen=5, fehler=[
                    ("P2", "Verbundbedingung vergessen – 12 · 12 = 144 Zeilen, jeder Fall mit jedem Ermittler."),
                    ("P7", "`Fall.LeiterNr = Ermittler.LeiterNr` – die Spalte heißt in Ermittler ErmittlerNr.")])),
        Auftrag(5, 3, "PAPIER", "Juri setzt *verdaechtig* „sparsam“ um. Trage Lasse Engelken (Person 42, Weserufer, "
                "verdächtig in Fall 1 und Fall 3, jeweils gering) und Jarne Kampen (Person 41, Hafenviertel, Fall 2, gering) "
                "in Juris Schema ein. **Begründe** an deinen Zeilen, warum das Schema die Beziehung nicht korrekt umsetzt, und nenne zwei "
                "Anomalien.", kompetenzen=["K42", "K08"], inhalt=a5, linien=6,
                loesung=Loesung(Tabelle(["PersonNr", "Vorname", "Nachname", "Stadtteil", "FallNr", "Grad"],
                                        [["42", "Lasse", "Engelken", "Weserufer", "1", "gering"],
                                         ["42", "Lasse", "Engelken", "Weserufer", "3", "gering"],
                                         ["41", "Jarne", "Kampen", "Hafenviertel", "2", "gering"]],
                                        breiten=[1.8, 1.6, 2.1, 2.4, 1.4, 1.5]),
                                "Lasse braucht zwei Zeilen: PersonNr 42 ist nicht mehr eindeutig (Primärschlüssel "
                                "verletzt) – oder FallNr müsste „1, 3“ enthalten (nicht atomar). Die Personendaten stehen "
                                "doppelt (Redundanz).",
                                "**Änderungsanomalie:** Zieht Lasse um, muss Stadtteil in allen seinen Zeilen geändert "
                                "werden, sonst widersprechen sich die Daten. **Löschanomalie:** Wird der Verdacht in Fall 2 "
                                "aufgehoben und die Zeile gelöscht, ist Jarne Kampen ganz verschwunden. **Einfügeanomalie:** "
                                "Eine Person ohne Verdacht lässt sich nur mit FallNr NULL speichern.",
                                "Korrekt: verdaechtig(<u>↑PersonNr</u>, <u>↑FallNr</u>, Grad) – R3.",
                                fehler=[("P10", "Nur „Redundanz“ genannt, ohne Bezug auf die eingetragenen Zeilen.")])),
    ])
    d.abi_ecke(mit_be("Gegeben ist das ER-Modell einer Erweiterung der Fall-Datenbank. **Überführen** Sie das Modell in ein "
               "Relationenschema in Kurzschreibweise und **begründen** Sie die Umsetzung des Beziehungstyps "
               "*arbeitetMitAn*.", 5), be=5, operator="überführen (erstellen) / begründen", afb="II/III",
               kompetenzen=["K41", "K42"], inhalt=[gross(er_erweiterung(), 11.5)], linien=8,
               loesung=Loesung(Schema("Ermittler(<u>ErmittlerNr</u>, Name, Dienstgrad, ↑VertretungNr)"),
                               Schema("Fall(<u>FallNr</u>, Titel, Eroeffnet, ↑LeiterNr)"),
                               Schema("arbeitetMitAn(<u>↑ErmittlerNr</u>, <u>↑FallNr</u>, Rolle)"),
                               "**Begründung:** *arbeitetMitAn* ist n:m – ein Ermittler arbeitet an mehreren Fällen mit, an "
                               "einem Fall arbeiten mehrere mit. In Ermittler oder Fall passte kein einzelner Wert (mehrere "
                               "Werte je Zelle bzw. wiederholte Zeilen → Anomalien), daher eigene Relation; das Paar ist "
                               "Primärschlüssel, Rolle hängt vom Paar ab und steht deshalb dort.",
                               "**Bewertung:** Ermittler mit rekursivem Fremdschlüssel (Rollenname) **(1 BE)**, Fall mit "
                               "↑LeiterNr **(1 BE)**, arbeitetMitAn mit zusammengesetztem Schlüssel und Rolle **(1 BE)**, "
                               "Begründung **(2 BE)**.",
                               fehler=[("P9", "Rolle in Ermittler; ↑FallNr in Ermittler für arbeitetMitAn."),
                                       ("P9", "Rekursiver Fremdschlüssel als zweites „ErmittlerNr“ benannt."),
                                       ("P10", "Begründung ohne Sachzusammenhang („weil n:m eine Tabelle braucht“).")]))
    d.spur_im_lab("5.3 „Vom Plan zur Datenbank“", "A5-E3-01 … A5-E3-08, danach Verhör 5 · Kletterhalle (20 Minuten)",
                  werkzeug="Schema-Werkbank, Bauplan-Tisch (Rückrichtung), Konsole auf der Fall-Datenbank",
                  text="A5-E3-08 fragt die Fall-Datenbank, wen Nora Falter kennt (Self-Join). Cold Case im Lab: "
                       "A5-E3-CC „Mentor-Kette“.")
    kette = ("SELECT e.Name AS Ermittler, m.Name AS Mentor, mm.Name AS MentorDesMentors\nFROM Ermittler e, Ermittler m, "
             "Ermittler mm\nWHERE e.MentorNr = m.ErmittlerNr\n  AND m.MentorNr = mm.ErmittlerNr\nORDER BY e.Name;")
    d.cold_case("Mentor-Kette", "Gib zu jedem Ermittler den Namen des Mentors und den Namen des Mentors des Mentors aus. "
                "Begründe, welche Ermittler in deinem Ergebnis fehlen und warum.", kompetenzen=["K24"], linien=6,
                loesung=Loesung(SQL(kette), ergebnis(kette),
                                "Ermittler wird **dreimal** mit Alias verwendet. Es fehlen Edda Lorenz und Ulf Witte (kein "
                                "Mentor, MentorNr NULL) sowie Juri Brandes, Maren Kolthoff und Helge Saathoff: Ihr Mentor "
                                "Edda Lorenz hat selbst keinen Mentor – der zweite Vergleich ist nie wahr."))
    name = ("5_Loesungen/DD_A5_AB3_Vom_Plan_zur_Datenbank_LSG.docx" if lsg
            else "2_Arbeitsblaetter/DD_A5_AB3_Vom_Plan_zur_Datenbank.docx")
    return d.speichern(ORDNER / name)


# ================================================================== Präsentation „Der Bauplan“ (DS 9–11)
#: Foliennummern je Abschnitt (für die Lehrerhinweise; wird von praesentation() gefüllt)
FOLIEN = {}

BRIEFING9A = ("Unsere Pinnwand ist voll, und Juri hat die rote Wolle aufgebraucht. Bevor wir eine Datenbank bauen, brauchen "
              "wir einen Bauplan – und Baupläne für Datenbanken heißen ER-Diagramme. Ich habe euch einen alten Plan aus dem "
              "Polizeiarchiv mitgebracht.")
BRIEFING9B = ("Rechtecke sind Entitätstypen, Rauten Beziehungstypen, Ellipsen Attribute, und unterstrichen sind die "
              "Schlüsselattribute. An den Kanten steht die Kardinalität: 1, n oder m. Eine Kardinalität rät man nicht, man "
              "begründet sie. Auch Beziehungstypen können Attribute haben – etwa den Grad, wie verdächtig eine Person in einem "
              "Fall ist. Gleich verbindet ihr Rollenkarten mit Fäden und lest die Kardinalität einfach ab.")
BRIEFING10A = ("Heute ist Architekturwettbewerb: Jedes Team entwirft den Bauplan für die Fall-Datenbank der Detektei. "
               "Ermittler leiten Fälle, erfahrene Ermittler betreuen neue, Personen sind in Fällen verdächtig, Personen kennen "
               "andere Personen, und eine Aussage macht eine Person in einem Fall an einem Ort.")
BRIEFING10B = ("Da stecken Fallstricke drin: Eine Beziehung kann einen Entitätstyp mit sich selbst verbinden. Zwischen "
               "denselben Entitätstypen kann es mehrere Beziehungen geben, und manche verbinden sogar drei. Nicht alles lässt "
               "sich mit 1, n und m ausdrücken. Am Ende geht ihr auf Galerie-Rundgang – der Siegerentwurf bekommt einen "
               "Ehrenplatz an der Pinnwand.")
BRIEFING11A = ("Ein Bauplan ist schön, aber ermitteln können wir nur mit Tabellen. Für die Überführung gibt es klare Regeln: "
               "Jeder Entitätstyp wird eine Tabelle, bei 1:n wandert der Fremdschlüssel auf die n-Seite, und eine "
               "n:m-Beziehung bekommt eine eigene Tabelle mit zusammengesetztem Schlüssel.")
BRIEFING11B = ("Im Abitur reicht die Umsetzung nicht, ihr müsst sie begründen. Manchmal geht es auch rückwärts: aus dem "
               "Schema das ER-Diagramm. Wenn das Schema steht, füllt Juri die Datenbank mit allem, was wir wissen – und dann "
               "fragen wir sie: Wer ist in allen vier Fällen verdächtig, und wen kennt diese Person?")


def _merke(p, schluessel):
    FOLIEN[schluessel] = p._nr


def praesentation():
    p = AktenPraesentation(5, "Der Bauplan", datenbank=DB)
    # ---------------------------------------------------------------- DS 9
    p.titelfolie("Der Bauplan", "Akte 5 · DS 9 „Baupläne lesen“ · DS 10 „Der Architekturwettbewerb“ · DS 11 „Vom Plan zur "
                 "Datenbank“ · nur eA",
                 notizen="Akte 5 nur in eA-Kursen (im gA-Kurs entfällt sie, Niveau gA blendet sie aus). Diese Präsentation "
                         "übernimmt die Inhalte der früheren Präsentation „ER-Modell“ mit den Korrekturen aus Spec §11.2: "
                         "Entitätstyp/Beziehungstyp statt „Entitätsmenge“, Schlüsselattribute im ER unterstrichen, Regeln für "
                         "rekursive und ternäre Beziehungen und für Beziehungsattribute bei 1:n, zusätzliches "
                         "Schlüsselattribut bei mehrfach vorkommenden Paaren.")
    _merke(p, "ds9_start")
    p.inhaltsfolie("Warm-up: Karten W-A4-05 … 08",
                   Liste(["Einzelarbeit, je Karte 2–3 Minuten, dann Vergleich im Team",
                          "Schlüsselkriterien · Ausschnitt-Falle · Verlustfrei? · Zusammengelegt"]),
                   "Wiederholung aus Akte 4 – die Ausschnitt-Falle brauchen wir heute wieder: Auch eine Kardinalität "
                   "beweist man nicht mit ein paar Zeilen.",
                   notizen="DS 9, Warm-up (0–10'): Warm-up-Karten A4 (Karten 05–08). Kurz auflösen; Brücke zur Stunde: "
                           "„Daten stützen, Begründung kommt aus dem Sachzusammenhang.“")
    p.briefing("Akte 5: Der Bauplan", BRIEFING9A,
               notizen="Briefing (10–20'). Foto/Skizze der vollen Pinnwand zeigen, falls vorhanden.")
    p.briefing("Akte 5: Der Bauplan (Fortsetzung)", BRIEFING9B,
               notizen="Die Begriffe des Briefings werden auf den nächsten Folien am Plan beschriftet.")
    p.inhaltsfolie("Der alte Bauplan aus dem Polizeiarchiv",
                   "Was sagt die **1** an der Kante bei Ermittler – und was die **n** bei Fall?", er_folie(er_leitet(), 28),
                   notizen="Impuls aus DS_PLANUNG. Erwartung: „Ein Ermittler leitet mehrere Fälle, jeder Fall hat genau einen "
                           "Leiter.“ Falls jemand die n dem Ermittler zuordnet: Leseschablone vorwegnehmen (Zahl steht beim "
                           "Entitätstyp, über den man spricht).")
    p.begriffe("Chen-Notation nach EH §7", [("Rechteck", "**Entitätstyp** – Singular: Ermittler, Fall"),
                                            ("Raute", "**Beziehungstyp** – Verb in lowerCamelCase: *leitet*"),
                                            ("Ellipse", "**Attribut** – Name, Titel"),
                                            ("unterstrichene Ellipse", "**Schlüsselattribut** – <u>FallNr</u>"),
                                            ("1 · n · m an der Kante", "**Kardinalität** (ternär zusätzlich p)")],
               inhalt=["Entitätstyp = alle Fälle · Entität = ein einzelner Fall. Im Abitur: Entitätstyp, nicht "
                       "„Entitätsmenge“; Schlüsselattribute werden im ER-Diagramm **immer** unterstrichen."],
               notizen="Korrektur gegenüber der alten Präsentation „ER-Modell“: jetzt Entitätstyp statt „Entitätsmenge“, und die "
                       "Schlüssel werden im ER markiert. Keine (min,max)- und keine Krähenfuß-Notation.")
    p.inhaltsfolie("Kardinalitäten 1:1 · 1:n · n:m", mengen_bild(MENGEN_BEISPIEL, "folie_mengen", breite_cm=30),
                   "Leseschablone – immer in **beide** Richtungen: „Ein Ermittler leitet mehrere Fälle; jeder Fall wird von "
                   "genau einem Ermittler geleitet.“",
                   notizen="Mengendiagramm = einzelne Entitäten als Punkte. Bei 1:n hat E3 keine Linie – auch das ist erlaubt "
                           "(die 1 sagt „höchstens/genau einer“ auf der anderen Seite, nicht „jeder Ermittler muss“). Bei n:m "
                           "haben Punkte auf **beiden** Seiten mehrere Linien.")
    p.inhaltsfolie("Beispiel Schritt für Schritt: Kardinalität bestimmen", er_folie(er_schule(), 17),
                   Liste(["Satz von links: „Ein Schüler besucht **genau eine** Klasse.“ → 1 an die Kante zu **Klasse**",
                          "Satz von rechts: „Eine Klasse wird von **mehreren** Schülern besucht.“ → n an die Kante zu "
                          "**Schueler**",
                          "Ergebnis lesen: Schueler n : 1 Klasse",
                          "Gegenprobe im Sachzusammenhang: Kann ein Schüler zwei Klassen besuchen? (Nein.)"],
                         nummeriert=True),
                   notizen="Eigenes Beispiel (Schule), nicht die Detektei. Typischer Fehler P9: die Zahl an die eigene Kante "
                           "schreiben („Ein Schüler – 1“ an Schueler). Regel: Die Zahl steht beim Entitätstyp, über den der "
                           "Satz eine Mengenangabe macht.")
    p.inhaltsfolie("Wo steht der Verdachtsgrad?",
                   "Person P3 ist in Fall F2 *hoch*, in Fall F5 *gering* verdächtig → der Grad hängt vom **Paar** ab → "
                   "Attribut am Beziehungstyp.", er_folie(er_verdacht(), 24),
                   notizen="Impuls: „Wo steht der Verdachtsgrad: bei der Person, beim Fall oder dazwischen?“ Erst abstimmen "
                           "lassen, dann mit dem Gegenbeispiel auflösen. Merkfrage: Ändert sich der Wert, wenn ich eine andere "
                           "Person oder einen anderen Fall nehme?")
    beleg = "SELECT FallNr, COUNT(*) AS Verdaechtige\nFROM verdaechtig\nGROUP BY FallNr\nHAVING COUNT(*) > 1;"
    p.beispiel_schritte("Daten stützen die Kardinalität", beleg,
                        ["**FROM** – die Beziehungstabelle verdaechtig(PersonNr, FallNr, Grad)",
                         "**GROUP BY** – eine Gruppe je Fall",
                         "**HAVING** – nur Fälle mit mehr als einer verdächtigen Person",
                         "Ergebnis: Ein Fall – mehrere Personen. Die andere Richtung prüft ihr im Lab (A5-E1-08)."],
                        ergebnis=ergebnis(beleg, max_zeilen=5),
                        notizen="SQL-Wiederholung aus Akte 3. Die Daten stützen n auf der Personen-Seite – beweisen tun sie "
                                "nichts (Ausschnitt-Falle). Die Gegenrichtung (Personen in mehreren Fällen) ist Lab-Aufgabe "
                                "A5-E1-08 und wird hier nicht gezeigt.")
    p.unplugged("Fäden-Kardinalitäten",
                ["**1:n** – Jede Fall-Karte nimmt einen Faden zu der Ermittlerin, die auf ihrer Karte steht. Wer hält wie "
                 "viele Fäden?",
                 "**n:m** – Personen spannen Fäden zu allen Fällen, in denen sie verdächtig sind.",
                 "**1:1** – Jeder Ermittler zu seinem Dienstausweis.",
                 "**Beziehungsattribut** – Wohin kommt das Etikett „Grad: hoch“ für Person 3 und Fall 2?",
                 "Impuls: Woran erkennt ihr in einer Tabelle, welche Kardinalität vorliegt?"],
                material="Rollenkarten, Wollfäden, Etiketten", dauer="15 MIN",
                notizen="Kopiervorlage DD_A5_Unplugged_Faeden-Kardinalitaeten (20 Rollenkarten + Etiketten). Ablauf und "
                        "erwartete Beobachtungen in den Lehrerhinweisen. Sicherung: Foto der Fadennetze; drei "
                        "Kardinalitäten mit Leseschablone an die Tafel. Antwort auf den Impuls: Kommt der Wert einer "
                        "Fremdschlüsselspalte mehrfach vor → n-Seite.")
    _merke(p, "ds9_lab")
    p.lab_auftrag("5.1", "Baupläne lesen", "A5-E1-01 … 08 (★ → ★★ → ★★★)",
                  werkzeug="Bauplan-Tisch (Modi lesen und ergänzen)",
                  hinweise=["Auf dem iPad: Listenmodus statt Zeichnen", "Parallel: AB 5.1, Aufträge 1–4 auf Papier",
                            "Schnelle: Cold Case A5-E1-CC und AB-Cold-Case „Dienstwagen“"],
                  notizen="Lab-Einsatz (35–80'). Herumgehen: Kardinalität an der falschen Kante (P9), Grad an Person (P9), "
                          "Erläutern ohne Sachzusammenhang (P10) sammeln.")
    p.merkfolie("Ermittler-Notiz", ["Rechteck = Entitätstyp · Raute = Beziehungstyp · Ellipse = Attribut · unterstrichen = "
                                    "Schlüsselattribut · 1/n/m = Kardinalität.",
                                    "Leseschablone in **beide** Richtungen; die Zahl steht beim Entitätstyp, über den du "
                                    "sprichst.",
                                    "Hängt ein Attribut vom Paar ab → an die Raute.",
                                    "Daten stützen eine Kardinalität nur – begründet wird im Sachzusammenhang.",
                                    "Fachbegriffe: Entitätstyp, Beziehungstyp – nie „Entitätsmenge“."],
                notizen="Fallbesprechung (80–90'): Merkkasten sichern.")
    p.abi_format("Gegeben ist ein ER-Diagramm zur Fallverwaltung einer Detektei. **Begründen** Sie die Kardinalitäten der "
                 "Beziehungstypen *leitet* und *verdaechtig* im Sachzusammenhang.", be=4, operator="begründen",
                 inhalt=[er_fallverwaltung()],
                 notizen="Abi-Ecke von AB 5.1 (4 BE): je Richtung 1 BE. Operator begründen = Sachverhalt auf Regeln und "
                         "Zusammenhänge zurückführen („weil …“).")
    p.loesungsfolie("Lösung Abi-Format (DS 9)",
                    "*leitet*: Jeder Fall wird von **genau einer** Person verantwortlich geleitet (1); eine Ermittlerin kann "
                    "**mehrere** Fälle gleichzeitig leiten (n).",
                    "*verdaechtig*: Eine Person kann in **mehreren** Fällen verdächtig sein (Serie); in einem Fall können "
                    "**mehrere** Personen verdächtig sein (n:m). Der Grad hängt vom Paar ab.",
                    notizen="Typisch: „1:n, weil an der Kante 1 und n steht“ (Ablesen statt Begründen, P10).")
    p.briefing("Cliffhanger", "Juri hat schon mal angefangen, den Bauplan der Detektei zu zeichnen – mit vierzehn "
               "Entitätstypen und keiner einzigen Kardinalität.", von="Edda Lorenz",
               notizen="Abschluss DS 9. Ausblick: Architekturwettbewerb – Teams bilden, Anforderungstext vorab lesen.")
    _merke(p, "ds9_ende")
    # ---------------------------------------------------------------- DS 10
    p.inhaltsfolie("Warm-up: Karten W-A5-01 … 04",
                   Liste(["Notation benennen · Mengenbild · Aus den Daten · Wohin mit der Note?",
                          "Einzelarbeit, dann Vergleich im Team; Karte 03 lösen wir gemeinsam"]),
                   notizen="DS 10, Warm-up (0–10'): Warm-up-Karten A5, Karten 01–04.")
    _merke(p, "ds10_start")
    p.loesungsfolie("Lösung Warm-up W-A5-03: Aus den Daten",
                    Tabelle(["LeserNr", "BuchNr"], [["1", "10"], ["1", "11"], ["2", "10"], ["3", "12"]]),
                    "Leser 1 leiht zwei Bücher **und** Buch 10 wird von zwei Lesern geliehen → **n:m**. Wer nur eine Richtung "
                    "prüft, kommt auf 1:n.",
                    notizen="Nachsatz: Die Daten stützen n:m; begründet wird es mit dem Sachzusammenhang (ein Buch wird im "
                            "Laufe der Zeit von vielen gelesen).")
    p.briefing("Der Architekturwettbewerb", BRIEFING10A, notizen="Briefing (10–20').")
    p.briefing("Der Architekturwettbewerb (Fortsetzung)", BRIEFING10B,
               notizen="Die drei neuen Bausteine folgen auf den nächsten Folien an Mini-Beispielen.")
    p.inhaltsfolie("Rekursive Beziehungen",
                   "Zwei Kanten zum **selben** Entitätstyp, Rollen an den Kanten – kein Entitätstyp „Mentor“.",
                   Bild(str(nebeneinander([er_bild(er_betreut(), 6, "f_betreut"), er_bild(er_kennt(), 6, "f_kennt")],
                                          "f_rekursiv")), breite_cm=20),
                   notizen="Impuls: „Wer betreut wen? Kann eine Raute zweimal an dasselbe Rechteck?“ Herkunft: Beispiel "
                           "*istPateVon* aus den EH → bei uns *betreut*. Korrektur gegenüber „ER-Modell“: rekursive "
                           "Beziehungen fehlten dort.")
    p.inhaltsfolie("Zwei Beziehungstypen zwischen denselben Entitätstypen",
                   "Zwei Rollen (Heim/Gast) = zwei Beziehungstypen. In der Detektei: Person *meldet* Fall und Person ist "
                   "*verdaechtig* in Fall.", er_folie(er_heim_gast(), 20),
                   notizen="Fehler, der hier vorkommt: eine Raute „spielt“ mit Attribut „Heim/Gast“ – dann wären Heim- und "
                           "Gastteam nicht unterscheidbar bzw. es bräuchte n:m.")
    p.inhaltsfolie("Ternär: eine Raute, drei Kanten",
                   "Reichen drei binäre Beziehungen? Nein – „wer, in welchem Fall, wo“ ist nur zu dritt eindeutig.",
                   er_folie(er_sagt_aus(), 20),
                   notizen="Impuls aus DS_PLANUNG: „Reichen drei binäre Beziehungen statt der ternären? Was geht verloren?“ "
                           "Gegenbeispiel: Gesa sagt zu Fall 1 und Fall 3 aus, es wurde im Fundbüro und auf dem Festplatz "
                           "ausgesagt – wo sagte sie zu welchem Fall aus? Kardinalitäten ternär: m, n, p.")
    p.inhaltsfolie("Beispiel Schritt für Schritt: vom Text zum Modell",
                   "„Im Tierheim wird jedes **Tier** (Nummer, Name, Art) von genau einer **Pflegekraft** (Nummer, Name) "
                   "versorgt. Erfahrene Pflegekräfte arbeiten neue ein. **Helfer** (Ausweisnummer, Name) führen mehrere "
                   "Tiere aus, ein Tier wird von mehreren Helfern ausgeführt.“",
                   Liste(["Substantive → Entitätstypen mit Schlüsselattribut", "Verben → Beziehungstypen",
                          "Mengenangaben → Kardinalitäten in beide Richtungen", "Rollen → rekursiver Beziehungstyp"],
                         nummeriert=True),
                   notizen="Eigenes Beispiel (nicht die Detektei). Gemeinsam an der Tafel markieren; Lösung auf der nächsten "
                           "Folie.")
    p.inhaltsfolie("Das Modell zum Tierheim", "Kontrolle: Jede Information des Textes ist im Modell abgehakt.",
                   er_folie(er_tierheim(), 30),
                   notizen="Varianten würdigen: *betreut* statt *versorgt* ist gleichwertig. Die Rollen Erfahrene/Neue an die "
                           "Kanten von *arbeitetEin*.")
    _merke(p, "ds10_wettbewerb")
    p.inhaltsfolie("Der Anforderungstext", ANFORDERUNGSTEXT,
                   notizen="Wettbewerb (20–50'): Text auch auf AB 5.2 (Auftrag 1) und als Plakat-Vorlage im Unplugged-Set.")
    p.unplugged("Architekturwettbewerb",
                ["**Rollen** im Team: Moderation · Zeichner:in · Textprüfer:in (hakt jede Information ab) · "
                 "Notationsprüfer:in (EH §7)",
                 "**Modellieren** (30') auf Plakat A2 oder im Bauplan-Tisch (Modus entwickeln, Export als Bild)",
                 "**Galerie** (15'): Jedes Team bewertet zwei fremde Entwürfe mit dem Bewertungsbogen",
                 "Pflichtkommentar: „Das würden wir übernehmen: …“",
                 "**Siegerehrung**: Punktekleber – das Siegerteam erhält den Stempel ARCHITEKTUR"],
                material="Plakate, Stifte, Bewertungsbögen, Klebepunkte", dauer="45 MIN",
                notizen="Kopiervorlage DD_A5_Unplugged_Architekturwettbewerb. Zeitansage nach 15': „Habt ihr schon alle "
                        "Beziehungstypen?“ Textprüfer:in verhindert, dass Teams sich in Details verlieren.")
    p.inhaltsfolie("Der Bewertungsbogen", Liste(KRITERIEN, nummeriert=True),
               notizen="Je Kriterium 1 Punkt. Die Kriterien sind zugleich die Kernpunkte der Referenzlösung.")
    _merke(p, "ds10_lab")
    p.lab_auftrag("5.2", "Der Architekturwettbewerb", "A5-E2-01 … 06 in der Stunde, 07/08 als Hausaufgabe",
                  werkzeug="Bauplan-Tisch (Modi ergänzen und entwickeln)",
                  hinweise=["Die Referenzlösung erscheint im Lab erst nach A5-E2-07",
                            "Parallel bzw. zu Hause: AB 5.2"],
                  notizen="Lab-Einsatz (65–85'), verkürzt wegen Wettbewerb.")
    p.loesungsfolie("Referenzlösung = Fall-Datenbank der Detektei", er_folie(er_referenz(), 30),
                    notizen="Fallbesprechung (85–90'): neben den Siegerentwurf legen; Abweichungen als Alternativen würdigen "
                            "(K40), z. B. aufgelöste Zwischenentität statt n:m. Erst nach dem Wettbewerb zeigen.")
    p.merkfolie("Ermittler-Notiz", ["Rekursiv: Beziehungstyp mit zwei Kanten zum selben Entitätstyp, Rollen an die Kanten.",
                                    "Zwei Rollen (Heim/Gast, Leitung/Verdacht) = zwei Beziehungstypen.",
                                    "Ternär: eine Raute, drei Kanten (m:n:p) – nur zu dritt eindeutig.",
                                    "Grenzen: „genau drei“, „höchstens 12“, „höherer Dienstgrad“ → Anmerkung zum Modell.",
                                    "Beurteilen: Alternativen benennen, Folgen abwägen, Urteil begründen."],
                notizen="Merkkasten DS 10 sichern.")
    p.abi_format("Gegeben ist das ER-Modell der Detektei (Ermittler *leitet* Fall). Die Detektei möchte zusätzlich "
                 "erfassen, welche Ermittler mit welcher Rolle an einem Fall mitarbeiten und welcher Ermittler welchen "
                 "anderen vertritt (höchstens eine feste Vertretung je Ermittler). **Erweitern** Sie das ER-Modell "
                 "entsprechend.", be=5, operator="erweitern", inhalt=[er_folie(er_leitet(), 20)],
                 notizen="Abi-Ecke von AB 5.2 (gekürzt). 2 BE arbeitetMitAn n:m mit Rolle, 2 BE vertritt rekursiv 1:n, "
                         "1 BE Notation.")
    p.loesungsfolie("Lösung Abi-Format (DS 10)", er_folie(er_erweiterung(), 26),
                    notizen="Typisch: Rolle an Ermittler (P9), Entitätstyp „Vertretung“ (P9).")
    p.briefing("Cliffhanger", "Der Siegerentwurf kommt auf den Bauplan-Tisch, und Edda sagt nur: „Schön – nächste Stunde "
               "wird daraus eine echte Datenbank, und dann fragen wir sie, wer Nora Falter kennt.“", von="Edda Lorenz",
               notizen="Abschluss DS 10. Hausaufgabe: AB 5.2, Lab A5-E2-07/08; Teamplakate für DS 11 aufheben.")
    _merke(p, "ds10_ende")
    # ---------------------------------------------------------------- DS 11
    p.inhaltsfolie("Warm-up: Karten W-A5-05 … 08",
                   Liste(["Rekursiv lesen · Ternär lesen · Text → Mini-ER · Grenzen"]),
                   notizen="DS 11, Warm-up (0–10'): Warm-up-Karten A5, Karten 05–08.")
    _merke(p, "ds11_start")
    p.briefing("Vom Plan zur Datenbank", BRIEFING11A, notizen="Briefing (10–20'): fragend-entwickelnd.")
    p.briefing("Vom Plan zur Datenbank (Fortsetzung)", BRIEFING11B)
    p.inhaltsfolie("Wohin kommt der Verweis?",
                   "Ermittler **1** – *leitet* – **n** Fall. Probiert beides mit Daten aus:",
                   Tabelle(["ErmittlerNr", "Name", "FallNr ?"], [["5", "Imke Tiemann", "7, 8, 11"], ["8", "Deniz Akgül", "12"]]),
                   Tabelle(["FallNr", "Titel", "LeiterNr"], [["7", "Der doppelte Bauantrag", "5"],
                                                             ["8", "Die verschwundene Gästeliste", "5"]]),
                   notizen="Impuls aus DS_PLANUNG: „Wohin kommt der Verweis bei leitet – in Ermittler oder in Fall?“ Oben: "
                           "mehrere Werte in einer Zelle (nicht atomar). Unten: ein Wert je Zeile → Fremdschlüssel auf der "
                           "n-Seite. Zweiter Impuls: „Warum kann Person keinen Fremdschlüssel FallNr bekommen, wenn sie in "
                           "mehreren Fällen verdächtig ist?“")
    p.inhaltsfolie("Überführungsregeln R1–R6 (Beispiel Schule)",
                   Tabelle(["Regel", "Umsetzung im Relationenschema"], [[r[0], r[2]] for r in REGELN_SCHULE],
                           breiten=[4.0, 26.0]), karte=False,
               notizen="Werkzeugkarte 5 im Ermittler-Handbuch. Korrekturen gegenüber „ER-Modell“: R5 (rekursiv) und R6 "
                       "(ternär) fehlten dort; Beziehungsattribute bei 1:n kommen in die Relation der n-Seite.")
    p.inhaltsfolie("Zwei Feinheiten, die im Abitur zählen",
                   "**Attribut am 1:n-Beziehungstyp** → in die Relation der **n-Seite**: Lehrkraft 1 – *leitet* (seit) – n "
                   "Kurs ergibt Kurs(<u>KursNr</u>, Titel, ↑LehrkraftNr, LeitetSeit).",
                   "**Paar kommt mehrfach vor** → zusätzliches Schlüsselattribut: Spieler bucht Platz jede Woche → "
                   "bucht(<u>↑SpielerNr</u>, <u>↑PlatzNr</u>, <u>Beginn</u>) oder eigene BuchungNr.",
                   "**Rekursiver Fremdschlüssel** heißt nach der Rolle: Schueler(<u>SchuelerNr</u>, Name, ↑PateNr).",
                   notizen="Diese drei Punkte waren in der alten Präsentation „ER-Modell“ falsch oder fehlten (Spec §11.2). "
                           "Prüffrage für den Schlüssel: Kann dieselbe Kombination zweimal vorkommen?")
    p.inhaltsfolie("Rückrichtung: Schema → ER",
                   Liste(["Relation mit eigenem Schlüssel → Entitätstyp",
                          "Fremdschlüssel in einer Entitätsrelation → 1:n (Fremdschlüssel auf der n-Seite)",
                          "Schlüssel aus zwei (drei) Fremdschlüsseln → n:m (ternär); übrige Attribute an die Raute",
                          "Fremdschlüssel werden **keine** Attribute im ER-Diagramm"]),
                   notizen="Lab A5-E3-06 und AB 5.3 Auftrag 3 (Stadtbücherei). Den Namen des Beziehungstyps erfinden die "
                           "Schüler aus dem Sachzusammenhang.")
    p.beispiel_schritte("Self-Join Schritt für Schritt", SELFJOIN,
                        ["**FROM** – Ermittler zweimal: `e` (betreut) und `m` (Mentor)",
                         "**WHERE** – Verbund über den rekursiven Fremdschlüssel MentorNr",
                         "**ORDER BY** – nach Mentor, dann nach Name",
                         "**SELECT** – Spaltenaliase machen die Rollen lesbar"],
                        ergebnis=ergebnis(SELFJOIN, max_zeilen=5),
                        notizen="10 Zeilen. Edda Lorenz und Ulf Witte erscheinen nicht als „Ermittler“ (MentorNr NULL). Ohne "
                                "Aliase: „Name ist mehrdeutig“ (P7). Tabellenalias ohne AS (EH).")
    p.unplugged("Regelkarten-Umzug",
                ["Für jedes Rechteck eures Wettbewerbsplakats entsteht eine Relationenkarte (R1), Schlüssel unterstrichen.",
                 "Für jede Raute die passende Regelkarte ziehen und „umziehen“: Fremdschlüssel rot mit ↑, ggf. neue Karte.",
                 "Probe: drei Beispielzeilen für *verdaechtig* – ohne Mehrfachwerte, ohne wiederholte Personendaten?",
                 "Austausch: „Welche Regel habt ihr hier angewendet – und warum nicht eine andere?“"],
                material="Regelkarten R1–R6, Relationenkarten, Teamplakat", dauer="15 MIN",
                notizen="Kopiervorlage DD_A5_Unplugged_Regelkarten (Vorderseite Regel, Rückseite Begründung, beidseitig "
                        "drucken). Sicherung: Werkzeugkarte 5.")
    _merke(p, "ds11_lab")
    p.lab_auftrag("5.3", "Vom Plan zur Datenbank", "A5-E3-01 … 08",
                  werkzeug="Schema-Werkbank, Bauplan-Tisch (Rückrichtung), Konsole (Fall-Datenbank)",
                  hinweise=["Danach: Verhör 5 · Kletterhalle (20 Minuten, ohne Tipps)", "Parallel: AB 5.3"],
                  notizen="Lab-Einsatz (35–80'). Verhör 5 im Kurs auf „auto“ oder „offen“ stellen. Typische Fehler: FK auf "
                          "der 1-Seite, n:m ohne eigene Relation, rekursiver FK gleich benannt wie der PK.")
    p.merkfolie("Ermittler-Notiz", ["R1 Entitätstyp → Relation · R2 1:n → Fremdschlüssel auf die n-Seite",
                                    "R3 n:m → eigene Relation, Schlüssel aus beiden Fremdschlüsseln",
                                    "R4 1:1 → Fremdschlüssel auf einer Seite · R5 rekursiv → Fremdschlüssel nach Rolle",
                                    "R6 ternär → eigene Relation mit drei Fremdschlüsseln; Paar/Trio mehrfach → "
                                    "zusätzliches Schlüsselattribut",
                                    "Self-Join: dieselbe Tabelle zweimal mit verschiedenen Aliasen"],
                notizen="Fallbesprechung (80–90'): mit Begründungen sichern (Werkzeugkarte 5).")
    p.abi_format("Gegeben ist das ER-Modell einer Erweiterung der Fall-Datenbank. **Überführen** Sie das Modell in ein "
                 "Relationenschema in Kurzschreibweise und **begründen** Sie die Umsetzung des Beziehungstyps "
                 "*arbeitetMitAn*.", be=5, operator="überführen / begründen", inhalt=[er_folie(er_erweiterung(), 20)],
                 notizen="Abi-Ecke von AB 5.3: 3 BE Relationen, 2 BE Begründung.")
    p.loesungsfolie("Lösung Abi-Format (DS 11)",
                    Schema("Ermittler(<u>ErmittlerNr</u>, Name, Dienstgrad, ↑VertretungNr)"),
                    Schema("Fall(<u>FallNr</u>, Titel, Eroeffnet, ↑LeiterNr)"),
                    Schema("arbeitetMitAn(<u>↑ErmittlerNr</u>, <u>↑FallNr</u>, Rolle)"),
                    "n:m → in keiner der beiden Relationen passt ein einzelner Wert; das Paar ist Schlüssel, die Rolle hängt "
                    "vom Paar ab.",
                    notizen="Typisch: Rolle in Ermittler (P9); Begründung ohne Sachzusammenhang (P10).")
    alle4 = ("SELECT Person.PersonNr, Person.Vorname, Person.Nachname\nFROM Person, verdaechtig\n"
             "WHERE Person.PersonNr = verdaechtig.PersonNr\n  AND verdaechtig.FallNr IN (1, 2, 3, 4)\n"
             "GROUP BY Person.PersonNr, Person.Vorname, Person.Nachname\nHAVING COUNT(*) = 4;")
    p.loesungsfolie("Beweisstück: Ermittlungsnetz (1/2)", SQL(alle4), ergebnis(alle4),
                    notizen="Fallbesprechung: die entscheidende Frage aus dem Briefing. Gruppiert wird nach dem Schlüssel "
                            "(Variante gleichnamig: eine zweite Nora Falter!).")
    netz = ("SELECT b.Vorname, b.Nachname, b.Stadtteil, k.Seit\nFROM Person p, kennt k, Person b\n"
            "WHERE p.PersonNr = k.PersonNr\n  AND k.BekannterNr = b.PersonNr\n  AND p.PersonNr = 13;")
    p.loesungsfolie("Beweisstück: Ermittlungsnetz (2/2)", SQL(netz), ergebnis(netz),
                    notizen="Bekannte an allen vier Tatorten (Fundbüro, StadtRad, Stadtfest, TuS) und Keno Bargmann aus der "
                            "IT der Stadtverwaltung – eine Insiderin. Lab A5-E3-08 fragt dasselbe über den Namen.")
    p.briefing("Cliffhanger", "Die Antwort ist eindeutig: Nora Falter – verdächtig in allen vier Fällen, mit Bekannten an "
               "jedem Tatort. Doch die Bürgermeisterin lässt ausrichten: Eine Nora Falter ist in Weserstedt nicht gemeldet.",
               von="Edda Lorenz", nur_lehrkraft=True,
               notizen="Abschluss DS 11 (verrät das Lab-Ergebnis, daher nicht im Handout). Steckbrief ergänzen: einzige "
                       "Verdächtige in allen vier Fällen · kennt Menschen an jedem Tatort · Insiderin · nicht gemeldet. "
                       "Hinweis Verhör 5 (ER, 20') und Video 6.")
    _merke(p, "ds11_ende")
    p.speichern(ORDNER / "1_Praesentation" / "DD_A5_Bauplan.pptx")
    return FOLIEN


# ================================================================== Unplugged: Hilfen für Karten
def _karten_raster(d, zellen_inhalt, spalten, zeilen_je_seite, hoehe_cm, rahmen=A.RAND_STARK, umbruch_vorher=False):
    """Füllt ein Raster aus Karten (Tabelle mit fester Zeilenhöhe, gestrichelte Schnittlinien) – wie Akte 1."""
    breite = d.breite / spalten
    for start in range(0, len(zellen_inhalt), spalten * zeilen_je_seite):
        if start or umbruch_vorher:
            d.seitenumbruch()
        seite = zellen_inhalt[start:start + spalten * zeilen_je_seite]
        n_zeilen = (len(seite) + spalten - 1) // spalten
        tab = d.doc.add_table(rows=n_zeilen, cols=spalten)
        A._tabelle_breiten(tab, [breite] * spalten)
        A._tabelle_ohne_rahmen(tab)
        for zi in range(n_zeilen):
            zeile = tab.rows[zi]
            trpr = zeile._tr.get_or_add_trPr()
            trpr.append(A._el("w:trHeight", val=int(Cm(hoehe_cm) / 635), hRule="exact"))
            A._zeile_nicht_teilen(zeile)
            for si in range(spalten):
                z = zeile.cells[si]
                A._zelle_rahmen(z, rahmen, 0.75, art="dashed")
                A._zelle_innenabstand(z, 0.35, 0.45, 0.3, 0.45)
                i = zi * spalten + si
                z.paragraphs[0].style = d.doc.styles["DD Tabellentext"]
                if i < len(seite) and seite[i] is not None:
                    seite[i](z, breite)


def _karte(d, etikett, gross_text, zeilen=(), farbe=A.SIEGELROT, flaeche=None, groesse=30, code=False, mitte=True):
    """Karte: Etikett oben, großer Text, darunter kleinere Zeilen (Mini-Auszeichnung)."""
    def fuellen(z, breite):
        if flaeche:
            A._zelle_flaeche(z, flaeche)
        z.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if mitte else WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = z.paragraphs[0]
        d._etikett(p, etikett, groesse=7.5, farbe=farbe, fett=True)
        p = z.add_paragraph(style="DD Tabellentext")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(6)
        if code:
            d._lauf_stil(p.add_run(gross_text), code=True, groesse=groesse, fett=True, farbe=farbe)
        else:
            d._lauf_stil(p.add_run(gross_text), schrift=A.SCHRIFT_KOPF, groesse=groesse)
        for t in zeilen:
            p = z.add_paragraph(style="DD Tabellentext")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            d._text(p, t, groesse=11)
    return fuellen


# ================================================================== Unplugged DS 9: Fäden-Kardinalitäten
ROLLEN_ERMITTLER = [("Edda Lorenz", "Leitung", "D-01"), ("Maren Kolthoff", "Senior", "D-02"),
                    ("Helge Saathoff", "Senior", "D-03"), ("Juri Brandes", "Ausbildung", "D-04")]
#: Fall-Karten: (Nr, Titel, Leitung) – Planspiel (Fall 6 wird hier von Maren Kolthoff geleitet)
ROLLEN_FAELLE = [(1, "Das Fundbüro", "Maren Kolthoff"), (2, "Spur der Räder", "Helge Saathoff"),
                 (3, "Die Stadtfest-Kasse", "Maren Kolthoff"), (4, "Chaos im Vereinsarchiv", "Edda Lorenz"),
                 (5, "Die vertauschten Koffer", "Helge Saathoff"), (6, "Falsche Rechnungen im Hafenamt", "Maren Kolthoff")]
#: Personen-Karten: (Nr, Name, [(Fall, Grad)]) – erfundene Namen (namen.json), Person 3: Fall 2 hoch, Fall 5 gering
ROLLEN_PERSONEN = [(1, "Anke Backhaus", [(1, "gering"), (3, "mittel")]), (2, "Arne Diekmann", [(1, "mittel")]),
                   (3, "Carla Eggers", [(2, "hoch"), (5, "gering")]),
                   (4, "Enno Gerdts", [(4, "gering"), (5, "mittel"), (6, "gering")]),
                   (5, "Fenja Haake", [(2, "gering"), (6, "hoch")]), (6, "Hauke Habben", [(3, "gering")])]


def faeden_auswertung():
    """Erwartete Beobachtungen (für die Lehrerhinweise): Fäden je Ermittler, je Fall, je Person."""
    je_ermittler = {e[0]: [f[0] for f in ROLLEN_FAELLE if f[2] == e[0]] for e in ROLLEN_ERMITTLER}
    je_fall = {f[0]: [p[0] for p in ROLLEN_PERSONEN if f[0] in [x[0] for x in p[2]]] for f in ROLLEN_FAELLE}
    assert je_ermittler["Juri Brandes"] == [] and max(len(v) for v in je_ermittler.values()) == 3
    assert any(len(v) > 1 for v in je_fall.values()) and any(len(p[2]) > 1 for p in ROLLEN_PERSONEN)
    return je_ermittler, je_fall


def unplugged_faeden():
    faeden_auswertung()
    d = AktenDokument("unplugged", 5, "Fäden-Kardinalitäten", fassung="schueler", ermittlung="5.1", ds=9,
                      material="Unplugged A5 · Fäden-Kardinalitäten", datenbank="keine",
                      untertitel="Kopiervorlage: 20 Rollenkarten zum Umhängen (4 Ermittler, 6 Fälle, 6 Personen, "
                                 "4 Dienstausweise) und Etiketten für Fäden – an den gestrichelten Linien schneiden")
    d.kopfband(name_datum=False)
    d.absatz("**So geht's:** Jede/r hängt sich eine Rollenkarte um – ihr seid die Entitäten. Wollfäden sind Beziehungen "
             "zwischen einzelnen Entitäten. Was auf eurer Karte steht, sagt euch, zu wem ihr einen Faden spannt.")
    d.inhalt(Liste(["**1:n · leitet:** Jede Fall-Karte nimmt einen Faden zu der Ermittlerin bzw. dem Ermittler, die bzw. "
                    "der auf ihrer Karte steht. Zählt: Wie viele Fäden hält jede Ermittler-Karte, wie viele jede Fall-Karte?",
                    "**n:m · verdaechtig:** Personen spannen Fäden zu allen Fällen, die auf ihrer Karte stehen. Gibt es auf "
                    "beiden Seiten Karten mit mehreren Fäden?",
                    "**1:1 · hat:** Jeder Ermittler verbindet sich mit seinem Dienstausweis.",
                    "**Beziehungsattribut:** Wohin kommt das Etikett „Grad: hoch“ für Person 3 und Fall 2 – an die Person, "
                    "an den Fall oder an den Faden?",
                    "**Impuls:** Woran erkennt ihr in einer Tabelle, welche Kardinalität vorliegt?"], nummeriert=True))
    d.notiz("Leseschablone", "„Ein Ermittler leitet **mehrere** Fälle; jeder Fall wird von **genau einem** Ermittler "
            "geleitet.“ → Ermittler 1 – *leitet* – n Fall. Immer in beide Richtungen lesen.")
    karten = []
    for name, grad, ausweis in ROLLEN_ERMITTLER:
        karten.append(_karte(d, "ROLLENKARTE · ENTITÄTSTYP ERMITTLER", name, [f"Dienstgrad: {grad}",
                                                                               f"Dienstausweis: **{ausweis}**"],
                             farbe=A.SIEGELROT, flaeche=A.ROT_FLAECHE))
    for nr, titel, leitung in ROLLEN_FAELLE:
        karten.append(_karte(d, "ROLLENKARTE · ENTITÄTSTYP FALL", f"Fall {nr}", [f"„{titel}“",
                                                                                f"geleitet von: **{leitung}**"],
                             farbe=A.BLAU, flaeche=A.BLAU_FLAECHE))
    for nr, name, faelle in ROLLEN_PERSONEN:
        karten.append(_karte(d, "ROLLENKARTE · ENTITÄTSTYP PERSON", f"Person {nr}", [name, "verdächtig in: **" + ", ".join(
            f"Fall {f} ({g})" for f, g in faelle) + "**"], farbe=A.GRUEN, flaeche=A.GRUEN_FLAECHE))
    for name, _, ausweis in ROLLEN_ERMITTLER:
        karten.append(_karte(d, "ROLLENKARTE · ENTITÄTSTYP DIENSTAUSWEIS", ausweis, [f"gehört: **{name}**"],
                             farbe=A.MESSING_TEXT, flaeche=A.MESSING_FLAECHE))
    _karten_raster(d, karten, spalten=2, zeilen_je_seite=4, hoehe_cm=6.0, umbruch_vorher=True)
    etiketten = ([_karte(d, "ETIKETT · BEZIEHUNGSTYP", b, farbe=A.TINTE_2, groesse=22) for b in ("leitet", "verdaechtig", "hat")]
                 + [_karte(d, "ETIKETT · KARDINALITÄT", k, farbe=A.SIEGELROT, groesse=34, code=True)
                    for k in ("1", "n", "1", "n", "m", "1", "1")]
                 + [_karte(d, "ETIKETT · ATTRIBUT", f"Grad: {g}", farbe=A.GRUEN, groesse=20)
                    for g in ("hoch", "gering", "mittel", "gering", "mittel")])
    _karten_raster(d, etiketten, spalten=3, zeilen_je_seite=5, hoehe_cm=4.6, umbruch_vorher=True)
    return d.speichern(ORDNER / "3_Unplugged" / "DD_A5_Unplugged_Faeden-Kardinalitaeten.docx")


# ================================================================== Unplugged DS 10: Architekturwettbewerb
TEAMROLLEN = [("Moderation", ["achtet auf die Zeit (Ansage nach 15')", "sorgt dafür, dass alle mitreden",
                              "entscheidet bei Streit nach Abstimmung"]),
              ("Zeichner:in", ["zeichnet das Plakat in Chen-Notation", "plant Platz für alle Entitätstypen",
                               "hält Kreuzungen gering"]),
              ("Textprüfer:in", ["hakt **jede** Information im Anforderungstext ab", "fragt: „Wo steht das im Modell?“",
                                 "meldet fehlende Attribute"]),
              ("Notationsprüfer:in", ["prüft EH §7: Rechteck, Raute, Ellipse", "Schlüsselattribute unterstrichen?",
                                      "Kardinalitäten 1/n/m an der richtigen Kante?"])]


def unplugged_wettbewerb():
    d = AktenDokument("unplugged", 5, "Architekturwettbewerb", fassung="schueler", ermittlung="5.2", ds=10,
                      material="Unplugged A5 · Architekturwettbewerb", datenbank="keine",
                      untertitel="Kopiervorlage: Anforderungstext (Plakat), Teamrollen, Bewertungsbogen (zwei je Team), "
                                 "Siegerstempel")
    d.kopfband(name_datum=False)
    d.absatz("**Ablauf (45 Minuten):** 1. *Modellieren (30')* – jedes Team entwirft auf einem Plakat (A2) oder im "
             "Bauplan-Tisch (Modus entwickeln) das ER-Modell zum Anforderungstext. 2. *Galerie (15')* – die Plakate hängen "
             "aus; jedes Team bewertet **zwei** fremde Entwürfe mit dem Bewertungsbogen. 3. *Siegerehrung* – Punktekleber; "
             "das Siegerteam erhält den Stempel ARCHITEKTUR.")
    d.ueberschrift("Anforderungstext (für das Plakat vergrößern)", 2)
    tab, z = d._kasten(d.doc, d.breite, rand=A.TINTE, staerke=2.0, flaeche=A.KARTE)
    p = z.paragraphs[0]
    d._etikett(p, "DATENDETEKTEI WESERSTEDT · AUSSCHREIBUNG ARCHITEKTURWETTBEWERB", groesse=8, farbe=A.SIEGELROT, fett=True)
    p = z.add_paragraph(style="DD Tabellentext")
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(8)
    p.paragraph_format.line_spacing = 1.3
    d._text(p, ANFORDERUNGSTEXT, groesse=15)
    d._nach_kasten(d.doc)
    d.seitenumbruch()
    d.ueberschrift("Teamrollen (je Team einmal)", 2)
    _karten_raster(d, [_karte(d, "TEAMROLLE", r, ["▸ " + a for a in aufg], farbe=A.BLAU, flaeche=A.BLAU_FLAECHE,
                              groesse=22) for r, aufg in TEAMROLLEN], spalten=2, zeilen_je_seite=2, hoehe_cm=5.2)
    d.ueberschrift("Siegerstempel (ausschneiden, auf den Detektei-Ausweis kleben)", 2)
    stempel = A.stempel_png("ARCHITEKTUR", TMP / "stempel_architektur.png")

    def stempelkarte(z, breite):
        z.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = z.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(stempel), width=Cm(4.6))
    _karten_raster(d, [stempelkarte] * 6, spalten=3, zeilen_je_seite=2, hoehe_cm=3.4)
    for _ in range(2):
        d.seitenumbruch()
        d.ueberschrift("Bewertungsbogen Architekturwettbewerb", 1)
        d.absatz("Bewertet wird der Entwurf von Team ______________________ durch Team ______________________")
        d.tabelle(["Nr.", "Kriterium (je 1 Punkt)", "✓", "Kommentar"],
                  [[str(i), k, "☐", ""] for i, k in enumerate(KRITERIEN, 1)] + [["", "**Summe**", "__ / 10", ""]],
                  breiten=[1.2, 8.4, 1.6, 5.8], groesse=11.5)
        d.absatz("**Das würden wir übernehmen:**")
        d.inhalt(A.Linien(3))
        d.absatz("**Das würden wir verbessern:**")
        d.inhalt(A.Linien(3))
    return d.speichern(ORDNER / "3_Unplugged" / "DD_A5_Unplugged_Architekturwettbewerb.docx")


# ================================================================== Unplugged DS 11: Regelkarten-Umzug
REGELKARTEN = [
    ("R1", "Entitätstyp", "Jeder Entitätstyp wird eine **Relation**. Das Schlüsselattribut wird **Primärschlüssel** "
     "(unterstreichen), die übrigen Attribute werden Spalten.",
     "Jede Entität braucht genau eine Zeile, die sich eindeutig ansprechen lässt – über den Primärschlüssel."),
    ("R2", "1 : n", "Der **Fremdschlüssel** kommt in die Relation der **n-Seite** (rot, mit ↑). Attribute des "
     "Beziehungstyps ebenfalls dorthin.",
     "Auf der n-Seite gehört zu jeder Zeile **genau ein** Partner – ein Wert passt in die Zelle. Auf der 1-Seite bräuchte "
     "eine Zeile mehrere Werte."),
    ("R3", "n : m", "Eine **eigene Relation** mit beiden Fremdschlüsseln; zusammen bilden sie den **Primärschlüssel**. "
     "Attribute des Beziehungstyps kommen hinzu.",
     "Auf keiner Seite passt ein einzelner Wert: sonst mehrere Werte in einer Zelle oder wiederholte Daten → Redundanz "
     "und Anomalien. Kommt ein Paar mehrfach vor → zusätzliches Schlüsselattribut."),
    ("R4", "1 : 1", "Der Fremdschlüssel kommt in **eine** der beiden Relationen – oder beide Relationen werden "
     "zusammengelegt.",
     "Auf beiden Seiten passt genau ein Wert; man wählt meist die Seite, die „später entsteht“ (der Ausweis zum "
     "Ermittler)."),
    ("R5", "rekursiv", "1:n → Fremdschlüssel auf die **eigene** Relation, benannt nach der **Rolle** (↑MentorNr). "
     "n:m → eigene Relation mit **zwei** Fremdschlüsseln auf dieselbe Relation.",
     "Es gelten R2 und R3 – nur zeigen die Fremdschlüssel auf dieselbe Relation. Der Rollenname verhindert zwei Spalten "
     "mit gleichem Namen."),
    ("R6", "ternär", "Eine **eigene Relation** mit **drei** Fremdschlüsseln im Primärschlüssel; Attribute des "
     "Beziehungstyps kommen hinzu.",
     "Nur das Trio hält fest, wer – in welchem Fall – wo. Kann dasselbe Trio mehrfach vorkommen (zwei Aussagen), "
     "kommt z. B. der Zeitpunkt in den Schlüssel."),
]


def unplugged_regelkarten():
    d = AktenDokument("unplugged", 5, "Regelkarten-Umzug", fassung="schueler", ermittlung="5.3", ds=11,
                      material="Unplugged A5 · Regelkarten", datenbank="keine",
                      untertitel="Kopiervorlage je Team: 6 Regelkarten (Seite 2 Vorderseite, Seite 3 Rückseite – beidseitig "
                                 "über die lange Kante drucken, laminieren) und leere Relationenkarten")
    d.kopfband(name_datum=False)
    d.inhalt(Liste(["Für jedes **Rechteck** eures Wettbewerbsplakats aus DS 10 entsteht eine Relationenkarte (R1), "
                    "Primärschlüssel unterstrichen.",
                    "Für jede **Raute** zieht ihr die passende Regelkarte und lasst die Beziehung „umziehen“: Fremdschlüssel "
                    "**rot mit ↑** auf eine Relationenkarte schreiben – oder eine neue Karte anlegen (n:m, ternär, rekursiv "
                    "n:m).",
                    "**Probe:** Tragt drei Beispielzeilen für *verdaechtig* ein. Passt jede Zeile ohne Mehrfachwerte und "
                    "ohne wiederholte Personendaten?",
                    "**Austausch:** Ein anderes Team fragt: „Welche Regel habt ihr hier angewendet – und warum nicht eine "
                    "andere?“ Die Antwort steht auf der Rückseite der Regelkarte."], nummeriert=True))
    d.notiz("Kurzschreibweise nach EH §7", "Relation(<u>Primärschlüssel</u>, Attribut, ↑Fremdschlüssel) – Primärschlüssel "
            "unterstrichen, Fremdschlüssel mit ↑ direkt vor dem Namen.")

    def vorne(r):
        return _karte(d, f"REGELKARTE {r[0]} · VORDERSEITE", r[1], [r[2]], farbe=A.SIEGELROT, groesse=26)

    def hinten(r):
        return _karte(d, f"REGELKARTE {r[0]} · RÜCKSEITE", "Warum?", [r[3]], farbe=A.BLAU, flaeche=A.BLAU_FLAECHE,
                      groesse=22)
    _karten_raster(d, [vorne(r) for r in REGELKARTEN], spalten=2, zeilen_je_seite=3, hoehe_cm=7.9, umbruch_vorher=True)
    # Rückseite: Spalten je Zeile getauscht (Duplex über die lange Kante)
    rueck = []
    for i in range(0, 6, 2):
        rueck += [hinten(REGELKARTEN[i + 1]), hinten(REGELKARTEN[i])]
    _karten_raster(d, rueck, spalten=2, zeilen_je_seite=3, hoehe_cm=7.9, umbruch_vorher=True)

    def relationenkarte(z, breite):
        p = z.paragraphs[0]
        d._etikett(p, "RELATIONENKARTE", groesse=7.5, farbe=A.TINTE_2, fett=True)
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(10)
        d._etikett(p, "RELATION ", groesse=8)
        d._lauf_stil(p.add_run("_" * 26), farbe=A.RAND_STARK)
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(4)
        d._text(p, "Primärschlüssel unterstreichen · Fremdschlüssel **rot** mit ↑ · Regel: R__", groesse=8.5,
                farbe=A.TINTE_2)
        for _ in range(4):
            d._schreiblinie(z, breite - 1.0)
    _karten_raster(d, [relationenkarte] * 6, spalten=2, zeilen_je_seite=3, hoehe_cm=7.9, umbruch_vorher=True)
    return d.speichern(ORDNER / "3_Unplugged" / "DD_A5_Unplugged_Regelkarten.docx")


# ================================================================== Lehrerhinweise
#: Stand der Präsentation, auf den sich die Folienangaben der Lehrerhinweise beziehen
FOLIEN_ERWARTET = {'ds9_start': 1, 'ds9_lab': 11, 'ds9_ende': 16, 'ds10_start': 17, 'ds10_wettbewerb': 25, 'ds10_lab': 28,
                   'ds10_ende': 34, 'ds11_start': 35, 'ds11_lab': 43, 'ds11_ende': 50}
VERLAUF_BREITEN = [2.3, 1.4, 7.2, 2.4, 3.7]


def lehrerhinweise():
    if not FOLIEN:
        praesentation()
    assert FOLIEN == FOLIEN_ERWARTET, f"Folienangaben der Lehrerhinweise anpassen: {FOLIEN}"
    d = AktenDokument("lehrerhinweise", 5, "Lehrerhinweise Akte 5", fassung="lehrkraft", material="Lehrerhinweise A5",
                      datenbank=DB, untertitel="Der Bauplan · DS 9–11 · nur eA (im gA-Kurs entfällt die Akte)")
    d.kopfband()
    d.ueberschrift("Überblick", 1)
    d.tabelle(["DS", "Thema", "Lab", "Warm-up", "Abi-Bezug"],
              [["9", "Baupläne lesen – Chen-Notation nach EH §7, Kardinalitäten (auch aus Daten), Attribute an "
                "Beziehungstypen", "A5-E1-01 … 08 + CC", "W-A4-05 … 08", "Kardinalitäten begründen (eA E22, E23N, E25); "
                "ER-Diagramm im Sachzusammenhang erläutern"],
               ["10", "Der Architekturwettbewerb – ER-Modell entwickeln und erweitern, rekursiv, zwei Beziehungstypen, "
                "ternär, Grenzen, beurteilen", "A5-E2-01 … 08 + CC", "W-A5-01 … 04", "ER ergänzen (eA 8/9) bzw. "
                "entwickeln (seit 2025); rekursive Beziehung (neu im EH-Beispiel, Prognose)"],
               ["11", "Vom Plan zur Datenbank – Überführung R1–R6 mit Begründung, Rückrichtung, Self-Join",
                "A5-E3-01 … 08 + CC; Verhör 5 (V5)", "W-A5-05 … 08", "Überführung (eA E21, E26), Rückrichtung "
                "(E21N, E25N), Self-Join (Prognose)"]],
              breiten=[0.9, 5.6, 3.2, 2.4, 4.9])
    d.absatz("Datenbank: **Detektei** – die Fall-Datenbank ist zugleich die Referenzlösung des Architekturwettbewerbs "
             "(DS 10) und der Überführung (DS 11). Im Lab steht ihre Datenbank-Kachel nur in A5-E3-08 und A5-E3-CC, sonst "
             "stehen Tabellenausschnitte direkt in den Aufgaben. Relationenschema:")
    for r in ["Ermittler(<u>ErmittlerNr</u>, Name, Dienstgrad, ↑MentorNr)", "Fall(<u>FallNr</u>, Titel, Eroeffnet, Status, ↑LeiterNr)",
              "Person(<u>PersonNr</u>, Vorname, Nachname, Jahrgang, Stadtteil)", "kennt(<u>↑PersonNr</u>, <u>↑BekannterNr</u>, Seit)",
              "verdaechtig(<u>↑PersonNr</u>, <u>↑FallNr</u>, Grad)", "Ort(<u>OrtNr</u>, Bezeichnung, Stadtteil)",
              "Aussage(<u>↑PersonNr</u>, <u>↑FallNr</u>, <u>↑OrtNr</u>, <u>Zeitpunkt</u>, Inhalt)",
              "Beweisstueck(<u>BeweisNr</u>, Bezeichnung, ↑FallNr, ↑OrtNr)"]:
        d.inhalt(Schema(r))
    d.ueberschrift("Material dieser Akte", 2)
    d.tabelle(["Datei", "Zweck"],
              [["1_Praesentation/DD_A5_Bauplan.pptx", "Fallbesprechung DS 9 (Folien 1–16), DS 10 (17–34), DS 11 (35–50) mit "
                "Sprechernotizen; Lösungsfolien nur in der Lehrkraftfassung"],
               ["2_Arbeitsblaetter/DD_A5_AB1_… AB2_… AB3_…", "AB 5.1, 5.2, 5.3 (Lösungen in 5_Loesungen, *_LSG)"],
               ["3_Unplugged/…Faeden-Kardinalitaeten", "DS 9: 20 Rollenkarten + Etiketten"],
               ["3_Unplugged/…Architekturwettbewerb", "DS 10: Anforderungstext (Plakat), Teamrollen, Bewertungsbogen, "
                "Siegerstempel"],
               ["3_Unplugged/…Regelkarten", "DS 11: Regelkarten R1–R6 (beidseitig) + Relationenkarten"],
               ["Warm-up/DD_Warm-up-Karten_A5.docx", "W-A5-01 … 04 (DS 10), 05 … 08 (DS 11), 09 … 12 (DS 12, eA)"]],
              breiten=[6.2, 10.8])
    d.ueberschrift("Korrekturen gegenüber der früheren Präsentation „ER-Modell“", 2)
    d.inhalt(Liste(["Fachbegriffe nach EH: **Entitätstyp** und **Beziehungstyp** statt „Entitätsmenge“.",
                    "Schlüsselattribute werden im ER-Diagramm **unterstrichen** (EH §7) – in allen Diagrammen dieser Akte.",
                    "Neue Regeln: **rekursive** Beziehungen (R5, Fremdschlüssel nach Rolle benannt) und **ternäre** "
                    "Beziehungen (R6); **Attribute an 1:n-Beziehungstypen** wandern in die Relation der n-Seite.",
                    "**Zusätzliches Schlüsselattribut**, wenn dasselbe Paar (Trio) mehrfach vorkommen kann (z. B. Zeitpunkt "
                    "in Aussage, Beginn bei einer Platzbuchung).",
                    "Notation nur Chen mit 1/n/m (ternär p) – keine (min,max)-, keine Krähenfuß-Notation."]))

    # ---------------------------------------------------------------- DS 9
    d.ueberschrift("DS 9 · Baupläne lesen (Ermittlung 5.1)", 1)
    d.ueberschrift("Ziele", 2)
    d.absatz("Die Schülerinnen und Schüler …")
    d.inhalt(Liste(["**benennen** die Elemente der Chen-Notation nach EH §7 und **interpretieren** ein ER-Diagramm im "
                    "Sachzusammenhang (K33, K47),",
                    "**bestimmen** Kardinalitäten 1:1, 1:n und n:m aus Texten, Mengendiagrammen und Tabellendaten und "
                    "**begründen** sie in beide Leserichtungen (K34),",
                    "**modellieren** Attribute an Beziehungstypen und **begründen**, warum sie zum Paar gehören (K36),",
                    "**verwenden** die Fachsprache der EH konsequent."], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt / Ablauf", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A4-05 … 08 (Schlüsselkriterien, Ausschnitt-Falle, Zerlegung, zusammengelegte "
                "Tabelle)", "EA → Team", "Warm-up-Karten A4, Folie 2"],
               ["Briefing", "10–20'", "Akte 5: „Bevor wir bauen, lernen wir Baupläne lesen.“ Chen-Notation am Plan "
                "Ermittler–leitet–Fall; Kardinalitäten, Leseschablone, Beziehungsattribut, Daten stützen", "LV, UG",
                "Folien 3–10"],
               ["Unplugged", "20–35'", "Fäden-Kardinalitäten (siehe unten)", "Plenum", "Rollenkarten, Fäden, Etiketten, "
                "Folie 11"],
               ["Lab", "35–80'", "Ermittlung 5.1 im Bauplan-Tisch (lesen/ergänzen, Listenmodus auf dem iPad); AB 5.1 "
                "parallel", "EA", "Geräte, AB 5.1, Folie 12"],
               ["Fallbesprechung", "80–90'", "Leseschablone in beide Richtungen, Merkkasten, Abi-Ecke (4 BE), Cliffhanger: "
                "Juris 14 Entitätstypen ohne Kardinalität; Teams für DS 10 bilden", "UG", "Folien 13–16"]],
              breiten=VERLAUF_BREITEN)
    d.ueberschrift("Unplugged-Ablauf: Fäden-Kardinalitäten (15 Minuten)", 2)
    je_ermittler, je_fall = faeden_auswertung()
    d.absatz("Material: Rollenkarten zum Umhängen (4 Ermittler, 6 Fälle, 6 Personen, 4 Dienstausweise – Kopiervorlage), "
             "Wollfäden in zwei Farben, Klebeetiketten. Was jede Karte verbinden soll, steht auf der Karte – die Lehrkraft "
             "moderiert nur. Bei kleinen Kursen übernehmen Personen zwei Karten.")
    d.inhalt(Liste([
        "**1:n · leitet:** Jede Fall-Karte nimmt einen Faden zu ihrer Leitung. Ablesen: " + "; ".join(
            f"{n} {len(f)} {'Faden' if len(f) == 1 else 'Fäden'}" for n, f in je_ermittler.items())
        + " – jeder Fall genau einen. Notation: Ermittler 1 – *leitet* – n Fall.",
        "**n:m · verdaechtig:** Personen spannen Fäden zu ihren Fällen (zweite Wollfarbe). Mehrere Fäden auf beiden Seiten: "
        + "; ".join(f"Fall {f}: {len(p)} Personen" for f, p in je_fall.items() if len(p) > 1) + "; Person 3 und 5 je zwei, "
        "Person 4 drei Fälle.",
        "**1:1 · hat:** Jeder Ermittler zu seinem Dienstausweis (D-01 … D-04).",
        "**Beziehungsattribut:** Das Etikett „Grad: hoch“ für Person 3 – Fall 2 kommt an den **Faden** – Person 3 ist in "
        "Fall 5 nur „gering“ verdächtig, Fall 2 hat noch eine andere verdächtige Person.",
        "**Impuls:** „Woran erkennt ihr in einer Tabelle, welche Kardinalität vorliegt?“ → Wert der Fremdschlüsselspalte "
        "kommt mehrfach vor → n-Seite."], nummeriert=True))
    d.absatz("Sicherung: Foto der Fadennetze; Tafel: drei Kardinalitäten mit Leseschablone. Hinweis: Juri hält keinen "
             "Faden – bei 1:n darf ein Ermittler auch keinen Fall leiten.")
    d.ueberschrift("Tafelbild / Merkkasten „Ermittler-Notiz“", 2)
    d.inhalt(gross(er_leitet(), 11), gross(er_verdacht(), 10),
             "**Leseschablone:** „Ein Ermittler leitet mehrere (n) Fälle; jeder Fall wird von genau einem (1) Ermittler "
             "geleitet.“ · „Eine Person kann in mehreren Fällen verdächtig sein, in einem Fall können mehrere Personen "
             "verdächtig sein (n:m); der Verdachtsgrad gehört zum Paar Person–Fall.“",
             "**Kardinalität aus Daten:** Kommt in `verdaechtig` dieselbe PersonNr mit verschiedenen FallNr vor **und** "
             "dieselbe FallNr mit verschiedenen PersonNr → n:m. Daten stützen nur, die Begründung kommt aus dem "
             "Sachzusammenhang. · Fachbegriffe: Entitätstyp, Beziehungstyp, Schlüsselattribut, Kardinalität.")
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["Kardinalität an der falschen Kante (**P9**) → Leseschablone: Die Zahl steht beim Entitätstyp, über "
                    "den man spricht („… leitet *n Fälle*“ → n an der Kante zu Fall).",
                    "Kardinalität aus wenigen Zeilen „bewiesen“ (vgl. **P8**) → Sachzusammenhang als Begründung.",
                    "Beziehungsattribut an einen Entitätstyp gehängt (**P9**) → Merkfrage „anderes Paar – anderer Wert?“.",
                    "Erläutern ohne Sachzusammenhang (**P10**) → „Erläutern ist kein vornehmes Wort für nennen.“",
                    "Altbegriffe aus älteren Materialien → konsequent korrigieren: Entitätstyp statt „Entitätsmenge“."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Case „Dienstwagen“ (AB) und A5-E1-CC (Lab); Anforderungstext für DS 10 vorab lesen "
                    "und Entitätstypen markieren.",
                    "*Langsame:* Listenmodus im Bauplan-Tisch, Werkzeugkarte 4, Fotos der Fadennetze als Erinnerung."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Rollenkarten-Set (Kopiervorlage, laminiert, mit Band), Wollfäden in zwei Farben, Etiketten.",
                    "Akte 5 nur in eA-Kursen freischalten; Bauplan-Tisch auf den Schülergeräten testen (Touch, Listenmodus).",
                    "Warm-up-Karten A4 (05–08) kopieren."]))

    # ---------------------------------------------------------------- DS 10
    d.ueberschrift("DS 10 · Der Architekturwettbewerb (Ermittlung 5.2)", 1)
    d.ueberschrift("Ziele", 2)
    d.absatz("Die Schülerinnen und Schüler …")
    d.inhalt(Liste(["**entwickeln** ein ER-Modell aus einem Anforderungstext mit vollständigen Schlüsselattributen und "
                    "Kardinalitäten (K35),",
                    "**modellieren** rekursive Beziehungen und mehrere Beziehungstypen zwischen denselben Entitätstypen, "
                    "**lesen** und **ergänzen** ternäre Beziehungen (K37, K38),",
                    "**erweitern** ein gegebenes ER-Modell nach neuen Anforderungen (K39),",
                    "**beurteilen** Modellierungen und **erläutern** Grenzen der Notation (K40)."], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt / Ablauf", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A5-01 … 04; Karte 03 gemeinsam auflösen", "EA → Team",
                "Warm-up-Karten A5, Folien 17–18"],
               ["Briefing", "10–20'", "Neue Bausteine: rekursiv: *istPateVon* → *betreut*; zwei Beziehungstypen "
                "(Heim/Gast); ternär: *sagtAus*; Beispiel Tierheim vom Text zum Modell", "LV, UG", "Folien 19–25"],
               ["Wettbewerb", "20–50'", "Teams modellieren die Fall-Datenbank nach dem Anforderungstext (Plakat A2 oder "
                "Bauplan-Tisch, Modus entwickeln)", "Teams", "Anforderungstext, Plakate, Teamrollen, Folien 26–27"],
               ["Galerie", "50–65'", "Galerie-Rundgang: jedes Team bewertet zwei fremde Entwürfe mit dem Bewertungsbogen; "
                "Siegerentwurf per Punktekleber", "Teams, Plenum", "Bewertungsbögen, Klebepunkte, Folie 28"],
               ["Lab", "65–85'", "Ermittlung 5.2: 01–06 in der Stunde, 07/08 als Hausaufgabe", "EA", "Geräte, AB 5.2, "
                "Folie 29"],
               ["Fallbesprechung", "85–90'", "Referenzlösung neben dem Siegerentwurf; Abweichungen als Alternativen würdigen "
                "(K40); Hausaufgabe AB 5.2, Lab 07/08", "UG", "Folien 30–34"]],
              breiten=VERLAUF_BREITEN)
    d.absatz("*Abweichung vom Standardtakt:* Wettbewerb und Galerie ersetzen den Großteil der Lab-Zeit (20' statt 45'); der "
             "Lab-Rest wird Hausaufgabe.")
    d.ueberschrift("Unplugged-Ablauf: Architekturwettbewerb mit Galerie-Rundgang (45 Minuten)", 2)
    d.inhalt(Liste(["*Modellierphase (30'):* Rollen im Team – Moderation, Zeichner:in, Textprüfer:in (hakt jede "
                    "Information im Text ab), Notationsprüfer:in (EH §7). Zeitansage nach 15'.",
                    "*Galerie (15'):* Plakate hängen aus; jedes Team bewertet zwei Entwürfe mit dem Bewertungsbogen (je "
                    "Kriterium 1 Punkt) und dem Pflichtkommentar „Das würden wir übernehmen: …“.",
                    "*Siegerehrung:* Punktekleber; das Siegerteam erhält den Stempel ARCHITEKTUR (Kopiervorlage).",
                    "*Sicherung:* Fotos aller Plakate in den Kurs-Ordner; Referenzlösung im Lab nach A5-E2-07, auf Folie "
                    "30 nur für die Lehrkraft. Teamplakate für DS 11 aufheben!"], nummeriert=True))
    d.absatz("Bewertungsbogen = Kernpunkte der Referenzlösung:")
    d.inhalt(Liste(KRITERIEN, nummeriert=True), gross(er_referenz(), 15.5))
    d.ueberschrift("Tafelbild / Merkkasten „Ermittler-Notiz“", 2)
    d.inhalt(Liste(["**Rekursiv** = Beziehungstyp mit zwei Kanten zum selben Entitätstyp; Rollen an die Kanten "
                    "(Ermittler 1 Mentor – *betreut* – n Neuling; Person n – *kennt* – m Person mit Seit).",
                    "**Zwei Rollen** (Heim/Gast, Leitung/Verdacht) = zwei Beziehungstypen.",
                    "**Ternär** (m:n:p): eine Raute, drei Kanten; „wer – in welchem Fall – wo“ ist nur zu dritt eindeutig.",
                    "**Grenzen der Notation:** „genau drei Ermittler je Team“, „höchstens 12 Teilnehmende“, „Mentor muss "
                    "höheren Dienstgrad haben“ → Anmerkung zum Modell.",
                    "**Beurteilen:** Alternativen benennen (Attribut vs. eigener Entitätstyp, n:m vs. aufgelöste "
                    "Zwischenentität), Folgen abwägen, Urteil begründen."]))
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["Rekursive Beziehung als zweiter Entitätstyp „Mentor“ modelliert (**P9**) → Rollen an Kanten.",
                    "Ternäre Beziehung in drei binäre zerlegt – Information „wer – wo – in welchem Fall“ geht verloren.",
                    "Attribute vergessen oder Schlüsselattribute nicht markiert (**P9**).",
                    "Galerie-Bewertung zu oberflächlich → Bewertungsbogen mit Pflichtkommentar.",
                    "Zeit: Teams verlieren sich in Details → Textprüfer:in hakt ab, Zeitansage nach 15'."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Cases (AB: Komplizenschaft in einem Fall – ternär und rekursiv; Lab A5-E2-CC), "
                    "AB 5.2 Auftrag 3 (Vorgängerfall), Rolle Notationsprüfer:in.",
                    "*Langsame:* Startgerüst mit vorgegebenen Entitätstypen, Listenmodus, Werkzeugkarte 4; im Team die Rolle "
                    "Textprüfer:in."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Plakate A2, Stifte, Klebepunkte, Bewertungsbögen (zwei je Team), Teamrollen, Siegerstempel "
                    "(Kopiervorlage Architekturwettbewerb); Anforderungstext als Folie 26.",
                    "Die Referenzlösung im Bauplan-Tisch erscheint erst nach dem Lösen von A5-E2-07 (serverseitig).",
                    "Warm-up-Karten A5 (01–04) kopieren."]))

    # ---------------------------------------------------------------- DS 11
    d.ueberschrift("DS 11 · Vom Plan zur Datenbank (Ermittlung 5.3)", 1)
    d.ueberschrift("Ziele", 2)
    d.absatz("Die Schülerinnen und Schüler …")
    d.inhalt(Liste(["**überführen** ein ER-Modell nach den Regeln R1–R6 (auch rekursiv, ternär, Beziehungsattribute) in ein "
                    "Relationenschema in Kurzschreibweise (K41),",
                    "**begründen** die Umsetzungsregeln, insbesondere warum n:m eine eigene Relation braucht und 1:n nicht, "
                    "und wann ein zusätzliches Schlüsselattribut nötig ist (K42),",
                    "**rekonstruieren** aus einem Relationenschema das ER-Diagramm (K43),",
                    "**entwerfen** Abfragen auf der Fall-Datenbank, auch mit Self-Join über rekursive Beziehungen (K23, "
                    "K24, K16)."], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt / Ablauf", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A5-05 … 08 (rekursiv, ternär, Text → Mini-ER, Grenzen)", "EA → Team",
                "Warm-up-Karten A5, Folie 35"],
               ["Briefing", "10–20'", "„Der Siegerentwurf steht – jetzt wird gebaut.“ Regeln fragend-entwickelnd: Wohin "
                "kommt der Verweis bei *leitet*? Warum kein FallNr in Person? Regeln R1–R6, Feinheiten, Rückrichtung, "
                "Self-Join", "UG", "Folien 36–42"],
               ["Unplugged", "20–35'", "Regelkarten-Umzug mit dem eigenen Wettbewerbsplakat (siehe unten)", "Teams",
                "Regelkarten, Relationenkarten, Plakate, Folie 43"],
               ["Lab", "35–80'", "Ermittlung 5.3: Schema-Werkbank (02, 04, 07), Rückrichtung (06), SQL mit Self-Join (08); "
                "AB 5.3 parallel", "EA", "Geräte, AB 5.3, Folie 44"],
               ["Fallbesprechung", "80–90'", "Merkkasten Überführungsregeln mit Begründungen (Werkzeugkarte 5), Abi-Ecke "
                "(5 BE), Beweisstück „Ermittlungsnetz“, Cliffhanger; Hinweis Verhör 5", "UG", "Folien 45–50"]],
              breiten=VERLAUF_BREITEN)
    d.ueberschrift("Unplugged-Ablauf: Regelkarten-Umzug (15 Minuten)", 2)
    d.inhalt(Liste(["Material je Team: 6 Regelkarten (Vorderseite Regel, Rückseite Begründung), das eigene Plakat aus "
                    "DS 10, leere Relationenkarten, rote und schwarze Stifte.",
                    "Für jedes Rechteck eine Relationenkarte (R1), Schlüssel unterstrichen.",
                    "Für jede Raute die passende Regelkarte ziehen und „umziehen“: Fremdschlüssel rot mit ↑ oder neue Karte "
                    "(n:m, ternär, rekursiv n:m).",
                    "Probe mit drei Beispielzeilen für *verdaechtig*: ohne Mehrfachwerte, ohne wiederholte Personendaten?",
                    "Austausch zweier Teams: „Welche Regel habt ihr hier angewendet – und warum nicht eine andere?“",
                    "Sicherung: Regelkarten bleiben laminiert im Team; Werkzeugkarte 5 im Ermittler-Handbuch."],
                   nummeriert=True))
    d.ueberschrift("Tafelbild / Merkkasten „Ermittler-Notiz“ (Detektei)", 2)
    d.tabelle(["Regel", "ER", "Umsetzung (Detektei)", "Begründung"],
              [["R1 Entitätstyp", "Fall", "Fall(<u>FallNr</u>, Titel, Eroeffnet, Status, …)", "Schlüsselattribut → Primärschlüssel"],
               ["R2 1:n", "Ermittler 1 – *leitet* – n Fall", "Fall(<u>FallNr</u>, …, ↑LeiterNr)", "ein Wert je Zeile auf der "
                "n-Seite; Beziehungsattribute ebenfalls dorthin"],
               ["R3 n:m", "Person n – *verdaechtig* – m Fall", "verdaechtig(<u>↑PersonNr</u>, <u>↑FallNr</u>, Grad)",
                "auf keiner Seite passt ein Wert; Paar ist Schlüssel; Paar mehrfach → zusätzliches Schlüsselattribut"],
               ["R4 1:1", "Ermittler 1 – *hat* – 1 Dienstausweis", "Dienstausweis(<u>AusweisNr</u>, …, ↑ErmittlerNr)",
                "FK auf einer Seite oder zusammenlegen"],
               ["R5 rekursiv", "*betreut* (1:n), *kennt* (n:m)", "Ermittler(…, ↑MentorNr) · kennt(<u>↑PersonNr</u>, "
                "<u>↑BekannterNr</u>, Seit)", "wie R2/R3, Fremdschlüssel nach Rolle benannt"],
               ["R6 ternär", "Person, Fall, Ort – *sagtAus*", "Aussage(<u>↑PersonNr</u>, <u>↑FallNr</u>, <u>↑OrtNr</u>, "
                "<u>Zeitpunkt</u>, Inhalt)", "Zeitpunkt im Schlüssel: Gesa sagte zu Fall 1 im Fundbüro zweimal aus"]],
              breiten=[2.2, 3.8, 6.0, 5.0], groesse=8.5)
    d.inhalt(SQL(SELFJOIN), ergebnis(SELFJOIN),
             "**Ermittler-Notiz:** Self-Join = dieselbe Tabelle zweimal mit verschiedenen Aliasen, Verbund über den "
             "rekursiven Fremdschlüssel. Rückrichtung: eigene Schlüssel → Entitätstyp; Fremdschlüssel in einer "
             "Entitätsrelation → 1:n; Schlüssel aus zwei/drei Fremdschlüsseln → n:m bzw. ternär.")
    d.ueberschrift("Beweisstück „Ermittlungsnetz“ (Fallbesprechung)", 2)
    alle4 = ("SELECT Person.PersonNr, Person.Vorname, Person.Nachname\nFROM Person, verdaechtig\n"
             "WHERE Person.PersonNr = verdaechtig.PersonNr\n  AND verdaechtig.FallNr IN (1, 2, 3, 4)\n"
             "GROUP BY Person.PersonNr, Person.Vorname, Person.Nachname\nHAVING COUNT(*) = 4;")
    d.inhalt(SQL(alle4), ergebnis(alle4),
             "Nora Falter ist als **einzige** Person in allen vier Fällen verdächtig; über *kennt* ist sie mit Menschen an "
             "allen Tatorten verbunden (Gesa Ohlendiek, Bent Hollwedel, Selin Aksoy, Wiebke Lüttjohann) und mit Keno "
             "Bargmann aus der IT der Stadtverwaltung – eine Insiderin. Cliffhanger: Eine Nora Falter ist in Weserstedt "
             "nicht gemeldet. (Variante gleichnamig: eine zweite Nora Falter – deshalb nach dem Schlüssel gruppieren.)")
    d.ueberschrift("Verhör 5 · Kletterhalle", 2)
    d.absatz("Pool von 15 Aufgaben (A5-V-01 … 15) zum Kontext Kletterhalle (nur ER und Schema, keine SQL-Aufgaben), "
             "Ziehung 8 Aufgaben (3 × ★, 3 × ★★, 2 × ★★★), 20 Minuten, Abimodus streng, kein Feedback während des Verhörs. "
             "Nur eA. Freischaltung im Kurs: „auto“ oder „offen“ (Hausaufgabe bis DS 12). Video 6 freigeben.")
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["Fremdschlüssel auf der 1-Seite (FallNr in Ermittler, **P9**) → Probedaten: Ein Ermittler mit drei "
                    "Fällen braucht drei Werte in einer Zelle.",
                    "n:m ohne eigene Relation oder Beziehungsrelation ohne zusammengesetzten Schlüssel (**P9**).",
                    "Beziehungsattribut bei 1:n in die falsche Relation (**P9**) → gehört auf die n-Seite.",
                    "Rekursiver Fremdschlüssel gleich benannt wie der Primärschlüssel → Rollenname (MentorNr).",
                    "Self-Join ohne zwei Aliase oder Spalten ohne Alias-Präfix (**P7**); Ermittler ohne Mentor verschwinden "
                    "(MentorNr NULL, **P5**)."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Case „Mentor-Kette“ (AB und Lab A5-E3-CC); eigene Wettbewerbslösung in der "
                    "Schema-Werkbank nachbauen; Ausblick Kür K1 (eigene Datenbank).",
                    "*Langsame:* Regelkarten liegen aus; Schema-Werkbank mit Startzustand aus dem ER; Werkzeugkarte 5."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Regelkarten-Set je Team (beidseitig gedruckt, laminiert), Relationenkarten; Teamplakate aus DS 10.",
                    "Verhör 5 im Kurs auf „auto“/„offen“; Video 6 freigeben; Warm-up-Karten A5 (05–08) kopieren."]))
    d.ueberschrift("Hinweise zum Niveau", 2)
    d.absatz("Akte 5 ist **nur eA** (alle Aufgaben, Arbeitsblätter und Verhör 5). Im gA-Kurs entfällt sie: Die Lehrkraft "
             "stellt „Niveau: gA“ ein, die Akte wird ausgeblendet, der Steckbrief-Hinweis der Akte wird beim Freischalten "
             "von Akte 6 automatisch aufgedeckt; das Warm-up in DS 12 nutzt dort W-A4-05 … 08. Die Arbeitsblätter sind "
             "nicht 1:1 die Lab-Aufgaben: Sie üben dieselben Kompetenzen mit anderen Beispielen (Mengendiagramme, "
             "Stadtbücherei, Entwurf „Hafenkante“, Komplizenschaft); die Abi-Ecken entsprechen den Operatoren begründen, "
             "erweitern und überführen/begründen (AFB II bzw. II/III).")
    return d.speichern(ORDNER / "4_Lehrerhinweise" / "DD_A5_Lehrerhinweise.docx")


# ================================================================== Warm-up-Karten
def warmup():
    karten = json.loads((REPO / "seed" / "warmup" / "warmup_a5.json").read_text(encoding="utf-8"))["karten"]
    return warmup_karten_dokument(karten, 5, "Warm-up-Karten Akte 5").speichern(MAT / "Warm-up" / "DD_Warm-up-Karten_A5.docx")


def main():
    for f in ("schueler", "loesung"):
        ab1(f)
        ab2(f)
        ab3(f)
    praesentation()
    unplugged_faeden()
    unplugged_wettbewerb()
    unplugged_regelkarten()
    lehrerhinweise()
    warmup()
    for p in sorted(ORDNER.rglob("*")):
        if p.is_file():
            print(f"✔ {p.relative_to(REPO)}")
    print("✔ materialien/Warm-up/DD_Warm-up-Karten_A5.docx")


if __name__ == "__main__":
    main()
