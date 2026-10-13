#!/usr/bin/env python3
"""Erstellungshilfe: Material Akte 3 „Die Stadtfest-Kasse“ (DS 5–6) im Akten-Design.

Erzeugt in materialien/Akte3_Stadtfest/:
  1_Praesentation/DD_A3_Stadtfest.pptx
  2_Arbeitsblaetter/DD_A3_AB1_Kassensturz.docx, DD_A3_AB2_Die_Statistik_luegt_nicht.docx
  3_Unplugged/DD_A3_Unplugged_GROUP_BY_Karten.docx (+ _LSG), DD_A3_Unplugged_Zeitlupe_Karten.docx (+ _LSG)
  4_Lehrerhinweise/DD_A3_Lehrerhinweise.docx
  5_Loesungen/DD_A3_AB1_Kassensturz_LSG.docx, DD_A3_AB2_Die_Statistik_luegt_nicht_LSG.docx
und materialien/Warm-up/DD_Warm-up-Karten_A3.docx.

Grundlage: docs/DS_PLANUNG.md (DS 5, DS 6), docs/STORY.md (5.5, 5.6, 4.3), seed/akten/akte3.json (Spur im Lab),
seed/datenbanken/stadtfest/LIESMICH.md, docs/KOMPETENZMATRIX.md (AB 3.1/…, AB 3.2/…).
Alle Ergebnistabellen werden aus der Stadtfest-DB berechnet (abfrage_ergebnis). Die Bon-Karten des Unplugged sind
eigene Daten (DS_PLANUNG: „eigene Daten“).

Die erzeugten docx/pptx sind danach maßgeblich (Spec §12); dieses Skript ist nur die Erstellungshilfe.
Aufruf: python3 quellen/material/akte3.py   (überschreibt die oben genannten Dateien)
"""
import json
import re
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "werkzeuge"))

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import akten_design as ad  # noqa: E402
from akten_design import (  # noqa: E402
    AktenDokument, AktenPraesentation, Auftrag, Bild, Ergebnis, Liste, Loesung, Notiz, Schema, SQL, Tabelle,
    abfrage_ergebnis, warmup_karten_dokument)

DB = "stadtfest"
ORDNER = REPO / "materialien" / "Akte3_Stadtfest"
WARMUP = REPO / "materialien" / "Warm-up"
TMP = Path(tempfile.mkdtemp(prefix="dd-akte3-"))

SCHEMA = [
    "Stand(<u>StandNr</u>, Name, Art, Platz)",
    "Helfer(<u>HelferNr</u>, Vorname, Nachname, Jahrgang, Verein)",
    "Schicht(<u>↑HelferNr, ↑StandNr, Beginn</u>, Ende)",
    "Produkt(<u>ProduktNr</u>, Bezeichnung, Preis, ↑StandNr)",
    "Bon(<u>BonNr</u>, ↑StandNr, ↑HelferNr, Zeitpunkt, Zahlart)",
    "Position(<u>↑BonNr, ↑ProduktNr</u>, Menge)",
]


# ------------------------------------------------------------------ Hilfen
def _runden(zelle):
    """Lange Kommazahlen (AVG) auf zwei Nachkommastellen – die SQL-Prüfung vergleicht auf die angezeigten Stellen."""
    if re.fullmatch(r"-?\d+,\d{3,}", zelle):
        return f"{float(zelle.replace(',', '.')):.2f}".replace(".", ",")
    return zelle


def ergebnis(sql, max_zeilen=None):
    kopf, zeilen = abfrage_ergebnis(sql, DB)
    zeilen = [[_runden(c) for c in z] for z in zeilen]
    if max_zeilen and len(zeilen) > max_zeilen:
        return Ergebnis(kopf, zeilen[:max_zeilen], abgeschnitten=True)
    return Ergebnis(kopf, zeilen)


def daten(sql):
    """Datentabelle (kein Ergebnisformat) aus der DB, z. B. Ausschnitte für Aufgaben."""
    return abfrage_ergebnis(sql, DB)


def umbrechen(text, breite=60):
    """Lange SELECT-/GROUP-BY-Zeilen nach Kommas umbrechen (Fortsetzung eingerückt) – passt in schmale Lösungskästen."""
    aus = []
    for zeile in text.split("\n"):
        m = re.match(r"(SELECT |GROUP BY )(.*)", zeile)
        if len(zeile) <= breite or not m:
            aus.append(zeile)
            continue
        einzug, teile, aktuell = " " * len(m.group(1)), m.group(2).split(", "), m.group(1)
        for i, t in enumerate(teile):
            stueck = t + ("," if i < len(teile) - 1 else "")
            if aktuell.strip() and len(aktuell) + len(stueck) + 1 > breite and aktuell not in (m.group(1), einzug):
                aus.append(aktuell.rstrip())
                aktuell = einzug
            aktuell += stueck + " "
        aus.append(aktuell.rstrip())
    return "\n".join(aus)


def S(text, breite=60, **kw):
    """SQL mit Semikolon am Ende (Schreibweise der Materialien), lange Zeilen umbrochen."""
    text = umbrechen(text.strip("\n").rstrip(";"), breite)
    return SQL(text + ";", **kw)


def schema_inhalt():
    return [Schema(s) for s in SCHEMA]


# ------------------------------------------------------------------ Bon-Karten (eigene Daten, Unplugged DS 5)
STAENDE_UNPLUGGED = {
    "Kaffeemobil": {"preis": 2.5, "helfer": ["Mika R.", "Lena K."],
                    "menge": [1, 1, 1, 2, 1, 1, 1, 1, 2, 1],
                    "zahlart": ["bar", "Karte", "bar", "bar", "Karte", "bar", "Karte", "bar", "bar", "Karte"]},
    "Grillhütte": {"preis": 3.5, "helfer": ["Jonas T.", "Emma S."],
                   "menge": [2, 1, 3, 1, 2, 1, 2, 1],
                   "zahlart": ["Karte", "bar", "bar", "Karte", "bar", "Karte", "bar", "bar"]},
    "Saftbar": {"preis": 3.0, "helfer": ["Finn B.", "Ida W."],
                "menge": [1, 2, 1, 1, 2, 1, 1],
                "zahlart": ["bar", "bar", "Karte", "bar", "bar", "Karte", "bar"]},
    "Crêpes-Stand": {"preis": 3.0, "helfer": ["Noah M.", "Lia H."],
                     "menge": [4, 3, 5, 4, 3],
                     "zahlart": ["Karte", "bar", "Karte", "Karte", "bar"]},
}


def bon_karten():
    """30 Bon-Karten, BonNr 101–130 reihum über die Stände verteilt (deterministisch)."""
    zeiger = {s: 0 for s in STAENDE_UNPLUGGED}
    karten, nr = [], 101
    while any(zeiger[s] < len(d["menge"]) for s, d in STAENDE_UNPLUGGED.items()):
        for s, d in STAENDE_UNPLUGGED.items():
            i = zeiger[s]
            if i >= len(d["menge"]):
                continue
            menge = d["menge"][i]
            karten.append({"bon": nr, "stand": s, "helfer": d["helfer"][i % 2], "zahlart": d["zahlart"][i],
                           "menge": menge, "betrag": menge * d["preis"]})
            zeiger[s] += 1
            nr += 1
    return karten


def euro(x):
    return f"{x:.2f}".replace(".", ",") + " €"


def auswertung_unplugged(karten, nur_karte=False):
    zeilen = {}
    for k in karten:
        if nur_karte and k["zahlart"] != "Karte":
            continue
        z = zeilen.setdefault(k["stand"], [0, 0, 0.0])
        z[0] += 1
        z[1] += k["menge"]
        z[2] += k["betrag"]
    return zeilen


def faecher_png(pfad):
    """Grafik „GROUP BY = Fächer bilden“ (Unplugged DS 5) mit eigenen Daten der Bon-Karten."""
    from PIL import Image, ImageDraw
    karten = bon_karten()
    auswahl = {"Kaffeemobil": 4, "Grillhütte": 3, "Crêpes-Stand": 2}
    faecher = {s: [k for k in karten if k["stand"] == s][:n] for s, n in auswahl.items()}
    alle = sorted([k for ks in faecher.values() for k in ks], key=lambda k: k["bon"])
    W, H = 1800, 800
    bild = Image.new("RGB", (W, H), "#" + ad.KARTE)
    d = ImageDraw.Draw(bild)
    rgb = lambda h: "#" + h  # noqa: E731
    f_etikett, f_karte, f_klein = ad._font("mono", 24), ad._font("text_halb", 27), ad._font("text", 23)
    f_kopf, f_erg = ad._font("kopf", 40), ad._font("mono", 28)

    def karte(x, y, k, b=176, mit_stand=True):
        h = 104 if mit_stand else 74
        d.rounded_rectangle([x, y, x + b, y + h], radius=8, fill=rgb(ad.PAPIER), outline=rgb(ad.RAND_STARK), width=2)
        d.text((x + 12, y + 6), f"Bon {k['bon']}", font=f_karte, fill=rgb(ad.TINTE))
        d.text((x + 12, y + 40), f"Menge {k['menge']}", font=f_klein, fill=rgb(ad.TINTE_2))
        if mit_stand:
            d.text((x + 12, y + 70), k["stand"], font=f_klein, fill=rgb(ad.TINTE_3))

    # 1. Reihe: ungeordnete Bon-Karten
    d.text((30, 14), "1 · ALLE BON-KARTEN (FROM Bon)", font=f_etikett, fill=rgb(ad.SIEGELROT))
    for i, k in enumerate(alle):
        karte(30 + i * 194, 52, k)
    # 2. Fächer
    d.text((30, 180), "2 · GROUP BY Stand – je Stand ein Fach", font=f_etikett, fill=rgb(ad.SIEGELROT))
    fach_b, abstand = 560, 30
    for i, (stand, ks) in enumerate(faecher.items()):
        x0, y0 = 30 + i * (fach_b + abstand), 218
        d.rounded_rectangle([x0, y0, x0 + fach_b, y0 + 250], radius=14, fill=rgb(ad.FLAECHE), outline=rgb(ad.MESSING), width=4)
        d.text((x0 + 18, y0 + 12), f"FACH {stand.upper()}", font=f_etikett, fill=rgb(ad.MESSING_TEXT))
        for j, k in enumerate(ks):
            karte(x0 + 18 + j * 134, y0 + 52, k, b=124, mit_stand=False)
        d.text((x0 + 18, y0 + 146), f"COUNT(*) = {len(ks)}", font=f_erg, fill=rgb(ad.TINTE))
        d.text((x0 + 18, y0 + 190), f"SUM(Menge) = {sum(k['menge'] for k in ks)}", font=f_erg, fill=rgb(ad.TINTE))
        d.text((x0 + 380, y0 + 150), "Zähler", font=f_klein, fill=rgb(ad.TINTE_3))
        d.text((x0 + 380, y0 + 194), "Summierer", font=f_klein, fill=rgb(ad.TINTE_3))
    # 3. Ergebniszeilen
    d.text((30, 500), "3 · JE FACH GENAU EINE ERGEBNISZEILE", font=f_etikett, fill=rgb(ad.SIEGELROT))
    spalten = [("Stand", 30), ("COUNT(*)", 470), ("SUM(Menge)", 770)]
    y = 540
    d.rectangle([30, y, 1100, y + 54], fill=rgb(ad.FLAECHE))
    for t, x in spalten:
        d.text((x + 14, y + 10), t, font=f_erg, fill=rgb(ad.TINTE))
    for stand, ks in faecher.items():
        y += 58
        d.rectangle([30, y, 1100, y + 54], outline=rgb(ad.RAND), width=2)
        for (t, x), w in zip(spalten, [stand, str(len(ks)), str(sum(k["menge"] for k in ks))]):
            d.text((x + 14, y + 8), w, font=f_kopf if x == 30 else f_erg, fill=rgb(ad.TINTE))
    d.text((1160, 560), "Meiste Bons:", font=f_karte, fill=rgb(ad.TINTE_2))
    d.text((1160, 596), "Kaffeemobil (COUNT)", font=f_kopf, fill=rgb(ad.TINTE))
    d.text((1160, 670), "Meiste Stück:", font=f_karte, fill=rgb(ad.TINTE_2))
    d.text((1160, 706), "Crêpes-Stand (SUM)", font=f_kopf, fill=rgb(ad.TINTE))
    bild.save(pfad)
    return pfad


def zeitlupe_png(pfad):
    """Leiste der Auswertungsreihenfolge (Zeitlupe) für AB 3.2 und Präsentation."""
    from PIL import Image, ImageDraw
    schritte = [("FROM", "Tabellen holen\nund verbinden"), ("WHERE", "Zeilen\nfiltern"), ("GROUP BY", "Fächer\nbilden"),
                ("HAVING", "Fächer\nfiltern"), ("SELECT", "Spalten\nberechnen"), ("ORDER BY", "sortieren"),
                ("LIMIT", "abschneiden")]
    W, H = 1800, 330
    bild = Image.new("RGB", (W, H), "#" + ad.KARTE)
    d = ImageDraw.Draw(bild)
    f_nr, f_kl, f_txt = ad._font("kopf", 46), ad._font("mono", 34), ad._font("text", 30)
    b, gap = 228, 22
    for i, (kl, txt) in enumerate(schritte):
        x = 20 + i * (b + gap)
        farbe = ad.SIEGELROT if kl in ("WHERE", "HAVING") else ad.TINTE
        d.rounded_rectangle([x, 20, x + b, 310], radius=14, fill="#" + ad.PAPIER, outline="#" + farbe, width=4)
        d.text((x + 16, 40), str(i + 1), font=f_nr, fill="#" + ad.SIEGELROT)
        d.text((x + 14, 100), kl, font=f_kl, fill="#" + farbe)
        d.multiline_text((x + 14, 170), txt, font=f_txt, fill="#" + ad.TINTE_2, spacing=6)
        if i < len(schritte) - 1:
            xa = x + b + 2
            d.polygon([(xa, 140), (xa + gap - 4, 150), (xa, 160)], fill="#" + ad.SIEGELROT)
    bild.save(pfad)
    return pfad


# ================================================================== AB 3.1 Kassensturz (DS 5)
def ab31(fassung):
    d = AktenDokument("ab", 3, "Kassensturz", fassung=fassung, ermittlung="3.1", ds=5, material="AB 3.1", datenbank=DB)
    d.kopfband()
    d.fall_einstieg("Nach dem Stadtfest am 29. August fehlen in der Festkasse 312,50 €. Die Stände haben jeden Verkauf als "
                    "Bon erfasst – Frauke Melchers vom Festausschuss hat uns alle Kassendaten gebracht. Edda Lorenz: "
                    "„Zählt nicht Zettel – zählt, was verkauft wurde.“")
    d.beweisstueck(
        1, "Aggregatfunktionen verdichten Spalten",
        "Eine **Aggregatfunktion** fasst die Werte einer Spalte über viele Zeilen zu **einem** Wert zusammen: "
        "`COUNT` (Anzahl), `SUM` (Summe), `AVG` (Durchschnitt), `MIN` (kleinster Wert), `MAX` (größter Wert). "
        "Ohne GROUP BY entsteht genau **eine** Ergebniszeile. Datenbank Stadtfest (Ausschnitt des Schemas):",
        Schema(SCHEMA[3]), Schema(SCHEMA[4]), Schema(SCHEMA[5]),
        S("SELECT COUNT(*) AS Produkte, MIN(Preis) AS Billigste,\n       MAX(Preis) AS Teuerste, AVG(Preis) AS Schnitt\nFROM Produkt"),
        ergebnis("SELECT COUNT(*) AS Produkte, MIN(Preis) AS Billigste,\n       MAX(Preis) AS Teuerste, AVG(Preis) AS Schnitt\nFROM Produkt"),
    )
    d.beweisstueck(
        2, "COUNT(*), COUNT(Spalte) oder SUM(Menge)?",
        "Die Bons 400–403 haben fünf Positionen mit den Mengen 3, 1, 3, 2 und 0 (Bon 403: Grillkäse, Menge 0 – siehe Auftrag 2):",
        Tabelle(["Ausdruck", "Wert", "Bedeutung"],
                [["`COUNT(*)`", "5", "Anzahl der **Zeilen** (Positionen) – auch die mit Menge 0"],
                 ["`SUM(Menge)`", "9", "Anzahl verkaufter **Stück**"],
                 ["`COUNT(Menge)`", "5", "Anzahl der Zeilen, in denen Menge **nicht NULL** ist"]],
                breiten=[3.6, 1.4, 11.8]),
        "Steht in einer Spalte NULL (z. B. `Helfer.Verein` bei Personen ohne Verein), zählt `COUNT(Verein)` diese Zeile "
        "nicht mit, `COUNT(*)` schon. Faustregel: „Anzahl Käufe/Bons“ → `COUNT(*)`, „Anzahl verkaufter Stück“ → `SUM(Menge)`.",
    )
    d.beweisstueck(
        3, "GROUP BY bildet Fächer",
        "`GROUP BY Stand` sortiert alle Zeilen in **Fächer** – ein Fach je Stand. Die Aggregatfunktion wird dann **je Fach** "
        "berechnet, und jedes Fach liefert **genau eine** Ergebniszeile (Bild aus dem Menschlichen GROUP BY, eigene Daten):",
        Bild(str(faecher_png(TMP / "faecher.png")), breite_cm=16.4),
    )
    d.beweisstueck(
        4, "Die GROUP-BY-Regel – und nach dem Schlüssel gruppieren",
        "**GROUP-BY-Regel:** Jede Spalte im SELECT ist entweder **aggregiert** oder steht in **GROUP BY**. "
        "Gruppiert wird nach dem **Schlüssel**; Namen kommen zusätzlich in GROUP BY, damit sie ausgegeben werden dürfen. "
        "Beispiel: Wie viele Schichten hatte jede Person?",
        S("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname,\n       COUNT(*) AS Schichten\nFROM Helfer, Schicht\n"
          "WHERE Helfer.HelferNr = Schicht.HelferNr\nGROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname\n"
          "ORDER BY Schichten DESC, Helfer.HelferNr"),
        ergebnis("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname,\n       COUNT(*) AS Schichten\nFROM Helfer, Schicht\n"
                 "WHERE Helfer.HelferNr = Schicht.HelferNr\nGROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname\n"
                 "ORDER BY Schichten DESC, Helfer.HelferNr", max_zeilen=3),
        Notiz("Läuft ≠ richtig", "SQLite führt auch Abfragen aus, die die GROUP-BY-Regel verletzen – und zeigt dann einen "
              "zufälligen Wert aus dem Fach. Im Abitur und in der Datendetektei gilt die Regel streng. Und: Am Stadtfest "
              "halfen **zwei Personen mit genau demselben Namen** – wer nur nach dem Namen gruppiert, legt sie in ein Fach."),
    )

    sql_a2 = "SELECT BonNr, COUNT(*) AS Positionen, SUM(Menge) AS Stueck\nFROM Position\nWHERE BonNr BETWEEN 400 AND 404\nGROUP BY BonNr"
    kopf_pos, zeilen_pos = daten("SELECT BonNr, ProduktNr, Menge FROM Position WHERE BonNr BETWEEN 400 AND 404 ORDER BY BonNr, ProduktNr")
    q4a = ("SELECT Stand.Name, Stand.Platz, COUNT(*) AS Bons\nFROM Stand, Bon\nWHERE Stand.StandNr = Bon.StandNr\n"
           "GROUP BY Stand.StandNr, Stand.Name")
    q4b = "SELECT Helfer.Verein, COUNT(*) AS Personen\nFROM Helfer\nGROUP BY Helfer.Verein"
    q4c = "SELECT Bon.Zahlart, Bon.Zeitpunkt, COUNT(*) AS Bons\nFROM Bon\nGROUP BY Bon.Zahlart"
    q4a_k = ("SELECT Stand.Name, Stand.Platz, COUNT(*) AS Bons\nFROM Stand, Bon\nWHERE Stand.StandNr = Bon.StandNr\n"
             "GROUP BY Stand.StandNr, Stand.Name, Stand.Platz")
    q4c_k = "SELECT Bon.Zahlart, MIN(Bon.Zeitpunkt) AS ErsterBon, COUNT(*) AS Bons\nFROM Bon\nGROUP BY Bon.Zahlart"
    q3 = ("SELECT Stand.Art, AVG(Produkt.Preis) AS Durchschnittspreis\nFROM Stand, Produkt\n"
          "WHERE Stand.StandNr = Produkt.StandNr\nGROUP BY Stand.Art")
    q5 = ("SELECT Helfer.Nachname,\n       SUM(Position.Menge * Produkt.Preis) AS Umsatz\nFROM Helfer, Bon, Position, Produkt\n"
          "WHERE Helfer.HelferNr = Bon.HelferNr\n  AND Bon.BonNr = Position.BonNr\n  AND Position.ProduktNr = Produkt.ProduktNr\n"
          "GROUP BY Helfer.Nachname")
    q5_beleg = ("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname,\n       SUM(Position.Menge * Produkt.Preis) AS Umsatz\n"
                "FROM Helfer, Bon, Position, Produkt\nWHERE Helfer.HelferNr = Bon.HelferNr\n  AND Bon.BonNr = Position.BonNr\n"
                "  AND Position.ProduktNr = Produkt.ProduktNr\n  AND Helfer.Nachname = 'Bödeker'\n"
                "GROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname")
    d.ermittlungsauftraege([
        Auftrag(1, 1, "PAPIER", "Frauke Melchers hat vier Fragen. Gib jeweils die passende Aggregatfunktion **mit Spalte** an "
                "und begründe bei b), warum COUNT hier nicht reicht.", kompetenzen=["K25"],
                inhalt=[Tabelle(["Frage des Festausschusses", "Aggregatfunktion (Spalte)"],
                                [["a) Wie viele Helferinnen und Helfer sind beim TuS gemeldet?", ""],
                                 ["b) Wie viele Waffeln wurden verkauft?", ""],
                                 ["c) Wann begann die früheste Schicht?", ""],
                                 ["d) Was kostet ein Produkt des Crêpes-Stands im Durchschnitt?", ""]],
                                breiten=[8.8, 4.6])],
                linien=1,
                loesung=Loesung("a) `COUNT(*)` über Helfer mit `WHERE Verein = 'TuS Weserstedt'` · b) `SUM(Position.Menge)` · "
                                "c) `MIN(Schicht.Beginn)` · d) `AVG(Produkt.Preis)`.",
                                "Zu b): Eine Position kann mehrere Waffeln enthalten (Menge 2 oder 3); `COUNT(*)` zählt nur "
                                "die Positionen, nicht die Stück.",
                                fehler=[("P4", "COUNT statt SUM bei Stückzahlen – zählt Positionen statt Waffeln.")])),
        Auftrag(2, 1, "PAPIER", "Die Bons 400 bis 404 enthalten diese acht Positionen. Gib das Ergebnis der Abfrage an "
                "(Reihenfolge egal). Bilde dazu in Gedanken Fächer.", kompetenzen=["K26", "K29"],
                inhalt=[Tabelle(kopf_pos, zeilen_pos, pk=("BonNr", "ProduktNr"), fk=("BonNr", "ProduktNr")),
                        S(sql_a2),
                        Tabelle(["BonNr", "Positionen", "Stueck"], [["", "", ""]] * 5) if fassung == "schueler" else None],
                loesung=Loesung(ergebnis(sql_a2),
                                "Fünf Fächer (je Bon eines) → fünf Ergebniszeilen. Bon 403 hat eine Position, aber 0 Stück.",
                                fehler=[("P4", "Bei Bon 403 „0 Positionen“ eintragen – COUNT(*) zählt die Zeile mit Menge 0 mit."),
                                        ("P3", "Eine Zeile je Position statt je Bon notiert – GROUP BY fasst zusammen.")])),
        Auftrag(3, 2, "LAB", "Wie teuer ist ein Produkt im Durchschnitt – getrennt nach der **Art** des Standes (Essen, "
                "Getränke, …)? Entwirf die Abfrage, prüfe sie in der Konsole und notiere das Ergebnis.",
                kompetenzen=["K26", "K25"], linien=4,
                loesung=Loesung(S(q3, 48), ergebnis(q3),
                                "Das Info-Zelt (Art Info) bietet keine Produkte an und fällt im Verbund heraus.",
                                fehler=[("P3", "`SELECT Stand.Name, AVG(…) … GROUP BY Stand.Art` – Name ist weder aggregiert "
                                               "noch gruppiert."),
                                        ("P2", "Verbundbedingung vergessen – jedes Produkt mit jedem Stand kombiniert.")])),
        Auftrag(4, 2, "PAPIER", "Alle drei Abfragen laufen in SQLite. Prüfe jede auf die GROUP-BY-Regel. Markiere die "
                "Spalte, die gegen die Regel verstößt, und korrigiere die fehlerhaften Abfragen.", kompetenzen=["K26"],
                inhalt=["**a)**", S(q4a, pruefung="fehlerhaft"), "**b)**", S(q4b), "**c)**", S(q4c, pruefung="fehlerhaft")],
                linien=3,
                loesung=Loesung("a) **Verstoß:** `Stand.Platz` ist weder aggregiert noch gruppiert. Korrektur:",
                                S(q4a_k, 48), ergebnis(q4a_k, max_zeilen=3),
                                "b) **korrekt** – Verein ist gruppiert. Personen ohne Verein bilden ein eigenes Fach (NULL):",
                                S(q4b), ergebnis(q4b),
                                "c) **Verstoß:** `Bon.Zeitpunkt` – im Fach „bar“ liegen Hunderte Zeitpunkte. Korrektur, z. B. "
                                "mit dem ersten Zeitpunkt je Zahlart (oder Zeitpunkt weglassen):",
                                S(q4c_k, 48), ergebnis(q4c_k),
                                fehler=[("P3", "Auch b) für falsch halten, weil „Verein NULL sein kann“ – NULL bildet "
                                               "einfach ein eigenes Fach.")])),
        Auftrag(5, 3, "PAPIER", "Juri will den Umsatz je Person berechnen und schreibt die folgende Abfrage. Sie hält die "
                "GROUP-BY-Regel ein. **Erkläre**, warum sie trotzdem zu falschen Ergebnissen führen kann, und gib an, "
                "wie GROUP BY richtig lauten muss.", kompetenzen=["K26"],
                inhalt=[S(q5)], linien=4,
                loesung=Loesung("Nachnamen sind nicht eindeutig: Alle Personen mit dem Nachnamen Bödeker landen in **einem** "
                                "Fach, ihre Umsätze werden addiert. Juris Abfrage liefert für Bödeker 402 € – tatsächlich sind "
                                "es drei Personen (zwei davon heißen sogar beide Lotta Bödeker):",
                                S(q5_beleg, 48), ergebnis(q5_beleg),
                                "Richtig: nach dem Schlüssel gruppieren und Namen zusätzlich aufnehmen – "
                                "`GROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname` (dann auch HelferNr ausgeben).",
                                fehler=[("P3", "„Dann eben nach Vor- und Nachname“ – die beiden Lotta Bödeker verschmelzen "
                                               "trotzdem (68,00 € + 67,50 €).")])),
    ])
    sql_abi = ("SELECT Stand.StandNr, Stand.Name, SUM(Position.Menge) AS Stueck\nFROM Stand, Bon, Position\n"
               "WHERE Stand.StandNr = Bon.StandNr\n  AND Bon.BonNr = Position.BonNr\nGROUP BY Stand.StandNr, Stand.Name")
    sql_abi_v = ("SELECT Stand.StandNr, Stand.Name, SUM(Position.Menge) AS Stueck\nFROM Stand, Produkt, Position\n"
                 "WHERE Stand.StandNr = Produkt.StandNr\n  AND Produkt.ProduktNr = Position.ProduktNr\n"
                 "GROUP BY Stand.StandNr, Stand.Name")
    d.abi_ecke("Gegeben ist die Datenbank Stadtfest (Relationenschema siehe Beweisstück 1 sowie "
               "Stand(<u>StandNr</u>, Name, Art, Platz)). **Entwerfen** Sie eine SQL-Abfrage, die für jeden Stand die "
               "Standnummer, den Namen und die Anzahl der dort verkauften Produkte (Stückzahl) ausgibt.",
               be=4, operator="entwerfen", afb="II", kompetenzen=["K24", "K25", "K26"], linien=5,
               loesung=Loesung(S(sql_abi), ergebnis(sql_abi),
                               "Je 1 BE: Tabellen und Verbundbedingungen · `SUM(Position.Menge)` statt `COUNT(*)` · "
                               "GROUP BY nach StandNr (und Name) · vollständige Ausgabe. Das Info-Zelt erscheint nicht "
                               "(kein Bon).",
                               varianten=[S(sql_abi_v)],
                               fehler=[("P4", "`COUNT(*)` liefert die Anzahl der Positionen, nicht die Stückzahl."),
                                       ("P3", "`GROUP BY Stand.Name` ohne StandNr bzw. StandNr im SELECT ohne GROUP BY."),
                                       ("P2", "Eine Verbundbedingung fehlt – Zeilen werden vervielfacht.")]))
    d.spur_im_lab("3.1 „Kassensturz“", "Einstieg A3-E1-04 (Sortierfächer), dann A3-E1-01 … A3-E1-08 · Cold Case A3-E1-CC",
                  werkzeug="Sortierfächer (aus jeder Konsole aufrufbar)")
    d.cold_case("Kartenzahler", "Welcher Anteil der Bons wurde mit Karte bezahlt? Überlege, wie du das mit **zwei** "
                "Abfragen ermitteln kannst, notiere beide und beschreibe, wie du aus den beiden Ergebnissen den Anteil "
                "berechnest.", kompetenzen=["K25", "K30"], linien=4,
                loesung=Loesung(S("SELECT COUNT(*) AS Kartenbons\nFROM Bon\nWHERE Zahlart = 'Karte'"),
                                ergebnis("SELECT COUNT(*) AS Kartenbons\nFROM Bon\nWHERE Zahlart = 'Karte'"),
                                S("SELECT COUNT(*) AS Bons\nFROM Bon"), ergebnis("SELECT COUNT(*) AS Bons\nFROM Bon"),
                                "Anteil = 200 : 554 ≈ 36 %. Mit einer Abfrage geht es gruppiert:",
                                S("SELECT Zahlart, COUNT(*) AS Bons\nFROM Bon\nGROUP BY Zahlart"),
                                ergebnis("SELECT Zahlart, COUNT(*) AS Bons\nFROM Bon\nGROUP BY Zahlart")))
    name = "DD_A3_AB1_Kassensturz"
    ziel = (ORDNER / "5_Loesungen" / f"{name}_LSG.docx") if fassung == "loesung" else (ORDNER / "2_Arbeitsblaetter" / f"{name}.docx")
    return d.speichern(ziel)


# ================================================================== AB 3.2 Die Statistik lügt nicht (DS 6)
def ab32(fassung):
    d = AktenDokument("ab", 3, "Die Statistik lügt nicht", fassung=fassung, ermittlung="3.2", ds=6, material="AB 3.2",
                      datenbank=DB)
    d.kopfband()
    d.fall_einstieg("Die Kasse rechnet richtig – und trotzdem fehlen 312,50 €. Am Kaffeemobil stehen Bons mit Positionen "
                    "der Menge 0. Edda Lorenz: „Daten lügen nicht. Abfragen manchmal schon.“")
    d.beweisstueck(
        1, "Die Zeitlupe: In dieser Reihenfolge wertet die Datenbank aus",
        "Geschrieben wird `SELECT … FROM … WHERE … GROUP BY … HAVING … ORDER BY … LIMIT`. Ausgewertet wird anders:",
        Bild(str(zeitlupe_png(TMP / "zeitlupe.png")), breite_cm=16.4),
        "Deshalb kann WHERE noch nichts über Gruppen wissen, und ein Alias aus dem SELECT ist in HAVING noch unbekannt "
        "(in ORDER BY dagegen schon). Nach EH steht in HAVING immer die Aggregatfunktion selbst.",
    )
    sql_b2 = ("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS Barbons\nFROM Stand, Bon\nWHERE Stand.StandNr = Bon.StandNr\n"
              "  AND Bon.Zahlart = 'bar'\nGROUP BY Stand.StandNr, Stand.Name\nHAVING COUNT(*) > 40\nORDER BY Barbons DESC")
    d.beweisstueck(
        2, "WHERE filtert Zeilen vorher – HAVING filtert Fächer nachher",
        "Welche Stände hatten mehr als 40 Barzahlungen?",
        S(sql_b2), ergebnis(sql_b2),
        Tabelle(["Schritt", "Was passiert mit den Bon-Karten?"],
                [["WHERE `Bon.Zahlart = 'bar'`", "Kartenzahler setzen sich hin, **bevor** Fächer gebildet werden (Zeilenbedingung)."],
                 ["GROUP BY", "Die übrigen Bons gehen in ihr Stand-Fach."],
                 ["HAVING `COUNT(*) > 40`", "Fächer mit höchstens 40 Bons werden **danach** aussortiert (Gruppenbedingung)."]],
                breiten=[5.2, 11.6]),
    )
    sql_b3 = ("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname,\n       COUNT(*) AS Schichten\nFROM Helfer, Schicht\n"
              "WHERE Helfer.HelferNr = Schicht.HelferNr\nGROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname\n"
              "ORDER BY Schichten DESC, Helfer.HelferNr")
    d.beweisstueck(
        3, "Top-n mit LIMIT – und das Gleichstand-Problem",
        "`ORDER BY … DESC LIMIT n` liefert die ersten n Zeilen nach der Sortierung. Wer hatte die meisten Schichten?",
        S(sql_b3), ergebnis(sql_b3, max_zeilen=5),
        "Mit `LIMIT 3` fiele eine der vier Personen mit 4 Schichten weg – welche, entscheidet der Zufall der Sortierung. "
        "Eine Top-n-Frage ist nur eindeutig, wenn an Platz n **kein Gleichstand** besteht. Sonst besser fragen: "
        "„alle Personen mit mindestens 4 Schichten“ (HAVING).",
    )
    sql_b4 = ("SELECT Bon.BonNr, Position.ProduktNr, Schicht.HelferNr, Schicht.Beginn\n"
              "FROM Bon, Position, Schicht\nWHERE Bon.BonNr = Position.BonNr\n  AND Bon.StandNr = Schicht.StandNr\n"
              "  AND Bon.BonNr = 486\n  AND Schicht.Beginn >= '2026-08-29 21:00'\nORDER BY Schicht.Beginn")
    d.beweisstueck(
        4, "Vervielfachung im Verbund",
        "Bon 486 (das Beweisstück) hat **eine** Position: Kaffee, Menge 0. Verbindet man zusätzlich Schicht nur über den "
        "Stand, passt jede Schicht des Kaffeemobils ab 21 Uhr zu diesem Bon:",
        S(sql_b4), ergebnis(sql_b4),
        "Aus einer Position werden drei Zeilen – `SUM(Produkt.Preis)` würde den Kaffee dreimal zählen (7,50 € statt 2,50 €). "
        "**Regel:** Nur die Tabellen verbinden, die für die Frage nötig sind, und jede Verbundbedingung prüfen. Ein "
        "Ergebnis, das „plausibel aussieht“, ist noch kein Beweis.",
    )
    d.beweisstueck(
        5, "Analysieren und Vorgehen beschreiben – Satzbausteine",
        "**Analyse im Sachzusammenhang** (nicht: „wählt aus, gruppiert, zählt …“):",
        Liste(["„Die Abfrage liefert **für jede/jeden …** (Gruppe mit Schlüssel und Name) **die/den …** (Aggregat mit "
               "Bedeutung) …“", "„… **berücksichtigt werden nur** … (WHERE, Zeilen) …“",
               "„… **ausgegeben werden nur** Gruppen, die/deren … (HAVING) …“", "„… **sortiert nach** … (ORDER BY), "
               "**nur die ersten** n (LIMIT).“"]),
        "**Vorgehen in Worten** (Abitur: „Beschreiben Sie Ihr Vorgehen …“): ① benötigte **Tabellen** → ② **Verbund** über "
        "Primär-/Fremdschlüssel → ③ **Zeilenfilter** (WHERE) → ④ **Gruppierung** nach dem Schlüssel → ⑤ **Gruppenfilter** "
        "(HAVING) → ⑥ **Sortierung/Begrenzung** und Ausgabe.",
    )

    sql_a2 = ("SELECT Helfer.Verein, COUNT(*) AS Bons\nFROM Helfer, Bon\nWHERE Helfer.HelferNr = Bon.HelferNr\n"
              "  AND Bon.Zeitpunkt LIKE '2026-08-28%'\nGROUP BY Helfer.Verein\nHAVING COUNT(*) >= 40\nORDER BY Bons DESC")
    sql_a3 = ("SELECT Stand.StandNr, Stand.Name, AVG(Produkt.Preis) AS Durchschnittspreis\nFROM Stand, Produkt\n"
              "WHERE Stand.StandNr = Produkt.StandNr\nGROUP BY Stand.StandNr, Stand.Name\nHAVING AVG(Produkt.Preis) < 3")
    sql_a4 = ("SELECT Stand.StandNr, Stand.Name, SUM(Position.Menge * Produkt.Preis) AS Umsatz\n"
              "FROM Stand, Bon, Position, Produkt\nWHERE Stand.StandNr = Bon.StandNr\n  AND Bon.BonNr = Position.BonNr\n"
              "  AND Position.ProduktNr = Produkt.ProduktNr\n  AND Bon.Zahlart = 'Karte'\n"
              "  AND Bon.Zeitpunkt LIKE '2026-08-29%'\nGROUP BY Stand.StandNr, Stand.Name\nORDER BY Umsatz DESC\nLIMIT 1")
    sql_a5 = ("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS Bons\nFROM Stand, Bon, Position\n"
              "WHERE Stand.StandNr = Bon.StandNr\n  AND Bon.BonNr = Position.BonNr\nGROUP BY Stand.StandNr, Stand.Name\n"
              "ORDER BY Bons DESC")
    sql_a5_k = ("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS Bons\nFROM Stand, Bon\nWHERE Stand.StandNr = Bon.StandNr\n"
                "GROUP BY Stand.StandNr, Stand.Name\nORDER BY Bons DESC")
    leer = "☐"
    d.ermittlungsauftraege([
        Auftrag(1, 1, "PAPIER", "Alle sechs Bedingungen stammen aus Abfragen, die nach Stand gruppieren. Kreuze an, ob "
                "sie in WHERE oder in HAVING gehören.", kompetenzen=["K27"],
                inhalt=[Tabelle(["Bedingung", "WHERE", "HAVING"],
                                [["a) `Helfer.Jahrgang < 1970`", leer, leer], ["b) `COUNT(*) > 20`", leer, leer],
                                 ["c) `Produkt.Preis >= 4`", leer, leer],
                                 ["d) `MIN(Bon.Zeitpunkt) < '2026-08-28 17:00'`", leer, leer],
                                 ["e) `Helfer.Verein = 'Shantychor Weserstedt'`", leer, leer],
                                 ["f) `SUM(Position.Menge * Produkt.Preis) > 500`", leer, leer]],
                                breiten=[9.2, 1.5, 1.5])],
                loesung=Loesung("WHERE: a), c), e) – sie prüfen eine einzelne Zeile. HAVING: b), d), f) – sie enthalten "
                                "eine Aggregatfunktion und beziehen sich auf ein ganzes Fach.",
                                fehler=[("P1", "d) in WHERE, weil „es um Zeitpunkte geht“ – entscheidend ist MIN(…).")])),
        Auftrag(2, 2, "PAPIER", "**Analysiere** die Abfrage: Beschreibe im Sachzusammenhang des Stadtfests, welche "
                "Information sie liefert. Nutze die Satzbausteine aus Beweisstück 5.", kompetenzen=["K28"],
                inhalt=[S(sql_a2)], linien=4,
                loesung=Loesung("Die Abfrage liefert für jeden **Verein** die Anzahl der **Bons, die Mitglieder dieses "
                                "Vereins am Freitag (28.08.) kassiert haben**. Ausgegeben werden nur Vereine mit "
                                "**mindestens 40** solchen Bons, absteigend nach Anzahl. Personen ohne Verein bilden eine "
                                "eigene Gruppe (NULL).", ergebnis(sql_a2),
                                fehler=[("P10", "Syntax nacherzählen („verbindet Helfer und Bon, gruppiert nach Verein …“) "
                                                "statt die Information zu beschreiben."),
                                        ("P5", "Die NULL-Gruppe übersehen.")])),
        Auftrag(3, 2, "LAB", "Welche Stände haben einen durchschnittlichen Produktpreis **unter 3 €**? Gib StandNr, Name "
                "und den Durchschnittspreis aus. Prüfe in der Konsole und notiere das Ergebnis.", kompetenzen=["K27", "K26"],
                linien=3,
                loesung=Loesung(S(sql_a3, 48), ergebnis(sql_a3),
                                fehler=[("P1", "`WHERE AVG(Produkt.Preis) < 3` – Aggregat in WHERE."),
                                        ("P1", "Alias in HAVING (`HAVING Durchschnittspreis < 3`) – läuft in SQLite, ist "
                                               "nach EH aber nicht vorgesehen.")])),
        Auftrag(4, 2, "PAPIER", "Edda fragt: „Welcher Stand hatte am Festsamstag (29.08.) den höchsten Umsatz mit "
                "**Kartenzahlungen**?“ **Beschreibe** dein Vorgehen in Worten – Tabellen, Verbund, Filter, Gruppierung, "
                "Sortierung. Eine fertige Abfrage ist nicht verlangt.", kompetenzen=["K30"], linien=4,
                loesung=Loesung(Liste(["Tabellen **Stand, Bon, Position, Produkt** (Umsatz = Menge × Preis).",
                                       "Verbund: Stand.StandNr = Bon.StandNr, Bon.BonNr = Position.BonNr, "
                                       "Position.ProduktNr = Produkt.ProduktNr.",
                                       "Zeilenfilter: Zahlart 'Karte' und Zeitpunkt am 2026-08-29.",
                                       "Gruppieren nach StandNr (und Name), je Stand `SUM(Position.Menge * Produkt.Preis)`.",
                                       "Absteigend nach dieser Summe sortieren, nur die erste Zeile ausgeben – vorher prüfen, "
                                       "ob an Platz 1 ein Gleichstand besteht."], nummeriert=True),
                                "Zur Kontrolle:", S(sql_a4, 48), ergebnis(sql_a4),
                                fehler=[("P10", "Nur „mit SQL die Daten abfragen“ – ohne Tabellen und Bedingungen."),
                                        ("P4", "Bons zählen statt Umsatz summieren.")])),
        Auftrag(5, 3, "PAPIER", "Juri soll die Anzahl der Bons je Stand ermitteln. Seine Abfrage liefert für das "
                "Kaffeemobil 316 Bons – Frauke Melchers hat aber nur 177 Bons in der Kasse. **Beurteile** die Abfrage "
                "und korrigiere sie. *(Schwerpunkt eA)*", kompetenzen=["K31", "K24"],
                inhalt=[S(sql_a5), ergebnis(sql_a5, max_zeilen=3)], linien=4,
                loesung=Loesung("Juri verbindet zusätzlich **Position**. Ein Bon mit drei Positionen erscheint dann in "
                                "drei Zeilen; `COUNT(*)` zählt **Positionen** statt Bons (Vervielfachung). Für die Frage "
                                "wird Position nicht gebraucht:", S(sql_a5_k, 48), ergebnis(sql_a5_k),
                                fehler=[("P2", "Unnötige Tabelle im Verbund vervielfacht Zeilen – das Ergebnis sieht "
                                               "plausibel aus, ist aber zu hoch.")])),
    ])
    sql_abi = ("SELECT Produkt.Bezeichnung, SUM(Position.Menge) AS Stueck\nFROM Bon, Produkt, Position\n"
               "WHERE Bon.BonNr = Position.BonNr\n  AND Position.ProduktNr = Produkt.ProduktNr\n"
               "  AND Bon.Zeitpunkt LIKE '2026-08-29%'\nGROUP BY Produkt.ProduktNr, Produkt.Bezeichnung\n"
               "HAVING SUM(Position.Menge) > 50\nORDER BY Stueck DESC")
    d.abi_ecke("Gegeben ist die folgende SQL-Abfrage zur Datenbank Stadtfest. **Analysieren** Sie die Abfrage, indem Sie "
               "beschreiben, welche Informationen sie im Sachzusammenhang liefert.", be=4, operator="analysieren", afb="II",
               kompetenzen=["K28"], inhalt=[S(sql_abi)], linien=4,
               loesung=Loesung("Die Abfrage liefert für jedes **Produkt**, das am **Samstag, 29.08.**, verkauft wurde, die "
                               "Bezeichnung und die **verkaufte Stückzahl** dieses Tages. Ausgegeben werden nur Produkte "
                               "mit **mehr als 50** verkauften Stück, das meistverkaufte zuerst.", ergebnis(sql_abi),
                               "Je 1 BE: Bezug je Produkt und Stückzahl (SUM der Menge, nicht Anzahl der Bons) · "
                               "Einschränkung auf den Samstag · Gruppenbedingung mehr als 50 Stück · Sortierung absteigend.",
                               fehler=[("P10", "„Die Abfrage verbindet drei Tabellen und gruppiert …“ – Syntax statt "
                                               "Sachzusammenhang."),
                                       ("P4", "„Anzahl der Verkäufe“ statt verkaufter Stück.")]))
    d.spur_im_lab("3.2 „Die Statistik lügt nicht“", "A3-E2-01 … A3-E2-08 · Cold Case A3-E2-CC · danach Verhör 3 (Schulkiosk)",
                  werkzeug="Zeitlupe (aus jeder Diagnose)",
                  text="A3-E2-06 und A3-E2-08 sind Analyse-Aufgaben: Deine Antwort bleibt auf deinem Gerät, du hakst die "
                       "Kernpunkte selbst ab.")
    sql_cc1 = ("SELECT Stand.StandNr, Stand.Name, SUM(Position.Menge * Produkt.Preis) AS Umsatz\n"
               "FROM Stand, Bon, Position, Produkt\nWHERE Stand.StandNr = Bon.StandNr\n  AND Bon.BonNr = Position.BonNr\n"
               "  AND Position.ProduktNr = Produkt.ProduktNr\nGROUP BY Stand.StandNr, Stand.Name")
    sql_cc2 = ("SELECT SUM(Position.Menge * Produkt.Preis) AS Gesamtumsatz\nFROM Position, Produkt\n"
               "WHERE Position.ProduktNr = Produkt.ProduktNr")
    sql_cc3 = ("SELECT Stand.Name, COUNT(*) AS Nullpositionen, SUM(Produkt.Preis) AS Warenwert\nFROM Stand, Produkt, Position\n"
               "WHERE Stand.StandNr = Produkt.StandNr\n  AND Produkt.ProduktNr = Position.ProduktNr\n  AND Position.Menge = 0\n"
               "GROUP BY Stand.StandNr, Stand.Name")
    d.cold_case("Die Kasse stimmt doch", "Zeige mit zwei Abfragen, dass die Umsätze je Stand und der Gesamtumsatz "
                "zusammenpassen. Erkläre dann, warum der Fehlbetrag in keiner der beiden Zahlen auftaucht – und mit welcher "
                "Abfrage man ihn sichtbar macht.", kompetenzen=["K25", "K31"], linien=3,
                loesung=Loesung(S(sql_cc1, 48), ergebnis(sql_cc1), S(sql_cc2), ergebnis(sql_cc2),
                                "Die Summe der Standumsätze ergibt genau 5023 €. Positionen mit Menge 0 tragen "
                                "0 × Preis = 0 € zum Umsatz bei – die Kasse „rechnet richtig“. Sichtbar wird der Fehlbetrag "
                                "erst über den **Warenwert** der Positionen mit Menge 0 (Listenpreis statt Menge × Preis):",
                                S(sql_cc3, 48), ergebnis(sql_cc3)))
    name = "DD_A3_AB2_Die_Statistik_luegt_nicht"
    ziel = (ORDNER / "5_Loesungen" / f"{name}_LSG.docx") if fassung == "loesung" else (ORDNER / "2_Arbeitsblaetter" / f"{name}.docx")
    return d.speichern(ziel)


# ================================================================== Präsentation
def praesentation():
    p = AktenPraesentation(3, "Die Stadtfest-Kasse", datenbank=DB)
    # ---------------------------------------------------------- DS 5
    p.titelfolie("Die Stadtfest-Kasse", "Akte 3 · DS 5 „Kassensturz“ · DS 6 „Die Statistik lügt nicht“",
                 notizen="DS 5 beginnt mit dem Warm-up (Karten W-A2-05 … 08, 10 min). Danach diese Folie: Akte 3 wird "
                         "geöffnet – im Lab die Akte für den Kurs freischalten.")
    p.briefing("Kassensturz",
               "Frauke Melchers, Kassenwartin des Festausschusses, hat uns alle Kassendaten des Stadtfests gebracht – und eine "
               "Zahl: 312,50 € Fehlbetrag. Die Datenbank ist groß: Stände, Helferinnen und Helfer, Schichten, Produkte, Bons "
               "und Positionen. Einzelne Zeilen anzuschauen bringt hier nichts; wir müssen verdichten – zählen, summieren, "
               "Durchschnitte bilden. Und es gibt eine eiserne Regel: Jede Spalte im SELECT ist entweder aggregiert oder "
               "steht im GROUP BY. Gleich sortiert ihr euch selbst in Fächer: Ihr seid die Datensätze.",
               notizen="Briefing vorlesen lassen (STORY 5.5b). Frage an den Kurs: Was heißt „verdichten“? Woher kennt ihr "
                       "Kassenbons mit mehreren Zeilen?")
    p.inhaltsfolie("Die Datenbank Stadtfest", *schema_inhalt(),
                   "Ein **Bon** hat mehrere **Positionen**; jede Position nennt ein Produkt und eine **Menge**.",
                   notizen="Schema gemeinsam lesen: Primärschlüssel unterstrichen, Fremdschlüssel mit ↑. Position hat einen "
                           "zusammengesetzten Primärschlüssel (BonNr, ProduktNr). Schicht: Beginn gehört zum Schlüssel, weil "
                           "dieselbe Person mehrmals am selben Stand arbeiten kann.")
    p.inhaltsfolie("Impuls: Zählen oder summieren?",
                   Liste(["Wie viele Bons gab es am Stadtfest?", "Wie viele Bratwürste wurden verkauft?",
                          "Ist das dieselbe Art von Zahl?", "Wir wollen die Anzahl der Bons **je Stand**. Wie viele "
                          "Ergebniszeilen erwartet ihr?"]),
                   notizen="Vermutungen sammeln, nicht auflösen. Ziel: Unterschied „Zeilen zählen“ (COUNT) vs. „Mengen "
                           "addieren“ (SUM) anbahnen. Erwartete Ergebniszeilen je Stand: so viele wie Stände mit Bons (7, "
                           "das Info-Zelt hat keine Kasse).")
    p.begriffe("Aggregatfunktionen", [("`COUNT(*)`", "Anzahl der Zeilen"), ("`COUNT(Spalte)`", "Anzahl der Werte ungleich NULL"),
                                      ("`SUM(Spalte)`", "Summe, z. B. verkaufte Stück `SUM(Menge)`"),
                                      ("`AVG(Spalte)`", "Durchschnitt"), ("`MIN` / `MAX`", "kleinster / größter Wert")],
               inhalt=["Ohne GROUP BY verdichtet eine Aggregatfunktion die ganze Tabelle zu **einer** Ergebniszeile."],
               notizen="Fachbegriff „Aggregatfunktion“ an die Tafel. COUNT(Spalte) mit Helfer.Verein (NULL) kurz zeigen.")
    sql1 = ("SELECT COUNT(*) AS Produkte, MIN(Preis) AS Billigste,\n       MAX(Preis) AS Teuerste, AVG(Preis) AS Schnitt\n"
            "FROM Produkt;")
    p.beispiel_schritte("Kennzahlen ohne GROUP BY", SQL(sql1),
                        ["**FROM** – alle 26 Zeilen der Tabelle Produkt", "**SELECT** – jede Aggregatfunktion verdichtet die "
                         "Spalte Preis zu einem Wert", "Ergebnis: **eine** Zeile"],
                        ergebnis=ergebnis(sql1),
                        notizen="Eigenes Beispiel (nicht die Lab-Aufgabe zu den Bons). Frage: Warum darf hier keine Spalte "
                                "Bezeichnung stehen? (Welche Bezeichnung sollte die eine Zeile zeigen?)")
    p.inhaltsfolie("GROUP BY bildet Fächer", Bild(str(faecher_png(TMP / "faecher_folie.png")), breite_cm=30),
                   notizen="Bild vor dem Unplugged nur kurz zeigen oder erst danach zur Sicherung. Kaffeemobil hat die "
                           "meisten Bons, der Crêpes-Stand die meisten Stück – COUNT und SUM ergeben andere Rangfolgen.")
    sql2 = ("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS Produkte,\n       MAX(Produkt.Preis) AS Hoechstpreis\n"
            "FROM Stand, Produkt\nWHERE Stand.StandNr = Produkt.StandNr\nGROUP BY Stand.StandNr, Stand.Name\n"
            "ORDER BY Produkte DESC, Stand.StandNr;")
    p.beispiel_schritte("Beispiel: Sortiment je Stand", SQL(sql2),
                        ["**FROM/WHERE** – Stand und Produkt verbinden", "**GROUP BY** – ein Fach je Stand (Schlüssel + Name)",
                         "**SELECT** – je Fach Anzahl und Höchstpreis", "**ORDER BY** – größtes Sortiment zuerst"],
                        ergebnis=ergebnis(sql2, max_zeilen=3),
                        notizen="GROUP-BY-Regel prüfen lassen: StandNr und Name stehen in GROUP BY, COUNT und MAX sind "
                                "Aggregate. Stand.Name allein würde reichen, solange Namen eindeutig sind – nach dem Schlüssel "
                                "gruppieren ist trotzdem die Regel.")
    p.unplugged("Menschliches GROUP BY",
                ["Jede/r bekommt eine Bon-Karte. `SELECT COUNT(*) FROM Bon`: alle stehen auf, eine Person zählt – eine Zahl.",
                 "`GROUP BY Stand`: alle gehen in die Ecke ihres Standes. Je Fach ein **Zähler** (COUNT) und ein **Summierer** "
                 "(SUM Betrag) – Ergebnis an die Tafel.",
                 "`SELECT Stand, Helfer, COUNT(*) … GROUP BY Stand`: Welchen Helfer meldet euer Fach? → GROUP-BY-Regel.",
                 "`SUM(Menge)` statt `COUNT(*)`: Welches Fach gewinnt jetzt?",
                 "`WHERE Zahlart = 'Karte'`: Barzahler setzen sich **vor** dem Gruppieren."],
                material="30 Bon-Karten, 4 Fächerschilder", dauer="15 MIN",
                notizen="Kopiervorlage 3_Unplugged/DD_A3_Unplugged_GROUP_BY_Karten (Bon-Karten 101–130, eigene Daten). "
                        "Fächerschilder in vier Raumecken. Erwartete Tafelwerte: siehe Lösung der Kopiervorlage. "
                        "Bei weniger als 30 Personen Karten mit hoher BonNr weglassen oder doppelt vergeben.")
    p.inhaltsfolie("Die GROUP-BY-Regel",
                   "**Jede Spalte im SELECT ist entweder aggregiert oder steht in GROUP BY.** Diese Abfrage läuft in SQLite – "
                   "trotzdem ist sie falsch:",
                   SQL("SELECT Stand.Name, Produkt.Bezeichnung, COUNT(*) AS Produkte\nFROM Stand, Produkt\n"
                       "WHERE Stand.StandNr = Produkt.StandNr\nGROUP BY Stand.StandNr, Stand.Name;", pruefung="fehlerhaft"),
                   "Welche Bezeichnung soll das Fach Kaffeemobil melden? Es hat sechs. Und: **nach dem Schlüssel** gruppieren – "
                   "am Stadtfest halfen zwei Personen mit genau demselben Namen.",
                   notizen="Läuft ≠ richtig. Die Plattform meldet den Regelverstoß (Parser-Hinweis). Bezug zum Unplugged-Schritt 3.")
    p.lab_auftrag("3.1", "Kassensturz", "Einstieg A3-E1-04 (Sortierfächer), dann A3-E1-01 … 08",
                  werkzeug="Sortierfächer",
                  hinweise=["AB 3.1 parallel bearbeiten (Aufträge 1–5, Abi-Ecke).",
                            "Schnelle: Cold Case „Der stille Stand“ im Lab, „Kartenzahler“ auf dem AB."],
                  notizen="Herumgehen: P3 (Name im SELECT, nicht in GROUP BY) und P4 (COUNT statt SUM) sammeln. Bei "
                          "Unsicherheit die Sortierfächer aus der Konsole öffnen lassen.")
    p.merkfolie("Ermittler-Notiz: Aggregation",
                ["Ohne GROUP BY: genau **eine** Ergebniszeile.", "GROUP BY bildet Fächer – je Fach **eine** Zeile.",
                 "**GROUP-BY-Regel:** jede SELECT-Spalte aggregiert oder in GROUP BY.",
                 "Nach dem **Schlüssel** gruppieren, Namen zusätzlich aufnehmen.",
                 "„Anzahl Bons“ = `COUNT(*)` · „verkaufte Stück“ = `SUM(Menge)` · `COUNT(Spalte)` übersieht NULL."],
                notizen="Merkkasten ins Heft. Tafelbild mit den drei Abfragen siehe Lehrerhinweise.")
    p.abi_format("Gegeben ist die Datenbank Stadtfest. **Entwerfen** Sie eine SQL-Abfrage, die für jeden Stand die "
                 "Standnummer, den Namen und die Anzahl der dort verkauften Produkte (Stückzahl) ausgibt.", be=4,
                 operator="entwerfen", notizen="Wie Abi-Ecke AB 3.1. 4 Minuten Einzelarbeit, dann Partner-Check mit der "
                                               "BE-Verteilung von der nächsten Folie.")
    p.loesungsfolie("Lösung Abi-Format (DS 5)",
                    SQL("SELECT Stand.StandNr, Stand.Name, SUM(Position.Menge) AS Stueck\nFROM Stand, Bon, Position\n"
                        "WHERE Stand.StandNr = Bon.StandNr\n  AND Bon.BonNr = Position.BonNr\nGROUP BY Stand.StandNr, Stand.Name;"),
                    "Je 1 BE: Verbund · `SUM(Position.Menge)` (nicht COUNT, P4) · GROUP BY nach StandNr und Name (P3) · Ausgabe.",
                    notizen="Auch richtig: über Produkt statt Bon verbinden (Produkt.StandNr).")
    p.briefing("Fallbesprechung", "Am Kaffeemobil gibt es Bons voller Positionen – aber auffällig viele davon mit der "
               "Menge 0. Wie finden wir nur die auffälligen Gruppen?",
               notizen="Cliffhanger (STORY 5.5c) und Impuls für DS 6 (HAVING). Nicht auflösen.")
    # ---------------------------------------------------------- DS 6
    p.titelfolie("Die Statistik lügt nicht", "Akte 3 · DS 6 · Ermittlung 3.2",
                 notizen="DS 6 beginnt mit dem Warm-up (Karten W-A3-01 … 04, 10 min).")
    p.briefing("Die Statistik lügt nicht",
               "Die Bons mit Menge 0 lassen mir keine Ruhe. Stellt euch vor: Jemand bekommt einen Kaffee, zahlt bar, und die "
               "Kasse bucht nichts – so verschwindet Geld, ohne dass die Kasse falsch rechnet. Wir müssen wissen, wann und bei "
               "wem: Dazu filtern wir Gruppen mit HAVING. Vorsicht beim Verbinden: Wer Bons, Positionen und Schichten in einen "
               "Topf wirft, kann Zeilen vervielfachen – dann lügt die Statistik eben doch. Und weil im Abitur gern eine "
               "fertige Abfrage vorgelegt wird, analysiert ihr heute auch Abfragen im Sachzusammenhang.",
               notizen="STORY 5.6b (gekürzt). Danach Hinweis: Am Ende der Stunde wartet Verhör 3.")
    p.inhaltsfolie("Impulse",
                   Liste(["Wie zeigen wir nur Personen mit **mehr als 3** Positionen der Menge 0? Kann das in WHERE stehen?",
                          "Juri hat Bons, Positionen und Schichten verbunden – plötzlich ist der Fehlbetrag **dreimal so hoch**. "
                          "Was ist passiert?"]),
                   notizen="Vermutungen festhalten (Tafelrand); Auflösung im Unplugged bzw. bei der Vervielfachung.")
    p.begriffe("Die Zeitlupe: Auswertungsreihenfolge",
               [("1 FROM", "Tabellen holen und verbinden"), ("2 WHERE", "Zeilen filtern (vor dem Gruppieren)"),
                ("3 GROUP BY", "Fächer bilden"), ("4 HAVING", "Fächer filtern (nach dem Gruppieren)"),
                ("5 SELECT", "Spalten und Aggregate berechnen, Aliase vergeben"), ("6 ORDER BY", "sortieren (Alias erlaubt)"),
                ("7 LIMIT", "nach n Zeilen abschneiden")],
               notizen="Warum kein Alias in HAVING? SELECT wird erst nach HAVING ausgewertet – nach EH steht in HAVING die "
                       "Aggregatfunktion selbst.")
    sql3 = ("SELECT Produkt.Bezeichnung, SUM(Position.Menge) AS Stueck\nFROM Bon, Position, Produkt\n"
            "WHERE Bon.BonNr = Position.BonNr\n  AND Position.ProduktNr = Produkt.ProduktNr\n"
            "  AND Bon.Zeitpunkt LIKE '2026-08-28%'\nGROUP BY Produkt.ProduktNr, Produkt.Bezeichnung\n"
            "HAVING SUM(Position.Menge) >= 40\nORDER BY Stueck DESC;")
    p.beispiel_schritte("WHERE und HAVING in einer Abfrage", SQL(sql3),
                        ["**WHERE** – nur Bons vom Freitag (Zeilenbedingung)", "**GROUP BY** – ein Fach je Produkt",
                         "**HAVING** – nur Fächer mit mindestens 40 Stück", "**SELECT/ORDER BY** – Bezeichnung und Stück, "
                         "meistverkauft zuerst"],
                        ergebnis=ergebnis(sql3),
                        notizen="Mit der Zeitlupe am Beamer Schritt für Schritt zeigen: Wie viele Zeilen bleiben nach WHERE, "
                                "wie viele Fächer nach GROUP BY, wie viele nach HAVING?")
    p.unplugged("Menschliche Zeitlupe",
                ["Sieben Freiwillige halten die Klauselkarten in **Schreibreihenfolge** hoch.",
                 "Das Plenum sortiert sie in **Auswertungsreihenfolge** um. Impuls: Warum kann WHERE noch nicht wissen, "
                 "wie viele Bons eine Gruppe hat?",
                 "Teams sortieren 12 Bedingungskarten auf zwei Stapel: „vor dem Gruppieren (WHERE)“ / „nach dem "
                 "Gruppieren (HAVING)“ – Wettlauf.",
                 "Auflösung durch die Klausel-Personen; Zusatzfrage: Alias in ORDER BY ja, in HAVING nein – warum?"],
                material="7 Klauselkarten, 12 Bedingungskarten", dauer="10 MIN",
                notizen="Kopiervorlage 3_Unplugged/DD_A3_Unplugged_Zeitlupe_Karten. Bedingungskarten je Team einmal "
                        "(laminiert). Sicherung: Reihenfolge als Leiste über der Tafel.")
    sql4 = ("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname,\n       COUNT(*) AS Schichten\nFROM Helfer, Schicht\n"
            "WHERE Helfer.HelferNr = Schicht.HelferNr\nGROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname\n"
            "ORDER BY Schichten DESC, Helfer.HelferNr;")
    p.beispiel_schritte("Top-n mit LIMIT – Vorsicht, Gleichstand", SQL(sql4),
                        ["`ORDER BY … DESC LIMIT 3` – die „drei fleißigsten“ Personen?", "Vier Personen haben je 4 Schichten.",
                         "`LIMIT 3` schneidet eine davon willkürlich ab.", "Besser: „alle mit mindestens 4 Schichten“ "
                         "(`HAVING COUNT(*) >= 4`)"],
                        ergebnis=ergebnis(sql4, max_zeilen=5),
                        notizen="Gleichstand an der Schnittkante (P6). Im Lab: A3-E2-05 (kein Gleichstand) und der Cold Case "
                                "„Gleichstand an der Kante“.")
    p.inhaltsfolie("Vervielfachung im Verbund",
                   "Bon 486 hat **eine** Position (Kaffee, Menge 0). Wird Schicht nur über den Stand verbunden, passt jede "
                   "Schicht des Kaffeemobils ab 21 Uhr – aus einer Position werden **drei Zeilen**:",
                   SQL("SELECT Bon.BonNr, Position.ProduktNr, Schicht.HelferNr, Schicht.Beginn\nFROM Bon, Position, Schicht\n"
                       "WHERE Bon.BonNr = Position.BonNr\n  AND Bon.StandNr = Schicht.StandNr\n  AND Bon.BonNr = 486\n"
                       "  AND Schicht.Beginn >= '2026-08-29 21:00'\nORDER BY Schicht.Beginn;"),
                   ergebnis("SELECT Bon.BonNr, Position.ProduktNr, Schicht.HelferNr, Schicht.Beginn\nFROM Bon, Position, Schicht\n"
                            "WHERE Bon.BonNr = Position.BonNr\n  AND Bon.StandNr = Schicht.StandNr\n  AND Bon.BonNr = 486\n"
                            "  AND Schicht.Beginn >= '2026-08-29 21:00'\nORDER BY Schicht.Beginn;"),
                   notizen="Merksatz: Eine Position, drei Zeilen – jede SUM wird dreimal so groß. Nur verbinden, was die Frage "
                           "braucht. Bezug zum Impuls („dreimal so hoch“). Im Lab A3-E2-07 „Dreimal zu viel“. Bon 486 ist das "
                           "Beweisstück der Akte (Kassenbon).")
    p.inhaltsfolie("Abfragen analysieren und Vorgehen beschreiben",
                   "**Analyse im Sachzusammenhang:** „Die Abfrage liefert für jede/jeden … die/den …, berücksichtigt "
                   "werden nur …, ausgegeben werden nur Gruppen, die …, sortiert nach …“ – nicht: „wählt aus und gruppiert“.",
                   "**Vorgehen in Worten:**",
                   Liste(["Tabellen", "Verbund", "Zeilenfilter (WHERE)", "Gruppierung", "Gruppenfilter (HAVING)",
                          "Sortierung/Begrenzung"], nummeriert=True),
                   notizen="Satzbausteine auch auf AB 3.2 (Beweisstück 5). Seit 2025 im Abitur: „Beschreiben Sie Ihr "
                           "Vorgehen …“.")
    p.lab_auftrag("3.2", "Die Statistik lügt nicht", "A3-E2-01 … 08 · Cold Case A3-E2-CC", werkzeug="Zeitlupe",
                  hinweise=["06 und 08: Analyse in eigenen Worten – die Antwort bleibt auf deinem Gerät.",
                            "Danach: Verhör 3 · Schulkiosk (20 min, gesiezt, wie im Abitur)."],
                  notizen="Bei jeder Fehldiagnose die Zeitlupe öffnen lassen. AB 3.2 parallel. Verhör 3 im Kurs auf "
                          "„offen“ stellen.")
    p.merkfolie("Ermittler-Notiz: WHERE, HAVING, LIMIT",
                ["FROM → WHERE → GROUP BY → HAVING → SELECT → ORDER BY → LIMIT",
                 "WHERE filtert **Zeilen vorher**, HAVING filtert **Fächer nachher** – Aggregat wiederholen, kein Alias.",
                 "Top-n mit LIMIT nur eindeutig ohne Gleichstand an Platz n.",
                 "Vervielfachung: nur Tabellen verbinden, die die Frage braucht.",
                 "Analyse: Information im Sachzusammenhang, nicht Syntax nacherzählen."],
                notizen="Merkkasten ins Heft.")
    sql_abi = ("SELECT Produkt.Bezeichnung, SUM(Position.Menge) AS Stueck\nFROM Bon, Produkt, Position\n"
               "WHERE Bon.BonNr = Position.BonNr\n  AND Position.ProduktNr = Produkt.ProduktNr\n"
               "  AND Bon.Zeitpunkt LIKE '2026-08-29%'\nGROUP BY Produkt.ProduktNr, Produkt.Bezeichnung\n"
               "HAVING SUM(Position.Menge) > 50\nORDER BY Stueck DESC;")
    p.abi_format("**Analysieren** Sie die Abfrage, indem Sie beschreiben, welche Informationen sie im Sachzusammenhang "
                 "liefert.", be=4, operator="analysieren", inhalt=[SQL(sql_abi)],
                 notizen="Wie Abi-Ecke AB 3.2. Erst schriftlich (4 min), dann zwei Antworten vorlesen und mit den "
                         "Kernpunkten vergleichen.")
    p.loesungsfolie("Lösung Abi-Format (DS 6)",
                    "Die Abfrage liefert für jedes **Produkt**, das am **Samstag, 29.08.**, verkauft wurde, die Bezeichnung "
                    "und die **verkaufte Stückzahl** dieses Tages – nur Produkte mit **mehr als 50** Stück, das "
                    "meistverkaufte zuerst.",
                    Tabelle(*abfrage_ergebnis(sql_abi, DB)),
                    "Je 1 BE: je Produkt + Stückzahl · Samstag · mehr als 50 · absteigend.",
                    notizen="Typischer Fehler P10: Syntax nacherzählen. P4: „Anzahl Verkäufe“ statt Stück.")
    sql_tafel = ("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname, COUNT(*) AS NullPositionen\n"
                 "FROM Bon, Position, Helfer\nWHERE Bon.BonNr = Position.BonNr\n  AND Bon.HelferNr = Helfer.HelferNr\n"
                 "  AND Position.Menge = 0\n  AND Bon.Zeitpunkt BETWEEN '2026-08-29 00:00' AND '2026-08-29 23:59'\n"
                 "GROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname\nHAVING COUNT(*) > 3\nORDER BY NullPositionen DESC;")
    p.loesungsfolie("Fallbesprechung: die auffällige Schicht", SQL(sql_tafel), ergebnis(sql_tafel),
                    notizen="Tafelbild DS 6. Erst nach der Lab-Phase zeigen (Lösung zu A3-E2-04). Analyse-Schema an dieser "
                            "Abfrage üben: „Die Abfrage liefert für jede Person …“.")
    p.briefing("Fallbesprechung",
               "Die Schicht von 22 bis 23 Uhr am Kaffeemobil hatte Helferin Nr. 23 – Nora Falter, Jahrgang 1988, gemeldet "
               "über den TuS Weserstedt. Und beim TuS stimmt seit Monaten die Beitragsliste nicht.",
               notizen="Abschluss (STORY 5.6c) erst ganz am Ende: Beweisstück 3 „Kassenbon“ an die Pinnwand. Steckbrief: "
                       "Jahrgang 1988 bestätigt, kennt Kassensysteme, Verbindung zum TuS. Überleitung zu Akte 4.",
               nur_lehrkraft=True)
    p.speichern(ORDNER / "1_Praesentation" / "DD_A3_Stadtfest.pptx")


# ================================================================== Lehrerhinweise
def lehrerhinweise():
    d = AktenDokument("lehrerhinweise", 3, "Lehrerhinweise Akte 3", fassung="lehrkraft", material="Lehrerhinweise A3",
                      datenbank=DB, untertitel="Die Stadtfest-Kasse · DS 5 „Kassensturz“ · DS 6 „Die Statistik lügt nicht“")
    d.kopfband()
    d.ueberschrift("Überblick", 1)
    d.tabelle(["DS", "Thema", "Lab", "Warm-up", "Abi-Bezug"],
              [["5", "Kassensturz: Aggregatfunktionen, COUNT oder SUM, GROUP BY, GROUP-BY-Regel",
                "A3-E1-01…08 (Sortierfächer), Cold Case A3-E1-CC", "W-A2-05…08 (Akte 2: Verbund, IS NULL)",
                "GROUP BY/HAVING in 16 von 19 Abituraufgaben 2021–2026"],
               ["6", "Die Statistik lügt nicht: HAVING, LIMIT, Vervielfachung, Analyse, Vorgehen in Worten",
                "A3-E2-01…08 (Zeitlupe), Cold Case A3-E2-CC; Verhör 3 · Schulkiosk", "W-A3-01…04 (Aggregation)",
                "Abfrage analysieren (14/19), Vorgehen in Worten (seit 2025), Top-n/LIMIT"]],
              breiten=[1.0, 5.0, 4.2, 3.2, 3.6])
    d.absatz("**Material:** Präsentation `DD_A3_Stadtfest.pptx` (DS 5: Folien 1–15, DS 6: Folien 16–30) · AB 3.1 und "
             "AB 3.2 mit Lösungen · Kopiervorlagen „Menschliches GROUP BY“ (30 Bon-Karten, 4 Fächerschilder) und "
             "„Menschliche Zeitlupe“ (7 Klauselkarten, 12 Bedingungskarten) · Warm-up-Karten A3 · Video 3.")
    d.ueberschrift("Datenbank Stadtfest", 2)
    d.inhalt(*schema_inhalt(),
             "Story-Fakten (STORY 4.3): Fest am Fr 28.08. (16–24 Uhr) und Sa 29.08.2026 (14–24 Uhr); 8 Stände, 26 Produkte, "
             "60 Helferinnen und Helfer, ≈ 550 Bons. Am Kaffeemobil (Stand 4) hat Helferin 23 „Nora Falter“ (Jahrgang 1988, "
             "TuS Weserstedt) am Samstag 22:00–23:00 kassiert: 52 Bons, alle bar, 110 Positionen mit Menge 0 im Listenwert "
             "von genau 312,50 €. Info-Zelt (Stand 8) und Kinderpunsch (Produkt 26) erscheinen in keiner Verbund-Auswertung. "
             "Zwei Helferinnen heißen Lotta Bödeker (HelferNr 26 und 30).",
             "Randfall-Varianten der Plattform: `null`, `gleichnamig`, `grenze` (Bon genau 23:00), `gleichstand`, `leer` – "
             "die Lab-Prüfung führt jede Lösung auch dort aus.")

    # ---------------------------------------------------------- DS 5
    d.seitenumbruch()
    d.ueberschrift("DS 5 · Ermittlung 3.1 „Kassensturz“", 1)
    d.ueberschrift("Ziele", 2)
    d.inhalt(Liste(["Die Schülerinnen und Schüler **entwerfen** Abfragen mit COUNT, SUM, AVG, MIN und MAX ohne GROUP BY "
                    "und **unterscheiden** COUNT(*), COUNT(Spalte) und SUM(Menge) (K25).",
                    "… **beschreiben** die Wirkung von GROUP BY als Bildung von Gruppen (Fächern) mit je einer Ergebniszeile (K26).",
                    "… **entwerfen** gruppierte Abfragen unter strikter Einhaltung der GROUP-BY-Regel und gruppieren nach dem "
                    "Schlüssel (K26).",
                    "… **geben** das Ergebnis einer gruppierten Abfrage auf einem kleinen Ausschnitt konkret **an** (K29)."],
                   nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A2-05…08 (Ermittlungsweg, Station zweimal, IS NULL, Fehlerzeile)", "EA → Team",
                "Warm-up-Karten A2"],
               ["Briefing", "10–20'", "Fehlbetrag 312,50 €, Schema Stadtfest; Impulse COUNT vs. SUM, „wie viele "
                "Ergebniszeilen?“; Aggregatfunktionen ohne GROUP BY", "LV, UG", "Folien 1–7"],
               ["Unplugged", "20–35'", "Menschliches GROUP BY (s. u.)", "Plenum, Teams", "Bon-Karten, Fächerschilder, Folie 9"],
               ["Lab", "35–80'", "Ermittlung 3.1, Einstieg Sortierfächer (A3-E1-04); AB 3.1 parallel", "EA",
                "iPad/Laptop, AB 3.1, Folie 11"],
               ["Fallbesprechung", "80–90'", "Merkkasten GROUP-BY-Regel + COUNT/SUM, Abi-Ecke (4 BE), Cliffhanger "
                "„Menge 0 am Kaffeemobil“", "UG", "Folien 12–15, Tafel"]],
              breiten=[3.0, 1.4, 6.8, 2.2, 3.6])
    d.ueberschrift("Unplugged-Ablauf: Menschliches GROUP BY (15 min)", 2)
    d.inhalt("**Material:** Kopiervorlage `DD_A3_Unplugged_GROUP_BY_Karten` – 30 Bon-Karten (BonNr 101–130, eigene Daten: "
             "Stand, Helfer, Zahlart, Menge, Betrag) und 4 Fächerschilder (Kaffeemobil, Grillhütte, Saftbar, Crêpes-Stand) "
             "für die Raumecken. Die Lösung der Kopiervorlage enthält alle Tafelwerte.",
             Liste(["Jede/r erhält eine Bon-Karte. `SELECT COUNT(*) FROM Bon` → alle stehen auf, eine Person zählt: eine Zahl, "
                    "eine Ergebniszeile.",
                    "`… GROUP BY Stand` → alle gehen in die Ecke ihres Standes. Je Fach ein **Zähler** (COUNT(*)) und ein "
                    "**Summierer** (SUM(Betrag)); Ergebnisse an die Tafel: eine Zeile je Fach.",
                    "`SELECT Stand, Helfer, COUNT(*) … GROUP BY Stand` → „Welchen Helfer soll euer Fach melden?“ Im Fach "
                    "stehen zwei Helfer → keine eindeutige Antwort → **GROUP-BY-Regel**.",
                    "`SUM(Menge)` statt `COUNT(*)` → Kaffeemobil hat die meisten Bons (10), der Crêpes-Stand die meisten "
                    "Stück (19): andere Rangfolge (P4).",
                    "`WHERE Zahlart = 'Karte'` **vor** dem Gruppieren: Barzahler setzen sich, bevor die Fächer gebildet "
                    "werden (Vorbereitung DS 6)."], nummeriert=True),
             "*Sicherung:* Ergebnistabelle an der Tafel bleibt als Tafelbild; Merksatz zur GROUP-BY-Regel.")
    d.ueberschrift("Tafelbild / Ermittler-Notiz", 2)
    t1 = "SELECT COUNT(*) AS Bons, MIN(Zeitpunkt) AS ErsterBon, MAX(Zeitpunkt) AS LetzterBon\nFROM Bon;"
    t2 = ("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS AnzahlBons\nFROM Bon, Stand\nWHERE Bon.StandNr = Stand.StandNr\n"
          "GROUP BY Stand.StandNr, Stand.Name\nORDER BY AnzahlBons DESC;")
    t3 = ("SELECT Produkt.ProduktNr, Produkt.Bezeichnung, SUM(Position.Menge) AS Stueck\nFROM Position, Produkt\n"
          "WHERE Position.ProduktNr = Produkt.ProduktNr\nGROUP BY Produkt.ProduktNr, Produkt.Bezeichnung;")
    d.inhalt(SQL(t1), ergebnis(t1), SQL(t2), ergebnis(t2, max_zeilen=3), SQL(t3),
             Notiz("Merksätze", "**GROUP-BY-Regel:** Jede Spalte im SELECT ist entweder aggregiert oder in GROUP BY "
                   "aufgezählt. · Nach dem **Schlüssel** gruppieren (zwei Helferinnen heißen „Lotta Bödeker“!) und Namen "
                   "zusätzlich in GROUP BY aufnehmen. · `COUNT(*)` zählt Zeilen, `COUNT(Spalte)` zählt Nicht-NULL-Werte, "
                   "`SUM(Menge)` addiert Stückzahlen."))
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["Name im SELECT, aber nicht in GROUP BY (P3) → SQLite liefert trotzdem ein Ergebnis! Die Plattform meldet "
                    "den Regelverstoß; Merksatz „läuft ≠ richtig“.",
                    "Gruppieren nur nach Name → Gleichnamige verschmelzen (P3; Prüf-Variante `gleichnamig`).",
                    "COUNT statt SUM für Stückzahlen (P4) bzw. `COUNT(Verein)` übersieht NULL.",
                    "Aggregat ohne GROUP BY mit weiterer Spalte (`SELECT Bezeichnung, MAX(Preis)`) → in SQLite „zufällige“ "
                    "Bezeichnung.",
                    "Verbund vergessen bei Produkt/Position (P2)."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Case „Kartenzahler“ (AB), Lab-Cold-Case „Der stille Stand“ (Info-Zelt fehlt im "
                    "Verbund), Vorgriff HAVING.",
                    "*Langsame:* Sortierfächer jederzeit aus der Konsole; Werkzeugkarte 2 (Auswertungsreihenfolge) im "
                    "Ermittler-Handbuch; ★★-Aufgaben mit Tipp."]))
    d.notiz("Vorbereitung DS 5", "Bon-Karten (30) und 4 Fächerschilder drucken (ggf. laminieren), Raumecken frei. "
            "Akte 3 im Kurs freischalten; Sortierfächer-Demo am Beamer testen. Warm-up-Karten A2 (W-A2-05…08) bereitlegen.")

    # ---------------------------------------------------------- DS 6
    d.seitenumbruch()
    d.ueberschrift("DS 6 · Ermittlung 3.2 „Die Statistik lügt nicht“", 1)
    d.ueberschrift("Ziele", 2)
    d.inhalt(Liste(["Die Schülerinnen und Schüler **unterscheiden** Zeilen- und Gruppenbedingungen und **entwerfen** "
                    "Abfragen mit WHERE und HAVING in einer Anweisung, ohne Alias in HAVING (K27).",
                    "… **entwerfen** Top-n-Abfragen mit ORDER BY und LIMIT und **erläutern** das Gleichstand-Problem (K22).",
                    "… **analysieren** eine gegebene Abfrage im Sachzusammenhang und **beschreiben** das Vorgehen zur "
                    "Beantwortung einer Frage in Worten (K28, K30).",
                    "… **beurteilen** die Korrektheit einer Abfrage, insbesondere die Vervielfachung beim Verbund über zwei "
                    "1:n-Beziehungen (K31, eA)."], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A3-01…04 (COUNT/SUM, Regel-Check, Fächer füllen, zwei Lotta Bödeker)",
                "EA → Team", "Warm-up-Karten A3"],
               ["Briefing", "10–20'", "Menge 0 am Kaffeemobil; Impulse „mehr als 3 – in WHERE?“ und „dreimal so hoch“; "
                "Zeitlupe, WHERE + HAVING", "LV, UG", "Folien 16–20"],
               ["Unplugged", "20–30'", "Menschliche Zeitlupe (s. u.)", "Plenum, Teams", "Klausel- und Bedingungskarten, Folie 21"],
               ["Lab", "30–75'", "Ermittlung 3.2; Zeitlupe bei jeder Fehldiagnose; 06 und 08 als `kernpunkte` (Freitext "
                "bleibt im Browser); AB 3.2 parallel", "EA", "Folien 22–25, AB 3.2"],
               ["Fallbesprechung", "75–90'", "Analyse-Schema an A3-E2-06; Tafelbild; Abschluss Nora Falter / TuS; Abi-Ecke "
                "(4 BE); Hinweis Verhör 3 (20')", "UG", "Folien 26–30"]],
              breiten=[3.0, 1.4, 6.8, 2.2, 3.6])
    d.ueberschrift("Unplugged-Ablauf: Menschliche Zeitlupe (10 min)", 2)
    d.inhalt("**Material:** Kopiervorlage `DD_A3_Unplugged_Zeitlupe_Karten` – 7 große Klauselkarten, 12 Bedingungskarten je "
             "Team (Lösung in der Lehrkraftfassung der Kopiervorlage).",
             Liste(["Sieben Freiwillige erhalten die Klauselkarten in Schreibreihenfolge. Das Plenum sortiert sie in "
                    "**Auswertungsreihenfolge**: FROM → WHERE → GROUP BY → HAVING → SELECT → ORDER BY → LIMIT. Impuls: "
                    "„Warum kann WHERE noch nicht wissen, wie viele Bons eine Gruppe hat?“",
                    "Teams sortieren ihre Bedingungskarten auf zwei Stapel „vor dem Gruppieren (WHERE)“ / „nach dem "
                    "Gruppieren (HAVING)“ – Wettlauf, Auflösung durch die Klausel-Personen.",
                    "Zusatzfrage: „Warum darf in ORDER BY der Alias stehen, in HAVING laut EH aber nicht?“ (SELECT wird erst "
                    "nach HAVING ausgewertet.)"], nummeriert=True),
             "*Sicherung:* Reihenfolge als Leiste über der Tafel; Werkzeugkarte 2.")
    d.ueberschrift("Tafelbild / Ermittler-Notiz", 2)
    t4 = ("SELECT Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname, COUNT(*) AS NullPositionen\nFROM Bon, Position, Helfer\n"
          "WHERE Bon.BonNr = Position.BonNr\n  AND Bon.HelferNr = Helfer.HelferNr\n  AND Position.Menge = 0\n"
          "  AND Bon.Zeitpunkt BETWEEN '2026-08-29 00:00' AND '2026-08-29 23:59'\n"
          "GROUP BY Helfer.HelferNr, Helfer.Vorname, Helfer.Nachname\nHAVING COUNT(*) > 3\nORDER BY NullPositionen DESC;")
    t5 = ("SELECT SUM(Produkt.Preis) AS Fehlbetrag\nFROM Bon, Position, Produkt\nWHERE Bon.BonNr = Position.BonNr\n"
          "  AND Position.ProduktNr = Produkt.ProduktNr\n  AND Bon.StandNr = 4\n  AND Position.Menge = 0\n"
          "  AND Bon.Zeitpunkt BETWEEN '2026-08-29 22:00' AND '2026-08-29 22:59';")
    t5_falsch = ("SELECT SUM(Produkt.Preis) AS Fehlbetrag\nFROM Bon, Position, Produkt, Schicht\nWHERE Bon.BonNr = Position.BonNr\n"
                 "  AND Position.ProduktNr = Produkt.ProduktNr\n  AND Bon.StandNr = Schicht.StandNr\n  AND Bon.StandNr = 4\n"
                 "  AND Position.Menge = 0\n  AND Bon.Zeitpunkt BETWEEN '2026-08-29 22:00' AND '2026-08-29 22:59'\n"
                 "  AND Schicht.Beginn >= '2026-08-29 21:00';")
    d.inhalt(SQL(t4), ergebnis(t4),
             "Der Fehlbetrag (Lösung A3-E2-07) – korrekt ohne Schicht:", SQL(t5), ergebnis(t5),
             "Juris Abfrage aus A3-E2-07 (Schicht nur über den Stand verbunden, drei Schichten ab 21 Uhr) verdreifacht die Summe:",
             SQL(t5_falsch), ergebnis(t5_falsch),
             Notiz("Merksätze", "**Auswertungsreihenfolge:** FROM → WHERE → GROUP BY → HAVING → SELECT → ORDER BY → LIMIT. · "
                   "WHERE filtert Zeilen vorher, HAVING Gruppen nachher (Aggregat wiederholen, kein Alias in HAVING). · "
                   "Top-n nur eindeutig ohne Gleichstand an Platz n. · Vervielfachung: nur verbinden, was die Frage braucht. · "
                   "Analyse: „Die Abfrage liefert für jede/n … die/den …, aber nur …, sortiert nach …“. · Vorgehen: Tabellen → "
                   "Verbund → Zeilenfilter → Gruppierung → Gruppenfilter → Sortierung/Begrenzung."))
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["Aggregat in WHERE (`WHERE COUNT(*) > 5`) oder Zeilenbedingung in HAVING (P1) → Sortierfächer, Zeitlupe.",
                    "Alias in HAVING (`HAVING Bons > 5`) → in SQLite lauffähig, laut EH nicht vorgesehen (Parser-Hinweis).",
                    "LIMIT bei Gleichstand – „die Top 3“ ist nicht eindeutig (P6) → Cold Case „Gleichstand an der Kante“.",
                    "Vervielfachung nicht bemerkt (P2): Das Ergebnis „sieht plausibel aus“.",
                    "Analyse als Nacherzählung der Syntax statt Sachzusammenhang (P10).",
                    "Stundengrenze: `BETWEEN '… 22:00' AND '… 23:00'` zählt einen Bon um 23:00 mit (Variante `grenze`)."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Cases (AB „Die Kasse stimmt doch“, Lab „Gleichstand an der Kante“), Verhör 3 noch in der "
                    "Stunde, eigene Analyse-Aufgabe für das Team formulieren.",
                    "*Langsame:* Satzbausteine (AB 3.2, Beweisstück 5), Zeitlupe, Werkzeugkarte 2; `kernpunkte`-Aufgaben "
                    "zuerst mündlich im Team.",
                    "*gA-Kurs:* unverändert; A3-E2-07 und AB 3.2 Auftrag 5 (K31) bleiben als Verbund-/Summenübung."]))
    d.notiz("Vorbereitung DS 6", "Klausel- und Bedingungskarten drucken; Zeitlupe-Demo (Abfrage aus dem Tafelbild) am Beamer "
            "testen. Verhör 3 · Schulkiosk im Kurs auf „auto“ oder „offen“ stellen (20 min, Abimodus streng, 6 Aufgaben).")
    d.ueberschrift("Verhör 3 · Schulkiosk (SQL-Gesamt)", 2)
    d.absatz("Eigene Kontext-DB (Produkt, Helfer, Schicht, Verkauf, Posten); gesiezt im Abitur-Stil; Pool A3-V-01…12, gezogen "
             "werden 2 × ★, 3 × ★★, 1 × ★★★ mit mindestens einer GROUP-BY-Aufgabe und einer Analyse- oder Fehleraufgabe. "
             "Rahmensatz: „Kassen lügen nicht, aber Abfragen manchmal – beweis mir, dass deine es nicht tun.“")
    d.speichern(ORDNER / "4_Lehrerhinweise" / "DD_A3_Lehrerhinweise.docx")


# ================================================================== Unplugged-Kopiervorlagen
def _karten_tabelle(d, zellen_texte, spalten, hoehe_cm, zeichner):
    """Schneidekarten: Tabelle mit fester Zeilenhöhe und gestrichelten Schnittlinien; zeichner(zelle, eintrag)."""
    zeilen = (len(zellen_texte) + spalten - 1) // spalten
    tab = d.doc.add_table(rows=zeilen, cols=spalten)
    ad._tabelle_breiten(tab, [d.breite / spalten] * spalten)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for zi in range(zeilen):
        zeile = tab.rows[zi]
        trpr = zeile._tr.get_or_add_trPr()
        trpr.append(ad._el("w:trHeight", val=int(Cm(hoehe_cm) / 635), hRule="exact"))
        ad._zeile_nicht_teilen(zeile)
        for si in range(spalten):
            z = zeile.cells[si]
            ad._zelle_rahmen(z, ad.RAND_STARK, 0.75, art="dashed")
            ad._zelle_innenabstand(z, 0.25, 0.3, 0.2, 0.3)
            z.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            i = zi * spalten + si
            z.paragraphs[0].style = d.doc.styles["DD Tabellentext"]
            if i < len(zellen_texte):
                zeichner(z, zellen_texte[i])
    return tab


def unplugged_group_by(fassung):
    karten = bon_karten()
    d = AktenDokument("unplugged", 3, "Menschliches GROUP BY", fassung=fassung, ermittlung="3.1", ds=5,
                      material="Unplugged 3.1", datenbank="keine",
                      untertitel="Kopiervorlage: 30 Bon-Karten und 4 Fächerschilder (eigene Daten)")
    d.kopfband(name_datum=False)
    if fassung == "schueler":
        d.absatz("**Vorbereitung:** Karten auf festes Papier drucken, an den gestrichelten Linien schneiden, mischen. Jede "
                 "Person zieht eine Karte. Die Fächerschilder (ab Seite 3) hängen in vier Raumecken. Ablauf: Präsentation "
                 "Folie „Menschliches GROUP BY“ bzw. Lehrerhinweise DS 5.")

        def bon(z, k):
            p = z.paragraphs[0]
            p.paragraph_format.tab_stops.add_tab_stop(Cm(d.breite / 3 - 0.7), ad.WD_TAB_ALIGNMENT.RIGHT)
            d._etikett(p, "BON", groesse=7, farbe=ad.SIEGELROT, fett=True)
            d._lauf_stil(p.add_run(f"\t{k['bon']}"), schrift=ad.SCHRIFT_KOPF, groesse=20, farbe=ad.TINTE)
            p = z.add_paragraph(style="DD Tabellentext")
            d._lauf_stil(p.add_run(k["stand"]), schrift=ad.SCHRIFT_KOPF, groesse=15, farbe=ad.TINTE)
            for feld, wert in (("Helfer", k["helfer"]), ("Zahlart", k["zahlart"]), ("Menge", str(k["menge"])),
                               ("Betrag", euro(k["betrag"]))):
                p = z.add_paragraph(style="DD Tabellentext")
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(2.2))
                d._etikett(p, feld.upper(), groesse=7)
                d._lauf_stil(p.add_run("\t" + wert), groesse=11, fett=feld in ("Menge", "Betrag"))

        _karten_tabelle(d, karten[:12], 3, 4.3, bon)
        d.seitenumbruch()
        _karten_tabelle(d, karten[12:], 3, 4.2, bon)
        for paar in (list(STAENDE_UNPLUGGED)[:2], list(STAENDE_UNPLUGGED)[2:]):
            d.seitenumbruch()

            def schild(z, stand):
                p = z.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                d._etikett(p, "FACH · GROUP BY Stand", groesse=12, farbe=ad.SIEGELROT, fett=True)
                p = z.add_paragraph(style="DD Tabellentext")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                d._lauf_stil(p.add_run(stand), schrift=ad.SCHRIFT_KOPF, groesse=60, farbe=ad.TINTE)
                p = z.add_paragraph(style="DD Tabellentext")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                d._lauf_stil(p.add_run("Zähler: COUNT(*) = ____    Summierer: SUM(Betrag) = ________"), groesse=13,
                             farbe=ad.TINTE_2)

            _karten_tabelle(d, paar, 1, 12.6, schild)
    else:
        d.absatz("Erwartete Tafelwerte für die 30 Bon-Karten (BonNr 101–130). Fehlen Karten, die Werte vorher anpassen. "
                 "Die SQL-Zeilen beziehen sich auf die vereinfachte Bon-Karte (Spalten Stand, Helfer, Zahlart, Menge, Betrag), "
                 "nicht auf die Stadtfest-Datenbank.")
        alle = auswertung_unplugged(karten)
        d.ueberschrift("Schritt 1 – ohne GROUP BY", 2)
        d.inhalt(SQL("SELECT COUNT(*) AS Bons, SUM(Betrag) AS Summe\nFROM Bon;", db="keine"),
                 Tabelle(["Bons", "Summe"], [[str(len(karten)), euro(sum(k["betrag"] for k in karten))]]),
                 "Eine Zahl je Aggregat – genau **eine** Ergebniszeile.")
        d.ueberschrift("Schritt 2 – GROUP BY Stand (Zähler und Summierer)", 2)
        d.inhalt(SQL("SELECT Stand, COUNT(*) AS Bons, SUM(Betrag) AS Summe\nFROM Bon\nGROUP BY Stand;", db="keine"),
                 Tabelle(["Stand", "Bons", "Summe"], [[s, str(z[0]), euro(z[2])] for s, z in alle.items()]))
        d.ueberschrift("Schritt 3 – Regelverstoß", 2)
        d.inhalt(SQL("SELECT Stand, Helfer, COUNT(*) AS Bons\nFROM Bon\nGROUP BY Stand;", db="keine", pruefung="fehlerhaft"),
                 Tabelle(["Fach", "Helfer im Fach"], [[s, " und ".join(STAENDE_UNPLUGGED[s]["helfer"])] for s in alle]),
                 "In jedem Fach stehen zwei Helfer – das Fach kann keinen eindeutigen Wert melden. **GROUP-BY-Regel:** Helfer "
                 "müsste aggregiert oder gruppiert werden (`GROUP BY Stand, Helfer` → acht Fächer).")
        d.ueberschrift("Schritt 4 – COUNT(*) oder SUM(Menge)?", 2)
        d.inhalt(SQL("SELECT Stand, COUNT(*) AS Bons, SUM(Menge) AS Stueck\nFROM Bon\nGROUP BY Stand;", db="keine"),
                 Tabelle(["Stand", "Bons", "Stueck"], [[s, str(z[0]), str(z[1])] for s, z in alle.items()]),
                 "Nach Bons liegt das Kaffeemobil vorn (viele Einzelkäufe), nach Stück der Crêpes-Stand (wenige große Käufe) "
                 "– P4.")
        karte = auswertung_unplugged(karten, nur_karte=True)
        d.ueberschrift("Schritt 5 – WHERE vor dem Gruppieren", 2)
        d.inhalt(SQL("SELECT Stand, COUNT(*) AS Kartenbons\nFROM Bon\nWHERE Zahlart = 'Karte'\nGROUP BY Stand;", db="keine"),
                 Tabelle(["Stand", "Kartenbons"], [[s, str(z[0])] for s, z in karte.items()]),
                 f"Insgesamt {sum(z[0] for z in karte.values())} Kartenzahler bleiben stehen; die Barzahler setzen sich, "
                 "**bevor** die Fächer gebildet werden. Vorgriff DS 6: Eine Bedingung wie „nur Fächer mit mehr als 5 Bons“ "
                 "kann erst **nach** dem Gruppieren geprüft werden (HAVING).")
    name = "DD_A3_Unplugged_GROUP_BY_Karten"
    return d.speichern(ORDNER / "3_Unplugged" / (f"{name}_LSG.docx" if fassung == "loesung" else f"{name}.docx"))


KLAUSELN = ["SELECT", "FROM", "WHERE", "GROUP BY", "HAVING", "ORDER BY", "LIMIT"]  # Schreibreihenfolge
AUSWERTUNG = ["FROM", "WHERE", "GROUP BY", "HAVING", "SELECT", "ORDER BY", "LIMIT"]
BEDINGUNGEN = [  # (Bedingung, WHERE/HAVING, Begründung)
    ("Position.Menge > 2", "WHERE", "prüft eine einzelne Position"),
    ("SUM(Position.Menge) > 100", "HAVING", "Summe über das ganze Fach"),
    ("Bon.Zahlart = 'bar'", "WHERE", "prüft einen einzelnen Bon"),
    ("COUNT(*) >= 50", "HAVING", "Anzahl der Zeilen im Fach"),
    ("Bon.Zeitpunkt BETWEEN '2026-08-29 22:00' AND '2026-08-29 22:59'", "WHERE", "Zeitpunkt eines einzelnen Bons"),
    ("AVG(Produkt.Preis) < 3", "HAVING", "Durchschnitt über das Fach"),
    ("Helfer.Verein IS NULL", "WHERE", "Eigenschaft einer einzelnen Person"),
    ("MAX(Produkt.Preis) > 4", "HAVING", "größter Wert im Fach"),
    ("Stand.Name = 'Bierwagen'", "WHERE", "prüft eine einzelne Zeile"),
    ("COUNT(*) = 1", "HAVING", "Fächer mit genau einer Zeile"),
    ("Produkt.Preis >= 3.5", "WHERE", "Preis eines einzelnen Produkts"),
    ("SUM(Position.Menge * Produkt.Preis) > 500", "HAVING", "Umsatz des ganzen Fachs"),
]


def unplugged_zeitlupe(fassung):
    d = AktenDokument("unplugged", 3, "Menschliche Zeitlupe", fassung=fassung, ermittlung="3.2", ds=6,
                      material="Unplugged 3.2", datenbank="keine",
                      untertitel="Kopiervorlage: 7 Klauselkarten und 12 Bedingungskarten (WHERE oder HAVING?)")
    d.kopfband(name_datum=False)
    if fassung == "schueler":
        d.absatz("**Vorbereitung:** Klauselkarten (Seiten 1–4) einmal drucken, Bedingungskarten (Seite 5) je Team einmal. "
                 "Die Klauselkarten werden zuerst in Schreibreihenfolge ausgegeben – die Auswertungsreihenfolge findet das "
                 "Plenum selbst.")

        def klausel(z, k):
            p = z.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            d._etikett(p, "KLAUSELKARTE · ZEITLUPE", groesse=11, farbe=ad.SIEGELROT, fett=True)
            p = z.add_paragraph(style="DD Tabellentext")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            d._lauf_stil(p.add_run(k), code=True, fett=True, groesse=66, farbe=ad.TINTE)
            p = z.add_paragraph(style="DD Tabellentext")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            d._lauf_stil(p.add_run("Auswertungsschritt Nr. ____"), groesse=14, farbe=ad.TINTE_2)

        _karten_tabelle(d, KLAUSELN[:1], 1, 10.4, klausel)
        for i in range(1, len(KLAUSELN), 2):
            d.seitenumbruch()
            _karten_tabelle(d, KLAUSELN[i:i + 2], 1, 12.6, klausel)
        d.seitenumbruch()
        d.ueberschrift("Bedingungskarten (je Team ein Satz)", 2)

        def bedingung(z, b):
            p = z.paragraphs[0]
            d._etikett(p, "BEDINGUNG", groesse=7, farbe=ad.SIEGELROT, fett=True)
            p = z.add_paragraph(style="DD Tabellentext")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            groesse = 13 if len(b[0]) <= 26 else 10
            d._lauf_stil(p.add_run(b[0]), code=True, fett=True, groesse=groesse, farbe=ad.TINTE)
            p = z.add_paragraph(style="DD Tabellentext")
            d._lauf_stil(p.add_run("vor dem Gruppieren?  nach dem Gruppieren?"), groesse=7.5, farbe=ad.TINTE_3)

        _karten_tabelle(d, BEDINGUNGEN, 3, 5.0, bedingung)
        d.absatz("Stapelschilder für die Teams: **„vor dem Gruppieren (WHERE)“** und **„nach dem Gruppieren (HAVING)“** – auf "
                 "die Rückseite eines Blattes schreiben.")
    else:
        d.ueberschrift("Klauselkarten", 2)
        d.tabelle(["Schreibreihenfolge", "Auswertungsreihenfolge (Zeitlupe)", "Was passiert?"],
                  [[f"{i + 1}. {s}", f"{i + 1}. {a}", t] for i, (s, a, t) in enumerate(zip(
                      KLAUSELN, AUSWERTUNG, ["Tabellen holen und verbinden", "Zeilen filtern", "Fächer bilden",
                                             "Fächer filtern", "Spalten/Aggregate berechnen, Aliase vergeben",
                                             "sortieren (Alias bekannt)", "abschneiden"]))],
                  breiten=[4.2, 5.4, 7.4])
        d.absatz("Impuls-Antwort: WHERE wird vor GROUP BY ausgewertet – zu diesem Zeitpunkt gibt es noch keine Fächer, also "
                 "kein COUNT je Gruppe. Alias: SELECT kommt nach HAVING, deshalb ist ein Alias in HAVING (nach EH) nicht "
                 "vorgesehen, in ORDER BY aber bekannt.")
        d.ueberschrift("Bedingungskarten – Lösung", 2)
        d.tabelle(["Bedingung", "Stapel", "Begründung"], [[f"`{b}`", s, g] for b, s, g in BEDINGUNGEN],
                  breiten=[8.4, 2.2, 6.4])
        d.absatz("Faustregel: Steht eine Aggregatfunktion in der Bedingung, gehört sie in HAVING (P1).")
    name = "DD_A3_Unplugged_Zeitlupe_Karten"
    return d.speichern(ORDNER / "3_Unplugged" / (f"{name}_LSG.docx" if fassung == "loesung" else f"{name}.docx"))


# ================================================================== Warm-up
def warmup():
    karten = json.loads((REPO / "seed" / "warmup" / "warmup_a3.json").read_text(encoding="utf-8"))["karten"]
    warmup_karten_dokument(karten, 3, "Warm-up-Karten Akte 3").speichern(WARMUP / "DD_Warm-up-Karten_A3.docx")


def main():
    erzeugt = [ab31("schueler"), ab31("loesung"), ab32("schueler"), ab32("loesung"),
               unplugged_group_by("schueler"), unplugged_group_by("loesung"),
               unplugged_zeitlupe("schueler"), unplugged_zeitlupe("loesung")]
    praesentation()
    lehrerhinweise()
    warmup()
    for p in erzeugt:
        print(f"✔ {p.relative_to(REPO)}")
    print("✔ Präsentation, Lehrerhinweise, Warm-up-Karten")


if __name__ == "__main__":
    main()
