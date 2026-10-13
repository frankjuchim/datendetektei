#!/usr/bin/env python3
"""Erstellungshilfe: Material Akte 2 „Spur der Räder“ (DS 3–4, DB stadtrad) im Akten-Design.

Erzeugt in materialien/Akte2_Raeder/:
  1_Praesentation/DD_A2_Raeder.pptx
  2_Arbeitsblaetter/DD_A2_AB1_Der_rote_Faden.docx, DD_A2_AB2_Nachtfahrten.docx
  3_Unplugged/DD_A2_Unplugged_Roter_Faden.docx, DD_A2_Unplugged_Ermittlungsweg.docx
  4_Lehrerhinweise/DD_A2_Lehrerhinweise.docx
  5_Loesungen/DD_A2_AB1_Der_rote_Faden_LSG.docx, DD_A2_AB2_Nachtfahrten_LSG.docx
und materialien/Warm-up/DD_Warm-up-Karten_A2.docx.

Grundlage: docs/DS_PLANUNG.md (DS 3, DS 4), docs/STORY.md 4.2/5.3/5.4, seed/datenbanken/stadtrad/LIESMICH.md,
seed/akten/akte2.json (Spur im Lab), seed/warmup/warmup_a2.json. Alle Ergebnistabellen werden aus der Datenbank
erzeugt (abfrage_ergebnis) und von werkzeuge/material_sql_pruefen.py nachgeprüft.

Die erzeugten docx/pptx sind danach maßgeblich; dieses Skript ist nur die Erstellungshilfe.
Aufruf: python3 quellen/material/akte2.py   (überschreibt die Dateien oben)
"""
import json
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "werkzeuge"))

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import akten_design as ad  # noqa: E402
from akten_design import (  # noqa: E402
    AktenDokument, AktenPraesentation, Auftrag, Bild, Ergebnis, Liste, Loesung, Notiz, Schema, SQL, Tabelle,
    abfrage_ergebnis, warmup_karten_dokument)

DB = "stadtrad"
ZIEL = REPO / "materialien" / "Akte2_Raeder"
TMP = Path(tempfile.mkdtemp(prefix="dd-akte2-"))

STATION = "Station(<u>StationID</u>, Name, Stadtteil, Stellplaetze)"
RAD = "Rad(<u>RadID</u>, Typ, Baujahr, ↑StationID, Status)"
KUNDE = "Kunde(<u>KundenID</u>, Vorname, Nachname, Jahrgang, Tarif, Stadtteil)"
AUSLEIHE = "Ausleihe(<u>AusleihID</u>, ↑RadID, ↑KundenID, ↑StartStation, ↑ZielStation, Start, Ende, DauerMin, Preis)"
REPARATUR = "Reparatur(<u>ReparaturID</u>, ↑RadID, Datum, Grund, Kosten)"
SCHEMA_ALLE = [Schema(STATION), Schema(RAD), Schema(KUNDE), Schema(AUSLEIHE), Schema(REPARATUR)]


def ergebnis(sql, max_zeilen=None, db=DB):
    """Ergebnistabelle aus der Datenbank (gekürzt nur mit ORDER BY, damit der Vergleich eindeutig bleibt)."""
    kopf, zeilen = abfrage_ergebnis(sql, db)
    if max_zeilen and len(zeilen) > max_zeilen:
        if "ORDER BY" not in sql:
            raise SystemExit(f"gekürzte Ergebnistabelle ohne ORDER BY: {sql}")
        return Ergebnis(kopf, zeilen[:max_zeilen], abgeschnitten=True)
    return Ergebnis(kopf, zeilen)


def anzahl(sql, db=DB):
    return len(abfrage_ergebnis(sql, db, max_zeilen=100000)[1])


# ================================================================== Schema-Pfeildiagramm (Pillow)
TABELLEN = {
    # name: (x, y, [(attribut, art)])   art: "pk" | "fk" | ""
    "Kunde": (20, 70, [("KundenID", "pk"), ("Vorname", ""), ("Nachname", ""), ("Jahrgang", ""), ("Tarif", ""),
                       ("Stadtteil", "")]),
    "Ausleihe": (390, 20, [("AusleihID", "pk"), ("RadID", "fk"), ("KundenID", "fk"), ("StartStation", "fk"),
                           ("ZielStation", "fk"), ("Start", ""), ("Ende", ""), ("DauerMin", ""), ("Preis", "")]),
    "Rad": (780, 20, [("RadID", "pk"), ("Typ", ""), ("Baujahr", ""), ("StationID", "fk"), ("Status", "")]),
    "Station": (780, 240, [("StationID", "pk"), ("Name", ""), ("Stadtteil", ""), ("Stellplaetze", "")]),
    "Reparatur": (390, 400, [("ReparaturID", "pk"), ("RadID", "fk"), ("Datum", ""), ("Grund", ""), ("Kosten", "")]),
}
# (von Tabelle, FK-Attribut, nach Tabelle, Seite des Austritts, Seite des Eintritts, x der senkrechten Führung)
PFEILE = [  # so geführt, dass sich keine Linien kreuzen
    ("Ausleihe", "KundenID", "Kunde", "l", "r", 345),
    ("Ausleihe", "RadID", "Rad", "r", "l", 705),
    ("Ausleihe", "StartStation", "Station", "r", "l", 760),
    ("Ausleihe", "ZielStation", "Station", "r", "l", 735),
    ("Rad", "StationID", "Station", "r", "r", 1075),
    ("Reparatur", "RadID", "Rad", "r", "r", 1095),
]
KB, KH, ZH = 270, 40, 28  # Kastenbreite, Kopfhöhe, Zeilenhöhe (Grundmaß)


def schema_png(pfad, markiert=(), pfeile=None, massstab=2.0):
    """Pfeildiagramm des StadtRad-Schemas: Tabellen als Karten (PK unterstrichen, ↑FK rot), Pfeile vom Fremdschlüssel
    zum Primärschlüssel. markiert = Tabellen, die hervorgehoben werden (Ermittlungsweg); pfeile = Teilmenge der Pfeile
    [(Tabelle, Attribut)], die hervorgehoben werden (sonst alle gleich)."""
    from PIL import Image, ImageDraw
    s = massstab
    w, h = int(1110 * s), int(610 * s)
    bild = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(bild)
    rgb = lambda c: "#" + c  # noqa: E731
    f_kopf = ad._font("kopf", 24 * s)
    f_mono = ad._font("mono_normal", 17 * s)
    f_mono_f = ad._font("mono", 17 * s)

    def zeile_y(name, attr):
        x, y, attrs = TABELLEN[name]
        i = [a for a, _ in attrs].index(attr)
        return y + KH + i * ZH + ZH / 2

    hervor = set(markiert)
    for name, (x, y, attrs) in TABELLEN.items():
        aktiv = not hervor or name in hervor
        rand = ad.SIEGELROT if (hervor and name in hervor) else (ad.RAND_STARK if aktiv else ad.RAND)
        kopf = ad.FLAECHE if aktiv else "F4F1EC"
        hh = KH + len(attrs) * ZH + 8
        d.rounded_rectangle([x * s, y * s, (x + KB) * s, (y + hh) * s], radius=8 * s, fill=rgb(ad.KARTE),
                            outline=rgb(rand), width=int((3 if name in hervor else 2) * s))
        d.rounded_rectangle([x * s + 2 * s, y * s + 2 * s, (x + KB) * s - 2 * s, (y + KH) * s], radius=7 * s, fill=rgb(kopf))
        d.text(((x + 14) * s, (y + 7) * s), name, font=f_kopf, fill=rgb(ad.TINTE if aktiv else ad.RAND_STARK))
        for i, (a, art) in enumerate(attrs):
            ty = (y + KH + i * ZH + 4) * s
            farbe = ad.TINTE if aktiv else ad.RAND_STARK
            tx = (x + 16) * s
            if art == "fk":
                pf = rgb(ad.SIEGELROT if aktiv else ad.RAND_STARK)
                d.text((tx, ty), "↑", font=f_mono_f, fill=pf)
                d.text((tx + f_mono.getlength("↑"), ty), a, font=f_mono, fill=pf)
            elif art == "pk":
                d.text((tx, ty), a, font=f_mono_f, fill=rgb(farbe))
                lb = f_mono_f.getlength(a)
                d.line([tx, ty + 22 * s, tx + lb, ty + 22 * s], fill=rgb(farbe), width=int(2 * s))
            else:
                d.text((tx, ty), a, font=f_mono, fill=rgb(farbe))

    for von, attr, nach, aus, ein, fx in PFEILE:
        aktiv = pfeile is None or (von, attr) in pfeile
        farbe = rgb(ad.SIEGELROT if aktiv else ad.RAND)
        breite = int((3.2 if (pfeile is not None and aktiv) else 2.2) * s)
        xv, _, _ = TABELLEN[von]
        xn, _, _ = TABELLEN[nach]
        y1 = zeile_y(von, attr)
        pk = TABELLEN[nach][2][0][0]
        y2 = zeile_y(nach, pk)
        x1 = xv if aus == "l" else xv + KB
        x2 = xn if ein == "l" else xn + KB
        punkte = [(x1, y1), (fx, y1), (fx, y2), (x2, y2)]
        d.line([(px * s, py * s) for px, py in punkte], fill=farbe, width=breite, joint="curve")
        d.ellipse([(x1 - 4) * s, (y1 - 4) * s, (x1 + 4) * s, (y1 + 4) * s], fill=farbe)
        richtung = 1 if x2 > fx else -1  # Pfeilspitze zeigt in den Zielkasten
        spitze = [(x2 * s, y2 * s), ((x2 - richtung * 13) * s, (y2 - 7) * s), ((x2 - richtung * 13) * s, (y2 + 7) * s)]
        d.polygon(spitze, fill=farbe)
    bild.save(pfad)
    return pfad


def schema_bild(name, breite_cm=16.5, **kw):
    return Bild(str(schema_png(TMP / f"{name}.png", **kw)), breite_cm=breite_cm)


# ================================================================== Eigene Daten des Unplugged „Roter Faden“
UP_STATIONEN = [(1, "Leuchtturm", "Hafenviertel"), (2, "Kornmarkt", "Altstadt"), (3, "Bootshaus", "Lindenhof"),
                (4, "Deichtor", "Deichkamp")]
UP_RAEDER = [(1, "Citybike", 1), (2, "Lastenrad", 2), (3, "E-Bike", 1), (4, "Citybike", 3), (5, "Lastenrad", 2),
             (6, "E-Bike", None)]


# ================================================================== SQL der Arbeitsblätter
SQL_HAFEN = ("SELECT Rad.RadID, Rad.Typ, Station.Name\nFROM Rad, Station\nWHERE Rad.StationID = Station.StationID\n"
             "  AND Station.Stadtteil = 'Hafenviertel';")
SQL_REP100 = ("SELECT Reparatur.Datum, Reparatur.Grund, Reparatur.Kosten, Rad.Typ\nFROM Reparatur, Rad\n"
              "WHERE Reparatur.RadID = Rad.RadID\n  AND Reparatur.Kosten > 100\nORDER BY Reparatur.Datum;")
SQL_START = ("SELECT Ausleihe.AusleihID, Station.Name\nFROM Ausleihe, Station\nWHERE Ausleihe.StartStation = Station.StationID\n"
             "ORDER BY Ausleihe.AusleihID;")
SQL_FALSCH_PRAEFIX = "SELECT Name\nFROM Rad, Station\nWHERE StationID = StationID;"
SQL_PRAEFIX_RICHTIG = ("SELECT Station.Name\nFROM Rad, Station\nWHERE Rad.StationID = Station.StationID\n"
                       "ORDER BY Station.Name;")
SQL_DOPPELT = ("SELECT Ausleihe.AusleihID, Station.Name\nFROM Ausleihe, Station\n"
               "WHERE Ausleihe.StartStation = Station.StationID\n   OR Ausleihe.ZielStation = Station.StationID;")
SQL_DOPPELT_ZAEHLEN = ("SELECT COUNT(*) AS Zeilen\nFROM Ausleihe, Station\nWHERE Ausleihe.StartStation = Station.StationID\n"
                       "   OR Ausleihe.ZielStation = Station.StationID;")

SQL_EBIKE = ("SELECT DISTINCT k.Vorname, k.Nachname\nFROM Ausleihe a, Kunde k, Rad r\nWHERE a.KundenID = k.KundenID\n"
             "  AND a.RadID = r.RadID\n  AND r.Typ = 'E-Bike'\nORDER BY k.Nachname, k.Vorname;")
SQL_VONNACH = ("SELECT a.AusleihID, a.Start, s1.Name AS Startort, s2.Name AS Zielort\nFROM Ausleihe a, Station s1, Station s2\n"
               "WHERE a.StartStation = s1.StationID\n  AND a.ZielStation = s2.StationID\n"
               "  AND a.Start LIKE '2026-09-01%'\nORDER BY a.Start;")
SQL_FISCHMARKT = ("SELECT DISTINCT k.Vorname, k.Nachname\nFROM Ausleihe a, Kunde k, Station s\nWHERE a.KundenID = k.KundenID\n"
                  "  AND a.StartStation = s.StationID\n  AND s.Name = 'Fischmarkt'\nORDER BY k.Nachname, k.Vorname;")
SQL_LUECKE = ("SELECT k.Nachname, r.Typ, s.Name, a.Start\nFROM Ausleihe a, Kunde k, Rad r, Station s\n"
              "WHERE a.KundenID = k.KundenID\n  AND ______________________\n  AND a.StartStation = s.StationID\n"
              "  AND a.Start LIKE '2026-09-01%';")
SQL_LUECKE_LSG = ("SELECT k.Nachname, r.Typ, s.Name, a.Start\nFROM Ausleihe a, Kunde k, Rad r, Station s\n"
                  "WHERE a.KundenID = k.KundenID\n  AND a.RadID = r.RadID\n  AND a.StartStation = s.StationID\n"
                  "  AND a.Start LIKE '2026-09-01%'\nORDER BY a.Start;")
SQL_REP_ALTSTADT = ("SELECT rp.ReparaturID, rp.Datum, rp.Grund, r.RadID, s.Name\nFROM Reparatur rp, Rad r, Station s\n"
                    "WHERE rp.RadID = r.RadID\n  AND r.StationID = s.StationID\n  AND s.Stadtteil = 'Altstadt'\n"
                    "ORDER BY rp.Datum;")
SQL_ROSSI = ("SELECT a.AusleihID, a.Start, s1.Name AS Startort, s2.Name AS Zielort\n"
             "FROM Ausleihe a, Kunde k, Station s1, Station s2\nWHERE a.KundenID = k.KundenID\n"
             "  AND a.StartStation = s1.StationID\n  AND a.ZielStation = s2.StationID\n"
             "  AND k.Vorname = 'Clara'\n  AND k.Nachname = 'Rossi';")
SQL_ROSSI_LAUFEND = ("SELECT a.AusleihID, a.Start, s.Name AS Startort\nFROM Ausleihe a, Kunde k, Station s\n"
                     "WHERE a.KundenID = k.KundenID\n  AND a.StartStation = s.StationID\n  AND k.Vorname = 'Clara'\n"
                     "  AND k.Nachname = 'Rossi'\n  AND a.Ende IS NULL;")
SQL_ABI22 = ("SELECT a.AusleihID, s1.Name AS Startstation, s2.Name AS Zielstation\nFROM Ausleihe a, Station s1, Station s2\n"
             "WHERE a.StartStation = s1.StationID\n  AND a.ZielStation = s2.StationID\n  AND s1.Stadtteil = 'Hafenviertel'\n"
             "  AND s2.Stadtteil != 'Hafenviertel'\nORDER BY a.AusleihID;")
SQL_KONTO = ("SELECT a.KundenID, a.AusleihID, a.RadID, a.Start\nFROM Ausleihe a\nWHERE a.Ende IS NULL\n"
             "ORDER BY a.KundenID, a.Start;")
SQL_KONTO_VERBUND = ("SELECT DISTINCT a1.KundenID\nFROM Ausleihe a1, Ausleihe a2\nWHERE a1.KundenID = a2.KundenID\n"
                     "  AND a1.AusleihID != a2.AusleihID\n  AND a1.Ende IS NULL\n  AND a2.Ende IS NULL;")


def NULLZEILEN():
    """Ausschnitt aus Ausleihe: eine beendete und zwei laufende Ausleihen (aus der DB)."""
    return abfrage_ergebnis("SELECT AusleihID, StartStation, ZielStation, Start, Ende\nFROM Ausleihe\n"
                            "WHERE AusleihID IN (287, 465, 466)\nORDER BY AusleihID;", DB)[1]


def zeilen_zusammenhalten(d, tab):
    """Schülerfassung: Auftragszeilen nicht über Seiten teilen (Text und Schreiblinien bleiben zusammen)."""
    if d.fassung == "schueler":
        for zeile in tab.rows:
            ad._zeile_nicht_teilen(zeile)


# ================================================================== AB 2.1
def ab21(fassung):
    d = AktenDokument("ab", 2, "Der rote Faden", fassung=fassung, ermittlung="2.1", ds=3, material="AB 2.1",
                      datenbank=DB, untertitel="Fremdschlüssel, Kreuzprodukt und Verbund über zwei Tabellen")
    d.kopfband()
    d.fall_einstieg("Bent Hollwedel, Betriebsleiter von StadtRad Weserstedt, vermisst fünf Lastenräder. Er hat der "
                    "Detektei seine Datenbank geöffnet – aber die Informationen sind auf mehrere Tabellen verteilt. Wer "
                    "wissen will, *wo* ein Rad steht oder *wer* es ausgeliehen hat, muss Tabellen miteinander verbinden.")
    d.beweisstueck(1, "Das StadtRad-Schema",
                   "Fünf Tabellen, in Kurzschreibweise: Primärschlüssel **unterstrichen**, Fremdschlüssel mit **↑** davor.",
                   *SCHEMA_ALLE,
                   "Jeder Pfeil zeigt vom **Fremdschlüssel** zum **Primärschlüssel**, auf den er verweist:",
                   schema_bild("schema_ab21", breite_cm=15.5))
    d.beweisstueck(2, "Fremdschlüssel mit anderem Namen",
                   "Ein Fremdschlüssel muss nicht so heißen wie der Primärschlüssel, auf den er verweist. Die Ausleihe "
                   "verweist gleich **zweimal** auf Station – über zwei Attribute mit eigenen Namen:",
                   Tabelle(["Fremdschlüssel", "verweist auf", "Bedeutung"],
                           [["Rad.StationID", "Station.StationID", "Station, an der das Rad gerade steht"],
                            ["Ausleihe.StartStation", "Station.StationID", "Station, an der die Fahrt begann"],
                            ["Ausleihe.ZielStation", "Station.StationID", "Station, an der die Fahrt endete"]],
                           breiten=[4.6, 4.0, 7.6]),
                   Notiz("Ein Fremdschlüssel darf leer sein",
                         "`Rad.StationID` ist NULL, wenn ein Rad in der Werkstatt, unterwegs oder vermisst ist; "
                         "`Ausleihe.ZielStation` ist NULL, solange eine Fahrt noch läuft."))
    kreuz = []
    stationen = {s[0]: s for s in UP_STATIONEN[:2]}
    for r in (UP_RAEDER[0], UP_RAEDER[1], UP_RAEDER[5]):
        for sid, sname, _ in stationen.values():
            passt = r[2] == sid
            kreuz.append([str(r[0]), r[1], "NULL" if r[2] is None else str(r[2]), str(sid), sname,
                          "✔ bleibt" if passt else "✘ fällt weg"])
    d.beweisstueck(3, "Kreuzprodukt und Verbund",
                   "`FROM Rad, Station` verbindet **jede** Zeile von Rad mit **jeder** Zeile von Station – das "
                   "**Kreuzprodukt**. Im Unplugged war jede Verbindung ein roter Faden. Ausschnitt mit 3 Rädern und "
                   "2 Stationen (eigene Beispieldaten):",
                   Tabelle(["Rad.RadID", "Rad.Typ", "Rad.StationID", "Station.StationID", "Station.Name", "Faden"],
                           kreuz, breiten=[2.2, 2.4, 2.9, 3.2, 2.9, 2.6]),
                   "**Zeilenzahl des Kreuzprodukts = Zeilen der einen Tabelle · Zeilen der anderen Tabelle** "
                   "(hier 3 · 2 = 6). Erst die **Verbundbedingung** `WHERE Rad.StationID = Station.StationID` lässt nur "
                   "die Fäden stehen, bei denen Fremdschlüssel und Primärschlüssel übereinstimmen. Rad 6 (StationID NULL) "
                   "behält keinen Faden: NULL ist mit nichts gleich.")
    d.beweisstueck(4, "Das Verbund-Rezept",
                   Liste(["Alle beteiligten Tabellen in **FROM** (durch Komma getrennt).",
                          "Je Tabellenpaar eine **Verbundbedingung** FK = PK in WHERE (n Tabellen → mindestens n − 1 "
                          "Bedingungen).",
                          "Gleichnamige Spalten mit **Tabellenpräfix** (`Rad.StationID`).",
                          "Erst dann weitere Bedingungen mit **AND** anhängen."], nummeriert=True),
                   "Beispiel: Welche Räder stehen gerade an einer Station im Hafenviertel?",
                   SQL(SQL_HAFEN), ergebnis(SQL_HAFEN),
                   "Im ganzen Hafenviertel steht nur noch ein einziges Rad – Hollwedel wird das interessieren.")

    auftraege = d.ermittlungsauftraege([
        Auftrag(1, 1, "PAPIER", "**Erläutere** die Kennzeichnungen in der Relation Reparatur.", kompetenzen=["K04", "K05"],
                inhalt=[Schema(REPARATUR)], linien=4,
                loesung=Loesung("`ReparaturID` ist unterstrichen: **Primärschlüssel** – identifiziert jede Reparatur "
                                "eindeutig. `↑RadID` ist ein **Fremdschlüssel**: verweist auf den Primärschlüssel `RadID` "
                                "der Tabelle Rad, also auf das reparierte Rad. Datum, Grund und Kosten sind gewöhnliche "
                                "Attribute. Folge: Ein Rad kann viele Reparaturen haben, jede Reparatur gehört zu genau "
                                "einem Rad.",
                                fehler=[("P8", "„RadID ist der Primärschlüssel von Reparatur“ – dann hätte jedes Rad "
                                                "höchstens eine Reparatur."),
                                        ("P7", "Pfeilrichtung verdreht: „Rad verweist auf Reparatur“ – der Pfeil zeigt "
                                               "immer zum Primärschlüssel.")])),
        Auftrag(2, 1, "PAPIER", "Die Tabelle Kunde hat 80, die Tabelle Ausleihe 469 Datensätze. **Bestimme**, wie viele "
                "Zeilen die Abfrage ohne WHERE liefert – und wie viele mit der Verbundbedingung "
                "`Ausleihe.KundenID = Kunde.KundenID` übrig bleiben. Begründe kurz.", kompetenzen=["K23", "K29"],
                inhalt=[SQL("SELECT *\nFROM Kunde, Ausleihe;", pruefung="fehlerhaft")], linien=3,
                loesung=Loesung("Ohne WHERE: Kreuzprodukt 80 · 469 = **37 520** Zeilen. Mit Verbundbedingung: **469** "
                                "Zeilen – jede Ausleihe hat genau eine Kundin bzw. einen Kunden (KundenID ist nie NULL), "
                                "also bleibt von ihren 80 Fäden genau einer.",
                                fehler=[("P2", "80 + 469 = 549 (addiert statt multipliziert) oder „80 Zeilen, eine je "
                                               "Kunde“ – der Verbund liefert eine Zeile je **Ausleihe**.")])),
        Auftrag(3, 2, "PAPIER", "Hollwedel möchte zu jeder Ausleihe den **Namen der Startstation** sehen. **Notiere** die "
                "passende Verbundbedingung und begründe, warum du Tabellenpräfixe verwendest, obwohl `StartStation` "
                "nur in Ausleihe vorkommt.", kompetenzen=["K03", "K16"], linien=4,
                loesung=Loesung("Verbundbedingung: `Ausleihe.StartStation = Station.StationID`. Der Fremdschlüssel heißt "
                                "anders als der Primärschlüssel, auf den er zeigt. Die Präfixe sind hier nicht zwingend, "
                                "machen den Verweis aber eindeutig lesbar; Pflicht sind sie bei gleichnamigen Spalten "
                                "(z. B. `StationID` in Rad und Station). Vollständige Abfrage:",
                                SQL(SQL_START), ergebnis(SQL_START, max_zeilen=4),
                                fehler=[("P7", "`Ausleihe.StationID = Station.StationID` – eine Spalte StationID gibt es "
                                               "in Ausleihe nicht."),
                                        ("P7", "`Ausleihe.ZielStation = …` – verbindet mit der falschen Station.")])),
        Auftrag(4, 2, "LAB", "Teure Reparaturen: Gib für alle Reparaturen mit Kosten über 100 € das Datum, den Grund, die "
                "Kosten und den **Typ des reparierten Rades** aus, nach Datum sortiert. Prüfe deine Abfrage in der "
                "Konsole und notiere sie.", kompetenzen=["K23"], linien=5,
                loesung=Loesung(SQL(SQL_REP100), ergebnis(SQL_REP100),
                                fehler=[("P2", "Verbundbedingung vergessen → jede teure Reparatur erscheint 45-mal "
                                               "(einmal je Rad)."),
                                        ("P7", "`WHERE RadID = RadID` ohne Präfix → „mehrdeutig“.")])),
        Auftrag(5, 3, "PAPIER", "Juri hat notiert, wo die Räder stehen – die Konsole meldet einen Fehler. **Erkläre**, "
                "warum die Abfrage nicht funktioniert, und **korrigiere** sie.", kompetenzen=["K16", "K23"],
                inhalt=[SQL(SQL_FALSCH_PRAEFIX, pruefung="fehlerhaft")], linien=5,
                loesung=Loesung("`StationID` gibt es in Rad **und** in Station – ohne Präfix ist der Name mehrdeutig, "
                                "die Datenbank weiß nicht, welche Spalte gemeint ist. Selbst mit gleichem Präfix auf "
                                "beiden Seiten (`Station.StationID = Station.StationID`) wäre die Bedingung immer wahr "
                                "→ Kreuzprodukt. Gemeint ist: Fremdschlüssel von Rad = Primärschlüssel von Station.",
                                SQL(SQL_PRAEFIX_RICHTIG), ergebnis(SQL_PRAEFIX_RICHTIG, max_zeilen=5),
                                "26 Zeilen: nur Räder, die an einer Station stehen.",
                                fehler=[("P7", "Nur ein Präfix ergänzt (`Rad.StationID = StationID`) – die rechte Seite "
                                               "bleibt mehrdeutig."),
                                        ("P2", "WHERE-Zeile gestrichen, um den Fehler „loszuwerden“ → Kreuzprodukt.")])),
    ])
    zeilen_zusammenhalten(d, auftraege)
    d.abi_ecke("Gegeben ist die folgende Relation aus der Datenbank eines Fahrradverleihs. **Erläutern** Sie die "
               "Kennzeichnungen in dieser Relation.", be=3, operator="erläutern", afb="I/II",
               kompetenzen=["K04", "K05", "K03"], inhalt=[Schema(AUSLEIHE)], linien=5,
               loesung=Loesung("Kernpunkte (je 1 BE):",
                               Liste(["`AusleihID` ist unterstrichen: Primärschlüssel – identifiziert jede Ausleihe "
                                      "eindeutig.",
                                      "Die mit ↑ gekennzeichneten Attribute sind Fremdschlüssel: `RadID` verweist auf "
                                      "den Primärschlüssel der Relation Rad, `KundenID` auf den der Relation Kunde.",
                                      "`StartStation` und `ZielStation` verweisen beide auf den Primärschlüssel "
                                      "`StationID` der Relation Station (anderer Name, gleicher Verweis); "
                                      "`ZielStation` ist NULL, solange die Fahrt läuft."]),
                               fehler=[("P7", "„StartStation verweist auf eine Tabelle StartStation“ – Name des FK ≠ "
                                              "Name der Zieltabelle."),
                                       ("P8", "„Alle unterstrichenen und markierten Attribute bilden zusammen den "
                                              "Primärschlüssel.“")]))
    d.spur_im_lab("2.1 „Der rote Faden“", "A2-E1-01 … A2-E1-08 · Cold Case A2-E1-CC",
                  werkzeug="Spurenbrett (Einstieg A2-E1-03)",
                  text="Starte mit dem **Spurenbrett** (A2-E1-03): Dort siehst du jeden Faden des Kreuzprodukts und welche "
                       "übrig bleiben. Dann Kurzschreibweise (01, 02, 04) und SQL (05–08).")
    d.cold_case("Doppelt verbunden", "Wie viele Zeilen liefert diese Abfrage auf der StadtRad-Datenbank? Vermute zuerst, "
                "prüfe dann in der Konsole. **Erkläre**, was die Abfrage eigentlich zeigt.",
                SQL(SQL_DOPPELT, pruefung="fehlerhaft"), kompetenzen=["K23", "K17"], linien=4,
                loesung=Loesung(f"**{anzahl(SQL_DOPPELT)} Zeilen.** Jede Ausleihe erscheint einmal für jede Station, die sie berührt: zweimal, wenn Start und "
                                "Ziel verschieden sind, einmal bei Rundfahrten (7) und bei laufenden Ausleihen ohne Ziel "
                                "(14): 2 · 469 − 7 − 14 = 917. Die Abfrage zeigt also „Ausleihe berührt Station“ – ohne "
                                "zu sagen, ob als Start oder als Ziel."))
    pfad = (ZIEL / "5_Loesungen" / "DD_A2_AB1_Der_rote_Faden_LSG.docx" if fassung == "loesung"
            else ZIEL / "2_Arbeitsblaetter" / "DD_A2_AB1_Der_rote_Faden.docx")
    return d.speichern(pfad)


# ================================================================== AB 2.2
def ab22(fassung):
    d = AktenDokument("ab", 2, "Nachtfahrten", fassung=fassung, ermittlung="2.2", ds=4, material="AB 2.2",
                      datenbank=DB, untertitel="Verbund über drei und vier Tabellen, Tabellenalias, NULL im Verbund")
    d.kopfband()
    d.fall_einstieg("Fünf Nächte, vom 8. bis 12. September, jeweils zwischen 01:00 und 02:00 Uhr: An der Station "
                    "Hafenkante startet eine Lastenrad-Ausleihe, die nie endet. Edda Lorenz legt den Ausdruck auf den "
                    "Tisch: „Folgt dem Faden durch alle Tabellen.“")
    d.beweisstueck(1, "Der Ermittlungsweg durchs Schema",
                   Liste(["**Frage** genau lesen: Was soll ausgegeben werden, welche Bedingung gilt?",
                          "**Attribute** sammeln: Welche Spalten brauche ich (Ausgabe *und* Bedingungen)?",
                          "**Tabellen** bestimmen, in denen diese Attribute stehen – plus die Tabellen *dazwischen*, die "
                          "sie verbinden.",
                          "**Verbundbedingungen** entlang der Pfeile notieren (n Tabellen → n − 1 Bedingungen).",
                          "**Filter** und **Sortierung** ergänzen."], nummeriert=True),
                   "Beispiel: *Vor- und Nachname aller Kundinnen und Kunden, die ein E-Bike ausgeliehen haben.* Name steht "
                   "in Kunde, Typ in Rad – verbunden sind sie nur über Ausleihe. Weg: Kunde ← Ausleihe → Rad.",
                   schema_bild("schema_ab22", breite_cm=14.5, markiert=("Kunde", "Ausleihe", "Rad"),
                               pfeile={("Ausleihe", "KundenID"), ("Ausleihe", "RadID")}))
    d.beweisstueck(2, "Tabellenalias",
                   "In FROM bekommt jede Tabelle einen kurzen Zweitnamen – **ohne AS**, wie in den Abituraufgaben: "
                   "`FROM Ausleihe a, Kunde k, Rad r`. Ab dann heißt die Tabelle in der ganzen Abfrage nur noch so "
                   "(`k.Nachname`, nicht `Kunde.Nachname`). Das Beispiel aus Beweisstück 1:",
                   SQL(SQL_EBIKE), ergebnis(SQL_EBIKE, max_zeilen=3),
                   Notiz("Zwei Arten von Alias",
                         "**Tabellenalias** in FROM ohne AS: `Station s1`. **Spaltenalias** in SELECT mit AS: "
                         "`s1.Name AS Startort` – er benennt nur die Spalte im Ergebnis."))
    d.beweisstueck(3, "Zwei Fremdschlüssel auf dieselbe Tabelle",
                   "Ausleihe verweist mit `StartStation` **und** `ZielStation` auf Station. Für beide Stationsnamen in "
                   "einer Zeile brauchst du Station **zweimal** – mit zwei Aliasen, als lägen zwei Kopien des "
                   "Stationsverzeichnisses auf dem Tisch. Alle Fahrten vom 1. September:",
                   SQL(SQL_VONNACH), ergebnis(SQL_VONNACH, max_zeilen=3),
                   "Mit nur **einer** Station-Tabelle (`a.StartStation = s.StationID AND a.ZielStation = s.StationID`) "
                   "müsste dieselbe Station Start **und** Ziel sein – übrig blieben nur Rundfahrten.")
    nullzeilen = NULLZEILEN()
    d.beweisstueck(4, "NULL im Verbund",
                   "Eine laufende Ausleihe hat noch kein `Ende` und keine `ZielStation` (NULL). Ausschnitt aus Ausleihe:",
                   Tabelle(["AusleihID", "StartStation", "ZielStation", "Start", "Ende"], nullzeilen,
                           pk=("AusleihID",), fk=("StartStation", "ZielStation"), breiten=[2.6, 3.0, 3.0, 3.8, 3.8]),
                   Liste(["Die Bedingung `a.ZielStation = s2.StationID` ist für NULL **nie wahr** – laufende Ausleihen "
                          "fallen beim Verbund über die Zielstation **heraus**, ohne Fehlermeldung.",
                          "Laufende Ausleihen findest du mit `a.Ende IS NULL` – **nicht** mit `= NULL` (das liefert "
                          "keine einzige Zeile)."]))

    auftraege = d.ermittlungsauftraege([
        Auftrag(1, 1, "PAPIER", "**Beschreibe** den Ermittlungsweg für die Frage: *„Vor- und Nachname aller Kundinnen und "
                "Kunden, die ein Rad an der Station Fischmarkt ausgeliehen haben.“* Nenne die Tabellen, die "
                "Verbundbedingungen und den Filter.", kompetenzen=["K24", "K30"], linien=4,
                loesung=Loesung("Ausgabe Vorname, Nachname → Kunde; „an der Station Fischmarkt ausgeliehen“ → Name der "
                                "**Startstation** → Station; verbunden über Ausleihe. Tabellen: Kunde, Ausleihe, Station "
                                "(Rad wird nicht gebraucht). Verbundbedingungen: `a.KundenID = k.KundenID` und "
                                "`a.StartStation = s.StationID`; Filter `s.Name = 'Fischmarkt'`; DISTINCT, weil manche "
                                "mehrfach dort gestartet sind.",
                                SQL(SQL_FISCHMARKT), ergebnis(SQL_FISCHMARKT, max_zeilen=4),
                                fehler=[("P2", "Weg über `Rad.StationID` – das ist die Station, an der das Rad **heute** "
                                               "steht, nicht die, an der es ausgeliehen wurde (falscher Pfad).")])),
        Auftrag(2, 1, "PAPIER", "Juri will Nachname, Radtyp, Startstation und Startzeit aller Ausleihen vom 1. September. "
                "In seiner Abfrage fehlt eine Verbundbedingung. **Ergänze** sie.", kompetenzen=["K24"],
                inhalt=[SQL(SQL_LUECKE, pruefung="aus")], linien=1,
                loesung=Loesung("`a.RadID = r.RadID` – vier Tabellen brauchen drei Verbundbedingungen; ohne sie würde "
                                "jede Ausleihe mit allen 45 Rädern kombiniert.",
                                SQL(SQL_LUECKE_LSG), ergebnis(SQL_LUECKE_LSG, max_zeilen=4),
                                fehler=[("P2", "`r.StationID = s.StationID` – verbindet das Rad mit der Station, an der "
                                               "es heute steht; Räder ohne Station fallen heraus, Zuordnung falsch.")])),
        Auftrag(3, 2, "LAB", "Hollwedel lässt die Räder in der Altstadt überprüfen. Gib für alle Reparaturen an Rädern, die "
                "**zurzeit** an einer Station im Stadtteil Altstadt stehen, ReparaturID, Datum, Grund, RadID und den "
                "Stationsnamen aus – nach Datum sortiert.", kompetenzen=["K24"], linien=5,
                loesung=Loesung(SQL(SQL_REP_ALTSTADT), ergebnis(SQL_REP_ALTSTADT),
                                fehler=[("P2", "Weg über Ausleihe (`StartStation`) – liefert Stationen früherer Fahrten, "
                                               "nicht den heutigen Standort; Zeilen vervielfachen sich."),
                                        ("P7", "`Stadtteil` ohne Präfix ist hier eindeutig, `RadID` nicht.")])),
        Auftrag(4, 2, "PAPIER", "**Notiere** eine Abfrage, die für alle Ausleihen der Kundin Clara Rossi die AusleihID, "
                "den Start sowie den Namen der Start- **und** der Zielstation ausgibt.", kompetenzen=["K24", "K16"],
                linien=7,
                loesung=Loesung(SQL(SQL_ROSSI), ergebnis(SQL_ROSSI),
                                fehler=[("P7", "Nur eine Station-Tabelle → Bedingung verlangt Start = Ziel, Ergebnis leer."),
                                        ("P7", "`Station.Name` in SELECT, obwohl in FROM die Aliase s1, s2 vergeben "
                                               "sind → Fehler.")])),
        Auftrag(5, 3, "PAPIER", "Clara Rossi ist gerade mit einem Rad unterwegs – trotzdem zeigt deine Abfrage aus Auftrag 4 "
                "diese Fahrt nicht. **Begründe** das und **notiere** eine Abfrage, die ihre laufenden Ausleihen mit "
                "AusleihID, Start und Name der Startstation ausgibt.", kompetenzen=["K18", "K24"], linien=6,
                loesung=Loesung("Bei einer laufenden Ausleihe ist `ZielStation` NULL. Die Verbundbedingung "
                                "`a.ZielStation = s2.StationID` ist dann für keine Station wahr – die Zeile fällt heraus. "
                                "Ohne Zielstation geht es: Station nur einmal (über StartStation) und `a.Ende IS NULL`.",
                                SQL(SQL_ROSSI_LAUFEND), ergebnis(SQL_ROSSI_LAUFEND),
                                fehler=[("P5", "`a.Ende = NULL` bzw. `a.ZielStation = NULL` → keine Zeile."),
                                        ("P5", "Zielstation weiter verbunden und nur `AND a.Ende IS NULL` ergänzt → leer.")])),
    ])
    zeilen_zusammenhalten(d, auftraege)
    d.abi_ecke("Gegeben ist das Relationenschema des Fahrradverleihs (Beweisstück 1). **Entwerfen** Sie eine "
               "SQL-Abfrage, die für alle Ausleihen, die an einer Station im Stadtteil Hafenviertel begonnen und an einer "
               "Station in einem anderen Stadtteil beendet wurden, die AusleihID sowie die Namen der Start- und der "
               "Zielstation ausgibt.", be=4, operator="entwerfen", afb="II/III", kompetenzen=["K24", "K16", "K17"],
               linien=7,
               loesung=Loesung(SQL(SQL_ABI22), ergebnis(SQL_ABI22, max_zeilen=5),
                               f"({anzahl(SQL_ABI22)} Zeilen; ORDER BY nicht verlangt.) Kernpunkte (je 1 BE): Station "
                               "zweimal mit zwei Aliasen in FROM · beide Verbundbedingungen (Start und Ziel) · Bedingung "
                               "Start im Hafenviertel **und** Ziel in anderem Stadtteil (`!=`) · korrekte Ausgabe "
                               "mit Präfixen.",
                               fehler=[("P7", "Station nur einmal → Start = Ziel."),
                                       ("P6", "`s2.Stadtteil != s1.Stadtteil` allein – erfasst auch Fahrten, die nicht im "
                                              "Hafenviertel begonnen haben (nur mit der Hafenviertel-Bedingung richtig).")],
                               varianten=["Statt `s2.Stadtteil != 'Hafenviertel'` auch `s2.Stadtteil != s1.Stadtteil` "
                                          "(zusammen mit `s1.Stadtteil = 'Hafenviertel'`)."]))
    d.spur_im_lab("2.2 „Nachtfahrten“", "A2-E2-01 … A2-E2-08 · Cold Case A2-E2-CC",
                  werkzeug="Schema-Browser, Zeitlupe",
                  text="Reihenfolge ★ → ★★★; A2-E2-07 führt dich zu den Nachtfahrten. Danach wartet **Verhör 2 · "
                       "Bürgerbus** (20 Minuten, im Abitur-Stil).")
    d.cold_case("Ein Konto, viele Räder", "Hat ein Kundenkonto **mehrere Räder gleichzeitig** ausgeliehen (laufende "
                "Ausleihen)? Finde es heraus und **beschreibe** dein Vorgehen. Ab DS 5 geht das mit GROUP BY eleganter.",
                kompetenzen=["K18", "K30"], linien=3,
                loesung=Loesung("Alle laufenden Ausleihen nach KundenID sortiert ausgeben und nach Wiederholungen suchen:",
                                SQL(SQL_KONTO), ergebnis(SQL_KONTO),
                                "Konto **77** hat fünf laufende Ausleihen (Lastenräder 41–45, die Nachtfahrten). Für "
                                "Schnelle: Ausleihe zweimal mit Aliasen – zwei verschiedene laufende Ausleihen desselben "
                                "Kontos:",
                                SQL(SQL_KONTO_VERBUND), ergebnis(SQL_KONTO_VERBUND)))
    pfad = (ZIEL / "5_Loesungen" / "DD_A2_AB2_Nachtfahrten_LSG.docx" if fassung == "loesung"
            else ZIEL / "2_Arbeitsblaetter" / "DD_A2_AB2_Nachtfahrten.docx")
    return d.speichern(pfad)


# ================================================================== Präsentation
SQL_P_LINDENHOF = ("SELECT Rad.RadID, Rad.Typ, Station.Name\nFROM Rad, Station\nWHERE Rad.StationID = Station.StationID\n"
                   "  AND Station.Stadtteil = 'Lindenhof'\nORDER BY Rad.RadID;")
SQL_P_VONNACH = ("SELECT a.AusleihID, s1.Name AS Von, s2.Name AS Nach\nFROM Ausleihe a, Station s1, Station s2\n"
                 "WHERE a.StartStation = s1.StationID\n  AND a.ZielStation = s2.StationID\n"
                 "  AND a.Start LIKE '2026-09-01%'\nORDER BY a.Start;")
SQL_P_ABI21 = "Rad(<u>RadID</u>, Typ, Baujahr, ↑StationID, Status)"
SQL_P_ABI22 = ("SELECT a.AusleihID, k.Nachname, s.Name AS Zielstation\nFROM Ausleihe a, Kunde k, Rad r, Station s\n"
               "WHERE a.KundenID = k.KundenID\n  AND a.RadID = r.RadID\n  AND a.ZielStation = s.StationID\n"
               "  AND r.Typ = 'Lastenrad'\nORDER BY a.AusleihID;")


def praesentation():
    p = AktenPraesentation(2, "Spur der Räder", datenbank=DB)
    # ---------------------------------------------------------------- DS 3
    p.titelfolie("Spur der Räder", "Akte 2 · DS 3–4 · Fremdschlüssel und Verbund",
                 notizen="DS 3 beginnt mit dem Warm-up (Karten W-A1-05…08, 10 min); Auflösung von Karte 06 (NULL-Falle) "
                         "an der Tafel. Dann Akte 2 im Lab öffnen lassen (Stadtplan: StadtRad-Station).")
    p.briefing("Neuer Auftrag: Der rote Faden",
               "Bent Hollwedel, Betriebsleiter von StadtRad Weserstedt, vermisst fünf Lastenräder – nachts einfach "
               "verschwunden. Seine Daten liegen nicht in einer Tabelle, sondern in mehreren: Stationen, Räder, "
               "Kundschaft, Ausleihen. Wer die Spur finden will, muss sie verbinden.", von="Edda Lorenz",
               notizen="Briefing vorlesen lassen (STORY 5.3b). Impulse: „Wo steht, an welcher Station Rad 17 steht?“ – "
                       "„Warum steht dort nicht der Stationsname?“ (Redundanz vermeiden: Name nur einmal in Station).")
    p.inhaltsfolie("Das StadtRad-Schema", Bild(str(schema_png(TMP / "schema_folie.png")), breite_cm=27),
                   notizen="Schema-Plakat (A3) aufhängen. Gemeinsam die Pfeile lesen: Vom Fremdschlüssel zum "
                           "Primärschlüssel. Frage: „StartStation und ZielStation – worauf verweisen sie, obwohl sie "
                           "anders heißen?“ Die vermissten Räder 41–45 stehen an keiner Station (StationID NULL).")
    p.begriffe("Primärschlüssel und Fremdschlüssel",
               [("Primärschlüssel", "unterstrichen – identifiziert jeden Datensatz eindeutig, nie NULL"),
                ("Fremdschlüssel", "mit ↑ – verweist auf den Primärschlüssel einer anderen Tabelle"),
                ("anderer Name", "`StartStation` verweist auf `Station.StationID`"),
                ("NULL im FK", "kein Verweis: Rad in der Werkstatt, Fahrt läuft noch")],
               inhalt=[Schema(AUSLEIHE)],
               notizen="Begriffe Fremdschlüssel und ↑-Kennzeichnung einführen. Wichtig: Der Pfeil steht vor dem "
                       "Fremdschlüssel und zeigt (gedanklich) zum Primärschlüssel. Ausleihe hat vier Fremdschlüssel.")
    p.beispiel_schritte("Verbund Schritt für Schritt", SQL(SQL_P_LINDENHOF),
                        ["**FROM Rad, Station** – Kreuzprodukt: jedes Rad mit jeder Station (Zeilen · Zeilen)",
                         "**WHERE Rad.StationID = Station.StationID** – nur Fäden, bei denen FK = PK",
                         "**AND Stadtteil = 'Lindenhof'** – erst jetzt weiter filtern",
                         "**SELECT** – RadID, Typ und Stationsname zeigen"],
                        ergebnis=ergebnis(SQL_P_LINDENHOF),
                        notizen="Reihenfolge der Auswertung betonen (FROM → WHERE → SELECT), wie in der Zeitlupe. "
                                "Präfix bei StationID ist Pflicht (gleichnamig in Rad und Station). Zeilenzahl des "
                                "Kreuzprodukts NICHT verraten – sie ist Lab-Aufgabe A2-E1-05.")
    p.unplugged("Roter Faden", [
        "Stationskarten links, Radkarten rechts auf den Tisch legen.",
        "`FROM Rad, Station`: Verbindet **jedes** Rad mit **jeder** Station durch einen Faden. Zählt die Fäden.",
        "`WHERE Rad.StationID = Station.StationID`: Entfernt alle Fäden, bei denen die Nummern nicht übereinstimmen.",
        "Wo ist Rad 6? Warum taucht Deichtor nicht auf? Was passiert ohne WHERE?",
        "Schnelle Teams: `AND Station.Stadtteil = 'Hafenviertel'` – welche Fäden bleiben?"],
        material="Karten + 24 Wollfäden je Team", dauer="15 MIN",
        notizen="Kopiervorlage DD_A2_Unplugged_Roter_Faden (4 Stations-, 6 Radkarten, Auftragskarte). Erwartung: "
                "6 · 4 = 24 Fäden, nach WHERE bleiben 5. Rad 6 (StationID NULL, Werkstatt) behält keinen Faden; Deichtor "
                "hat kein Rad. Zusatz: Räder 1 und 3 (Leuchtturm). Sicherung: Foto nach Schritt 3; Tafel „Verbund = "
                "Kreuzprodukt + Filter FK = PK“.")
    p.lab_auftrag("2.1", "Der rote Faden", "A2-E1-01 … 08 (★ → ★★ → ★★★) · Cold Case A2-E1-CC",
                  werkzeug="Spurenbrett (Einstieg A2-E1-03)",
                  hinweise=["Erst das Spurenbrett, dann Kurzschreibweise, dann SQL.",
                            "AB 2.1 parallel: Aufträge 1–3 und 5 auf Papier, 4 im Lab."],
                  notizen="Herumgehen; typische Fehler sammeln: fehlende Verbundbedingung (P2), StationID ohne Präfix "
                          "(P7), FK mit anderem Namen nicht erkannt (P7). Bei riesigen Ergebnissen: Zeilenzahl = "
                          "Produkt der Tabellengrößen → Spurenbrett.")
    p.merkfolie("Das Verbund-Rezept", [
        "Alle beteiligten Tabellen in **FROM**.",
        "Je Tabellenpaar eine **Verbundbedingung** FK = PK (n Tabellen → n − 1 Bedingungen).",
        "Gleichnamige Spalten mit **Tabellenpräfix**: `Rad.StationID`.",
        "Erst dann weitere Bedingungen mit **AND**.",
        "Ohne Verbundbedingung entsteht das Kreuzprodukt; Zeilen mit NULL im Fremdschlüssel fallen heraus."],
        notizen="Merkkasten ins Heft (Werkzeugkarte 3 im Ermittler-Handbuch).")
    p.abi_format("Gegeben ist die Relation Rad der Datenbank eines Fahrradverleihs. **Erläutern** Sie die "
                 "Kennzeichnungen in dieser Relation.", be=2, operator="erläutern", inhalt=[Schema(SQL_P_ABI21)],
                 notizen="Zwei Minuten Einzelarbeit, dann Partner-Check mit den Kernpunkten (Lösungsfolie). Die "
                         "Abi-Ecke auf AB 2.1 (Ausleihe, 3 BE) als Hausaufgabe bzw. in der Fallbesprechung.")
    p.loesungsfolie("Lösung Abi-Format (Rad)",
                    Liste(["`RadID` unterstrichen: Primärschlüssel, identifiziert jedes Rad eindeutig (1 BE).",
                           "`↑StationID`: Fremdschlüssel, verweist auf den Primärschlüssel StationID der Relation "
                           "Station – die Station, an der das Rad steht; NULL, wenn es an keiner Station steht (1 BE)."]),
                    notizen="Erst nach der Einzelarbeit zeigen. Typischer Fehler P7: Pfeilrichtung verdreht.")
    p.briefing("Cliffhanger", "Fünf Ausleihen, alle gestartet an der Station Hafenkante, alle mitten in der Nacht – "
               "und bei keiner steht eine Ende-Zeit.", von="Juri Brandes",
               notizen="Abschluss DS 3 (STORY 5.3c). Nicht auflösen – in DS 4 geht es weiter.")
    # ---------------------------------------------------------------- DS 4
    p.briefing("Nachtfahrten", "Bent Hollwedel war heute Morgen hier, sichtlich nervös: Die Räder bleiben verschwunden, "
               "die Versicherung stellt Fragen. Wir wissen, wo die Fahrten begannen – aber nicht, wer sie gestartet hat. "
               "Plant euren Weg durchs Schema, bevor ihr tippt.", von="Edda Lorenz",
               notizen="DS 4: Warm-up W-A2-01…04 vorher (10 min). Briefing (STORY 5.4b). Impuls: „Welche Tabellen braucht "
                       "ihr für ‚Name des Kunden + Name der Startstation‘?“")
    p.inhaltsfolie("Der Ermittlungsweg durchs Schema",
                   "Frage → Attribute → Tabellen → Verbundbedingungen → Filter → Sortierung",
                   Bild(str(schema_png(TMP / "schema_weg.png", markiert=("Kunde", "Ausleihe", "Station"),
                                       pfeile={("Ausleihe", "KundenID"), ("Ausleihe", "StartStation")})), breite_cm=21),
                   notizen="Weg für „Name des Kunden + Name der Startstation“ am Plakat zeigen: Kunde ← Ausleihe → Station. "
                           "Rad wird nicht gebraucht. Drei Tabellen → zwei Verbundbedingungen.")
    p.begriffe("Tabellenalias und Spaltenalias",
               [("Tabellenalias (ohne AS)", "`FROM Ausleihe a, Kunde k` – gilt in der ganzen Abfrage"),
                ("Spaltenalias (mit AS)", "`s1.Name AS Startort` – benennt die Ergebnisspalte"),
                ("Pflicht", "wenn eine Tabelle zweimal vorkommt (`Station s1, Station s2`)")],
               inhalt=["Beispiel: Nachname und Startzeit jeder Ausleihe – mit Aliasen kurz und lesbar.",
                       SQL("SELECT k.Nachname, a.Start\nFROM Ausleihe a, Kunde k\nWHERE a.KundenID = k.KundenID;")],
               notizen="Alias ohne AS wie in den Ergänzenden Hinweisen. Häufiger Fehler: In FROM Alias vergeben, in "
                       "SELECT trotzdem `Station.Name` schreiben → Fehler.")
    p.beispiel_schritte("Zwei Pfeile auf Station", SQL(SQL_P_VONNACH),
                        ["**Station s1, Station s2** – zwei Kopien des Stationsverzeichnisses",
                         "**a.StartStation = s1.StationID** – s1 liefert den Start",
                         "**a.ZielStation = s2.StationID** – s2 liefert das Ziel",
                         "**AS Von / AS Nach** – zwei Namensspalten unterscheidbar machen"],
                        ergebnis=ergebnis(SQL_P_VONNACH, max_zeilen=4),
                        notizen="Unplugged-Erfahrung aufgreifen („Kopie von Station“). Gegenprobe: Mit nur einer "
                                "Station-Tabelle bleiben nur Rundfahrten (Start = Ziel).")
    p.inhaltsfolie("NULL im Verbund",
                   Tabelle(["AusleihID", "StartStation", "ZielStation", "Ende"],
                           [[z[0], z[1], z[2], z[4]] for z in NULLZEILEN()], pk=("AusleihID",),
                           fk=("StartStation", "ZielStation")),
                   Liste(["Verbund über `ZielStation`: Ausleihen 465 und 466 verschwinden – ohne Fehlermeldung.",
                          "Laufende Ausleihen: `Ende IS NULL` – nicht `= NULL`."]),
                   notizen="Frage an den Kurs: „Was passiert mit einer nie beendeten Fahrt, wenn ihr ihre Zielstation "
                           "verbinden wollt?“ Erst vermuten lassen, dann auflösen.")
    p.unplugged("Ermittlungsweg durchs Schema", [
        "Fragekarte ziehen. Nur die benötigten Tabellenkarten in eine Reihe legen.",
        "Faden über die Verweise spannen; an jeden Übergang die Verbundbedingung schreiben.",
        "Braucht ihr eine Tabelle zweimal? „Kopie von …“-Karte anlegen und Aliase vergeben.",
        "Karte 3: Es gibt zwei Wege zu Station – welcher passt zur Frage?",
        "Karte 4 (★★★): Was passiert mit Fahrten, die noch unterwegs sind?"],
        material="Tabellen-, Kopie- und Fragekarten, Faden", dauer="15 MIN",
        notizen="Kopiervorlage DD_A2_Unplugged_Ermittlungsweg. Lösungen der Fragekarten in den Lehrerhinweisen. "
                "Sicherung: Teams fotografieren ihre Wege; Lösung zu Karte 2 als SQL an der Tafel.")
    p.lab_auftrag("2.2", "Nachtfahrten", "A2-E2-01 … 08 (★ → ★★★) · Cold Case A2-E2-CC",
                  werkzeug="Schema-Browser, Zeitlupe",
                  hinweise=["07/08 = die Nachtfahrten des verdächtigen Kontos.",
                            "Danach: Verhör 2 · Bürgerbus (20 Minuten, Abitur-Stil)."],
                  notizen="Verhör 2 im Kurs auf „auto“ oder „offen“ stellen. Typische Fehler: eine Station-Tabelle für "
                          "Start und Ziel (P7), falscher Pfad über Rad.StationID (P2), laufende Ausleihen verschwinden "
                          "(P5, Prüf-Variante +null).")
    p.merkfolie("Ermittler-Notiz: Mehrere Tabellen", [
        "Ermittlungsweg: Frage → Attribute → Tabellen → Verbundbedingungen → Filter → Sortierung.",
        "Zwei Fremdschlüssel auf dieselbe Tabelle → Tabelle zweimal in FROM, zwei Aliase (ohne AS).",
        "Laufende Ausleihen: `Ende IS NULL` – ein Verbund über ZielStation lässt sie verschwinden.",
        "Vier Tabellen → mindestens drei Verbundbedingungen."],
        notizen="Merkkasten ins Heft. Verweis auf Werkzeugkarte 3.")
    p.abi_format("Gegeben ist das Relationenschema des Fahrradverleihs (Schema-Plakat). **Entwerfen** Sie eine SQL-Abfrage, die für "
                 "alle Ausleihen von Lastenrädern die AusleihID, den Nachnamen der Kundin bzw. des Kunden und den Namen "
                 "der Zielstation ausgibt.", be=4, operator="entwerfen",
                 notizen="Einzelarbeit 4 min, Partner-Check mit den Kernpunkten. Die Abi-Ecke auf AB 2.2 (Hafenviertel) "
                         "ist die schwerere Variante.")
    p.loesungsfolie("Lösung Abi-Format (Lastenräder)", SQL(SQL_P_ABI22), ergebnis(SQL_P_ABI22, max_zeilen=1),
                    notizen="Kernpunkte je 1 BE: vier Tabellen in FROM · drei Verbundbedingungen (Zielstation über "
                            "ZielStation!) · Filter Lastenrad · Ausgabe. Hinweis: Laufende Ausleihen fehlen, weil "
                            "ZielStation NULL ist – das ist hier richtig, sie haben keine Zielstation.")
    p.briefing("Das Konto 77", "Das Kundenkonto 77 gehört einer „Nora Falter“, Jahrgang 1988 – N. F. wie im Fundbuch. "
               "Und während wir noch staunen, meldet der Festausschuss: In der Stadtfest-Kasse fehlen 312,50 €.",
               von="Edda Lorenz", nur_lehrkraft=True,
               notizen="Story-Auflösung Akte 2 (STORY 5.4c) erst nach A2-E2-07 zeigen. Beweisstück 2 „Schloss-Code“ "
                       "erscheint im Lab. Steckbrief ergänzen: Tarnname Nora Falter, Jahrgang 1988 (unbestätigt), "
                       "nachts unterwegs. Cliffhanger zu Akte 3.")
    return p.speichern(ZIEL / "1_Praesentation" / "DD_A2_Raeder.pptx")


# ================================================================== Unplugged-Kopiervorlagen
def _karten_raster(d, spalten, hoehe_cm, karten):
    """Karten als Tabelle mit gestrichelten Schnittlinien; karten = Funktionen f(zelle, breite_cm)."""
    breite = d.breite / spalten
    zeilen = (len(karten) + spalten - 1) // spalten
    tab = d.doc.add_table(rows=zeilen, cols=spalten)
    ad._tabelle_breiten(tab, [breite] * spalten)
    ad._tabelle_ohne_rahmen(tab)
    for zi in range(zeilen):
        zeile = tab.rows[zi]
        trpr = zeile._tr.get_or_add_trPr()
        trpr.append(ad._el("w:trHeight", val=int(Cm(hoehe_cm) / 635), hRule="exact"))
        ad._zeile_nicht_teilen(zeile)
        for si in range(spalten):
            z = zeile.cells[si]
            ad._zelle_rahmen(z, ad.RAND_STARK, 0.75, art="dashed")
            ad._zelle_innenabstand(z, 0.45, 0.5, 0.3, 0.5)
            z.paragraphs[0].style = d.doc.styles["DD Tabellentext"]
            i = zi * spalten + si
            if i < len(karten):
                karten[i](z, breite)
    d._nach_kasten(d.doc)
    return tab


def _kartenkopf(d, z, etikett, titel, groesse=30, farbe=ad.SIEGELROT):
    p = z.paragraphs[0]
    d._etikett(p, etikett, groesse=9, farbe=farbe, fett=True)
    p = z.add_paragraph(style="DD Tabellentext")
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(8)
    d._lauf_stil(p.add_run(titel), schrift=ad.SCHRIFT_KOPF, groesse=groesse)


def _wert(d, z, name, wert, art="", groesse=18):
    """Zeile „Attribut: Wert“ in großer Schrift; art pk (unterstrichen) / fk (↑, rot)."""
    p = z.add_paragraph(style="DD Tabellentext")
    p.paragraph_format.space_after = Pt(6)
    if art == "pk":
        d._lauf_stil(p.add_run(name), code=True, unter=True, fett=True, groesse=groesse * 0.75)
    elif art == "fk":
        d._lauf_stil(p.add_run("↑" + name), code=True, fett=True, groesse=groesse * 0.75, farbe=ad.SIEGELROT)
    else:
        d._lauf_stil(p.add_run(name), code=True, groesse=groesse * 0.75, farbe=ad.TINTE_2)
    d._lauf_stil(p.add_run("   " + wert), groesse=groesse, fett=True,
                 farbe=ad.TINTE_3 if wert == "NULL" else ad.TINTE, kursiv=wert == "NULL")


def unplugged_roter_faden():
    d = AktenDokument("unplugged", 2, "Roter Faden", fassung="schueler", ermittlung="2.1", ds=3,
                      material="Unplugged 2.1", untertitel="Kopiervorlage: je Team 4 Stationskarten, 6 Radkarten, "
                                                           "1 Auftragskarte, 24 rote Wollfäden (50 cm)")
    d.kopfband(name_datum=False)
    d.inhalt("**So geht’s:** Seiten 1–3 auf festes Papier (160 g) kopieren und an den gestrichelten Linien schneiden. "
             "Eigene Beispieldaten – nicht die Lab-Datenbank. Wollfäden vorher auf 50 cm schneiden.")

    def station(s):
        def f(z, b):
            _kartenkopf(d, z, "TABELLE STATION", s[1], groesse=34)
            _wert(d, z, "StationID", str(s[0]), "pk", groesse=26)
            _wert(d, z, "Name", s[1])
            _wert(d, z, "Stadtteil", s[2])
        return f

    def rad(r):
        def f(z, b):
            _kartenkopf(d, z, "TABELLE RAD", f"Rad {r[0]}", groesse=30, farbe=ad.BLAU)
            _wert(d, z, "RadID", str(r[0]), "pk", groesse=22)
            _wert(d, z, "Typ", r[1])
            _wert(d, z, "StationID", "NULL" if r[2] is None else str(r[2]), "fk", groesse=22)
            if r[2] is None:
                p = z.add_paragraph(style="DD Tabellentext")
                d._text(p, "*in der Werkstatt*", groesse=12, farbe=ad.TINTE_2)
        return f

    d.ueberschrift("Stationskarten (links auf den Tisch)", 2)
    _karten_raster(d, 2, 9.6, [station(s) for s in UP_STATIONEN])
    d.seitenumbruch()
    d.ueberschrift("Radkarten (rechts auf den Tisch)", 2)
    _karten_raster(d, 2, 7.9, [rad(r) for r in UP_RAEDER])
    d.seitenumbruch()
    d.ueberschrift("Auftragskarte und Protokoll", 2)

    def auftrag(z, b):
        _kartenkopf(d, z, "AUFTRAGSKARTE · ROTER FADEN", "Vom Kreuzprodukt zum Verbund", groesse=22, farbe=ad.GRUEN)
        schritte = [
            "Legt die Stationskarten links, die Radkarten rechts auf den Tisch.",
            "`FROM Rad, Station`: Verbindet **jedes** Rad mit **jeder** Station durch einen roten Faden. Zählt die Fäden.",
            "`WHERE Rad.StationID = Station.StationID`: Entfernt jeden Faden, bei dem die beiden Nummern nicht "
            "übereinstimmen. Zählt erneut.",
            "Beantwortet die Fragen im Protokoll.",
            "Schnelle Teams: Zusätzlich `AND Station.Stadtteil = 'Hafenviertel'` – welche Fäden bleiben?"]
        for i, t in enumerate(schritte, 1):
            p = z.add_paragraph(style="DD Tabellentext")
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(5)
            d._lauf_stil(p.add_run(f"{i}.  "), fett=True, groesse=12.5)
            d._text(p, t, groesse=12.5)

    _karten_raster(d, 1, 8.6, [auftrag])
    d.inhalt("**Protokoll des Teams** ____________________________________")
    d.tabelle(["Frage", "Unsere Antwort"],
              [["Fäden nach FROM (Kreuzprodukt)", ""], ["Fäden nach WHERE (Verbund)", ""],
               ["Welches Rad hat keinen Faden mehr? Warum?", ""], ["Welche Station hat keinen Faden mehr? Warum?", ""],
               ["Was zeigt das Ergebnis ohne WHERE – und warum ist es Unsinn?", ""],
               ["Zusatz: Fäden mit Stadtteil = 'Hafenviertel'", ""]],
              breiten=[8.0, d.breite - 8.0], groesse=10.5)
    for zeile in d.doc.tables[-1].rows[1:]:
        trpr = zeile._tr.get_or_add_trPr()
        trpr.append(ad._el("w:trHeight", val=int(Cm(1.25) / 635), hRule="atLeast"))
    return d.speichern(ZIEL / "3_Unplugged" / "DD_A2_Unplugged_Roter_Faden.docx")


UP_TABELLEN = [
    ("Station", [("StationID", "pk"), ("Name", ""), ("Stadtteil", ""), ("Stellplaetze", "")]),
    ("Rad", [("RadID", "pk"), ("Typ", ""), ("Baujahr", ""), ("StationID", "fk"), ("Status", "")]),
    ("Kunde", [("KundenID", "pk"), ("Vorname", ""), ("Nachname", ""), ("Jahrgang", ""), ("Tarif", ""), ("Stadtteil", "")]),
    ("Ausleihe", [("AusleihID", "pk"), ("RadID", "fk"), ("KundenID", "fk"), ("StartStation", "fk"), ("ZielStation", "fk"),
                  ("Start", ""), ("Ende", ""), ("DauerMin", ""), ("Preis", "")]),
    ("Reparatur", [("ReparaturID", "pk"), ("RadID", "fk"), ("Datum", ""), ("Grund", ""), ("Kosten", "")]),
]
UP_FRAGEN = [
    (1, "★", "Vor- und Nachname aller Kundinnen und Kunden, die ein **Lastenrad** ausgeliehen haben."),
    (2, "★★", "Name der **Start- und** der **Zielstation** jeder Ausleihe."),
    (3, "★★", "Welche Kundinnen und Kunden haben Räder ausgeliehen, die **heute** an der Station *Deichtor* stehen? "
              "Vorsicht: Es gibt zwei Wege zu Station."),
    (4, "★★★", "Wo sind die Räder gerade, die **noch unterwegs** sind? Was verrät die Datenbank – und was nicht?"),
]


def unplugged_ermittlungsweg():
    d = AktenDokument("unplugged", 2, "Ermittlungsweg durchs Schema", fassung="schueler", ermittlung="2.2", ds=4,
                      material="Unplugged 2.2", untertitel="Kopiervorlage: je Team 5 Tabellenkarten, 2 Kopie-Karten, "
                                                           "4 Fragekarten, ein langer Faden, Folienstift")
    d.kopfband(name_datum=False)
    d.inhalt("**So geht’s:** Tabellenkarten (A5) und Kopie-Karten auf festes Papier kopieren, an der gestrichelten Linie "
             "halbieren, möglichst laminieren (dann mit Folienstift beschreibbar). Primärschlüssel sind unterstrichen, "
             "Fremdschlüssel rot mit ↑.")

    def tabelle(name, attrs):
        def f(z, b):
            _kartenkopf(d, z, "TABELLENKARTE", name, groesse=40, farbe=ad.TINTE_2)
            for a, art in attrs:
                p = z.add_paragraph(style="DD Tabellentext")
                p.paragraph_format.space_after = Pt(3)
                g = 20 if len(attrs) <= 6 else 17
                if art == "pk":
                    d._lauf_stil(p.add_run(a), code=True, unter=True, fett=True, groesse=g)
                elif art == "fk":
                    d._lauf_stil(p.add_run("↑" + a), code=True, fett=True, groesse=g, farbe=ad.SIEGELROT)
                else:
                    d._lauf_stil(p.add_run(a), code=True, groesse=g)
        return f

    def kopie(z, b):
        _kartenkopf(d, z, "TABELLENKARTE · KOPIE", "Kopie von ______________", groesse=34, farbe=ad.TINTE_2)
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(18)
        d._lauf_stil(p.add_run("Alias in FROM:  ________"), code=True, groesse=20)
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(18)
        d._text(p, "Braucht ihr eine Tabelle zweimal, weil **zwei Fremdschlüssel** auf sie verweisen? Dann legt diese "
                   "Karte als zweite Kopie dazu und gebt beiden Karten verschiedene Aliase.", groesse=14)

    karten = [tabelle(n, a) for n, a in UP_TABELLEN] + [kopie, kopie]
    # Seite 1: Kopf + eine Karte; danach je Seite zwei A5-Karten
    gruppen = [karten[:1]] + [karten[i:i + 2] for i in range(1, len(karten), 2)]
    for i, gruppe in enumerate(gruppen):
        if i:
            d.seitenumbruch()
        _karten_raster(d, 1, 12.0, gruppe)
    d.seitenumbruch()
    d.ueberschrift("Fragekarten", 2)

    def frage(nr, sterne, text):
        def f(z, b):
            _kartenkopf(d, z, f"FRAGEKARTE {nr}  {sterne}", "Plant den Ermittlungsweg", groesse=22, farbe=ad.SIEGELROT)
            p = z.add_paragraph(style="DD Tabellentext")
            d._text(p, text, groesse=15)
            p = z.add_paragraph(style="DD Tabellentext")
            p.paragraph_format.space_before = Pt(10)
            d._text(p, "Welche Karten? Welche Verbundbedingungen? Welcher Filter?", groesse=11, farbe=ad.TINTE_2)
        return f

    _karten_raster(d, 2, 11.6, [frage(*f) for f in UP_FRAGEN])
    return d.speichern(ZIEL / "3_Unplugged" / "DD_A2_Unplugged_Ermittlungsweg.docx")


# ================================================================== Lehrerhinweise
SQL_LH_FRAGE1 = ("SELECT DISTINCT k.Vorname, k.Nachname\nFROM Kunde k, Ausleihe a, Rad r\nWHERE a.KundenID = k.KundenID\n"
                 "  AND a.RadID = r.RadID\n  AND r.Typ = 'Lastenrad'\nORDER BY k.Nachname, k.Vorname;")
SQL_LH_FRAGE3 = ("SELECT DISTINCT k.Vorname, k.Nachname\nFROM Kunde k, Ausleihe a, Rad r, Station s\n"
                 "WHERE a.KundenID = k.KundenID\n  AND a.RadID = r.RadID\n  AND r.StationID = s.StationID\n"
                 "  AND s.Name = 'Marktplatz'\nORDER BY k.Nachname, k.Vorname;")


def lehrerhinweise():
    d = AktenDokument("lehrerhinweise", 2, "Lehrerhinweise Akte 2", fassung="lehrkraft", material="Lehrerhinweise A2",
                      datenbank=DB, untertitel="Spur der Räder · DS 3–4 · Fremdschlüssel, Kurzschreibweise, Verbund")
    d.kopfband()
    d.ueberschrift("Überblick", 1)
    d.tabelle(["DS", "Thema", "Lab", "Warm-up", "Abi-Bezug"],
              [["3", "Der rote Faden – Fremdschlüssel, Kreuzprodukt, Verbund über zwei Tabellen (K03, K04, K05, K23, K16)",
                "A2-E1-01…08, Cold Case A2-E1-CC (Spurenbrett)", "W-A1-05…08 (LIKE, NULL, BETWEEN, LIMIT)",
                "Kennzeichnungen im Relationenschema erläutern; Verbund in 19/19 Abiturklausuren 2021–2026"],
               ["4", "Nachtfahrten – Verbund über 3–4 Tabellen, Tabellenalias, zwei FKs auf dieselbe Tabelle, IS NULL "
                     "(K24, K18, K22)",
                "A2-E2-01…08, Cold Case A2-E2-CC; danach Verhör 2 · Bürgerbus (20')", "W-A2-01…04",
                "SQL entwerfen über mehrere Tabellen (jeder Termin); Prognose 2027/28: zwei FKs auf dieselbe Tabelle, "
                "IS NULL, „Vorgehen in Worten“ (seit 2025)"]],
              breiten=[1.0, 5.0, 3.6, 3.1, 4.3], groesse=8.5)
    d.inhalt("**Material:** Präsentation `DD_A2_Raeder.pptx` (Folien 1–11 DS 3, 12–22 DS 4) · AB 2.1 und AB 2.2 mit "
             "Lösungen · Unplugged-Kopiervorlagen „Roter Faden“ und „Ermittlungsweg“ · Warm-up-Karten A1 (DS 3) und "
             "A2 (DS 4; W-A2-05…08 eignen sich als Warm-up für DS 5) · Schema-Plakat StadtRad (A3, Folie 3 drucken).",
             "**Datenbank:** StadtRad – " + ", ".join(s.text.split("(")[0] for s in SCHEMA_ALLE) +
             "; Story-Fakten: Nachtfahrten 303–307 (Konto 77 „Nora Falter“, Lastenräder 41–45, ab Hafenkante, Ende NULL); "
             "insgesamt 14 laufende Ausleihen; 19 Räder ohne Station.")
    d.inhalt(*SCHEMA_ALLE)

    # ---------------------------------------------------------------- DS 3
    d.ueberschrift("DS 3 · Der rote Faden (Ermittlung 2.1)", 1)
    d.ueberschrift("Ziele", 2)
    d.inhalt(Liste(["**erläutern** Primär- und Fremdschlüssel sowie die Kennzeichnungen der Kurzschreibweise – auch für "
                    "Fremdschlüssel mit abweichendem Namen (StartStation → StationID)",
                    "**beschreiben** den Verbund als Kreuzprodukt mit anschließender Filterung (FK = PK) und "
                    "**bestimmen** die Zeilenzahl eines Kreuzprodukts",
                    "**entwerfen** Verbundabfragen über zwei Tabellen im impliziten Stil mit qualifizierten Spaltennamen",
                    "**erklären**, warum Datensätze mit NULL im Fremdschlüssel im Verbund nicht erscheinen"],
                   nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A1-05…08; Auflösung Karte 06 (NULL-Falle) an der Tafel", "EA → Team",
                "Warm-up-Karten A1"],
               ["Briefing", "10–20'", "Akte 2 öffnen; Schema in Kurzschreibweise; Impulse: Wo steht Rad 17? Warum steht "
                                      "dort nicht der Stationsname? Worauf verweisen StartStation/ZielStation?",
                "LV, UG", "Folien 1–5, Schema-Plakat"],
               ["Unplugged", "20–35'", "Roter Faden: Kreuzprodukt legen, Verbundbedingung kappt Fäden", "Teams",
                "Kopiervorlage Roter Faden, Wollfäden, Folie 6"],
               ["Lab", "35–80'", "Ermittlung 2.1, Einstieg Spurenbrett (A2-E1-03); AB 2.1 parallel", "EA",
                "iPad/Laptop, AB 2.1, Folie 7"],
               ["Fallbesprechung", "80–90'", "Verbund-Rezept (Merkkasten); Abi-Format (Rad, 2 BE) bzw. Abi-Ecke AB 2.1 "
                                             "(3 BE): zwei Teamlösungen vergleichen; Cliffhanger", "UG",
                "Folien 8–11"]],
              breiten=[2.7, 1.3, 6.8, 2.4, 3.8], groesse=8.5)
    d.ueberschrift("Unplugged „Roter Faden“ (15 min)", 2)
    d.inhalt(Liste([
        "Material je Team: 4 Stationskarten (Leuchtturm, Kornmarkt, Bootshaus, Deichtor – Deichtor ohne Rad), 6 Radkarten "
        "(Rad 6 mit StationID NULL, „in der Werkstatt“), Auftragskarte, Protokoll, 24 rote Wollfäden (50 cm), Klebepunkte. "
        "Eigene Daten, nicht die Lab-DB.",
        "Schritt 1 `FROM Rad, Station`: jedes Rad mit jeder Station verbinden – 6 · 4 = **24 Fäden** = 24 Zeilen.",
        "Schritt 2 `WHERE Rad.StationID = Station.StationID`: Es bleiben **5 Fäden** (Rad 1 und 3 → Leuchtturm, 2 und 5 → "
        "Kornmarkt, 4 → Bootshaus).",
        "Impulse: *Wo ist Rad 6?* (NULL ist mit nichts gleich → kein Faden) – *Warum taucht Deichtor nicht auf?* (kein Rad "
        "verweist darauf) – *Was passiert ohne WHERE?* (24 Zeilen, „Rad 1 steht überall“).",
        "Zusatz `AND Station.Stadtteil = 'Hafenviertel'`: Räder 1 und 3 (Leuchtturm).",
        "Sicherung: Foto nach Schritt 2; Tafel „Verbund = Kreuzprodukt + Filter FK = PK“. Danach dasselbe im Spurenbrett "
        "des Labs (A2-E1-03)."]))
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste([
        "Verbundbedingung vergessen → riesige Ergebnistabelle (P2); Diagnose „Zeilenzahl = Produkt der Tabellengrößen“ → "
        "Spurenbrett.",
        "`WHERE StationID = StationID` ohne Präfix → „mehrdeutig“ (P7); vgl. AB 2.1 Auftrag 5:"]),
        SQL(SQL_FALSCH_PRAEFIX, pruefung="fehlerhaft"),
        Liste(["Fremdschlüssel mit anderem Namen nicht erkannt (`Ausleihe.StartStation = Station.StationID`) (P7).",
               "PK/FK-Richtung verwechselt („Station verweist auf Rad“) → der Pfeil zeigt immer zum Primärschlüssel.",
               "NULL im Fremdschlüssel: Räder in der Werkstatt „verschwinden“ – kein Fehler der Abfrage, sondern "
               "Verbundlogik (Lab-Cold-Case A2-E1-CC: 19 Räder ohne Station)."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Case „Doppelt verbunden“ (AB 2.1), Lab-Cold-Case „Räder ohne Heimat“, Ausblick auf "
                    "drei Tabellen (Kunde–Ausleihe–Rad).",
                    "*Langsame:* Spurenbrett frei nutzen, Verbund-Rezept (Werkzeugkarte 3) als Schablone, ★-Aufgaben mit "
                    "Tipp; AB-Auftrag 4 mit Beweisstück 4 als Vorlage."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Unplugged-Set „Roter Faden“ je Team kopieren und schneiden; 24 Wollfäden je Team vorher schneiden.",
                    "Akte 2 im Kurs freischalten; Schema-Plakat StadtRad (A3) aufhängen.",
                    "Warm-up-Karten A1 (W-A1-05…08) bereitlegen."]))

    # ---------------------------------------------------------------- DS 4
    d.ueberschrift("DS 4 · Nachtfahrten (Ermittlung 2.2)", 1)
    d.ueberschrift("Ziele", 2)
    d.inhalt(Liste(["**planen** den Ermittlungsweg durchs Schema (Tabellen, Verbundbedingungen) und **beschreiben** ihn "
                    "in Worten",
                    "**entwerfen** Verbundabfragen über drei und vier Tabellen mit Tabellenaliasen",
                    "**entwerfen** Abfragen, in denen eine Tabelle zweimal (mit zwei Aliasen) vorkommt, weil zwei "
                    "Fremdschlüssel auf sie verweisen (Start-/Zielstation)",
                    "**erläutern**, warum laufende Ausleihen (ZielStation NULL) beim Verbund über die Zielstation "
                    "herausfallen, und **berücksichtigen** `IS NULL` korrekt"], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A2-01…04 (Kurzschreibweise, Kreuzprodukt, Verbund, FK mit anderem Namen)",
                "EA → Team", "Warm-up-Karten A2"],
               ["Briefing", "10–20'", "Hollwedel: Versicherung stellt Fragen. Frage → Tabellen; zwei Pfeile auf Station → "
                                      "Tabellenalias einführen; NULL im Verbund", "LV, UG", "Folien 12–16, Schema-Plakat"],
               ["Unplugged", "20–35'", "Ermittlungsweg durchs Schema (4 Fragekarten)", "Teams",
                "Kopiervorlage Ermittlungsweg, Faden, Folie 17"],
               ["Lab", "35–80'", "Ermittlung 2.2 (★ → ★★★; 07/08 = Nachtfahrten); AB 2.2 parallel", "EA",
                "iPad/Laptop, AB 2.2, Folie 18"],
               ["Fallbesprechung", "80–90'", "Story-Auflösung (Konto 77 „Nora Falter“); Merkkasten Alias und IS NULL; "
                                             "Abi-Format (4 BE); Hinweis Verhör 2 (20')", "UG", "Folien 19–22"]],
              breiten=[2.7, 1.3, 6.8, 2.4, 3.8], groesse=8.5)
    d.ueberschrift("Unplugged „Ermittlungsweg durchs Schema“ (15 min)", 2)
    d.inhalt("Material je Team: 5 Tabellenkarten (A5, PK unterstrichen, ↑FK rot), 2 leere Karten „Kopie von …“, "
             "4 Fragekarten, ein langer Faden, Folienstift. Erwartete Wege:")
    d.tabelle(["Fragekarte", "Karten (Weg)", "Verbundbedingungen / Pointe"],
              [["1 ★ Lastenrad-Kundschaft", "Kunde – Ausleihe – Rad",
                "a.KundenID = k.KundenID, a.RadID = r.RadID; Filter r.Typ = 'Lastenrad'; DISTINCT"],
               ["2 ★★ Start- und Zielstation", "Station (Kopie s1) – Ausleihe – Station (Kopie s2)",
                "a.StartStation = s1.StationID, a.ZielStation = s2.StationID – eine Station-Karte reicht nicht"],
               ["3 ★★ Räder heute am Deichtor", "Kunde – Ausleihe – Rad – Station",
                "Weg über **Rad.StationID** (Standort heute), nicht über StartStation (Ort der Ausleihe): zwei Pfade, "
                "zwei Bedeutungen. In den Kartendaten erfunden; in der StadtRad-DB gibt es Deichtor nicht – siehe SQL "
                "unten mit Marktplatz."],
               ["4 ★★★ Räder unterwegs", "Ausleihe – Station (nur Start)",
                "ZielStation und Ende sind NULL: Ein Verbund über ZielStation verliert genau diese Ausleihen. Die DB "
                "kennt nur die Startstation – wo die Räder jetzt sind, weiß sie nicht (Nachtfahrten!)."]],
              breiten=[3.6, 4.2, d.breite - 7.8], groesse=8.5)
    d.inhalt("Fragekarte 1 und 3 als SQL (zum Zeigen nach der Sicherung; Karte 3 mit einer echten Station):",
             SQL(SQL_LH_FRAGE1), ergebnis(SQL_LH_FRAGE1, max_zeilen=3),
             SQL(SQL_LH_FRAGE3), ergebnis(SQL_LH_FRAGE3, max_zeilen=3),
             "Sicherung: Teams fotografieren ihre Wege; an der Tafel die Lösung zu Karte 2 als SQL (Folie 15).")
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste([
        "Eine Station-Tabelle für Start und Ziel → `StartStation = StationID AND ZielStation = StationID` liefert nur "
        "Rundfahrten (P7/P2).",
        "Aliase in SELECT nicht verwendet, nachdem sie in FROM vergeben wurden (`Station.Name` trotz `Station s1`) → Fehler.",
        "Verbund über den falschen Pfad (Rad.StationID statt Ausleihe.StartStation) – syntaktisch korrekt, inhaltlich "
        "falsch (P2).",
        "Laufende Ausleihen verschwinden unbemerkt (P5) → Prüf-Variante `+null` zeigt „passt nur zufällig“.",
        "Zu wenige Verbundbedingungen bei vier Tabellen (Zähler n − 1)."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Case „Ein Konto, viele Räder“ (AB 2.2), Lab-Cold-Case „Rundfahrten“, Verhör 2 schon "
                    "in der Stunde.",
                    "*Langsame:* Schema-Karten aus dem Unplugged liegen weiter aus; Werkzeugkarte 3; Spurenbrett mit drei "
                    "Tabellen; AB 2.2 Auftrag 2 vor Auftrag 4."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Schema-Karten-Set (5 Tabellen + 2 Kopie-Karten) und Fragekarten je Team kopieren, laminieren.",
                    "Verhör 2 im Kurs auf „auto“ oder „offen“ stellen.",
                    "Folie 22 (Story-Auflösung) erst nach A2-E2-07 zeigen – sie fehlt im Handout."]))
    d.notiz("Verhör 2 · Bürgerbus", "Kontext-DB Bürgerbus (Haltestelle, Fahrer, Bus, Fahrt mit zwei Fremdschlüsseln auf "
            "Haltestelle). 20 Minuten, 6 Aufgaben (2 ★, 3 ★★, 1 ★★★), Abimodus streng, gesiezt. Mindestens eine Aufgabe "
            "zu zwei FKs auf dieselbe Tabelle.")
    return d.speichern(ZIEL / "4_Lehrerhinweise" / "DD_A2_Lehrerhinweise.docx")


# ================================================================== Warm-up-Karten
def warmup():
    karten = json.loads((REPO / "seed" / "warmup" / "warmup_a2.json").read_text(encoding="utf-8"))["karten"]
    pfad = REPO / "materialien" / "Warm-up" / "DD_Warm-up-Karten_A2.docx"
    return warmup_karten_dokument(karten, 2, "Warm-up-Karten Akte 2").speichern(pfad)


def main():
    erzeugt = [ab21("schueler"), ab21("loesung"), ab22("schueler"), ab22("loesung"), praesentation(),
               unplugged_roter_faden(), unplugged_ermittlungsweg(), lehrerhinweise(), warmup()]
    for p in erzeugt:
        print(f"✔ {Path(p).relative_to(REPO)}")


if __name__ == "__main__":
    main()
