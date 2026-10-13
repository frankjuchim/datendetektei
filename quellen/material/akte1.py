#!/usr/bin/env python3
"""Erstellungshilfe: Material Akte 1 „Das Fundbüro“ (DS 1–2) im Akten-Design (Phase 7, Task 2).

Erzeugt unter materialien/Akte1_Fundbuero/ (Spec §12.1) und materialien/Warm-up/:
  1_Praesentation/DD_A1_Fundbuero.pptx                         Fallbesprechung DS 1 + DS 2 (Notizen, Lösungsfolien)
  2_Arbeitsblaetter/DD_A1_AB1_Der_erste_Auftrag.docx            AB 1.1 (Schülerfassung)
  2_Arbeitsblaetter/DD_A1_AB2_Das_verschwundene_Fundstueck.docx AB 1.2 (Schülerfassung)
  5_Loesungen/…_LSG.docx                                        Lösungen mit typischen Fehlern (P-Nummern)
  3_Unplugged/DD_A1_Unplugged_Karteikarten.docx                 Karteikarten-Filter (DS 1, 30 Karten, eigene Daten)
  3_Unplugged/DD_A1_Unplugged_Platzhalter-Fahndung.docx         Fahndung mit Platzhaltern + NULL-Stuhl (DS 2)
  3_Unplugged/DD_A1_Unplugged_Detektei-Ausweise.docx            Detektei-Ausweise (Kick-off DS 1)
  4_Lehrerhinweise/DD_A1_Lehrerhinweise.docx                    Überblick, Ziele, Verlauf, Unplugged, Schwierigkeiten
  Warm-up/DD_Warm-up-Karten_A1.docx                             aus seed/warmup/warmup_a1.json

Grundlage: docs/DS_PLANUNG.md (DS 1, DS 2), docs/STORY.md 5.1/5.2, seed/akten/akte1.json (Spur im Lab),
seed/datenbanken/fundbuero/LIESMICH.md. Alle Ergebnistabellen kommen direkt aus der Fundbuero-DB (abfrage_ergebnis).
Die erzeugten docx/pptx sind danach maßgeblich (Spec §12); dieses Skript ist nur die Erstellungshilfe.

Aufruf: python3 quellen/material/akte1.py   (überschreibt die oben genannten Dateien)
"""
import json
import sqlite3
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "werkzeuge"))

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import akten_design as A  # noqa: E402
from akten_design import (  # noqa: E402
    AktenDokument, AktenPraesentation, Auftrag, Ergebnis, Liste, Loesung, Notiz, Schema, SQL, Tabelle,
    abfrage_ergebnis, warmup_karten_dokument)

MAT = REPO / "materialien"
ORDNER = MAT / "Akte1_Fundbuero"
DB = "fundbuero"
FUNDSACHE = ("Fundsache(<u>FundNr</u>, Bezeichnung, Kategorie, Farbe, Fundort, Stadtteil, Funddatum, Wert, Finder, "
             "AbgeholtAm)")


def ergebnis(sql, max_zeilen=None, db=DB, breiten=None):
    """Ergebnistabelle direkt aus der DB (optional gekürzt mit „…“)."""
    kopf, zeilen = abfrage_ergebnis(sql, db)
    if max_zeilen and len(zeilen) > max_zeilen:
        return Ergebnis(kopf, zeilen[:max_zeilen], abgeschnitten=True, breiten=breiten)
    return Ergebnis(kopf, zeilen, breiten=breiten)


def auftraege(d, liste):
    """Ermittlungsaufträge; jede Zeile bleibt auf einer Seite (kein Umbruch mitten im Auftrag)."""
    tab = d.ermittlungsauftraege(liste)
    if d.fassung != "schueler":
        return tab  # Lösungsfassung: lange Lösungskästen dürfen umbrechen
    for zeile in tab.rows[1:]:
        A._zeile_nicht_teilen(zeile)
    return tab


def loesung_sql(sql, *rest, max_zeilen=None, fehler=(), varianten=()):
    """Lösungskasten: SQL + Ergebnistabelle aus der DB + weitere Inhalte."""
    return Loesung(SQL(sql), ergebnis(sql, max_zeilen), *rest, fehler=fehler, varianten=varianten)


# ================================================================== Daten für AB 1.1 (Ausschnitt aus der Fundbuero-DB)
AUSSCHNITT_NR = (3, 11, 31, 38, 39, 73, 80)
AUSSCHNITT_SQL = ("SELECT FundNr, Bezeichnung, Kategorie, Farbe, Stadtteil, Funddatum, Wert\nFROM Fundsache\n"
                  f"WHERE FundNr IN ({', '.join(map(str, AUSSCHNITT_NR))})\nORDER BY FundNr;")


def ausschnitt():
    kopf, zeilen = abfrage_ergebnis(AUSSCHNITT_SQL, DB)
    # Plausibilität der Aufgaben 2 und 3 (Ausschnitt-Falle): Bezeichnung im Ausschnitt eindeutig, Funddatum nicht
    bez = [z[1] for z in zeilen]
    datum = [z[5] for z in zeilen]
    assert len(set(bez)) == len(bez), "Bezeichnung soll im Ausschnitt eindeutig sein"
    assert len(set(datum)) < len(datum), "Funddatum soll im Ausschnitt doppelt vorkommen"
    return Tabelle(kopf, zeilen, pk=("FundNr",), breiten=[1.5, 3.0, 2.2, 1.6, 3.0, 2.3, 1.3])


# ================================================================== AB 1.1 „Der erste Auftrag“ (DS 1)
def ab1(fassung):
    lsg = fassung == "loesung"
    d = AktenDokument("ab", 1, "Der erste Auftrag", fassung=fassung, ermittlung="1.1", ds=1, material="AB 1.1",
                      datenbank=DB, untertitel="Tabelle, Primärschlüssel, SELECT – FROM – WHERE – ORDER BY")
    d.kopfband()
    d.fall_einstieg("Gesa Ohlendiek vom Fundbüro Weserstedt bittet die Detektei, ihr Fundbuch zu prüfen – "
                    "„irgendetwas stimmt da nicht“. Bevor ihr ermittelt, müsst ihr das Fundbuch lesen und befragen können.")
    d.beweisstueck(1, "Aufbau einer Tabelle",
                   "Das Fundbuch ist eine **Tabelle** (Relation) mit dem Namen Fundsache. Jede Zeile ist ein "
                   "**Datensatz** (eine Fundsache), jede Spalte ein **Attribut** (z. B. Kategorie), jeder einzelne Eintrag "
                   "ein **Attributwert** (z. B. 'Dokumente'). `NULL` heißt: Hier ist nichts eingetragen. Ausschnitt:",
                   ausschnitt(),
                   "In der Kurzschreibweise steht der Tabellenname vor der Klammer, der Primärschlüssel ist unterstrichen:",
                   Schema(FUNDSACHE))
    d.beweisstueck(2, "Der Primärschlüssel",
                   "Ein **Primärschlüssel** bezeichnet jeden Datensatz eindeutig. Ein Attribut taugt nur dann, wenn es",
                   Liste(["**eindeutig** ist – für **alle möglichen** Daten, auch für die Fundsachen von morgen,",
                          "**nie NULL** ist – jeder Datensatz hat einen Wert,",
                          "**minimal** ist – nichts Überflüssiges enthält."], nummeriert=True),
                   Notiz("Ein Ausschnitt beweist nichts",
                         "Dass ein Wert in ein paar Zeilen nur einmal vorkommt, heißt nicht, dass er immer eindeutig ist. "
                         "Frag dich: *Kann morgen eine zweite Fundsache mit diesem Wert abgegeben werden?*"))
    d.beweisstueck(3, "Das Abfrage-Gerüst",
                   "Eine Abfrage beantwortet vier Fragen: *Welche Spalten?* (SELECT) – *Aus welcher Tabelle?* (FROM) – "
                   "*Welche Zeilen?* (WHERE) – *In welcher Reihenfolge?* (ORDER BY, ASC aufsteigend ist Standard, DESC "
                   "absteigend). Je Klausel eine Zeile, Schlüsselwörter groß:",
                   SQL("SELECT Bezeichnung, Stadtteil, Wert\nFROM Fundsache\nWHERE Kategorie = 'Schlüssel'\n"
                       "  AND Wert > 25\nORDER BY Wert DESC;"),
                   ergebnis("SELECT Bezeichnung, Stadtteil, Wert\nFROM Fundsache\nWHERE Kategorie = 'Schlüssel'\n"
                            "  AND Wert > 25\nORDER BY Wert DESC;"),
                   "Mehrstufig sortieren: `ORDER BY Stadtteil, Wert DESC` sortiert erst nach Stadtteil; nur bei gleichem "
                   "Stadtteil entscheidet der Wert (absteigend).")
    d.beweisstueck(4, "Vergleiche und Verknüpfungen",
                   Tabelle(["Operator", "Bedeutung", "Beispiel"],
                           [["`=`  `!=`", "gleich, ungleich", "`Kategorie != 'Kleidung'`"],
                            ["`<`  `>`", "kleiner, größer", "`Wert > 100`"],
                            ["`<=`  `>=`", "kleiner/größer **oder gleich**", "`Wert >= 20` (20 gehört dazu)"],
                            ["`AND`", "beide Bedingungen müssen gelten", "`Kategorie = 'Schmuck' AND Wert > 50`"],
                            ["`OR`", "mindestens eine Bedingung gilt", "`Stadtteil = 'Altstadt' OR Stadtteil = 'Lindenhof'`"],
                            ["`NOT`", "kehrt eine Bedingung um", "`NOT Stadtteil = 'Hafenviertel'`"]],
                           breiten=[2.6, 5.2, 8.6]),
                   Notiz("Klammern retten Fälle",
                         "**AND bindet stärker als OR** – wie Punkt vor Strich. `A OR B AND C` bedeutet `A OR (B AND C)`. "
                         "Wenn du etwas anderes meinst, setze Klammern: `(A OR B) AND C`. Text steht in `'…'`, Zahlen ohne "
                         "Hochkommata."))

    auftrag1_inhalt = [] if lsg else [Tabelle(["Fachbegriff", "Beispiel aus dem Ausschnitt"],
                                              [["Tabelle", ""], ["Datensatz", ""], ["Attribut", ""], ["Attributwert", ""]],
                                              breiten=[3.2, 9.6])]
    auftrag3_inhalt = [Tabelle(["Bedingung", "FundNr der Treffer"],
                               [["`Kategorie = 'Sonstiges' OR Farbe = 'schwarz' AND Wert < 10` (wie notiert)", ""],
                                ["`(Kategorie = 'Sonstiges' OR Farbe = 'schwarz') AND Wert < 10`", ""]],
                               breiten=[9.4, 3.4])]
    probe = ("SELECT FundNr\nFROM Fundsache\nWHERE FundNr IN (3, 11, 31, 38, 39, 73, 80)\n"
             "  AND (Kategorie = 'Sonstiges' OR Farbe = 'schwarz' AND Wert < 10);")
    probe_k = ("SELECT FundNr\nFROM Fundsache\nWHERE FundNr IN (3, 11, 31, 38, 39, 73, 80)\n"
               "  AND (Kategorie = 'Sonstiges' OR Farbe = 'schwarz')\n  AND Wert < 10;")
    sql4 = "SELECT Bezeichnung, Funddatum\nFROM Fundsache\nWHERE Stadtteil = 'Lindenhof'\nORDER BY Funddatum DESC;"
    sql5 = ("SELECT Bezeichnung, Kategorie, Stadtteil, Wert\nFROM Fundsache\nWHERE (Wert >= 100 OR Kategorie = 'Schmuck')\n"
            "  AND Stadtteil != 'Hafenviertel'\nORDER BY Stadtteil, Wert DESC;")
    auftraege(d, [
        Auftrag(1, 1, "PAPIER", "Beschrifte im Fundbuch-Ausschnitt aus Beweisstück 1 je ein Beispiel für **Tabelle**, "
                "**Datensatz**, **Attribut** und **Attributwert** (Pfeile) und trage es in die Tabelle ein.",
                kompetenzen=["K01"], inhalt=auftrag1_inhalt,
                loesung=Loesung("**Tabelle:** Fundsache (der ganze Ausschnitt mit Kopfzeile) · **Datensatz:** eine Zeile, "
                                "z. B. 38 | Kuscheltier | Spielzeug | schwarz | … · **Attribut:** eine Spalte, z. B. Farbe · "
                                "**Attributwert:** ein einzelner Eintrag, z. B. 'Bahnhofsviertel' oder 7.",
                                fehler=[("P7", "Attribut und Attributwert verwechselt („Attribut = schwarz“) – das "
                                               "Attribut ist die Spalte Farbe, 'schwarz' ist ein Wert darin.")])),
        Auftrag(2, 1, "PAPIER", "Entscheide begründet, ob **FundNr**, **Bezeichnung** oder **Funddatum** Primärschlüssel "
                "der Tabelle Fundsache sein kann. Prüfe alle drei Kriterien aus Beweisstück 2.",
                kompetenzen=["K02"], linien=4,
                loesung=Loesung("**FundNr:** geeignet – das Fundbüro vergibt jede Eingangsnummer genau einmal (eindeutig "
                                "auch für künftige Fundsachen), jede Fundsache bekommt eine (nie NULL), ein einzelnes Attribut "
                                "(minimal).",
                                "**Bezeichnung:** ungeeignet – im Ausschnitt zwar eindeutig, aber morgen kann ein zweiter "
                                "Regenschirm abgegeben werden; im ganzen Fundbuch gibt es z. B. mehrere Schals.",
                                "**Funddatum:** ungeeignet – an einem Tag werden mehrere Sachen gefunden; schon im Ausschnitt "
                                "haben FundNr 38 und 39 dasselbe Funddatum.",
                                fehler=[("P8", "„Im Ausschnitt kommt jede Bezeichnung nur einmal vor, also ist sie "
                                               "Primärschlüssel.“ – ein Ausschnitt beweist keine Eindeutigkeit."),
                                        ("P8", "Nur „eindeutig“ geprüft; „nie NULL“ und „minimal“ fehlen in der Begründung.")])),
        Auftrag(3, 2, "PAPIER", "Gib an, welche Datensätze des Ausschnitts die Bedingung liefert – einmal wie notiert, "
                "einmal mit Klammern um die OR-Bedingung. Erkläre den Unterschied in einem Satz.",
                kompetenzen=["K17", "K29"], inhalt=auftrag3_inhalt, linien=2,
                loesung=Loesung("Wie notiert: **31, 38, 73, 80** – AND bindet stärker, gemeint ist also "
                                "`Kategorie = 'Sonstiges' OR (Farbe = 'schwarz' AND Wert < 10)`: jede Fundsache der "
                                "Kategorie Sonstiges (auch die Sonnenbrille für 39 €) und dazu das schwarze Kuscheltier für 7 €.",
                                "Mit Klammern: **31, 38, 80** – die Sonnenbrille (73) fällt heraus, weil jetzt für alle "
                                "Zeilen Wert < 10 gelten muss.",
                                "Probe in der Konsole (der Ausschnitt per FundNr eingegrenzt):",
                                SQL(probe), ergebnis(probe), SQL(probe_k), ergebnis(probe_k),
                                fehler=[("P6", "OR vor AND ausgewertet – dann fehlt in der ersten Zeile die 73 bzw. "
                                               "es werden beide Fälle gleich beantwortet.")])),
        Auftrag(4, 2, "PAPIER", "Notiere eine Abfrage, die von allen Fundsachen aus dem Stadtteil Lindenhof die "
                "Bezeichnung und das Funddatum ausgibt – die neuesten Funde zuerst.",
                kompetenzen=["K15", "K21"], linien=5,
                loesung=loesung_sql(sql4, "(12 Zeilen, gekürzt.) Ein Datum im Format 'JJJJ-MM-TT' sortiert sich wie eine "
                                    "Zahl: DESC = neueste zuerst.", max_zeilen=4,
                                    fehler=[("P5", "`WHERE Stadtteil = Lindenhof` ohne Hochkommata – SQL sucht dann eine "
                                                   "Spalte namens Lindenhof.")])),
        Auftrag(5, 3, "LAB", "Gesucht sind alle Fundsachen mit einem Wert ab 100 € oder der Kategorie 'Schmuck', aber "
                "nicht aus dem Hafenviertel – mit Bezeichnung, Kategorie, Stadtteil und Wert, sortiert nach Stadtteil und "
                "innerhalb eines Stadtteils absteigend nach Wert. Entwirf die Abfrage auf Papier und prüfe sie in der "
                "Konsole (freie Ermittlung).", kompetenzen=["K17", "K21"], linien=6,
                loesung=loesung_sql(sql5, fehler=[
                    ("P6", "Ohne Klammern: `Wert >= 100 OR Kategorie = 'Schmuck' AND Stadtteil != 'Hafenviertel'` – dann "
                           "erscheinen die Armbanduhr (149 €) und das Fernglas (180 €) aus dem Hafenviertel."),
                    ("P6", "`Wert > 100` statt `>= 100` – Ladekabel (100 €) und Armbanduhr (100 €) fehlen.")],
                    varianten=[SQL(sql5.replace("  AND Stadtteil != 'Hafenviertel'", "  AND NOT Stadtteil = 'Hafenviertel'"))])),
    ])
    abi = "SELECT Bezeichnung, Wert\nFROM Fundsache\nWHERE Stadtteil = 'Altstadt'\n  AND Wert > 20;"
    if not lsg:
        d.seitenumbruch()
    d.abi_ecke("Gegeben ist die Tabelle Fundsache (siehe Beweisstück 1). **Entwerfen** Sie eine SQL-Abfrage, die "
               "Bezeichnung und Wert aller Fundsachen ausgibt, die im Stadtteil Altstadt gefunden wurden und einen Wert "
               "von mehr als 20 € haben.", be=2, operator="entwerfen", afb="II", kompetenzen=["K15", "K17"], linien=6,
               loesung=loesung_sql(abi, "**Bewertung:** 1 BE SELECT/FROM mit den richtigen Spalten, 1 BE vollständige "
                                   "Bedingung (Text in Hochkommata, `>` statt `>=`). Der Ball und der Helm mit genau 20 € "
                                   "gehören nicht dazu.",
                                   fehler=[("P6", "`Wert >= 20` – „mehr als“ heißt `>`."),
                                           ("P5", "`Stadtteil = Altstadt` ohne Hochkommata.")]))
    d.spur_im_lab("1.1 „Der erste Auftrag“", "A1-E1-01 … A1-E1-08 (★ → ★★ → ★★★)",
                  werkzeug="Zeitlupe (WHERE-Schritt) bei Klammerfehlern",
                  text="Erst **Ausführen**, dann **Prüfen**. Wer fertig ist: Cold Case A1-E1-CC „Die Nachtfalter-Liste“ im Lab.")
    d.cold_case("Die Nachtfalter-Liste", "Im Feld Finder stehen Namen wie 'D. Harjes' – manchmal auch NULL. Welche "
                "Bedingungen bräuchtest du, um Finder-Einträge zu finden, die gar keine Personennamen sind? Notiere Ideen "
                "– in DS 2 bekommst du das passende Werkzeug.", kompetenzen=["K15"], linien=4,
                loesung=Loesung("Ideen genügen, z. B.: „Finder beginnt mit einem Buchstaben, dann Punkt …“, „Finder enthält "
                                "einen Tiernamen“, „Finder ist leer (NULL)“. In DS 2 mit `LIKE` und `IS NULL`; im Lab führt "
                                "die sortierte Finder-Liste (A1-E1-CC) zu 'N. Falter' – schnell gesprochen: *Nachtfalter*."))
    name = "5_Loesungen/DD_A1_AB1_Der_erste_Auftrag_LSG.docx" if lsg else "2_Arbeitsblaetter/DD_A1_AB1_Der_erste_Auftrag.docx"
    return d.speichern(ORDNER / name)


# ================================================================== AB 1.2 „Das verschwundene Fundstück“ (DS 2)
def ab2(fassung):
    lsg = fassung == "loesung"
    d = AktenDokument("ab", 1, "Das verschwundene Fundstück", fassung=fassung, ermittlung="1.2", ds=2, material="AB 1.2",
                      datenbank=DB, untertitel="LIKE, BETWEEN, IN, IS NULL – Literale, Berechnungen, DISTINCT, LIMIT")
    d.kopfband()
    d.fall_einstieg("Frau Ohlendiek erinnert sich nur an Bruchstücke: etwas aus Messing, im Juni oder Juli gefunden, "
                    "mehrere hundert Euro wert. Edda Lorenz: „Raten ist in diesem Haus verboten. Zählen nicht.“")
    like = "SELECT FundNr, Bezeichnung\nFROM Fundsache\nWHERE Bezeichnung LIKE 'Fahrrad%';"
    d.beweisstueck(1, "Textmuster mit LIKE",
                   "Mit `LIKE` vergleichst du Text mit einem **Muster**: `%` steht für beliebig viele Zeichen (auch "
                   "keins), `_` für **genau ein** Zeichen.",
                   Tabelle(["Muster", "bedeutet", "passt auf", "passt nicht auf"],
                           [["`'Fahrrad%'`", "beginnt mit Fahrrad", "Fahrradhelm, Fahrradschloss", "Luftpumpe"],
                            ["`'%jacke'`", "endet auf jacke", "Regenjacke, Jacke", "Jackentasche"],
                            ["`'%karte%'`", "enthält karte", "Monatskarte, Kartenetui", "Notizbuch"],
                            ["`'Sch_l'`", "Sch, ein Zeichen, l", "Schal", "Schaal, Schlüssel"]],
                           breiten=[2.6, 4.0, 5.2, 4.6]),
                   SQL(like), ergebnis(like),
                   "Groß- und Kleinschreibung spielt bei LIKE in SQLite (für A–Z) keine Rolle: `'%jacke'` findet auch 'Jacke'.")
    zwischen = ("SELECT Bezeichnung, Kategorie, Wert\nFROM Fundsache\nWHERE Kategorie IN ('Spielzeug', 'Wertsachen')\n"
                "  AND Wert BETWEEN 7 AND 20;")
    d.beweisstueck(2, "BETWEEN, IN und IS NULL",
                   Tabelle(["Bedingung", "bedeutet"],
                           [["`Wert BETWEEN 7 AND 20`", "7 ≤ Wert ≤ 20 – **beide Grenzen gehören dazu**"],
                            ["`Kategorie IN ('Spielzeug', 'Wertsachen')`", "Kategorie ist einer der Werte der Liste"],
                            ["`Finder IS NULL`", "im Feld Finder ist nichts eingetragen (anonym abgegeben)"],
                            ["`NOT AbgeholtAm IS NULL`", "im Feld AbgeholtAm steht ein Datum (schon abgeholt)"]],
                           breiten=[6.6, 9.8]),
                   SQL(zwischen), ergebnis(zwischen),
                   Notiz("„= NULL“ ist nie wahr",
                         "NULL heißt: *weiß keiner*. Ein Vergleich mit NULL ergibt weder wahr noch falsch – die Zeile fällt "
                         "heraus, auch bei `Farbe != 'rot'`. Prüfe Lücken **nur** mit `IS NULL`."))
    d.beweisstueck(3, "Literale richtig notieren",
                   Tabelle(["Art", "richtig", "falsch"],
                           [["Text", "`'Hafenviertel'`", "`Hafenviertel` (wird als Spaltenname gelesen)"],
                            ["Datum", "`'2026-09-30'` (JJJJ-MM-TT, als Text)", "`2026-09-30` (Rechnung 2026 − 9 − 30), `'30.09.2026'`"],
                            ["Zahl", "`12.5` (Dezimalpunkt)", "`12,5` oder `'12.5'`"],
                            ["ungleich", "`!=`", "`<>` (nicht im Abitur)"]],
                           breiten=[2.0, 6.4, 8.0]),
                   "Weil ein Datum im Format 'JJJJ-MM-TT' notiert ist, kannst du es vergleichen und sortieren wie eine Zahl. "
                   "Achtung: `7 / 2` ergibt bei ganzen Zahlen `3` (ganzzahlige Division).")
    lohn = ("SELECT Bezeichnung, Wert, Wert * 0.05 AS Finderlohn\nFROM Fundsache\nWHERE Kategorie = 'Taschen'\n"
            "ORDER BY Finderlohn DESC\nLIMIT 2;")
    distinct = "SELECT DISTINCT Stadtteil\nFROM Fundsache\nWHERE Kategorie = 'Wertsachen';"
    d.beweisstueck(4, "Berechnen, umbenennen, Doppeltes weglassen, begrenzen",
                   "Spalten kann man berechnen und mit `AS` benennen; nach dem Spaltenalias darf man sortieren. `LIMIT n` "
                   "zeigt nur die ersten n Zeilen – sinnvoll nur nach ORDER BY:",
                   SQL(lohn), ergebnis(lohn),
                   "`DISTINCT` lässt doppelte Zeilen weg – hier: In welchen Stadtteilen wurden Wertsachen gefunden?",
                   SQL(distinct), ergebnis(distinct),
                   Notiz("LIMIT mit Vorsicht", "„Die zwei teuersten“ mit `LIMIT 2` stimmt nur, wenn an der Schnittkante kein "
                         "Gleichstand ist – sonst fällt eine gleich teure Zeile zufällig weg. Ohne ORDER BY ist LIMIT reiner Zufall."))

    a1 = [] if lsg else [Tabelle(["Muster", "passt (zwei Wörter)", "passt nicht (ein Wort)"],
                                 [["`'Sch%'`", "", ""], ["`'%ring'`", "", ""], ["`'_a%'`", "", ""]], breiten=[2.6, 6.2, 4.2])]
    a2 = [Tabelle(["Juris Bedingung (gemeint: …)", "korrigiert"],
                  [["a) `WHERE Farbe = NULL` (gemeint: Farbe nicht erfasst)", ""],
                   ["b) `WHERE Funddatum >= 2026-09-01` (gemeint: ab dem 1. September 2026)", ""],
                   ["c) `WHERE Kategorie <> 'Kleidung'` (gemeint: nicht Kleidung)", ""],
                   ["d) `WHERE Wert > 12,50` (gemeint: mehr als 12,50 €)", ""],
                   ["e) `WHERE AbgeholtAm = '05.09.2026'` (gemeint: am 5. September 2026 abgeholt)", ""]],
                  breiten=[6.8, 5.5])]
    a3 = [Tabelle(["FundNr", "Bezeichnung", "Wert"],
                  [["201", "Stofftasche", "9.50"], ["202", "Mütze", "10"], ["203", "Schal", "20"], ["204", "Handschuh", "20.50"],
                   ["205", "Regenschirm", "NULL"]], breiten=[1.8, 3.4, 1.8]),
          "Bedingung: `WHERE Wert BETWEEN 10 AND 20`"]
    sql4 = "SELECT DISTINCT Kategorie\nFROM Fundsache\nWHERE Stadtteil = 'Hafenviertel'\nORDER BY Kategorie;"
    kette = ("SELECT Bezeichnung, Fundort, Wert, AbgeholtAm, Finder\nFROM Fundsache\nWHERE Bezeichnung LIKE '%Messing%'\n"
             "  AND Funddatum BETWEEN '2026-06-01' AND '2026-07-31'\n  AND (Wert >= 400 OR Wert IS NULL)\n"
             "  AND NOT AbgeholtAm IS NULL;")
    auftraege(d, [
        Auftrag(1, 1, "PAPIER", "Gib zu jedem Muster zwei passende Wörter und ein unpassendes Wort an.",
                kompetenzen=["K19"], inhalt=a1,
                loesung=Loesung("Zum Beispiel: `'Sch%'` Schal, Schlüsselbund – Tasche · `'%ring'` Ohrring, Silberring – "
                                "Ringbuch · `'_a%'` Kamera, Matze – Helm (zweites Zeichen muss ein a sein).",
                                fehler=[("P6", "`'%ring'` mit „Ringbuch“ – `%ring` verlangt, dass das Wort auf ring endet."),
                                        ("P6", "`_` als „beliebig viele Zeichen“ gelesen: „Schal“ passt nicht auf `'_a%'`.")])),
        Auftrag(2, 1, "PAPIER", "Juri hat beim Protokollieren fünf Bedingungen falsch notiert. Korrigiere sie.",
                kompetenzen=["K18", "K32"], inhalt=a2,
                loesung=Loesung("a) `WHERE Farbe IS NULL` · b) `WHERE Funddatum >= '2026-09-01'` · "
                                "c) `WHERE Kategorie != 'Kleidung'` (im Abitur `!=` statt `<>`) · d) `WHERE Wert > 12.5` · "
                                "e) `WHERE AbgeholtAm = '2026-09-05'`",
                                fehler=[("P5", "a) `= NULL` ist nie wahr – die Abfrage liefert stillschweigend keine Zeile."),
                                        ("P5", "b) ohne Hochkommata rechnet SQL 2026 − 9 − 1 = 2016 und vergleicht mit dieser Zahl.")])),
        Auftrag(3, 2, "PAPIER", "Gib an, welche Datensätze des folgenden (erfundenen) Ausschnitts die Bedingung liefert, "
                "und begründe die Grenzfälle.", kompetenzen=["K18", "K29"], inhalt=a3, linien=3,
                loesung=Loesung("**202 und 203.** BETWEEN schließt beide Grenzen ein (10 und 20 gehören dazu); 9.50 und 20.50 "
                                "liegen außerhalb; beim Regenschirm ist der Wert unbekannt (NULL) – der Vergleich ist nicht wahr.",
                                fehler=[("P6", "Grenzen ausgeschlossen (nur 203 bzw. keine) – BETWEEN ist inklusiv."),
                                        ("P5", "NULL als 0 gelesen.")])),
        Auftrag(4, 2, "LAB", "Liste in der Konsole alle **verschiedenen** Kategorien auf, die im Hafenviertel vorkommen "
                "(alphabetisch). Notiere die Abfrage und die Anzahl der Zeilen.", kompetenzen=["K15"], linien=4,
                loesung=loesung_sql(sql4, "8 Kategorien.", fehler=[
                    ("P7", "DISTINCT vergessen – dann erscheint z. B. Kleidung dreimal (eine Zeile je Fundsache).")])),
        Auftrag(5, 3, "PAPIER", "Beschreibe in Worten, mit welchen drei Abfragen du im Lab (Indizienkette A1-E2-06 … 08) "
                "das verschwundene Fundstück eingekreist hast: Welche Bedingung kam in welchem Schritt dazu, und warum?",
                kompetenzen=["K30"], linien=7,
                loesung=Loesung("1. Alle Fundsachen, deren Bezeichnung „Messing“ enthält (`LIKE '%Messing%'` – mit % auf beiden "
                                "Seiten, weil Messing auch mitten im Wort stehen kann): vier Treffer.",
                                "2. Zusätzlich Funddatum im Juni oder Juli (`BETWEEN '2026-06-01' AND '2026-07-31'`, Grenzen "
                                "inklusive) und Wert mindestens 400 € oder unbekannt (`Wert >= 400 OR Wert IS NULL` in Klammern): "
                                "nur noch der Messing-Kompass.",
                                "3. Zusätzlich `NOT AbgeholtAm IS NULL`: Er gilt als abgeholt – und Fundort, Abholdatum und "
                                "Finder werden sichtbar:",
                                SQL(kette), ergebnis(kette, breiten=[2.8, 3.0, 1.2, 2.2, 1.9]),
                                fehler=[("P6", "`LIKE 'Messing%'` – verliert Bezeichnungen, in denen Messing nicht am Anfang steht."),
                                        ("P5", "„Wert unbekannt“ vergessen: ohne `OR Wert IS NULL` fällt ein Messing-Fund "
                                               "mit unbekanntem Wert stillschweigend weg.")])),
    ])
    abi = ("SELECT Bezeichnung, Funddatum\nFROM Fundsache\nWHERE AbgeholtAm IS NULL\n"
           "  AND Funddatum BETWEEN '2026-09-01' AND '2026-09-30'\nORDER BY Funddatum DESC;")
    d.abi_ecke("Gegeben ist die Tabelle Fundsache. **Entwerfen** Sie eine SQL-Abfrage, die Bezeichnung und Funddatum aller "
               "noch nicht abgeholten Fundsachen ausgibt, die im September 2026 gefunden wurden. Die neuesten Funde sollen "
               "zuerst erscheinen.", be=3, operator="entwerfen", afb="I/II", kompetenzen=["K18", "K32", "K21"], linien=6,
               loesung=loesung_sql(abi, "**Bewertung:** 1 BE `AbgeholtAm IS NULL`, 1 BE Zeitraum mit korrekten "
                                   "Datumsliteralen, 1 BE Sortierung absteigend (sowie SELECT/FROM).",
                                   varianten=[SQL(abi.replace("BETWEEN '2026-09-01' AND '2026-09-30'",
                                                              "LIKE '2026-09-%'")),
                                              SQL(abi.replace("BETWEEN '2026-09-01' AND '2026-09-30'",
                                                              ">= '2026-09-01'\n  AND Funddatum <= '2026-09-30'"))],
                                   fehler=[("P5", "`AbgeholtAm = NULL` – liefert keine Zeile."),
                                           ("P5", "Datum ohne Hochkommata oder im Format '01.09.2026'."),
                                           ("P6", "`Funddatum < '2026-09-30'` – der 30. September fehlt.")]))
    d.spur_im_lab("1.2 „Das verschwundene Fundstück“", "A1-E2-01 … A1-E2-08 (Indizienkette 06–08)",
                  werkzeug="Zeitlupe (WHERE-Schritt)",
                  text="Danach wartet **Verhör 1 · Stadtbibliothek** (20 Minuten, ohne Tipps) – im Lab, sobald es "
                       "freigeschaltet ist. Cold Case im Lab: A1-E2-CC „Der Schatten-Finder“.")
    stufe1 = ("SELECT Bezeichnung, Wert, Wert * 0.05 AS Finderlohn\nFROM Fundsache\nWHERE Wert <= 500\n"
              "ORDER BY Finderlohn DESC;")
    stufe2 = ("SELECT Bezeichnung, Wert, 25 + (Wert - 500) * 0.03 AS Finderlohn\nFROM Fundsache\nWHERE Wert > 500;")
    d.cold_case("Finderlohn-Stufen", "Der gesetzliche Finderlohn beträgt 5 % des Werts bis 500 € und vom Mehrwert über "
                "500 € noch 3 %. Entwirf eine Abfrage, die für alle Fundsachen bis 500 € Bezeichnung, Wert und Finderlohn "
                "ausgibt (höchster Finderlohn zuerst). Überlege, warum der Rest mit unseren SQL-Mitteln umständlich wird.",
                kompetenzen=["K20"], linien=5,
                loesung=Loesung(SQL(stufe1), ergebnis(stufe1, max_zeilen=3),
                                "Für Fundsachen über 500 € gilt 25 € + 3 % vom Mehrwert – das braucht eine zweite Abfrage "
                                "(in der Fundbuero-DB gibt es keine solche Fundsache, das Ergebnis ist leer):",
                                SQL(stufe2),
                                "Eine Fallunterscheidung in *einer* Abfrage kennt das Abitur-SQL nicht – deshalb zwei Abfragen. "
                                "Fundsachen ohne bekannten Wert (NULL) fehlen in beiden."))
    name = ("5_Loesungen/DD_A1_AB2_Das_verschwundene_Fundstueck_LSG.docx" if lsg
            else "2_Arbeitsblaetter/DD_A1_AB2_Das_verschwundene_Fundstueck.docx")
    return d.speichern(ORDNER / name)


# ================================================================== Präsentation „Das Fundbüro“ (DS 1 + DS 2)
BRIEFING1A = ("Moin, Ermittlerteam – ich bin Edda Lorenz, 31 Jahre Kriminalpolizei und ab heute eure Chefin. Die "
              "Datendetektei Weserstedt hat ihren ersten Auftrag: Gesa Ohlendiek vom Fundbüro bittet uns, ihr Fundbuch zu "
              "prüfen, weil dort „etwas nicht stimmt“. In der Stadt reden die Leute seit Wochen über kleine gezeichnete "
              "Nachtfalter an merkwürdigen Orten, und das Weserstedter Tageblatt nennt den Unbekannten schon „das Phantom“.")
BRIEFING1B = ("Ob das zusammenhängt, weiß ich nicht – und Raten ist in diesem Haus verboten. Bevor wir suchen, lernen wir "
              "lesen: Das Fundbuch ist eine Tabelle, jede Zeile ein Datensatz, jede Spalte ein Attribut, und die FundNr ist "
              "der Primärschlüssel. Gleich bekommt jede und jeder von euch eine Karteikarte, denn heute seid ihr selbst die "
              "Datensätze. Eure Detektei-Ausweise holt ihr euch über das Portal – damit meldet ihr euch im Lab an.")
BRIEFING2A = ("Gesa Ohlendiek hat uns ihre Notizen gegeben, und sie sind – freundlich gesagt – lückenhaft: „irgendwas mit "
              "Messing“, „Juni oder Juli“, „mehrere hundert Euro“. Mit einem einfachen Gleichheitszeichen kommen wir da "
              "nicht weit; wir brauchen Werkzeuge für ungefähre Angaben: LIKE für Textmuster, BETWEEN für Zeiträume, IN für "
              "Wertelisten.")
BRIEFING2B = ("Und wir müssen mit Lücken umgehen: Wo im Fundbuch nichts steht, steht NULL, und NULL prüft man mit IS NULL. "
              "Merkt euch: Text gehört in einfache Hochkommata, ein Datum schreibt man als 'JJJJ-MM-TT', und „ungleich“ heißt "
              "im Abitur !=. Im Lab kreist ihr das Fundstück mit einer Kette aus drei Abfragen ein. Wer fertig ist, stellt "
              "sich dem ersten Verhör – meinem Test, ob ihr das auch in einem fremden Archiv könnt.")


def praesentation():
    p = AktenPraesentation(1, "Das Fundbüro", datenbank=DB)
    # ---------------------------------------------------------------- DS 1
    p.titelfolie("Das Fundbüro", "Akte 1 · DS 1 „Der erste Auftrag“ · DS 2 „Das verschwundene Fundstück“",
                 notizen="DS 1, Kick-off (0–10'): Statt Warm-up. Teams zu 3–4 bilden, Detektei-Ausweise (Kopiervorlage "
                         "3_Unplugged) ausfüllen lassen: Teamname, Ermittlername. Portal-Login vorbereiten (Kursliste mit "
                         "persönlichen Codes bereithalten).")
    p.briefing("Der erste Auftrag", BRIEFING1A,
               notizen="Briefing vorlesen (lassen). Kurz nachfragen: Was ist ein Fundbüro, was steht wohl im Fundbuch?")
    p.briefing("Der erste Auftrag (Fortsetzung)", BRIEFING1B,
               notizen="Die Fachbegriffe im Briefing (Tabelle, Datensatz, Attribut, Primärschlüssel) werden gleich am "
                       "Fundbuch beschriftet – hier nur ankündigen.")
    p.inhaltsfolie("Impuls", "**Wo werden über euch Daten in Tabellen gespeichert – und woran erkennt man euch dort "
                   "eindeutig?**", Liste(["Schule: Kursliste, Notenverwaltung …", "Verein, Bibliothek, Handyvertrag …",
                                          "Erkennungsmerkmal: Name? Geburtsdatum? Nummer?"]),
                   notizen="Ideen sammeln (2–3 min). Ziel: „Name reicht nicht – es gibt Namensgleiche“ → Überleitung zum "
                           "Primärschlüssel. Antworten an der Tafel stehen lassen.")
    kopf, zeilen = abfrage_ergebnis("SELECT FundNr, Bezeichnung, Kategorie, Farbe, Stadtteil, Funddatum, Wert\n"
                                    "FROM Fundsache\nWHERE FundNr BETWEEN 2 AND 9;", DB)
    p.inhaltsfolie("Das Fundbuch (Ausschnitt)", Tabelle(kopf, zeilen, pk=("FundNr",)),
                   "Was ist hier die Tabelle, ein Datensatz, ein Attribut, ein Attributwert?",
                   notizen="Briefing-Phase (10–20'): Begriffe am Ausschnitt sammeln und beschriften (Beamer + Tafel). Impuls: "
                           "„Welche Spalte würdet ihr nehmen, um genau eine Fundsache zu bezeichnen?“ – „Im Ausschnitt ist "
                           "auch die Bezeichnung eindeutig. Reicht das?“ (Ausschnitt-Falle anbahnen: Es gibt im ganzen "
                           "Fundbuch mehrere Regenjacken, Schals …). Noch nicht auflösen.")
    p.begriffe("Begriffe am Fundbuch", [("Tabelle (Relation)", "Fundsache – das ganze Fundbuch"),
                                         ("Datensatz", "eine Zeile = eine Fundsache"),
                                         ("Attribut", "eine Spalte, z. B. Kategorie"),
                                         ("Attributwert", "ein Eintrag, z. B. 'Kleidung'"),
                                         ("Primärschlüssel", "FundNr: eindeutig für **alle möglichen** Daten, nie NULL, minimal")],
               inhalt=[Schema(FUNDSACHE)],
               notizen="Tafelbild Teil 1 (Merkkasten). Kurzschreibweise: Tabellenname, in Klammern die Attribute, "
                       "Primärschlüssel unterstrichen (später: Fremdschlüssel mit ↑).")
    p.inhaltsfolie("Primärschlüssel: Ein Ausschnitt beweist nichts",
                   Liste(["**eindeutig** – für alle möglichen Daten, auch für die Fundsachen von morgen",
                          "**nie NULL** – jeder Datensatz hat einen Wert", "**minimal** – nichts Überflüssiges"],
                         nummeriert=True),
                   "Gegenfrage: *Kann morgen eine zweite Fundsache „Regenschirm“ abgegeben werden?*",
                   notizen="Typischer Fehler P8: „Im Ausschnitt eindeutig, also Schlüssel.“ Gegenbeispiel-Frage stellen lassen. "
                           "Funddatum: an einem Tag mehrere Funde. FundNr: vom Fundbüro fortlaufend vergeben.")
    p.unplugged("Karteikarten-Filter",
                ["Jede/r zieht eine Karteikarte – ihr seid die Datensätze der Tabelle Fundsache.",
                 "WHERE = aufstehen, wer passt · SELECT = nur die genannten Spalten vorlesen.",
                 "Vergleiche testen: `>` und `>=` – wer steht mit genau 20 € zusätzlich auf?",
                 "AND und OR live: erst per Handzeichen vermuten, dann auflösen – danach mit Klammern.",
                 "ORDER BY: Die Stehenden bilden eine Reihe – auch zweistufig (Stadtteil, dann Wert)."],
                material="30 Karteikarten (Kopiervorlage)", dauer="15 MIN",
                notizen="Abfragen und erwartete Ergebnisse: Lehrerhinweise, Abschnitt „Unplugged-Ablauf DS 1“. Schild "
                        "„FROM Fundsache“ an die Tafel. Die Abfragen nacheinander anschreiben bzw. von der nächsten Folie "
                        "zeigen. Sicherung: Foto der Reihe; Merksätze „SELECT wählt Spalten – WHERE wählt Zeilen – "
                        "ORDER BY ordnet“, „AND vor OR → im Zweifel Klammern“.")
    p.inhaltsfolie("Unplugged: die Abfragen",
                   SQL("SELECT Bezeichnung\nFROM Fundsache\nWHERE Kategorie = 'Schirm';", db="keine"),
                   SQL("SELECT Bezeichnung\nFROM Fundsache\nWHERE Wert > 20\n  AND Stadtteil = 'Altstadt';", db="keine"),
                   SQL("SELECT Bezeichnung\nFROM Fundsache\nWHERE Kategorie = 'Tasche'\n   OR Kategorie = 'Rucksack'\n"
                       "  AND Wert > 50;", db="keine"),
                   karte=False,
                   notizen="Die Karten sind eigene Daten (nicht die Lab-DB) – daher Kategorien wie 'Schirm', 'Tasche'. "
                           "Abfrage 2 danach mit >= wiederholen (Daunenjacke und Haustürschlüssel, je 20 €, stehen dann "
                           "zusätzlich). Abfrage 3: Die Handtasche für 15 € steht auch – AND bindet stärker! Danach mit "
                           "Klammern um die OR-Bedingung. Anschließend ORDER BY Wert DESC (Reihe) und ORDER BY Stadtteil, "
                           "Wert DESC.")
    bsp = ("SELECT Bezeichnung, Stadtteil, Wert\nFROM Fundsache\nWHERE (Kategorie = 'Kleidung'\n    OR Kategorie = 'Dokumente')\n"
           "  AND Wert > 30\nORDER BY Stadtteil, Wert DESC;")
    k, z = abfrage_ergebnis(bsp, DB)
    p.beispiel_schritte("Beispiel Schritt für Schritt", bsp,
                        ["**FROM** – alle 80 Zeilen der Tabelle Fundsache",
                         "**WHERE** – Kleidung oder Dokumente (Klammern!) und Wert über 30 €",
                         "**ORDER BY** – erst nach Stadtteil, bei Gleichstand teuerste zuerst",
                         "**SELECT** – nur Bezeichnung, Stadtteil und Wert zeigen"],
                        ergebnis=Ergebnis(k, z[:5], abgeschnitten=len(z) > 5),
                        notizen=f"Tafelbild Teil 2. Reihenfolge der Auswertung betonen: FROM → WHERE → ORDER BY → SELECT "
                                f"(Werkzeug Zeitlupe im Lab). Ergebnis: {len(z)} Zeilen, hier gekürzt. Frage: Was passiert "
                                f"ohne Klammern? (Dann alle Kleidungsstücke, auch billige, plus teure Dokumente.)")
    p.lab_auftrag("1.1", "Der erste Auftrag", "A1-E1-01 … 08 (★ → ★★ → ★★★)", werkzeug="Zeitlupe (WHERE-Schritt)",
                  hinweise=["Portal → Datendetektei → Stadtplan → Fundbüro → Ermittlung 1.1",
                            "Erst „Ausführen“, dann „Prüfen“ – bei Fehlern die Zeitlupe öffnen.",
                            "Parallel: AB 1.1, Aufträge 1–4 auf Papier."],
                  notizen="Lab-Einsatz (35–80'): Optional Video 0 (3 min) für die Bedienung. Herumgehen, typische Fehler "
                          "sammeln (P5 Text ohne Hochkommata, P6 Klammern) für die Fallbesprechung. Schnelle: Cold Case "
                          "A1-E1-CC, AB-Cold-Case.")
    p.merkfolie("Ermittler-Notiz", ["SELECT wählt Spalten – WHERE wählt Zeilen – ORDER BY ordnet (ASC ist Standard).",
                                    "AND bindet stärker als OR – Klammern setzen!",
                                    "Text steht in `'…'`, Zahlen ohne Hochkommata.",
                                    "ORDER BY mehrstufig: erst nach dem ersten Kriterium, bei Gleichstand nach dem zweiten.",
                                    "Primärschlüssel: eindeutig für alle möglichen Daten, nie NULL, minimal."],
                notizen="Fallbesprechung (80–90'): Merkkasten an der Tafel sichern, ins Heft übernehmen lassen.")
    p.abi_format("Gegeben ist die Tabelle Fundsache. **Entwerfen** Sie eine SQL-Abfrage, die Bezeichnung und Wert aller "
                 "Fundsachen ausgibt, die im Stadtteil Altstadt gefunden wurden und einen Wert von mehr als 20 € haben.",
                 be=2, operator="entwerfen",
                 notizen="Abi-Ecke von AB 1.1: 2 Minuten Einzelarbeit, dann Kurzbewertung im Team (1 BE SELECT/FROM, "
                         "1 BE Bedingung). Operator „entwerfen“ = eine Lösung selbstständig erstellen.")
    abi1 = "SELECT Bezeichnung, Wert\nFROM Fundsache\nWHERE Stadtteil = 'Altstadt'\n  AND Wert > 20;"
    p.loesungsfolie("Lösung Abi-Format (DS 1)", SQL(abi1), ergebnis(abi1, max_zeilen=4),
                    notizen="Erst nach der Einzelarbeit zeigen (6 Zeilen, gekürzt). Ball und Helm (je genau 20 €) gehören nicht dazu: „mehr als“ "
                            "heißt >.")
    p.briefing("Cliffhanger", "Kurz vor Feierabend ruft Gesa Ohlendiek noch einmal an: „Laut Fundbuch wurde etwas "
               "Wertvolles abgeholt – aber in meinem Ausgabebuch steht keine Unterschrift.“", von="Fundbüro",
               notizen="Abschluss-Satz DS 1. Hausaufgabe: Video 1 (SQL-Grundlagen), Rest von AB 1.1.")
    # ---------------------------------------------------------------- DS 2
    p.inhaltsfolie("Warm-up: Karten W-A1-01 … 04",
                   Liste(["Einzelarbeit, je Karte 2–3 Minuten", "Dann Vergleich im Team", "Karte 02 (Klammer-Check) lösen wir gemeinsam"]),
                   "Begriffe am Fundbuch · Klammer-Check · Zweistufig sortieren · Schlüsselfrage",
                   notizen="DS 2, Warm-up (0–10'): Warm-up-Karten A1 (Papier) austeilen, optional Leitner im Lab. "
                           "Karte 02 an der Tafel auflösen (nächste Folie, nur Lehrkraft).")
    wu = ("SELECT FundNr\nFROM Fundsache\nWHERE FundNr IN (4, 9, 10, 31, 73, 80)\n"
          "  AND (Kategorie = 'Sonstiges' OR Farbe = 'rot' AND Wert > 10);")
    wu_k = ("SELECT FundNr\nFROM Fundsache\nWHERE FundNr IN (4, 9, 10, 31, 73, 80)\n"
            "  AND (Kategorie = 'Sonstiges' OR Farbe = 'rot')\n  AND Wert > 10;")
    p.loesungsfolie("Lösung Warm-up W-A1-02: Klammer-Check",
                    "Wie notiert (AND zuerst): **10, 31, 73, 80** · mit Klammern um OR: **10, 73**",
                    SQL(wu), ergebnis(wu),
                    notizen="Die FundNr-Liste grenzt die Probe auf die sechs Karten-Zeilen ein. Mit Klammern: "
                            f"{', '.join(r[0] for r in abfrage_ergebnis(wu_k, DB)[1])}. Bezug zum NULL-Thema: FundNr 31 "
                            "hat keine Farbe (NULL) – steht trotzdem wegen 'Sonstiges'.")
    p.briefing("Das verschwundene Fundstück", BRIEFING2A,
               notizen="Briefing (10–20'). Gesa Ohlendieks Notizen an die Tafel: „Messing“, „Juni oder Juli“, „mehrere "
                       "hundert Euro“ – sie bleiben dort für die Fallbesprechung stehen.")
    p.briefing("Das verschwundene Fundstück (Fortsetzung)", BRIEFING2B,
               notizen="Impulse: „Wie findet ihr alles aus Messing, wenn ihr nicht wisst, ob ‚Messingleuchter‘ oder "
                       "‚Kompass (Messing)‘ eingetragen ist?“ – „Was bedeutet ein leeres Feld AbgeholtAm?“")
    p.begriffe("Neue Werkzeuge", [("`LIKE 'Schl%'`", "Textmuster: `%` beliebig viele Zeichen, `_` genau eins"),
                                  ("`BETWEEN 10 AND 20`", "Bereich – beide Grenzen inklusive"),
                                  ("`IN ('a', 'b')`", "Werteliste, z. B. `Kategorie IN ('Schmuck', 'Elektronik')`"),
                                  ("`IS NULL`", "kein Eintrag – niemals `= NULL`"),
                                  ("`DISTINCT`", "doppelte Zeilen weglassen"),
                                  ("`… AS Finderlohn`", "Spaltenalias, z. B. für `Wert * 0.05`"),
                                  ("`LIMIT 3`", "nur die ersten 3 Zeilen (nach ORDER BY)")],
               notizen="Tafelbild DS 2 (Merkkasten). Jedes Werkzeug mit einem Beispiel an der Tafel, Schüler formulieren "
                       "die Bedeutung selbst.")
    p.begriffe("Literale – so schreibt man Werte", [("Text", "`'Hafenviertel'` – in einfachen Hochkommata"),
                                                    ("Datum", "`'2026-09-30'` – Text im Format JJJJ-MM-TT, darum vergleichbar"),
                                                    ("Zahl", "`0.05` – Dezimalpunkt, ohne Hochkommata"),
                                                    ("ungleich", "`!=` – im Abitur statt `<>`"),
                                                    ("Division", "`7 / 2` ergibt bei ganzen Zahlen `3`")],
               notizen="Typischer Fehler: Datum ohne Hochkommata – 2026-09-01 wird als Rechnung 2026 − 9 − 1 gelesen! "
                       "Deutsches Datumsformat '01.09.2026' sortiert falsch.")
    like = "SELECT FundNr, Bezeichnung, Wert\nFROM Fundsache\nWHERE Bezeichnung LIKE '%uhr%'\n  AND Wert IS NULL;"
    p.beispiel_schritte("Beispiel: Uhren ohne Wert", like,
                        ["**FROM** – alle Fundsachen", "**WHERE** – „uhr“ irgendwo in der Bezeichnung (`%` vorn und hinten)",
                         "… **und** kein Wert eingetragen (`IS NULL`)", "**SELECT** – FundNr, Bezeichnung, Wert"],
                        ergebnis=ergebnis(like),
                        notizen="Gegenprobe im Gespräch: Was liefert `Wert = NULL`? (Keine Zeile.) Was liefert `LIKE 'uhr'`? "
                                "(Nur die Bezeichnung „uhr“ – ohne Platzhalter wirkt LIKE wie =.)")
    lohn = ("SELECT Bezeichnung, Wert, Wert * 0.05 AS Finderlohn\nFROM Fundsache\nWHERE Kategorie IN ('Taschen', 'Spielzeug')\n"
            "ORDER BY Finderlohn DESC\nLIMIT 3;")
    p.beispiel_schritte("Beispiel: Finderlohn der Top 3", lohn,
                        ["**WHERE** – nur Taschen und Spielzeug (IN)", "**SELECT** berechnet 5 % als Spalte Finderlohn",
                         "**ORDER BY** über den Alias – höchster Finderlohn zuerst", "**LIMIT 3** – nur die ersten drei Zeilen"],
                        ergebnis=ergebnis(lohn),
                        notizen="LIMIT nur, wenn an der Schnittkante kein Gleichstand ist (hier 3. und 4. Platz verschieden). "
                                "Alias in WHERE (WHERE Finderlohn > 5) geht in SQLite, im Abitur nicht – Berechnung wiederholen.")
    p.unplugged("Fahndung mit Platzhaltern",
                ["Teams ordnen jeder Musterkarte alle passenden Wortkarten zu (Wettlauf, 4 Minuten).",
                 "Streitfälle klären: Passt „Uhrenkette“ zu `'%uhr'`? Und zu `'%uhr%'`?",
                 "NULL-Stuhl: Eine Person hält die NULL-Karte als Farbe. `WHERE Farbe = NULL` – wer steht auf?",
                 "Dann `WHERE Farbe IS NULL` und `WHERE Farbe != 'rot'` – wer bleibt sitzen?",
                 "BETWEEN 10 AND 20 mit den Wertkarten 10 · 15 · 20 · 21: Wer gehört dazu?"],
                material="Wort-/Musterkarten je Team, NULL-Karte", dauer="10 MIN",
                notizen="Kopiervorlage DD_A1_Unplugged_Platzhalter-Fahndung; Lösungen in den Lehrerhinweisen. Sicherung: "
                        "„% beliebig viele, _ genau ein Zeichen“ · „NULL nur mit IS NULL“ · „BETWEEN schließt beide Grenzen ein“.")
    p.lab_auftrag("1.2", "Das verschwundene Fundstück", "A1-E2-01 … 05 (Werkzeuge), dann Indizienkette 06 → 07 → 08",
                  werkzeug="Zeitlupe (WHERE-Schritt)",
                  hinweise=["Jede Abfrage der Indizienkette baut auf der vorigen auf – Bedingungen ergänzen, nicht neu raten.",
                            "Parallel: AB 1.2. Wer fertig ist: Verhör 1 · Stadtbibliothek (20 Minuten)."],
                  notizen="Lab-Einsatz (30–75'). Verhör 1 im Kurs auf „auto“ (ab 60 % der Akte-1-Punkte) oder „offen“ "
                          "stellen, wenn es Hausaufgabe bis DS 3 ist. Typische Fehler sammeln: = NULL, Datum ohne "
                          "Hochkommata, LIKE 'Messing%' statt '%Messing%'.")
    p.merkfolie("Ermittler-Notiz", ["`%` = beliebig viele Zeichen (auch keins), `_` = genau ein Zeichen.",
                                    "NULL heißt „weiß keiner“ – nur `IS NULL` findet es, `= NULL` ist nie wahr.",
                                    "BETWEEN schließt beide Grenzen ein.",
                                    "Text `'…'` · Datum `'JJJJ-MM-TT'` · Zahl mit Punkt · ungleich `!=`.",
                                    "LIMIT nur nach ORDER BY und ohne Gleichstand an der Schnittkante."],
                notizen="Fallbesprechung (75–90'). Merkkasten Literale/NULL sichern.")
    p.abi_format("Gegeben ist die Tabelle Fundsache. **Entwerfen** Sie eine SQL-Abfrage, die Bezeichnung und Funddatum aller "
                 "noch nicht abgeholten Fundsachen ausgibt, die im September 2026 gefunden wurden. Die neuesten Funde sollen "
                 "zuerst erscheinen.", be=3, operator="entwerfen",
                 notizen="Abi-Ecke von AB 1.2 im Team, Kurzvergleich. 1 BE IS NULL, 1 BE Zeitraum mit korrekten Literalen, "
                         "1 BE Sortierung.")
    abi2 = ("SELECT Bezeichnung, Funddatum\nFROM Fundsache\nWHERE AbgeholtAm IS NULL\n"
            "  AND Funddatum BETWEEN '2026-09-01' AND '2026-09-30'\nORDER BY Funddatum DESC;")
    p.loesungsfolie("Lösung Abi-Format (DS 2)", SQL(abi2), ergebnis(abi2, max_zeilen=4),
                    notizen="Ergebnis: 6 Zeilen (gekürzt). Auch richtig: Funddatum >= '2026-09-01' AND Funddatum <= '2026-09-30' bzw. LIKE '2026-09-%'. "
                            "Typisch falsch: AbgeholtAm = NULL; < '2026-09-30' (der 30. fehlt).")
    s1 = "SELECT FundNr, Bezeichnung, Funddatum, Wert\nFROM Fundsache\nWHERE Bezeichnung LIKE '%Messing%';"
    s3 = ("SELECT Bezeichnung, Fundort, Wert, AbgeholtAm, Finder\nFROM Fundsache\nWHERE Bezeichnung LIKE '%Messing%'\n"
          "  AND Funddatum BETWEEN '2026-06-01' AND '2026-07-31'\n  AND (Wert >= 400 OR Wert IS NULL)\n"
          "  AND NOT AbgeholtAm IS NULL;")
    p.loesungsfolie("Indizienkette rekonstruiert (1/2)", SQL(s1), ergebnis(s1),
                    notizen="Fallbesprechung: „Welche Bedingung hat den entscheidenden Datensatz verraten?“ Schritt 1: vier "
                            "Messing-Funde, zwei außerhalb Juni/Juli, einer mit unbekanntem Wert.")
    p.loesungsfolie("Indizienkette rekonstruiert (2/2)", SQL(s3), ergebnis(s3),
                    notizen="Schritte 2 und 3: Zeitraum + Wert (mindestens 400 € oder unbekannt), dann „abgeholt“. Übrig: "
                            "Messing-Kompass, 450 €, Weserpromenade, abgeholt am 14.09. – Finder 'N. Falter'.")
    p.briefing("Cliffhanger", "Im Feld Finder steht „N. Falter“ – sprecht es einmal schnell aus, dann wisst ihr, wer im "
               "Fundbuch war: das Phantom mit dem Nachtfalter.", nur_lehrkraft=True,
               notizen="Abschluss-Satz DS 2 (verrät das Lab-Ergebnis, daher nicht im Handout). Steckbrief an der Pinnwand "
                       "beginnen: Initialen „N. F.“, Signatur Nachtfalter, Zugriff auf städtische Daten. Hinweis Verhör 1 "
                       "(20 Minuten, bis DS 3 im Lab).")
    p.speichern(ORDNER / "1_Praesentation" / "DD_A1_Fundbuero.pptx")


# ================================================================== Unplugged DS 1: Karteikarten-Filter (eigene Daten)
# FundNr, Bezeichnung, Kategorie, Farbe, Stadtteil, Funddatum, Wert – erfunden, nicht die Lab-DB (DS_PLANUNG DS 1)
KARTEN = [
    (101, "Stockschirm", "Schirm", "schwarz", "Altstadt", "2026-09-21", 12),
    (102, "Taschenschirm", "Schirm", "rot", "Hafenviertel", "2026-09-21", 8),
    (103, "Kinderschirm", "Schirm", "gelb", "Lindenhof", "2026-09-22", 6),
    (104, "Handtasche", "Tasche", "braun", "Altstadt", "2026-09-22", 15),
    (105, "Ledertasche", "Tasche", "schwarz", "Weserufer", "2026-09-23", 70),
    (106, "Laptoptasche", "Tasche", "grau", "Bahnhofsviertel", "2026-09-23", 45),
    (107, "Wanderrucksack", "Rucksack", "grün", "Mühlenfeld", "2026-09-24", 80),
    (108, "Schulrucksack", "Rucksack", "blau", "Nordheide", "2026-09-24", 30),
    (109, "Daunenjacke", "Kleidung", "blau", "Altstadt", "2026-09-25", 20),
    (110, "Wollmütze", "Kleidung", "rot", "Hafenviertel", "2026-09-25", 9),
    (111, "Regenmantel", "Kleidung", "gelb", "Altstadt", "2026-09-26", 35),
    (112, "Schlüsselbund", "Schlüssel", "silber", "Lindenhof", "2026-09-26", 25),
    (113, "Fahrradschlüssel", "Schlüssel", "schwarz", "Hafenviertel", "2026-09-27", 15),
    (114, "Haustürschlüssel", "Schlüssel", "gold", "Altstadt", "2026-09-28", 20),
    (115, "Smartphone", "Elektronik", "schwarz", "Bahnhofsviertel", "2026-09-28", 140),
    (116, "Kopfhörer", "Elektronik", "weiß", "Altstadt", "2026-09-29", 60),
    (117, "Tablet", "Elektronik", "silber", "Nordheide", "2026-09-30", 180),
    (118, "Silberring", "Schmuck", "silber", "Hafenviertel", "2026-09-30", 55),
    (119, "Halskette", "Schmuck", "gold", "Altstadt", "2026-10-01", 90),
    (120, "Armbanduhr", "Schmuck", "silber", "Weserufer", "2026-10-01", 75),
    (121, "Teddybär", "Spielzeug", "braun", "Lindenhof", "2026-10-02", 10),
    (122, "Fußball", "Spielzeug", "weiß", "Mühlenfeld", "2026-10-03", 18),
    (123, "Monatskarte", "Dokumente", "blau", "Bahnhofsviertel", "2026-10-03", 40),
    (124, "Notizbuch", "Dokumente", "grün", "Altstadt", "2026-10-04", 5),
    (125, "Sporttasche", "Tasche", "rot", "Mühlenfeld", "2026-10-05", 55),
    (126, "Kinderrucksack", "Rucksack", "rot", "Lindenhof", "2026-10-06", 22),
    (127, "Sonnenbrille", "Sonstiges", "braun", "Weserufer", "2026-10-07", 25),
    (128, "Trinkflasche", "Sonstiges", "blau", "Nordheide", "2026-10-08", 12),
    (129, "Fernglas", "Wertsachen", "schwarz", "Hafenviertel", "2026-10-09", 120),
    (130, "Regenschirm", "Schirm", "blau", "Bahnhofsviertel", "2026-10-10", 15),
]
SPALTEN = ["FundNr", "Bezeichnung", "Kategorie", "Farbe", "Stadtteil", "Funddatum", "Wert"]

#: Abfragen des Karteikarten-Filters (DS_PLANUNG DS 1, Unplugged Schritt 2–5)
UNPLUGGED_ABFRAGEN = [
    ("1 · Projektion und Selektion", "SELECT Bezeichnung\nFROM Fundsache\nWHERE Kategorie = 'Schirm';"),
    ("2a · Grenze mit >", "SELECT Bezeichnung\nFROM Fundsache\nWHERE Wert > 20\n  AND Stadtteil = 'Altstadt';"),
    ("2b · Grenze mit >=", "SELECT Bezeichnung\nFROM Fundsache\nWHERE Wert >= 20\n  AND Stadtteil = 'Altstadt';"),
    ("3a · AND vor OR", "SELECT Bezeichnung\nFROM Fundsache\nWHERE Kategorie = 'Tasche'\n   OR Kategorie = 'Rucksack'\n"
                        "  AND Wert > 50;"),
    ("3b · mit Klammern", "SELECT Bezeichnung\nFROM Fundsache\nWHERE (Kategorie = 'Tasche' OR Kategorie = 'Rucksack')\n"
                          "  AND Wert > 50;"),
    ("4 · Reihe bilden", "SELECT Bezeichnung, Wert\nFROM Fundsache\nWHERE Wert >= 20\n  AND Stadtteil = 'Altstadt'\n"
                         "ORDER BY Wert DESC;"),
    ("5 · zweistufig", "SELECT Bezeichnung, Stadtteil, Wert\nFROM Fundsache\nWHERE Kategorie = 'Schmuck'\n"
                       "   OR Kategorie = 'Elektronik'\nORDER BY Stadtteil, Wert DESC;"),
]


def karten_db():
    con = sqlite3.connect(":memory:")
    con.execute("CREATE TABLE Fundsache (FundNr INTEGER PRIMARY KEY, Bezeichnung TEXT, Kategorie TEXT, Farbe TEXT, "
                "Stadtteil TEXT, Funddatum TEXT, Wert REAL)")
    con.executemany("INSERT INTO Fundsache VALUES (?, ?, ?, ?, ?, ?, ?)", KARTEN)
    return con


def unplugged_ergebnisse():
    """Erwartete Ergebnisse der Unplugged-Abfragen (für die Lehrerhinweise) – berechnet mit SQLite."""
    con = karten_db()
    aus = []
    for titel, sql in UNPLUGGED_ABFRAGEN:
        cur = con.execute(sql)
        kopf = [c[0] for c in cur.description]
        zeilen = [[A.zelle_text(v) for v in z] for z in cur.fetchall()]
        aus.append((titel, sql, kopf, zeilen))
    # Plausibilität (Blueprint): 3a enthält die billige Handtasche, 3b nicht; 2b hat genau zwei Zeilen mehr als 2a
    e = {t: [z[0] for z in zz] for t, _, _, zz in aus}
    assert "Handtasche" in e["3a · AND vor OR"] and "Handtasche" not in e["3b · mit Klammern"]
    assert len(e["2b · Grenze mit >="]) == len(e["2a · Grenze mit >"]) + 2
    return aus


def _karten_raster(d, zellen_inhalt, spalten, zeilen_je_seite, hoehe_cm, rahmen=A.RAND_STARK):
    """Füllt ein Raster aus Karten (Tabelle mit fester Zeilenhöhe, gestrichelte Schnittlinien)."""
    breite = d.breite / spalten
    for start in range(0, len(zellen_inhalt), spalten * zeilen_je_seite):
        if start:
            d.seitenumbruch()
        seite = zellen_inhalt[start:start + spalten * zeilen_je_seite]
        n_zeilen = (len(seite) + spalten - 1) // spalten
        tab = d.doc.add_table(rows=n_zeilen, cols=spalten)
        A._tabelle_breiten(tab, [breite] * spalten)
        A._tabelle_ohne_rahmen(tab)
        for zi in range(n_zeilen):
            zeile = tab.rows[zi]
            trpr = zeile._tr.get_or_add_trPr()
            trpr.append(A._el("w:trHeight", val=int(Cm(hoehe_cm) / 635), hRule="exact"))
            A._zeile_nicht_teilen(zeile)
            for si in range(spalten):
                z = zeile.cells[si]
                A._zelle_rahmen(z, rahmen, 0.75, art="dashed")
                A._zelle_innenabstand(z, 0.35, 0.45, 0.3, 0.45)
                i = zi * spalten + si
                z.paragraphs[0].style = d.doc.styles["DD Tabellentext"]
                if i < len(seite):
                    seite[i](z, breite)


def _karteikarte(d, k):
    def fuellen(z, breite):
        p = z.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Cm(breite - 1.0), A.WD_TAB_ALIGNMENT.RIGHT)
        d._etikett(p, "FUNDBUCH · KARTEIKARTE", groesse=7.5, farbe=A.SIEGELROT, fett=True)
        d._etikett(p, f"\tNR. {k[0]}", groesse=8, farbe=A.TINTE_2, fett=True)
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(10)
        d._lauf_stil(p.add_run(k[1]), schrift=A.SCHRIFT_KOPF, groesse=26)
        tab = z.add_table(rows=len(SPALTEN), cols=2)
        A._tabelle_breiten(tab, [3.3, breite - 4.4])
        for i, (name, wert) in enumerate(zip(SPALTEN, k)):
            a, b = tab.rows[i].cells
            for c in (a, b):
                A._zelle_rahmen(c, A.RAND, 0.5, seiten=("bottom",))
                A._zelle_innenabstand(c, 0.1, 0.1, 0.1, 0.1)
            A._zelle_flaeche(a, A.FLAECHE)
            pa = a.paragraphs[0]
            pa.style = d.doc.styles["DD Tabellentext"]
            d._text(pa, f"<u>{name}</u>" if name == "FundNr" else name, code=True, groesse=10)
            pb = b.paragraphs[0]
            pb.style = d.doc.styles["DD Tabellentext"]
            d._text(pb, str(wert), groesse=13, fett=name in ("Kategorie", "Stadtteil", "Wert"))
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(8)
        d._etikett(p, "TABELLE FUNDSACHE · DATENDETEKTEI WESERSTEDT", groesse=6.5)
    return fuellen


def unplugged_karteikarten():
    d = AktenDokument("unplugged", 1, "Karteikarten-Filter", fassung="schueler", ermittlung="1.1", ds=1,
                      material="Unplugged A1 · Karteikarten", datenbank="keine",
                      untertitel="Kopiervorlage: 30 Karteikarten (vier je Seite, an den gestrichelten Linien schneiden, "
                                 "laminieren) und das Tafelschild „FROM Fundsache“")
    d.kopfband(name_datum=False)
    d.absatz("**So geht's:** Jede Karte ist ein Datensatz der Tabelle Fundsache (eigene Beispieldaten, nicht das Fundbuch "
             "im Lab). Jede/r zieht eine Karte. Die Lehrkraft ruft eine Abfrage auf: Wer die WHERE-Bedingung erfüllt, "
             "**steht auf** und liest **nur die Spalten aus dem SELECT** vor. Bei ORDER BY bilden alle Stehenden eine Reihe.")
    d.notiz("Merksätze", "SELECT wählt Spalten – WHERE wählt Zeilen – ORDER BY ordnet. · AND bindet stärker als OR – im "
            "Zweifel Klammern.")
    d.ueberschrift("Tafelschild (auf A4 quer vergrößern)", 2)
    tab, z = d._kasten(d.doc, d.breite, rand=A.TINTE, staerke=2.0, flaeche=A.KARTE)
    p = z.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(18), Pt(18)
    d._lauf_stil(p.add_run("FROM Fundsache"), code=True, groesse=44, fett=True)
    d._nach_kasten(d.doc)
    d.ueberschrift("Die Tabelle in Kurzschreibweise", 2)
    d.inhalt(Schema("Fundsache(<u>FundNr</u>, Bezeichnung, Kategorie, Farbe, Stadtteil, Funddatum, Wert)"))
    d.seitenumbruch()
    _karten_raster(d, [_karteikarte(d, k) for k in KARTEN], spalten=2, zeilen_je_seite=2, hoehe_cm=12.6)
    return d.speichern(ORDNER / "3_Unplugged" / "DD_A1_Unplugged_Karteikarten.docx")


# ================================================================== Unplugged DS 2: Fahndung mit Platzhaltern
WOERTER = ["Schal", "Schlüssel", "Schlüsselbund", "Mütze", "Matze", "Taschenuhr", "Armbanduhr", "Uhrenkette",
           "Tischtennisschläger", "Kamera", "Laptop", "Sonnenbrille"]
MUSTER = ["'Sch%'", "'%uhr'", "'%uhr%'", "'M_tze'", "'_a%'", "'Schl_ssel'"]


def fahndung_loesung():
    """Zuordnung Muster → Wörter, berechnet mit SQLite-LIKE (wie in der Plattform)."""
    con = sqlite3.connect(":memory:")
    con.execute("CREATE TABLE Wort (Text TEXT)")
    con.executemany("INSERT INTO Wort VALUES (?)", [(w,) for w in WOERTER])
    aus = []
    for m in MUSTER:
        treffer = [r[0] for r in con.execute(f"SELECT Text FROM Wort WHERE Text LIKE {m}")]
        aus.append((m, treffer))
    zuordnung = dict(aus)
    assert zuordnung["'%uhr'"] == ["Taschenuhr", "Armbanduhr"], zuordnung["'%uhr'"]
    assert "Uhrenkette" in zuordnung["'%uhr%'"]
    assert zuordnung["'M_tze'"] == ["Mütze", "Matze"]
    return aus


def _wortkarte(d, text, etikett, farbe, flaeche=None, groesse=24, code=False):
    def fuellen(z, breite):
        if flaeche:
            A._zelle_flaeche(z, flaeche)
        z.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = z.paragraphs[0]
        d._etikett(p, etikett, groesse=7, farbe=farbe, fett=True)
        p = z.add_paragraph(style="DD Tabellentext")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(8)
        if code:
            d._lauf_stil(p.add_run(text), code=True, groesse=groesse, fett=True, farbe=farbe)
        else:
            d._lauf_stil(p.add_run(text), schrift=A.SCHRIFT_KOPF, groesse=groesse)
    return fuellen


def unplugged_fahndung():
    fahndung_loesung()
    d = AktenDokument("unplugged", 1, "Fahndung mit Platzhaltern", fassung="schueler", ermittlung="1.2", ds=2,
                      material="Unplugged A1 · Platzhalter-Fahndung", datenbank="keine",
                      untertitel="Kopiervorlage je Team: 12 Wortkarten, 6 Musterkarten · für den Kurs: NULL-Karte und "
                                 "4 Wertkarten (NULL-Stuhl, BETWEEN)")
    d.kopfband(name_datum=False)
    d.absatz("**So geht's (4 Minuten, Wettlauf):** Legt jede Musterkarte (rot) aus und ordnet ihr **alle** Wortkarten zu, "
             "die zu dem Muster passen – eine Wortkarte darf zu mehreren Mustern gehören (dann notiert ihr sie). "
             "`%` steht für beliebig viele Zeichen (auch keins), `_` für genau ein Zeichen. Groß-/Kleinschreibung zählt "
             "bei LIKE (für A–Z) nicht.")
    d.tabelle(["Muster", "passende Wortkarten"], [[f"`{m}`", ""] for m in MUSTER], breiten=[4.8, 12.2], groesse=18)
    d.notiz("Merksätze", "`%` beliebig viele, `_` genau ein Zeichen · NULL nur mit IS NULL · BETWEEN schließt beide Grenzen ein.")
    d.seitenumbruch()
    d.ueberschrift("Wortkarten (12 je Team)", 2)
    _karten_raster(d, [_wortkarte(d, w, "WORTKARTE · BEZEICHNUNG", A.TINTE_2, groesse=24 if len(w) <= 14 else 17)
                       for w in WOERTER],
                   spalten=3, zeilen_je_seite=4, hoehe_cm=5.4)
    d.seitenumbruch()
    d.ueberschrift("Musterkarten (6 je Team)", 2)
    _karten_raster(d, [_wortkarte(d, f"LIKE {m}", "MUSTERKARTE", A.SIEGELROT, flaeche=A.ROT_FLAECHE, groesse=20, code=True)
                       for m in MUSTER], spalten=2, zeilen_je_seite=3, hoehe_cm=5.6)
    d.seitenumbruch()
    d.ueberschrift("NULL-Stuhl und BETWEEN (einmal für den Kurs)", 2)
    _karten_raster(d, [_wortkarte(d, "NULL", "FARBE · NICHT ERFASST", A.TINTE_2, flaeche=A.FLAECHE, groesse=72, code=True)],
                   spalten=1, zeilen_je_seite=1, hoehe_cm=9.0)
    _karten_raster(d, [_wortkarte(d, w, "WERTKARTE · WERT", A.BLAU, flaeche=A.BLAU_FLAECHE, groesse=54, code=True)
                       for w in ("10", "15", "20", "21")], spalten=2, zeilen_je_seite=2, hoehe_cm=6.4)
    return d.speichern(ORDNER / "3_Unplugged" / "DD_A1_Unplugged_Platzhalter-Fahndung.docx")


# ================================================================== Detektei-Ausweise (Kick-off DS 1)
def unplugged_ausweise():
    d = AktenDokument("unplugged", 1, "Detektei-Ausweise", fassung="schueler", ermittlung="1.1", ds=1,
                      material="Unplugged A1 · Detektei-Ausweise", untertitel="Kopiervorlage für den Kick-off: acht Ausweise "
                      "je Seite – ausschneiden, ausfüllen, falten oder in eine Hülle stecken")
    d.kopfband(name_datum=False)

    def ausweis(z, breite):
        A._zelle_flaeche(z, A.PAPIER)
        p = z.paragraphs[0]
        d._etikett(p, "DATENDETEKTEI WESERSTEDT · ERMITTLER-AUSWEIS", groesse=7, farbe=A.SIEGELROT, fett=True)
        for feld in ("Ermittlername", "Team", "Kurs"):
            p = z.add_paragraph(style="DD Tabellentext")
            p.paragraph_format.space_before = Pt(9)
            d._etikett(p, feld.upper() + " ", groesse=7)
            d._lauf_stil(p.add_run("_" * (26 if feld != "Kurs" else 14)), farbe=A.RAND_STARK)
        p = z.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(9)
        d._etikett(p, "ANMELDUNG IM LAB: PORTAL → DATENDETEKTEI", groesse=6.5, farbe=A.TINTE_2)
        p = z.add_paragraph(style="DD Tabellentext")
        d._text(p, "*„Erst die Frage, dann die Abfrage.“* – E. Lorenz", groesse=8.5, farbe=A.TINTE_2)

    _karten_raster(d, [ausweis] * 8, spalten=2, zeilen_je_seite=4, hoehe_cm=5.4)
    d.absatz("Hinweis: Den persönlichen Code für das Portal nicht auf den Ausweis schreiben.")
    return d.speichern(ORDNER / "3_Unplugged" / "DD_A1_Unplugged_Detektei-Ausweise.docx")


# ================================================================== Lehrerhinweise
def lehrerhinweise():
    d = AktenDokument("lehrerhinweise", 1, "Lehrerhinweise Akte 1", fassung="lehrkraft", material="Lehrerhinweise A1",
                      datenbank=DB, untertitel="Das Fundbüro · DS 1–2 · gA und eA unverändert")
    d.kopfband()
    d.ueberschrift("Überblick", 1)
    d.tabelle(["DS", "Thema", "Lab", "Warm-up", "Abi-Bezug"],
              [["1", "Der erste Auftrag – Tabelle, Primärschlüssel, SELECT/WHERE/ORDER BY", "A1-E1-01 … 08 + CC",
                "– (Kick-off)", "Schlüsseleignung (14/19 Abiturklausuren); SQL entwerfen in jeder Klausur"],
               ["2", "Das verschwundene Fundstück – LIKE, BETWEEN, IN, IS NULL, Literale, AS, DISTINCT, LIMIT",
                "A1-E2-01 … 08 + CC; Verhör 1 (V1)", "W-A1-01 … 04", "Top-n/LIMIT (E22N, E26); Literale und NULL in fast "
                "jeder SQL-Teilaufgabe"]],
              breiten=[0.9, 5.4, 3.4, 2.4, 4.9])
    d.absatz("Datenbank: **Fundbuero** – eine Tabelle mit 80 Fundsachen (Randfall-Varianten +null, +gleichnamig, +grenze, "
             "+gleichstand prüfen die Lab-Lösungen mit). Kurzschreibweise:")
    d.inhalt(Schema(FUNDSACHE))
    d.ueberschrift("Material dieser Akte", 2)
    d.tabelle(["Datei", "Zweck"],
              [["1_Praesentation/DD_A1_Fundbuero.pptx", "Fallbesprechung DS 1 (Folien 1–15) und DS 2 (Folien 16–31), "
                "Sprechernotizen; Lösungsfolien nur in der Lehrkraftfassung"],
               ["2_Arbeitsblaetter/DD_A1_AB1_…, DD_A1_AB2_…", "AB 1.1 und AB 1.2 (Lösungen in 5_Loesungen, *_LSG)"],
               ["3_Unplugged/…Karteikarten", "Karteikarten-Filter DS 1: 30 Karten + Tafelschild"],
               ["3_Unplugged/…Platzhalter-Fahndung", "DS 2: Wort-/Musterkarten je Team, NULL-Karte, Wertkarten"],
               ["3_Unplugged/…Detektei-Ausweise", "Kick-off DS 1 (Teamname, Ermittlername)"],
               ["Warm-up/DD_Warm-up-Karten_A1.docx", "W-A1-01 … 04 zu Beginn von DS 2, W-A1-05 … 08 zu Beginn von DS 3"]],
              breiten=[6.2, 10.8])

    # ---------------------------------------------------------------- DS 1
    d.ueberschrift("DS 1 · Der erste Auftrag (Ermittlung 1.1)", 1)
    d.ueberschrift("Ziele", 2)
    d.absatz("Die Schülerinnen und Schüler …")
    d.inhalt(Liste(["**erläutern** den Aufbau einer relationalen Tabelle mit den Fachbegriffen Tabelle, Datensatz, Attribut, "
                    "Attributwert und Primärschlüssel am Fundbuch (K01, K47),",
                    "**beurteilen**, ob ein Attribut als Primärschlüssel geeignet ist (eindeutig, nie NULL, minimal), und "
                    "erkennen, dass ein Datenausschnitt keine Eindeutigkeit beweist (K02),",
                    "**entwerfen** SQL-Abfragen mit SELECT, FROM, WHERE (Vergleiche, AND/OR/NOT mit Klammern) und ORDER BY, "
                    "auch mehrstufig (K15, K17, K21),",
                    "**geben** das Ergebnis einer Abfrage auf einem kleinen Ausschnitt konkret **an** (K29)."], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt / Ablauf", "Sozialform", "Material"],
              [["Kick-off", "0–10'", "Briefing Edda Lorenz (Folien 2–3), Teams zu 3–4, Detektei-Ausweise; Impuls „Wo werden "
                "über euch Daten in Tabellen gespeichert – woran erkennt man euch eindeutig?“", "Plenum, Teams",
                "Folien 1–4, Ausweise"],
               ["Briefing", "10–20'", "Fundbuch-Ausschnitt: Begriffe sammeln und beschriften, Primärschlüssel; "
                "Ausschnitt-Falle anbahnen", "LV, UG", "Folien 5–7, Tafel"],
               ["Unplugged", "20–35'", "Karteikarten-Filter (siehe unten)", "Plenum", "30 Karteikarten, Folien 8–9"],
               ["Lab", "35–80'", "Ermittlung 1.1 (★ → ★★ → ★★★), optional Video 0; AB 1.1 Aufträge 1–4 parallel; Fehler "
                "P5/P6 sammeln", "EA", "Geräte, AB 1.1, Folien 10–11"],
               ["Fallbesprechung", "80–90'", "Merkkasten sichern, Abi-Ecke (2 BE) im Team, Cliffhanger; Hausaufgabe: "
                "Video 1, AB-Rest", "UG", "Folien 12–15"]],
              breiten=[2.3, 1.4, 7.2, 2.4, 3.7])
    d.ueberschrift("Unplugged-Ablauf: Karteikarten-Filter (15 Minuten)", 2)
    d.absatz("Material: 30 laminierte Karteikarten (Kopiervorlage, eigene Daten – nicht die Lab-DB), Tafelschild "
             "„FROM Fundsache“. Jede/r zieht eine Karte (kleine Kurse: zwei). Die Lehrkraft schreibt die Abfragen nacheinander "
             "an; wer passt, steht auf und liest **nur** die SELECT-Spalten vor. Bei weniger als 30 Personen die Karten so "
             "auswählen, dass die Karten 104, 105, 107, 109, 111, 114, 116 und 119 im Spiel sind.")
    for titel, sql, kopf, zeilen in unplugged_ergebnisse():
        d.ueberschrift(f"Abfrage {titel}", 3)
        d.inhalt(SQL(sql, db="keine"), Tabelle(kopf, zeilen, groesse=8.5))
    d.notiz("Impulse im Unplugged", "Vor 3a per Handzeichen vermuten lassen, ob die Handtasche für 15 € aufsteht. Bei 4 "
            "stehen Daunenjacke und Haustürschlüssel mit je 20 € gleich – „Wer kommt zuerst?“ (Reihenfolge bei Gleichstand "
            "ist nicht festgelegt) → zweites Sortierkriterium. "
            "Sicherung: Foto der Reihe; Tafel: „SELECT wählt Spalten – WHERE wählt Zeilen – ORDER BY ordnet“, „AND vor OR "
            "→ im Zweifel Klammern“.")
    d.ueberschrift("Tafelbild / Merkkasten „Ermittler-Notiz“", 2)
    d.inhalt(Schema(FUNDSACHE),
             Tabelle(["Fachbegriff", "im Fundbuch"], [["Tabelle (Relation)", "Fundsache"], ["Datensatz", "eine Zeile = eine Fundsache"],
                                                      ["Attribut", "eine Spalte, z. B. Kategorie"],
                                                      ["Attributwert", "ein Eintrag, z. B. 'Kleidung'"],
                                                      ["Primärschlüssel", "FundNr: eindeutig für alle möglichen Daten, nie NULL, minimal"]],
                     breiten=[4.0, 13.0]),
             SQL("SELECT Bezeichnung, Stadtteil, Wert\nFROM Fundsache\nWHERE (Kategorie = 'Kleidung' OR Kategorie = 'Dokumente')\n"
                 "  AND Wert > 30\nORDER BY Stadtteil, Wert DESC;"),
             "**Ermittler-Notiz:** AND bindet stärker als OR – Klammern setzen! · Text in `'…'` · ORDER BY mehrstufig: erst "
             "nach dem ersten Kriterium, bei Gleichstand nach dem zweiten · ASC ist Standard.")
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["Portal-Login scheitert (persönlicher Code vergessen) → Kursliste mit Codes bereithalten; notfalls `/gast` "
                    "(Eisdiele) zum Ausprobieren – zählt nicht.",
                    "Spalten- und Zeilenauswahl verwechselt (SELECT vs. WHERE) → Bild aus dem Unplugged: „aufstehen = WHERE, "
                    "vorlesen = SELECT“.",
                    "AND/OR ohne Klammern (**P6**) → Zeitlupe zeigt die verworfenen Zeilen; Merkkasten.",
                    "Text ohne Hochkommata (**P5**) → deutsche Fehlermeldung „Die Spalte Kleidung gibt es nicht …“ gemeinsam deuten:",
                    "Primärschlüssel aus dem Ausschnitt „bewiesen“ (**P8**) → Gegenbeispiel-Frage „Kann es morgen eine zweite "
                    "Fundsache *Regenschirm* geben?“"]),
             SQL("SELECT Bezeichnung\nFROM Fundsache\nWHERE Kategorie = Kleidung;", pruefung="fehlerhaft"))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* ★★★-Aufgaben, Cold Case (AB und Lab A1-E1-CC), freie Ermittlung in der Konsole (eigene "
                    "Fragen an das Fundbuch), AB 1.1 Auftrag 5.",
                    "*Langsame:* ★-Aufgaben zuerst, Tipps (kosten Punkte), Zeitlupe, Ermittler-Handbuch Werkzeugkarte 1; "
                    "Partner-Check im Team (das Lab bleibt Einzelarbeit)."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Kurs im Portal vorhanden, Datendetektei für den Kurs freigeschaltet, Akte 1 offen, Niveau eA/gA eingestellt.",
                    "Karteikarten-Set (30 Karten) gedruckt und laminiert; Detektei-Ausweise; Beamer mit Präsentation Akte 1.",
                    "Geräte geladen, WLAN getestet (die Konsole lädt die Fundbuero-DB); Video-0-Link im Kurs bereit."]))

    # ---------------------------------------------------------------- DS 2
    d.ueberschrift("DS 2 · Das verschwundene Fundstück (Ermittlung 1.2)", 1)
    d.ueberschrift("Ziele", 2)
    d.absatz("Die Schülerinnen und Schüler …")
    d.inhalt(Liste(["**entwerfen** Bedingungen mit LIKE (`_`, `%`), BETWEEN, IN und IS NULL und **begründen**, warum "
                    "`= NULL` nie zutrifft (K18, K19),",
                    "**notieren** Literale korrekt (Text `'…'`, Datum `'JJJJ-MM-TT'`, Dezimalzahl mit Punkt) und verwenden "
                    "`!=` statt `<>` (K32),",
                    "**entwerfen** Abfragen mit Berechnungen, Spaltenalias (AS), DISTINCT und LIMIT (K15, K16, K20, K22),",
                    "**kreisen** ein Fundstück über eine Kette aufeinander aufbauender Abfragen **ein** und **beschreiben** "
                    "das Vorgehen in Worten (K30)."], nummeriert=True))
    d.ueberschrift("Verlauf", 2)
    d.tabelle(["Phase", "Zeit", "Inhalt / Ablauf", "Sozialform", "Material"],
              [["Warm-up", "0–10'", "Karten W-A1-01 … 04 in Einzelarbeit, Vergleich im Team; Karte 02 (Klammer-Check) an "
                "der Tafel auflösen", "EA → Team", "Warm-up-Karten A1, Folien 16–17"],
               ["Briefing", "10–20'", "Gesa Ohlendieks Notizen: „Messing“, „Juni oder Juli“, „mehrere hundert Euro“; neue "
                "Werkzeuge LIKE, BETWEEN, IN, IS NULL, DISTINCT, AS, LIMIT, Literale an Beispielen", "LV, UG", "Folien 18–23"],
               ["Unplugged", "20–30'", "Fahndung mit Platzhaltern + NULL-Stuhl (siehe unten)", "Teams",
                "Wort-/Musterkarten, NULL-Karte, Folie 24"],
               ["Lab", "30–75'", "Ermittlung 1.2: 01–05 (Werkzeuge), dann Indizienkette 06–08; AB 1.2 parallel", "EA",
                "Geräte, AB 1.2, Folie 25"],
               ["Fallbesprechung", "75–90'", "Indizienkette gemeinsam rekonstruieren („Welche Bedingung hat den entscheidenden "
                "Datensatz verraten?“), Merkkasten Literale/NULL, Abi-Ecke (3 BE), Cliffhanger, Hinweis Verhör 1", "UG",
                "Folien 26–31"]],
              breiten=[2.3, 1.4, 7.2, 2.4, 3.7])
    d.ueberschrift("Unplugged-Ablauf: Fahndung mit Platzhaltern (10 Minuten)", 2)
    d.inhalt(Liste(["Teams ordnen jeder Musterkarte alle passenden Wortkarten zu (Wettlauf, 4 Minuten); Lehrkraft prüft "
                    "stichprobenartig. Streitfall: Passt „Uhrenkette“ zu `'%uhr'`? Nein – `%uhr` verlangt das Ende „uhr“; "
                    "zu `'%uhr%'` passt sie (Groß/Klein ist bei LIKE in SQLite für A–Z egal).",
                    "**NULL-Stuhl:** Eine Person erhält die NULL-Karte als „Farbe“. `WHERE Farbe = NULL` – wer steht auf? "
                    "Niemand, auch NULL nicht. Dann `WHERE Farbe IS NULL` (nur NULL steht). Ebenso `WHERE Farbe != 'rot'` – "
                    "NULL bleibt sitzen!",
                    "`BETWEEN 10 AND 20` mit den Wertkarten 10/15/20/21 → 10, 15 und 20 gehören dazu, 21 nicht.",
                    "Sicherung: „`%` beliebig viele, `_` genau ein Zeichen“ · „NULL nur mit IS NULL“ · „BETWEEN schließt "
                    "beide Grenzen ein“."], nummeriert=True))
    d.absatz("Lösung der Zuordnung (mit SQLite-LIKE geprüft):")
    d.tabelle(["Muster", "passende Wortkarten"], [[f"`{m}`", ", ".join(t) or "–"] for m, t in fahndung_loesung()],
              breiten=[3.2, 13.8])
    d.absatz("Ohne Muster bleiben „Tischtennisschläger“ und „Sonnenbrille“ (Ablenker).")
    d.ueberschrift("Tafelbild / Merkkasten „Ermittler-Notiz“", 2)
    lohn = ("SELECT Bezeichnung, Wert, Wert * 0.05 AS Finderlohn\nFROM Fundsache\nWHERE Kategorie IN ('Schmuck', 'Elektronik')\n"
            "ORDER BY Finderlohn DESC\nLIMIT 3;")
    d.inhalt(Tabelle(["Bedingung", "Bedeutung"],
                     [["`WHERE Bezeichnung LIKE 'Schl%'`", "`%` = beliebig viele Zeichen (auch keins)"],
                      ["`WHERE Bezeichnung LIKE 'M_tze'`", "`_` = genau ein Zeichen"],
                      ["`WHERE Funddatum BETWEEN '2026-09-01' AND '2026-09-30'`", "Grenzen inklusive"],
                      ["`WHERE Kategorie IN ('Schmuck', 'Elektronik')`", "Werteliste"],
                      ["`WHERE AbgeholtAm IS NULL`", "niemals `= NULL`"],
                      ["`WHERE Kategorie != 'Kleidung'`", "im Abitur `!=`"]], breiten=[9.0, 8.0]),
             SQL(lohn), Ergebnis(*abfrage_ergebnis(lohn, DB)),
             "**Ermittler-Notiz Literale:** Text `'Hafenviertel'` · Datum `'2026-09-30'` (Text im Format JJJJ-MM-TT, deshalb "
             "sortier- und vergleichbar) · Zahl `0.05` ohne Hochkommata · `7 / 2` ergibt bei Ganzzahlen `3` · LIMIT nur, "
             "wenn an der Schnittkante kein Gleichstand ist.")
    d.ueberschrift("Die Indizienkette (Fallbesprechung)", 2)
    d.absatz("Auf der Haupt-DB bleibt schon nach Schritt 2 (A1-E2-07) nur FundNr 36 übrig; Schritt 3 (A1-E2-08) bestätigt "
             "„abgeholt“ und deckt Fundort und Finder auf. Die Variante +null (Messingleuchter mit unbekanntem Wert) prüft im "
             "Lab das `OR Wert IS NULL`.")
    kette = ("SELECT Bezeichnung, Fundort, Wert, AbgeholtAm, Finder\nFROM Fundsache\nWHERE Bezeichnung LIKE '%Messing%'\n"
             "  AND Funddatum BETWEEN '2026-06-01' AND '2026-07-31'\n  AND (Wert >= 400 OR Wert IS NULL)\n"
             "  AND NOT AbgeholtAm IS NULL;")
    d.inhalt(SQL(kette), Ergebnis(*abfrage_ergebnis(kette, DB)),
             "Story-Auflösung (Abschluss-Satz): „N. Falter“ – schnell gesprochen *Nachtfalter*. Wer `AbgeholtAm` setzen kann, "
             "hat Zugriff auf die Verwaltungsdaten der Stadt; im Ausgabebuch fehlt die Unterschrift. Steckbrief: Initialen "
             "„N. F.“, Signatur Nachtfalter. Cold Case A1-E2-CC: Nur dieser eine Finder-Eintrag enthält „falter“.")
    d.ueberschrift("Verhör 1 · Stadtbibliothek", 2)
    d.absatz("Pool von 12 Aufgaben (A1-V-01 … 12) auf der Kontext-DB Stadtbibliothek, Ziehung 6 Aufgaben, 20 Minuten, "
             "Abimodus streng, kein Feedback während des Verhörs. Freischaltung im Kurs: „auto“ (ab 60 % der Akte-1-Punkte) "
             "oder „offen“, wenn es Hausaufgabe bis DS 3 ist.")
    d.ueberschrift("Typische Schwierigkeiten", 2)
    d.inhalt(Liste(["`= NULL` / `!= NULL` (**P5**) → an den NULL-Stuhl erinnern; Parser-Hinweis NULL_VERGLEICH.",
                    "Datum ohne Hochkommata (`2026-09-01` wird als Rechnung 2026 − 9 − 1 gelesen!) bzw. deutsches "
                    "Datumsformat (**P5**).",
                    "`%` und `_` vertauscht, `LIKE 'uhr'` ohne Platzhalter (**P6**).",
                    "BETWEEN-Grenzen falsch eingeschätzt, `>` statt `>=` (**P6**).",
                    "LIMIT ohne ORDER BY („die drei teuersten“) → Ergebnis zufällig.",
                    "Alias in WHERE (`WHERE Finderlohn > 5`) → in SQLite möglich, im Abitur nicht vorgesehen: Berechnung "
                    "wiederholen."]))
    d.ueberschrift("Differenzierung", 2)
    d.inhalt(Liste(["*Schnelle:* Cold Case „Finderlohn-Stufen“ (AB), Lab-Cold-Case „Der Schatten-Finder“, Verhör 1 schon in der Stunde.",
                    "*Langsame:* Indizienkette mit Tipps; Werkzeugkarte 1; die Musterkarten aus dem Unplugged liegen aus."]))
    d.ueberschrift("Vorbereitung", 2)
    d.inhalt(Liste(["Wort-/Musterkarten je Team (Unplugged-Set A1), NULL-Karte, Wertkarten.",
                    "Warm-up-Karten A1 (Karten 01–04) kopieren.",
                    "Verhör 1 im Kurs auf „auto“ oder „offen“ stellen (siehe oben)."]))
    d.ueberschrift("Hinweise zum Niveau", 2)
    d.absatz("Akte 1 ist für gA und eA identisch (alle Aufgaben gA/eA). Die Arbeitsblätter sind nicht 1:1 die Lab-Aufgaben: "
             "Sie üben dieselben Kompetenzen mit anderen Fragen; die Abi-Ecken entsprechen dem Operator „entwerfen“ (AFB II "
             "bzw. I/II).")
    return d.speichern(ORDNER / "4_Lehrerhinweise" / "DD_A1_Lehrerhinweise.docx")


# ================================================================== Warm-up-Karten
def warmup():
    karten = json.loads((REPO / "seed" / "warmup" / "warmup_a1.json").read_text(encoding="utf-8"))["karten"]
    warmup_karten_dokument(karten, 1, "Warm-up-Karten Akte 1").speichern(MAT / "Warm-up" / "DD_Warm-up-Karten_A1.docx")


def main():
    for f in ("schueler", "loesung"):
        ab1(f)
        ab2(f)
    praesentation()
    unplugged_karteikarten()
    unplugged_fahndung()
    unplugged_ausweise()
    lehrerhinweise()
    warmup()
    for p in sorted(ORDNER.rglob("*")):
        if p.is_file():
            print(f"✔ {p.relative_to(REPO)}")
    print(f"✔ materialien/Warm-up/DD_Warm-up-Karten_A1.docx")


if __name__ == "__main__":
    main()
