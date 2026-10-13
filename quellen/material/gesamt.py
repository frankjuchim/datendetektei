#!/usr/bin/env python3
"""Erstellungshilfe: Gesamtmaterial der Datendetektei Weserstedt (materialien/00_Gesamt/) im Akten-Design.

Erzeugt (Spec §11.4, §12.1; Plan Phase 7 Task 9):
  materialien/00_Gesamt/DD_Gesamtuebersicht.docx   – DS-Plan, eA/gA-Pfad, Kompetenz- und Abi-Bezugsmatrix,
                                                     Materialliste, Checkliste vor dem Einsatz (Lehrkraft)
  materialien/00_Gesamt/DD_Abi-Kompass.docx        – 19 Abituraufgaben 2021–2026 als Verweise, Zuordnung zu Akten
                                                     und Kompetenzen, EH-2025-Hinweise (Lehrkraft)
  materialien/00_Gesamt/DD_Ermittler-Handbuch.docx – 8 Werkzeugkarten, zwei je A4-Seite (Schülerinnen und Schüler)

Datenquellen (werden beim Lauf gelesen, damit die Übersicht zu den Planungsdateien passt):
  docs/DS_PLANUNG.md (§1 Überblick, Grundkurs-Pfad), docs/KOMPETENZMATRIX.md (§1, §6),
  Spec §2.4/§4/§5 (Abi-Befunde, Belege je Kompetenz und Falle), seed/arbeitsblaetter.json (Materialliste).
Keine Originaltexte von Abituraufgaben – nur Kennungen (Jahr, Termin, Niveau, Teilaufgabe).

Die erzeugten docx sind danach maßgeblich (Spec §12); dieses Skript ist nur die Erstellungshilfe. Nach dem Mergen
weiterer Akten-Materialien die Materialliste mit diesem Skript neu erzeugen (sie liest seed/arbeitsblaetter.json).
Aufruf: python3 quellen/material/gesamt.py
"""
import json
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "werkzeuge"))

from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_TAB_ALIGNMENT  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import akten_design as A  # noqa: E402
from akten_design import (  # noqa: E402
    AktenDokument, ER, Ergebnis, Liste, Schema, SQL, Tabelle, abfrage_ergebnis)

MAT = REPO / "materialien"
ZIEL = MAT / "00_Gesamt"
SPEC = REPO / "docs" / "superpowers" / "specs" / "2026-10-12-datendetektei-design.md"
DS_PLANUNG = REPO / "docs" / "DS_PLANUNG.md"
MATRIX = REPO / "docs" / "KOMPETENZMATRIX.md"


# ================================================================== Markdown-Tabellen lesen
def md_tabelle(pfad, kopf_beginn):
    """Zeilen (Zellenlisten) der ersten Markdown-Tabelle, deren Kopfzeile mit kopf_beginn anfängt."""
    zeilen = pfad.read_text(encoding="utf-8").splitlines()
    for i, z in enumerate(zeilen):
        if z.startswith(kopf_beginn):
            aus = []
            for t in zeilen[i + 2:]:
                if not t.startswith("|"):
                    break
                aus.append([c.strip() for c in t.strip().strip("|").split("|")])
            return aus
    raise SystemExit(f"Tabelle „{kopf_beginn}“ in {pfad.name} nicht gefunden")


def ohne_fett(t):
    return t.replace("**", "")


def kuerze(liste_text, n=3):
    """„A, B, C, D, E“ → „A, B, C … (5)“ – für schmale Spalten."""
    teile = [x.strip() for x in liste_text.split(",") if x.strip() and x.strip() != "–"]
    if not teile:
        return "–"
    if len(teile) <= n:
        return ", ".join(teile)
    return ", ".join(teile[:n]) + f" … ({len(teile)})"


def anzahl(liste_text):
    teile = [x.strip() for x in liste_text.split(",") if x.strip() and x.strip() != "–"]
    return str(len(teile)) if teile else "–"


def ds_nummern(text):
    """„1, 8“ → {1, 8}; „1–6“ → {1..6}; „K1“/„alle“ → leer."""
    aus = set()
    for teil in text.split(","):
        teil = teil.strip()
        m = re.fullmatch(r"(\d+)\s*[–-]\s*(\d+)", teil)
        if m:
            aus.update(range(int(m.group(1)), int(m.group(2)) + 1))
        elif teil.isdigit():
            aus.add(int(teil))
    return aus


DS_AKTE = {1: 1, 2: 1, 3: 2, 4: 2, 5: 3, 6: 3, 7: 4, 8: 4, 9: 5, 10: 5, 11: 5, 12: 6, 13: 7}
AKTEN_NAMEN = {1: "Das Fundbüro", 2: "Spur der Räder", 3: "Die Stadtfest-Kasse", 4: "Chaos im Vereinsarchiv",
               5: "Der Bauplan (eA)", 6: "Datenspuren", 7: "Die Ermittlerprüfung"}

# Kompetenzmatrix (KOMPETENZMATRIX §1): ID → Felder
_KM_KOPF = ("id", "kompetenz", "niveau", "einf", "uebung", "lab", "ab", "warmup", "verhoer", "gp", "kl", "video")
KM = {}
for _z in md_tabelle(MATRIX, "| ID | Kompetenz | Niveau | Einführung DS"):
    _e = dict(zip(_KM_KOPF, _z))
    _e["id"] = ohne_fett(_e["id"])
    KM[_e["id"]] = _e


def akten_der_kompetenz(k):
    e = KM.get(k)
    if not e:
        return set()
    return {DS_AKTE[d] for d in ds_nummern(e["einf"]) if d in DS_AKTE}


# Spec §4: Belege je Kompetenz („Abi E21, E26“) · Spec §5: Belege je Falle („E22 c, E23 c“)
KENNUNG = re.compile(r"\b([EG])(2[1-6])(N?)\b(?:\s+([a-h]\d?)\b)?")


def _belege():
    je_kompetenz, je_falle = {}, {}
    for z in md_tabelle(SPEC, "| ID | Kompetenz | Niveau | Quelle | DS |"):
        if len(z) < 5 or not z[0].startswith("K"):
            continue
        for m in KENNUNG.finditer(z[3]):
            je_kompetenz.setdefault("".join(m.group(1, 2, 3)), set()).add(z[0])
    for z in md_tabelle(SPEC, "| Nr. | Kurztitel | Erläuterung | Beleg (Abi) |"):
        falle = ohne_fett(z[0])
        for m in KENNUNG.finditer(z[3]):
            je_falle.setdefault("".join(m.group(1, 2, 3)), []).append((falle, m.group(4)))
    return je_kompetenz, je_falle


BELEG_K, BELEG_P = _belege()


# ================================================================== Tabellen-Hilfe
def tab(d, kopf, zeilen, breiten=None, groesse=9.0, behaelter=None):
    """Datentabelle wie d.tabelle, aber alle Zellen linksbündig; Zeilen, deren erste Zelle mit „**“ beginnt und deren
    übrige Zellen leer sind, werden zu einer durchgehenden Zwischenüberschrift verbunden. ↔ wird ausgeschrieben
    (Barlow hat das Zeichen nicht)."""
    zeilen = [[re.sub(r"^SELECT …", "Abfrage SELECT …", str(c).replace(" ↔ ", " vs. ").replace("↔", " vs. ")
                      .replace("Self-Join", "Selbstverbund")) for c in z] for z in zeilen]
    t = d._datentabelle(behaelter if behaelter is not None else d.doc,
                        Tabelle(kopf, zeilen, breiten=breiten, groesse=groesse), d.breite)
    for zi, z in enumerate(zeilen, start=1):
        row = t.rows[zi]
        for c in row.cells:
            for p in c.paragraphs:
                p.alignment = None
        if z[0].startswith("**") and not any(x.strip() for x in z[1:]):
            erste = row.cells[0].merge(row.cells[-1])
            for p in list(erste.paragraphs)[1:]:
                if not p.text:
                    p._p.getparent().remove(p._p)
            A._zelle_flaeche(erste, A.FLAECHE_2)
    return t


# ================================================================== Gesamtübersicht
def gesamtuebersicht():
    d = AktenDokument("gesamt", 0, "Gesamtübersicht Datenbanken", fassung="lehrkraft", material="Gesamtübersicht",
                      querformat=True, untertitel="Datendetektei Weserstedt – Ermitteln mit Datenbanken · Informatik "
                      "Leistungskurs (eA) mit Grundkurs-Pfad (gA) · Niedersachsen, Abitur ab 2027")
    B = d.breite
    d.kopfband(name_datum=False)
    d.absatz("Der Kurs ist das Ermittlerteam der **Datendetektei Weserstedt**. Chefin Edda Lorenz übergibt in sieben "
             "**Akten** Aufträge mit eigener Datenbank; jede Doppelstunde ist eine **Ermittlung** im Lab, jedes gewonnene "
             "**Verhör** bringt ein Beweisstück an die Pinnwand. Seit Wochen hinterlässt ein Unbekannter – „das Phantom“ – "
             "einen Nachtfalter in den Datenbeständen der Stadt; in Akte 6 wird es enttarnt, und die Frage „Durfte die "
             "Detektei das?“ führt in die Datenschutz-Erörterung. Alle Personen, Orte und Daten sind frei erfunden.")
    d.inhalt(Liste([
        "**1** Takt einer Doppelstunde · **2** Plan der Doppelstunden · **3** eA- und gA-Pfad · **4** Kompetenzmatrix im "
        "Überblick und Fallen P1–P10 · **5** Abi-Bezugsmatrix · **6** Materialliste je Akte · **7** Checkliste vor dem "
        "Einsatz"]))

    d.ueberschrift("1 · Takt einer Doppelstunde (90 Minuten)", 1)
    tab(d, ["Phase", "Zeit", "Inhalt", "Sozialform"], [
        ["Warm-up", "10'", "Papierkarte (Wiederholung der Vorstunde, abi-nah); digital im Leitner-Trainer ohne Punkte",
         "Einzelarbeit"],
        ["Briefing / Unplugged", "15–20'", "Fall-Einstieg (Präsentation), neue Begriffe, haptische Aktivität", "Plenum, Teams"],
        ["Lab-Einsatz", "45'", "Ermittlung im Lab (Stufen ★/★★/★★★), Arbeitsblatt parallel; Zeitlupe, Spurenbrett, "
         "Sortierfächer als Verstehenshilfen", "Einzelarbeit (keine Team-Punkte im Lab)"],
        ["Fallbesprechung", "10–15'", "Sicherung (Ermittler-Notiz), Abi-Ecke, Ausblick auf die nächste Akte",
         "Plenum"],
    ], breiten=[4.0, 2.0, 13.7, 6.0])
    d.absatz("Feste **Teams** (3–4) für Unplugged, Peer-Check und Wettbewerbe; im Lab arbeitet jede und jeder für sich. "
             "Sonder-Doppelstunden: DS 1 (Kick-off ohne Warm-up), DS 10 (Architekturwettbewerb), DS 12 (Debatte), "
             "DS 13 (Generalprobe).")

    # ---------------------------------------------------------- 2 DS-Plan
    d.ueberschrift("2 · Plan der Doppelstunden (DS 1–13, Kür K1/K2)", 1)
    zeilen = []
    for z in md_tabelle(DS_PLANUNG, "| DS | Akte | Thema | Kompetenzen"):
        ds, akte, thema, komp, unpl, lab, ab, verhoer, video = z[:9]
        zeilen.append([ds.replace("DS ", ""), akte, thema, komp.replace(" · prozessbezogen", " ·"), unpl, lab,
                       ab, verhoer, video])
    for z in zeilen:  # JOIN-Hinweise in Planungstexten stehen nicht im Plan; Kür klar kennzeichnen
        if z[0].startswith("K1"):
            z[2] += " (Kür, nicht im Abitur)"
    tab(d, ["DS", "Akte", "Thema", "Kompetenzen", "Unplugged", "Lab", "AB", "Verhör", "Video"], zeilen,
              breiten=[1.5, 2.6, 3.2, 4.2, 3.6, 3.3, 2.1, 3.1, 2.1], groesse=7.5)
    d.absatz("Verhöre V1–V5 dauern 20 Minuten, V6 15 Minuten, V7 (Lab-Generalprobe) 30 Minuten; Freischaltung "
             "automatisch ab 60 % der Punkte der Akte (Kurs-Einstellung auto/offen/gesperrt). Warm-up-Karten wiederholen "
             "die Vorstunde (W-A1 in DS 2–3, W-A2 in DS 4–5 usw.; W-A6/W-A7 als Hausaufgabe vor DS 13 bzw. im Warm-up "
             "von K2). Die **Klausur** (eA und gA, mit Nachschreibklausur) liegt außerhalb der 15 Doppelstunden.")

    # ---------------------------------------------------------- 3 eA/gA
    d.ueberschrift("3 · eA- und gA-Pfad", 1)
    tab(d, ["Pfad", "Doppelstunden", "Prüfungen", "Hinweise"], [
        ["**eA** (Leistungskurs)", "DS 1–13 (13 DS) + Kür K1 (nach DS 11) und K2 (vor DS 13) = 15 DS",
         "Generalprobe eA (30 BE, ≈ 70'), Klausur eA Aufgabe A + B (60 BE, auf 45 BE kürzbar), Nachschreibklausur",
         "Kür-Inhalte sind ausdrücklich nicht abiturrelevant (K1: CREATE/INSERT, Ausblick auf andere Verbund-"
         "Schreibweisen; Normalformen als Cold Case in DS 8)."],
        ["**gA** (Grundkurs-Pfad)", "DS 1–8, 12, 13 (10 DS) + optional K2", "Generalprobe gA (25 BE, ≈ 60', ohne ER, "
         "mit Datenschutz), Klausur gA = Aufgabe A (30 BE) + Aufgabe B-gA (15 BE)",
         "Im Kurs „Niveau: gA“ einstellen → Akte 5 und eA-Aufgaben sind ausgeblendet; Ränge beziehen sich auf die "
         "freigeschalteten Akten."],
    ], breiten=[3.6, 6.2, 7.6, 8.3])
    d.ueberschrift("Besonderheiten im gA-Kurs", 3)
    tab(d, ["gA-DS", "entspricht", "Besonderheit"],
              [[a, b, c] for a, b, c in md_tabelle(DS_PLANUNG, "| gA-DS | entspricht |")], breiten=[2.0, 2.6, 21.1])

    # ---------------------------------------------------------- 4 Kompetenzmatrix
    d.seitenumbruch()
    d.ueberschrift("4 · Kompetenzmatrix im Überblick (DB-K01 … DB-K48)", 1)
    d.absatz("Niveau **gA/eA** = beide, **eA** = nur Leistungskurs, **Kür** = nicht abiturverbindlich. Zahlen = Anzahl "
             "der zugeordneten Lab-Aufgaben, Warm-up-Karten bzw. Verhör-Aufgaben; AB-Bezüge als „AB x.y/Nr.“. Jede "
             "abiturverbindliche Kompetenz ist mindestens einer DS, einer Lab-Aufgabe und einer Prüfungsaufgabe "
             "zugeordnet (Kür: K06, K13, K48).")
    gruppen = [("A – Relationales Modell und Schlüssel", 1, 6), ("B – Redundanz, Anomalien, Schemaqualität", 7, 14),
               ("C – SQL (Umfang EH §8)", 15, 32), ("D – ER-Modellierung (Chen, EH §7)", 33, 43),
               ("E – Datenschutz, Informatik und Gesellschaft", 44, 46), ("F – Prozessbezogen", 47, 48)]
    zeilen = []
    for titel, von, bis in gruppen:
        zeilen.append([f"**{titel}**", "", "", "", "", "", "", "", "", ""])
        for i in range(von, bis + 1):
            e = KM[f"K{i:02d}"]
            zeilen.append([e["id"], e["kompetenz"], e["niveau"], e["einf"], anzahl(e["lab"]), kuerze(e["ab"], 3),
                           anzahl(e["warmup"]), anzahl(e["verhoer"]), e["gp"], kuerze(e["kl"], 3)])
    tab(d, ["ID", "Kompetenz", "Niveau", "DS", "Lab", "Arbeitsblatt / Abi-Ecke", "Warm-up", "Verhör",
               "Generalprobe", "Klausur"], zeilen,
              breiten=[1.1, 8.6, 1.4, 1.2, 1.0, 4.6, 1.5, 1.4, 2.5, 3.4], groesse=7.5)

    d.seitenumbruch()
    d.ueberschrift("Fallen P1–P10 (überall dieselben Nummern)", 2)
    fallen = []
    for z in md_tabelle(SPEC, "| Nr. | Kurztitel | Erläuterung | Beleg (Abi) |"):
        fallen.append([ohne_fett(z[0]), z[1], z[2]])
    tab(d, ["Nr.", "Falle", "Erläuterung"], fallen, breiten=[1.3, 4.6, 19.8], groesse=8)
    d.absatz("Die Fallen erscheinen im Lab (Feld „Fallen“, Profil „Deine Fallen“), im Fallen-Profil des Kurses "
             "(Lehrkraft-Bereich → Auswertung), auf Werkzeugkarte 8 des Ermittler-Handbuchs und in allen "
             "Erwartungshorizonten.")

    # ---------------------------------------------------------- 5 Abi-Bezug
    d.ueberschrift("5 · Abi-Bezugsmatrix", 1)
    d.absatz("Aufgabentypen der 19 vorliegenden Abiturklausuren 2021–2026 (gA und eA) × Einführung, Lab und Prüfung. "
             "Zahlen = Anzahl der Lab- bzw. Verhör-Aufgaben. Die Zuordnung der einzelnen Klausuren steht im "
             "**Abi-Kompass**; Originaltexte der Abituraufgaben sind nicht Teil des Materials.")
    zeilen = []
    for z in md_tabelle(MATRIX, "| Abituraufgaben-Typ |"):
        typ, beleg, komp, einf, lab_ds, lab, verh, pruef = z[:8]
        lab_n = lab.split(":")[0] if ":" in lab else lab
        verh_n = verh.split(":")[0] if ":" in verh else verh
        zeilen.append([ohne_fett(typ), beleg, komp, einf, lab_ds, lab_n, verh_n, kuerze(pruef, 5)])
    tab(d, ["Abituraufgaben-Typ", "Beleg Abi 2021–26", "Kompetenzen", "Einf. DS", "Lab in DS", "Lab",
               "Verhör", "Generalprobe / Klausur"], zeilen,
              breiten=[4.6, 4.4, 2.6, 2.0, 2.6, 1.2, 1.7, 6.6], groesse=8)

    # ---------------------------------------------------------- 6 Materialliste
    d.seitenumbruch()
    d.ueberschrift("6 · Materialliste je Akte", 1)
    d.absatz("Maßgeblich sind die docx/pptx; die PDFs stehen nach Anmeldung in der Plattform (Lehrkraft-Bereich → "
             "Materialien). **Schüler-PDF** = Arbeitsblaetter_PDF (alle Angemeldeten), **Lehrkraft-PDF** = "
             "Loesungen_Portal bzw. Lehrkraft_PDF (nur Lehrkräfte). Präsentationen gibt es als Handout (ohne "
             "Lösungsfolien und Notizen) und als Lehrkraftfassung mit Sprechernotizen.")
    materialliste(d, B)

    # ---------------------------------------------------------- 7 Checkliste
    d.seitenumbruch()
    d.ueberschrift("7 · Checkliste vor dem Einsatz", 1)
    d.absatz("Einmal vor Akte 1 durchgehen (ca. 45 Minuten); die mit ◆ markierten Punkte vor jeder Akte erneut prüfen.")
    pruef = [
        ["Portal", "Lehrkraft-Konto im Portal InfoLabs anmelden; Freischaltung Lehrkraft × Datendetektei durch den "
         "Admin des Portals prüfen.", "Portal InfoLabs: Plattform Datendetektei wird angezeigt"],
        ["Kurs", "Kurs im Portal anlegen; Kurscode notieren; Schülerinnen und Schüler treten mit dem Kurscode bei und "
         "erhalten ihren persönlichen Code (sicher aufbewahren lassen).", "Portal (Kursverwaltung); ein Konto je Schüler:in"],
        ["Freischaltung", "Kurs × Plattform Datendetektei freischalten; einmal selbst per Single Sign-on aus dem "
         "Portal in die Datendetektei wechseln.", "Portal (Kursverwaltung, Plattformen des Kurses)"],
        ["Niveau", "Im Lehrkraft-Bereich der Datendetektei je Kurs „Niveau eA/gA“ einstellen (gA blendet Akte 5 und "
         "eA-Aufgaben aus).", "Datendetektei → Lehrkraft → Kurs → Einstellungen"],
        ["◆ Akten", "Akte(n) der nächsten Stunden freischalten; Verhör-Modus wählen (auto = ab 60 % der Punkte, "
         "offen, gesperrt); Kür-Welt K1 erst nach DS 11.", "Einstellungen → Akten / Verhör"],
        ["Vorschau", "Ermittlungen und Verhör-Pools in der Vorschau ansehen (mit Lösungen) – nur für Sie.",
         "Lehrkraft → Vorschau"],
        ["Geräte", "iPad/Laptop je Schüler:in, aktueller Browser; Test-Anmeldung mit einem Schülerkonto (Lab, "
         "Konsole „Ausführen“, Zeitlupe).", "Lab → Akte 1 → Ermittlung 1.1"],
        ["◆ WLAN / SSE", "Live-Modus im Schul-WLAN testen: Testrunde starten, mit einem zweiten Gerät über die Seite "
         "„Live“ beitreten, eine Aufgabe freischalten (Server-Sent Events dürfen vom Proxy nicht gepuffert werden).",
         "Lehrkraft → Live → Testrunde; spätestens vor K2"],
        ["◆ Board", "Beamer-Board aktivieren (Kurs-Einstellung) und Board-Link am Beamer öffnen; Scoreboard an/aus "
         "nach Kurskultur.", "Lehrkraft → Board; Link mit eigenem Board-Schlüssel"],
        ["◆ Materialien", "Arbeitsblätter, Warm-up-Karten (Schülerfassung) und Unplugged-Kopiervorlagen der Akte "
         "drucken; Detektei-Ausweise (DS 1), Karteikarten, Fäden, Fächerschilder, Umschläge vorbereiten.",
         "Lehrkraft → Materialien; Liste in Abschnitt 6"],
        ["Handbuch", "Ermittler-Handbuch (8 Werkzeugkarten) je Team einmal drucken und laminieren.",
         "Gesamtmaterial → Ermittler-Handbuch"],
        ["Datenschutz", "Kurs über die Datenschutzerklärung informieren: gespeichert werden nur Anzeigename, Kurs und "
         "Lernstand; SQL-Entwürfe bleiben auf dem Gerät.", "Seite „Datenschutz“ der Plattform, Portal"],
        ["Prüfungen", "Generalprobe und Klausur (eA/gA) mit Erwartungshorizont rechtzeitig sichten; Klausur nicht "
         "vorab im Kurs zeigen.", "Lehrkraft → Materialien (nur Lehrkraft)"],
    ]
    tab(d, ["erledigt", "Bereich", "Was ist zu tun?", "Wo / Hinweis"], [[""] + z for z in pruef],
              breiten=[2.0, 2.8, 13.3, 7.6], groesse=8.5)
    d.notiz("Abschluss der Einheit", "Nach DS 13: Endstand festschreiben, Urkunden „Detektei-Lizenz“ im Lehrkraft-Bereich "
            "drucken (Lehrkraft → Urkunden). Live-Runden speichern keine Antworten; im Lab bleiben nur Anzeigename, "
            "Kurs-ID und Lernstand.")
    return d.speichern(ZIEL / "DD_Gesamtuebersicht.docx")


PDF_ZIEL = {  # art → (Schüler-PDF?, Lehrkraft-PDF?)
    "ab": ("ja", "Lösung"), "kuer": ("ja", "Lösung"), "unplugged": ("ja", "–"), "handbuch": ("ja", "–"),
    "warmup": ("ohne Lösungen", "mit Lösungen"), "praesentation": ("Handout", "mit Notizen"),
    "lehrerhinweise": ("–", "ja"), "gesamt": ("–", "ja"), "generalprobe": ("–", "ja + Erwartungshorizont"),
    "klausur": ("–", "ja + Erwartungshorizont"), "video": ("–", "Folien"),
}
ART_NAME = {"ab": "Arbeitsblatt", "kuer": "Kür", "unplugged": "Unplugged", "handbuch": "Handbuch", "warmup": "Warm-up",
            "praesentation": "Präsentation", "lehrerhinweise": "Lehrerhinweise", "gesamt": "Gesamt",
            "generalprobe": "Generalprobe", "klausur": "Klausur", "video": "Video"}
GEPLANT = {  # Spec §11, §12.1 – Platzhalter, bis die Akte in seed/arbeitsblaetter.json eingetragen ist
    "A5": ["AB 5.1 Baupläne lesen, AB 5.2 Der Architekturwettbewerb, AB 5.3 Vom Plan zur Datenbank (je mit Lösung)",
           "Präsentation Der Bauplan · Lehrerhinweise Akte 5 · Warm-up-Karten A5",
           "Unplugged: Rollenkarten Fäden-Kardinalitäten, Bewertungsbogen Architekturwettbewerb, Regelkarten-Umzug"],
    "A6": ["AB 6.1 Das Phantom (mit Lösung) · Präsentation Datenspuren · Lehrerhinweise Akte 6 · Warm-up-Karten A6",
           "Unplugged: „Steh auf, wenn …“, Rollenkarten-Debatte „Durfte die Detektei das?“"],
    "A7": ["Generalprobe eA und gA mit Erwartungshorizont · Klausur eA (A + B) und gA (A + B-gA) mit "
           "Erwartungshorizont · Nachschreibklausur · Warm-up-Karten A7"],
    "Videos": ["Folien und Drehbücher der Erklärvideos 0–7 (Lehrkraft)"],
}


def materialliste(d, breite):
    z = json.loads((REPO / "seed" / "arbeitsblaetter.json").read_text(encoding="utf-8"))
    for g in z["gruppen"]:
        d.ueberschrift(g["titel"], 3)
        if not g["materialien"]:
            d.inhalt(Liste([f"*geplant (noch nicht eingetragen):* {t}" for t in GEPLANT.get(g["gruppe"], ["–"])]))
            continue
        zeilen = []
        for x in g["materialien"]:
            titel = f"AB {x['nr']} · {x['titel']}" if x.get("nr") and x["art"] == "ab" else x["titel"]
            s, l = PDF_ZIEL.get(x["art"], ("–", "–"))
            zeilen.append([titel, ART_NAME.get(x["art"], x["art"]), x.get("niveau", "beide"), s, l, x["datei"]])
        tab(d, ["Material", "Art", "Niveau", "Schüler-PDF", "Lehrkraft-PDF", "Dateiname (PDF)"], zeilen,
                  breiten=[8.0, 2.6, 1.6, 2.4, 3.6, 7.5], groesse=8)


# ================================================================== Abi-Kompass
KLAUSUREN = [  # Kennung, Niveau, Jahr, Termin – Kennungen wie in Spec §4/§5 und KOMPETENZMATRIX §6
    ("E21", "eA", 2021, "Haupttermin"), ("E21N", "eA", 2021, "Nachtermin"), ("E22", "eA", 2022, "Haupttermin"),
    ("E22N", "eA", 2022, "Nachtermin"), ("E23", "eA", 2023, "Haupttermin"), ("E23N", "eA", 2023, "Nachtermin"),
    ("E25", "eA", 2025, "Haupttermin"), ("E25N", "eA", 2025, "Nachtermin"), ("E26", "eA", 2026, "Haupttermin"),
    ("G21", "gA", 2021, "Haupttermin"), ("G21N", "gA", 2021, "Nachtermin"), ("G22", "gA", 2022, "Haupttermin"),
    ("G23", "gA", 2023, "Haupttermin"), ("G23N", "gA", 2023, "Nachtermin"), ("G24", "gA", 2024, "Haupttermin"),
    ("G24N", "gA", 2024, "Nachtermin"), ("G25", "gA", 2025, "Haupttermin"), ("G25N", "gA", 2025, "Nachtermin"),
]
KURZ = {  # Kurzbezeichnung der Kompetenz für die Kompass-Tabelle
    "K02": "Schlüsseleignung", "K05": "Kennzeichnungen im Schema", "K07": "Redundanz", "K09": "Schema auf Anomalien",
    "K10": "Redundanz und Anomalie", "K11": "verlustfrei zerlegen", "K12": "Redundanzfreiheit begründen",
    "K14": "nicht-atomare Werte", "K17": "Bedingungslogik", "K20": "Berechnungen", "K22": "Top-n / LIMIT",
    "K24": "Verbund ≥ 3 Tabellen", "K27": "HAVING", "K29": "Ergebnis angeben", "K30": "Vorgehen in Worten",
    "K31": "Fehler finden", "K34": "Kardinalitäten", "K35": "ER entwickeln", "K38": "ternär / zwei Rollen",
    "K40": "Modellierung beurteilen", "K41": "ER → Schema", "K42": "Umsetzung begründen", "K43": "Schema → ER",
    "K45": "Datenschutz anwenden", "K46": "Datenanalyse, Risiken",
}
FALLE_NAME = {"P1": "WHERE vs. HAVING", "P2": "Verbund", "P3": "GROUP-BY-Regel", "P4": "COUNT vs. SUM",
              "P5": "NULL/Literale", "P6": "Bedingungslogik", "P7": "Benennung", "P8": "Schlüssel", "P9": "ER/Überführung",
              "P10": "Operator"}
EH_HINWEIS = {  # Spec §2.4: „Ältere Lösungen verletzen die GROUP-BY-Regel (z. B. 2023 eA, 2025 gA)“ – Termin nicht benannt
    "E23": "GROUP-BY-Regel prüfen (Befund „2023 eA“, Termin offen)",
    "E23N": "GROUP-BY-Regel prüfen (Befund „2023 eA“, Termin offen)",
    "G25": "GROUP-BY-Regel prüfen (Befund „2025 gA“, Termin offen)",
    "G25N": "GROUP-BY-Regel prüfen (Befund „2025 gA“, Termin offen)",
}


def abi_kompass():
    d = AktenDokument("gesamt", 0, "Abi-Kompass Datenbanken", fassung="lehrkraft", material="Abi-Kompass",
                      untertitel="Abituraufgaben 2021–2026 (gA/eA) → Akten, Kompetenzen, Übungen · EH-2025-Hinweise")
    d.kopfband(name_datum=False)
    d.notiz("Nur Verweise – keine Originaltexte",
            "Die Abituraufgaben sind urheberrechtlich geschützt (Niedersächsisches Kultusministerium). Dieser Kompass "
            "nennt nur **Jahr, Termin, Niveau und – wo belegt – die Teilaufgabe**; die Aufgaben selbst liegen Ihnen "
            "über die Schule vor. Übungen im gleichen Stil sind die Abi-Ecken der Arbeitsblätter, die Verhöre, die "
            "Generalprobe und die Klausuren der Datendetektei.")
    d.ueberschrift("1 · Lesehilfe", 1)
    tab(d, ["Kennung", "Bedeutung"], [
        ["E21, E22N …", "**E** = erhöhtes Anforderungsniveau (eA), **G** = grundlegendes Anforderungsniveau (gA); "
                        "Ziffern = Abiturjahr (21 = 2021); **N** = Nachtermin, ohne N = Haupttermin"],
        ["E22 c", "Teilaufgabe c der Datenbank-Aufgabe dieser Klausur"],
        ["Kompetenz", "DB-K01 … DB-K48 aus dem Kompetenzraster (Gesamtübersicht, Abschnitt 4)"],
        ["Falle", "P1–P10 (Werkzeugkarte 8): typische Fehler, die in der Teilaufgabe naheliegen"],
        ["Zuordnung offen", "Die Planungsunterlagen belegen die Zuordnung nicht eindeutig – bitte an der Originalaufgabe "
                            "ergänzen (nicht erfunden)."],
    ], breiten=[3.0, 14.0])

    d.ueberschrift("2 · Befunde im Überblick (19 Klausuren: 9 eA, 10 gA)", 1)
    tab(d, ["Aufgabentyp", "Häufigkeit", "Akte", "Kompetenzen"], [
        ["Verbund über mehrere Tabellen", "19 von 19", "2 (DS 3–4)", "K23, K24"],
        ["Anomalien (seit 2023 oft: „zusammengelegte Tabelle“ untersuchen)", "18 von 19", "4 (DS 7–8)",
         "K07–K10"],
        ["GROUP BY (mit HAVING)", "16 von 19", "3 (DS 5–6)", "K26, K27"],
        ["Schlüsseleignung beurteilen", "14 von 19", "1, 4 (DS 1, 8)", "K02"],
        ["Gegebene Abfrage analysieren („… welche Informationen die Abfrage im Sachzusammenhang liefert“)",
         "14 von 19", "3 (DS 6)", "K28"],
        ["ER-Modell (ergänzen/erweitern; seit 2025 auch entwickeln, Rückrichtung, ternär, zwei Rollen)",
         "9 von 9 eA, 0 von 10 gA (ergänzen/erweitern 8 von 9)", "5 (DS 9–11, nur eA)", "K33–K43"],
    ], breiten=[8.0, 3.6, 2.6, 2.8])
    d.inhalt(Liste([
        "Festes SQL-Muster: **2–3 Abfragen entwerfen + 1 gegebene Abfrage analysieren**.",
        "Seit 2025 viele kleine Teilaufgaben (a–g, 2–8 BE), neu: **Vorgehen in Worten beschreiben** (E25N, G25N).",
        "Welche einzelnen Klausuren die Häufigkeiten „18 von 19“, „16 von 19“ und „14 von 19“ bilden, ist in den "
        "Planungsunterlagen nicht aufgeschlüsselt – deshalb stehen diese Typen unten nur dort, wo eine Kompetenz oder "
        "Falle die Klausur ausdrücklich nennt.",
    ]))

    d.ueberschrift("3 · Zuordnung der Klausuren", 1)
    d.absatz("Spalte „belegt“: Kompetenzen und Fallen, für die Spec §4/§5 bzw. die Abi-Bezugsmatrix diese Klausur "
             "ausdrücklich als Beleg nennen. Daraus abgeleitet: die Akten (Einführungs-DS der Kompetenz). Für **alle** "
             "Klausuren gilt zusätzlich: Verbund (Akte 2); für alle eA-Klausuren: ER-Teil (Akte 5).")
    zeilen = []
    for kennung, niveau, jahr, termin in KLAUSUREN:
        komp = sorted(BELEG_K.get(kennung, set()))
        fallen = BELEG_P.get(kennung, [])
        akten = set()
        for k in komp:
            akten |= akten_der_kompetenz(k)
        for f, _ in fallen:
            akten |= {"P1": {3}, "P2": {2}, "P3": {3}, "P4": {3}, "P5": {1}, "P6": {1}, "P7": {2}, "P8": {1, 4},
                      "P9": {5}, "P10": set()}[f]
        teile = sorted({t for _, t in fallen if t})
        belegt = "; ".join([f"{k} {KURZ.get(k, KM[k]['kompetenz'][:24])}" for k in komp] +
                           [f"{f}{' (' + t + ')' if t else ''} {FALLE_NAME[f]}" for f, t in fallen])
        zeilen.append([f"**{kennung}**", niveau, f"{jahr} · {termin}", ", ".join(teile) or "–",
                       belegt or "Zuordnung offen", ", ".join(str(a) for a in sorted(akten)) or "–",
                       EH_HINWEIS.get(kennung, "Lösung an EH 2025 §8 prüfen" if jahr < 2025 else "–")])
    zeilen.append(["**?**", "gA", "Jahr/Termin nicht benannt", "–",
                   "Zuordnung offen: 19. Klausur (10. gA-Klausur) – in den Planungsunterlagen ohne Kennung", "–", "–"])
    tab(d, ["Kennung", "Niv.", "Jahr · Termin", "Teil", "belegt (Kompetenzen, Fallen)", "Akten",
               "EH-2025-Hinweis"], zeilen, breiten=[1.6, 1.0, 2.6, 1.1, 6.4, 1.3, 3.0], groesse=7.5)
    d.absatz("Akten: 1 Fundbüro · 2 Spur der Räder · 3 Stadtfest-Kasse · 4 Vereinsarchiv · 5 Bauplan (eA) · "
             "6 Datenspuren. Eine Klausur des Jahres 2024 auf eA-Niveau liegt nicht vor.")

    d.ueberschrift("4 · Vom Abiturtyp zur Übung in der Datendetektei", 1)
    zeilen = []
    for z in md_tabelle(MATRIX, "| Abituraufgaben-Typ |"):
        typ, beleg, komp, einf, lab_ds, lab, verh, pruef = z[:8]
        akten = sorted({DS_AKTE[n] for n in ds_nummern(einf) if n in DS_AKTE})
        zeilen.append([ohne_fett(typ), beleg, komp, ", ".join(map(str, akten)),
                       f"Lab {lab.split(':')[0]} · Verhör {verh.split(':')[0]}", kuerze(pruef, 4)])
    tab(d, ["Typ", "Beleg", "Kompetenzen", "Akte", "Training", "Generalprobe / Klausur"], zeilen,
              breiten=[3.6, 3.2, 2.4, 1.1, 2.6, 4.1], groesse=7.5)
    d.absatz("Abi-Ecken der Arbeitsblätter zu diesen Typen: siehe Kompetenzmatrix (Spalte „Arbeitsblatt / Abi-Ecke“ der "
             "Gesamtübersicht). Lab- und Verhör-IDs zeigt die Vorschau im Lehrkraft-Bereich.")

    d.ueberschrift("5 · EH-2025-Hinweise: Wir lehren den Stil 2026", 1)
    d.absatz("Ab dem Abitur 2027 gelten die Ergänzenden Hinweise 2025 (§7 Notation, §8 SQL-Umfang). Ältere "
             "Musterlösungen weichen davon teilweise ab; beim Einsatz alter Abituraufgaben im Unterricht die Lösung "
             "vorher anpassen:")
    tab(d, ["Regel (EH 2025 §8)", "Was das heißt", "Bezug"], [
        ["GROUP-BY-Regel", "Jede Spalte im SELECT ist aggregiert **oder** in GROUP BY aufgezählt; nach dem Schlüssel "
         "gruppieren und Namen zusätzlich aufnehmen (Gleichnamige!).", "Ältere Lösungen verletzen sie (Befund: "
         "2023 eA, 2025 gA – Termin jeweils offen); P3"],
        ["Kein Alias in HAVING", "Aggregat in HAVING wiederholen: `HAVING COUNT(*) > 3`; Alias nur in ORDER BY "
         "(und GROUP BY).", "Ältere Lösungen nutzen Alias in HAVING – betroffene Klausuren: Zuordnung offen; P1"],
        ["Verbund nur implizit", "`FROM a, b WHERE a.fk = b.pk`; keine andere Verbund-Schreibweise, keine "
         "Unterabfragen, kein UNION.", "alle Klausuren (Verbund 19/19); P2"],
        ["`!=` statt `<>`", "Ungleich wird als `!=` geschrieben.", "P6"],
        ["Literale", "Text in `'…'`, Datum `'JJJJ-MM-TT'`, Zahlen ohne Hochkommata; `IS NULL` statt `= NULL`.", "P5"],
        ["LIMIT", "Top-n mit `ORDER BY … DESC LIMIT n` – nur eindeutig ohne Gleichstand an Platz n.", "E22N, E26; K22"],
        ["Notation §7", "PK unterstrichen, FK mit vorangestelltem ↑; ER in Chen-Notation mit 1/n/m (ternär p); "
         "Entitätstyp, Beziehungstyp, Schlüsselattribut.", "alle eA-Klausuren; P9"],
    ], breiten=[3.4, 8.6, 5.0], groesse=8.5)
    d.ueberschrift("Beispiel: GROUP-BY-Regel und HAVING (eigene Daten, Stadtfest-DB)", 3)
    d.absatz("So nicht (Name weder gruppiert noch aggregiert, Alias in HAVING):")
    d.inhalt(SQL("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS Bons\nFROM Stand, Bon\n"
                 "WHERE Stand.StandNr = Bon.StandNr\nGROUP BY Stand.StandNr\nHAVING Bons > 100;",
                 db="stadtfest", pruefung="fehlerhaft"))
    d.absatz("Stil 2026:")
    sql = ("SELECT Stand.StandNr, Stand.Name, COUNT(*) AS Bons\nFROM Stand, Bon\nWHERE Stand.StandNr = Bon.StandNr\n"
           "GROUP BY Stand.StandNr, Stand.Name\nHAVING COUNT(*) > 100\nORDER BY Bons DESC;")
    d.inhalt(SQL(sql, db="stadtfest"), Ergebnis(*abfrage_ergebnis(sql, "stadtfest")))

    d.ueberschrift("6 · Prognose 2027/28", 1)
    d.inhalt(Liste([
        "**Rekursive Beziehung** (neu im EH-Beispiel *istPateVon*, bisher nie geprüft) → Akte 5 DS 10/11, Selbstverbund "
        "über den rekursiven Fremdschlüssel (K37, K24).",
        "**Zwei Fremdschlüssel auf dieselbe Tabelle** (Tabelle zweimal mit Alias) → Akte 2 DS 4 (K24).",
        "**IS NULL, IN, `!=`, LIMIT** im SQL-Teil → Akte 1 DS 2, Akte 3 DS 6 (K18, K22).",
        "**Konsequente GROUP-BY-Regel** in den Erwartungshorizonten → Akte 3 DS 5 (K26, P3).",
        "ER **von Grund auf entwickeln**, Kardinalitäten begründen, ternär m:n:p, zwei Rollenbeziehungen → Akte 5.",
    ]))
    d.ueberschrift("7 · Offene Zuordnungen (bitte an den Originalaufgaben ergänzen)", 1)
    d.inhalt(Liste([
        "Die **19. Klausur** (10. gA-Klausur) ist in den Planungsunterlagen nicht mit Jahr/Termin benannt.",
        "Welche Klausuren die Typen **Anomalien (18)**, **GROUP BY (16)**, **Schlüsseleignung (14)** und "
        "**Abfrage analysieren (14)** enthalten, ist nur in der Summe belegt.",
        "**GROUP-BY-Verstoß** in älteren Lösungen: „2023 eA“ und „2025 gA“ – Haupt- oder Nachtermin offen.",
        "**Alias in HAVING** in älteren Lösungen: betroffene Klausuren nicht benannt.",
        "Teilaufgaben-Buchstaben nur dort, wo Spec §5 sie nennt; für die übrigen Belege offen.",
    ]))
    return d.speichern(ZIEL / "DD_Abi-Kompass.docx")


# ================================================================== Ermittler-Handbuch (8 Werkzeugkarten)
def _karte(d, zelle, nr, titel, rechts, spalten, breite, unten=()):
    """Füllt eine Kartenzelle: Kopf „WERKZEUGKARTE n · …“ + Titel, darunter eine oder zwei Spalten mit Inhalten."""
    p = zelle.paragraphs[0]
    p.style = d.doc.styles["DD Tabellentext"]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(breite - 1.5), WD_TAB_ALIGNMENT.RIGHT)
    d._etikett(p, f"ERMITTLER-HANDBUCH · WERKZEUGKARTE {nr}", groesse=7, farbe=A.SIEGELROT, fett=True)
    d._etikett(p, f"\t{rechts}", groesse=7, farbe=A.TINTE_2, fett=True)
    p = zelle.add_paragraph(style="DD Tabellentext")
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(4)
    d._lauf_stil(p.add_run(titel), schrift=A.SCHRIFT_KOPF, groesse=16)
    innen = breite - 1.4
    if len(spalten) == 1:
        d._fuelle(zelle, spalten[0], innen)
        return
    anteile = [s[0] for s in spalten]
    breiten = [innen * a / sum(anteile) for a in anteile]
    tab = zelle.add_table(rows=1, cols=len(spalten))
    A._tabelle_breiten(tab, breiten)
    A._tabelle_ohne_rahmen(tab)
    for i, (_, inhalt) in enumerate(spalten):
        z = tab.rows[0].cells[i]
        A._zelle_innenabstand(z, 0, 0.05 if i == 0 else 0.25, 0, 0.25 if i == 0 else 0.05)
        z.paragraphs[0].style = d.doc.styles["DD Tabellentext"]
        erster = z.paragraphs[0]
        d._fuelle(z, inhalt, breiten[i] - 0.3)
        if not erster.text and erster._p.getnext() is not None:  # leeren Startabsatz entfernen
            erster._p.getparent().remove(erster._p)
    if unten:
        d._fuelle(zelle, list(unten), innen)


def er_schule():
    """Mini-ER „Schule“ für Werkzeugkarte 4 (eigene Daten, Layout wie Bauplan-Tisch)."""
    def ent(i, name, attr):
        return {"id": i, "name": name, "attribute": [{"name": a, "schluessel": n == 0} for n, a in enumerate(attr)]}
    graph = {"entitaeten": [ent("e1", "Lehrkraft", ["Kuerzel", "Name"]), ent("e2", "Kurs", ["KursNr", "Fach"]),
                            ent("e3", "Schueler", ["SchuelerNr", "Name"])],
             "beziehungen": [{"id": "b1", "name": "leitet", "attribute": [],
                              "enden": [{"entitaet": "e1", "kard": "1"}, {"entitaet": "e2", "kard": "n"}]},
                             {"id": "b2", "name": "belegt", "attribute": [{"name": "Note"}],
                              "enden": [{"entitaet": "e2", "kard": "m"}, {"entitaet": "e3", "kard": "n"}]}]}
    layout = {"e1": {"x": 156, "y": 150}, "e1.0": {"x": 92, "y": 60}, "e1.1": {"x": 220, "y": 60},
              "e2": {"x": 616, "y": 150}, "e2.0": {"x": 552, "y": 60}, "e2.1": {"x": 680, "y": 60},
              "e3": {"x": 1076, "y": 150}, "e3.0": {"x": 1012, "y": 60}, "e3.1": {"x": 1140, "y": 60},
              "b1": {"x": 386, "y": 150}, "b2": {"x": 846, "y": 150}, "b2.0": {"x": 846, "y": 236}}
    return graph, layout


def karten_inhalte():
    """Die acht Werkzeugkarten: (nr, titel, rechts, spalten) – spalten = [(anteil, [Inhalte…]), …]."""
    k = []
    # 1 SQL nach EH §8
    k.append((1, "SQL nach EH §8", "GA/EA", [
        (1.05, [
            "**So ist jede Abfrage gebaut** (eckige Klammern = optional):",
            SQL("SELECT [DISTINCT] * | Spalte [AS Alias], …\nFROM Tabelle [Alias], …\n[WHERE Bedingung]\n"
                "[GROUP BY Spalte, …\n  [HAVING Gruppenbedingung]]\n[ORDER BY Spalte [ASC | DESC], …]\n[LIMIT Anzahl]",
                pruefung="aus"),
            "**Beispiel** (Schulbibliothek, eigene Daten):",
            SQL("SELECT Titel, Jahr\nFROM Buch\nWHERE (Fach = 'Kunst' OR Fach = 'Mathe')\n  AND Jahr >= 2020\n"
                "ORDER BY Jahr DESC, Titel;", db="keine"),
        ]),
        (1.0, [
            "**Bausteine im Abitur**",
            Liste(["Vergleich `= != < > <= >=` · `NOT` · `AND` vor `OR` → Klammern",
                   "`BETWEEN 10 AND 20` (Grenzen inklusive) · `IN ('A', 'B')`",
                   "`LIKE 'M%'` (% = beliebig viele Zeichen), `LIKE 'M_ier'` (_ = genau eins)",
                   "`IS NULL` / `IS NOT NULL` – nie `= NULL`",
                   "Rechnen `+ - * /` · Aggregate `COUNT SUM AVG MIN MAX`",
                   "Literale: Datum `'2026-09-30'`, Text `'Hafen'`, Zahl `4.5`"]),
            "**Stil:** Schlüsselwörter groß, je Klausel eine Zeile, Tabellenalias ohne AS, Spaltenalias mit AS.",
            "**Nicht im Abitur:** JOIN … ON, Unterabfragen, UNION, CREATE/INSERT; nie `<>`, sondern `!=`.",
        ]),
    ]))
    # 2 Auswertungsreihenfolge
    k.append((2, "Auswertungsreihenfolge (Zeitlupe)", "GA/EA", [
        (1.0, [
            "Die Datenbank liest eine Abfrage **nicht** von oben nach unten, sondern in dieser Reihenfolge:",
            Tabelle(["Schritt", "Klausel", "macht …"], [
                ["1", "FROM", "alle Zeilen der Tabellen (bei mehreren: jede mit jeder)"],
                ["2", "WHERE", "behält passende **Zeilen**"],
                ["3", "GROUP BY", "bildet Gruppen gleicher Werte"],
                ["4", "HAVING", "behält passende **Gruppen**"],
                ["5", "SELECT", "wählt/berechnet Spalten, Alias entsteht"],
                ["6", "ORDER BY", "sortiert (Alias erlaubt)"],
                ["7", "LIMIT", "schneidet nach n Zeilen ab"]], breiten=[1.7, 2.2, 4.8], groesse=9.3),
        ]),
        (0.95, [
            "**Daraus folgt**",
            Liste(["WHERE kennt noch keine Gruppen → Bedingungen mit COUNT/SUM gehören in **HAVING** (P1).",
                   "HAVING läuft vor SELECT → **kein Alias in HAVING**, Aggregat wiederholen.",
                   "ORDER BY läuft nach SELECT → Alias erlaubt: `ORDER BY Anzahl DESC`.",
                   "LIMIT zuletzt → Top-n nur eindeutig ohne Gleichstand an Platz n.",
                   "Ohne GROUP BY liefert ein Aggregat genau **eine** Zeile."]),
            "**Zeitlupe im Lab:** Konsole → „Zeitlupe“ zeigt jeden Schritt mit Zwischentabelle und markiert, welche "
            "Zeilen wegfallen.",
        ]),
    ]))
    # 3 Verbund-Rezept
    k.append((3, "Verbund-Rezept", "GA/EA", [
        (0.85, [
            "**Ermittlungsweg:** Frage → benötigte Attribute → Tabellen → Verbundbedingungen → Filter → Gruppierung → "
            "Sortierung.",
            Liste(["alle beteiligten Tabellen in FROM (Alias ohne AS)",
                   "je verbundenem Tabellenpaar eine Bedingung **FK = PK** – n Tabellen brauchen mindestens n − 1",
                   "gleichnamige Spalten immer mit Präfix: `s.Name`",
                   "erst danach weitere Bedingungen mit AND"], nummeriert=True),
            "**Zwei Fremdschlüssel auf dieselbe Tabelle** (Heim und Gast): Tabelle zweimal mit zwei Aliasen.",
        ]),
        (1.2, [
            Schema("Spiel(<u>SpielNr</u>, ↑HeimTeam, ↑GastTeam, Datum)"),
            Schema("Team(<u>TeamNr</u>, Name, Schule)"),
            SQL("SELECT s.SpielNr, h.Name AS Heim, g.Name AS Gast\nFROM Spiel s, Team h, Team g\n"
                "WHERE s.HeimTeam = h.TeamNr\n  AND s.GastTeam = g.TeamNr\n  AND s.Datum = '2026-10-01';", db="keine"),
            "**Ohne Schritt 2** entsteht das Kreuzprodukt (jede Zeile mit jeder, P2). **Vervielfachung:** Nur Tabellen "
            "verbinden, die die Frage braucht – zwei 1:n-Wege nebeneinander verfälschen COUNT/SUM.",
        ]),
    ]))
    # 4 ER-Notation EH §7
    k.append((4, "ER-Notation nach EH §7 (Chen)", "NUR EA", [
        (1.0, [
            Tabelle(["Symbol", "Bedeutung"], [
                ["Rechteck", "**Entitätstyp** (Singular): Schüler"],
                ["Raute", "**Beziehungstyp** (Verb, lowerCamelCase): belegt"],
                ["Ellipse", "**Attribut**; Schlüsselattribut <u>unterstrichen</u>"],
                ["1 / n / m an der Kante", "**Kardinalität** (ternär: m, n, p)"],
                ["Ellipse an der Raute", "Attribut des Beziehungstyps (z. B. Note)"]], breiten=[3.4, 5.3], groesse=9),
            "**Fachbegriffe:** Entitätstyp, Beziehungstyp, Schlüsselattribut, Kardinalität – nie „Entitätsmenge“; "
            "keine (min,max)- oder Krähenfuß-Notation.",
        ]),
        (1.0, [
            "**Lesen in beide Richtungen** (Beispiel Schule):",
            Liste(["*Lehrkraft —1— leitet —n— Kurs:* Eine Lehrkraft leitet mehrere Kurse, jeder Kurs wird von genau "
                   "einer Lehrkraft geleitet.",
                   "*Schüler —n— belegt —m— Kurs* (Note): Schüler belegen mehrere Kurse, Kurse haben mehrere Schüler; "
                   "die Note gehört zum Paar.",
                   "**Rekursiv:** Schüler —1— *istPateVon* —n— Schüler (Rollen an die Kanten: Pate, Patenkind).",
                   "**Zwei Rollen:** Team *spieltHeim* / *spieltGast* Spiel = zwei Beziehungstypen."]),
            "**Grenzen:** „genau drei“, „höchstens 12“ lassen sich mit 1/n/m nicht ausdrücken → Anmerkung zum Modell.",
        ]),
    ], [ER(*er_schule(), breite_cm=15.5)]))
    # 5 Überführungsregeln
    k.append((5, "Überführungsregeln ER → Relationenschema", "NUR EA", [
        (1.0, [
            Tabelle(["Regel", "Umsetzung"], [
                ["R1 Entitätstyp", "eigene Relation; Schlüsselattribut → Primärschlüssel"],
                ["R2 1:n", "Fremdschlüssel auf der **n-Seite** (Beziehungsattribute ebenfalls dorthin)"],
                ["R3 n:m", "eigene Relation, Schlüssel aus **beiden Fremdschlüsseln** (+ Beziehungsattribute); kommen "
                           "Paare mehrfach vor → zusätzliches Schlüsselattribut"],
                ["R4 1:1", "Fremdschlüssel auf einer Seite (oder zusammenlegen)"],
                ["R5 rekursiv", "wie R2/R3, Fremdschlüssel verweist auf die eigene Relation; Name nach Rolle"],
                ["R6 ternär", "eigene Relation mit drei Fremdschlüsseln"]], breiten=[2.6, 6.1], groesse=9),
        ]),
        (1.0, [
            "**Beispiel Schule:**",
            Schema("Lehrkraft(<u>Kuerzel</u>, Name)"),
            Schema("Kurs(<u>KursNr</u>, Fach, ↑Kuerzel)"),
            Schema("Schueler(<u>SchuelerNr</u>, Name, ↑PateNr)"),
            Schema("belegt(<u>↑SchuelerNr, ↑KursNr</u>, Note)"),
            "**Begründen:** 1:n braucht keine eigene Relation – auf der n-Seite steht genau ein Wert. Bei n:m passt auf "
            "keiner Seite ein einzelner Wert (sonst nicht-atomare Werte oder Redundanz → Anomalien).",
            "**Rückrichtung:** Relation mit eigenem Schlüssel → Entitätstyp; FK → 1:n; Schlüssel aus zwei FKs → n:m.",
        ]),
    ]))
    # 6 Schlüssel- und Anomalie-Check
    k.append((6, "Schlüssel- und Anomalie-Check", "GA/EA", [
        (1.0, [
            "**Schlüssel-Check** – ein Primärschlüssel ist …",
            Liste(["eindeutig für **alle möglichen** Daten (ein Ausschnitt kann das nur widerlegen, nie beweisen – P8)",
                   "nie NULL",
                   "minimal (kein Attribut weglassbar)",
                   "zusammengesetzt, wenn erst die Kombination eindeutig ist – Achtung bei Mehrfachvorkommen",
                   "künstlich (laufende Nr.), wenn kein natürlicher Schlüssel sicher ist"]),
            "**Redundanz:** dieselbe Information steht mehrfach – Ursache der Änderungsanomalie.",
        ]),
        (1.05, [
            Tabelle(["Anomalie", "erläutern: Art, Datensatz, Folge"], [
                ["Einfüge-", "Neues lässt sich nicht eintragen, weil anderes fehlt (Abteilung ohne Mitglied)."],
                ["Änderungs-", "Ein Wert steht mehrfach; wird eine Stelle vergessen, widerspricht sich die Tabelle."],
                ["Lösch-", "Mit einer Zeile verschwindet mehr als gewollt (letztes Mitglied → Abteilung weg)."]],
                breiten=[2.2, 6.2], groesse=9),
            "**Zusammengelegte Tabelle:** Hängt ein Nicht-Schlüssel-Attribut nur von einem Teil des Schlüssels oder "
            "von einem anderen Nicht-Schlüssel-Attribut ab? → Anomalien. Sonst: **keine Anomalie möglich**.",
            "**Zerlegen (eA):** verlustfrei (Verbund ergibt das Original) und redundanzfrei (jede Information genau einmal).",
        ]),
    ]))
    # 7 Operatoren
    k.append((7, "Operatoren in Datenbank-Aufgaben", "ABITUR", [
        (1.0, [
            Tabelle(["Operator", "Das wird erwartet"], [
                ["nennen, angeben", "ohne Begründung aufzählen bzw. Ergebnis hinschreiben"],
                ["beschreiben", "Sachverhalt in eigenen Worten wiedergeben (auch: Vorgehen in Worten)"],
                ["erläutern", "beschreiben **und** an einem Beispiel/Datensatz verständlich machen"],
                ["analysieren, untersuchen", "Bestandteile und Wirkung herausarbeiten – bei Abfragen: welche "
                 "Information im **Sachzusammenhang**"],
                ["begründen", "Aussage durch Argumente stützen („weil …“)"]], breiten=[3.1, 5.6], groesse=9),
        ]),
        (1.0, [
            Tabelle(["Operator", "Das wird erwartet"], [
                ["entwerfen, entwickeln", "eigene Lösung erstellen (SQL-Abfrage, ER-Modell)"],
                ["ergänzen, erweitern", "Gegebenes vervollständigen, ohne Bestehendes zu verfälschen"],
                ["erstellen (auch „überführen“)", "z. B. aus dem ER-Modell ein Relationenschema"],
                ["vergleichen", "Gemeinsamkeiten **und** Unterschiede kriterienorientiert"],
                ["beurteilen", "begründetes Urteil mit Kriterien, Alternativen benennen"],
                ["erörtern", "Pro und Contra abwägen, eigene Position begründen (Datenschutz)"]],
                breiten=[3.1, 5.6], groesse=9),
            "Operatoren sind in Aufgaben **fett**; die BE-Zahl zeigt den erwarteten Umfang (P10).",
        ]),
    ]))
    # 8 Abi-Check
    k.append((8, "Abi-Check P1–P10 und Zeitbudget", "ABITUR", [
        (1.15, [
            Tabelle(["Falle", "Prüffrage vor der Abgabe"], [
                ["P1", "Zeilenbedingung in WHERE, Gruppenbedingung (COUNT/SUM …) in HAVING?"],
                ["P2", "Jede Tabelle verbunden (n − 1 Bedingungen)? Nichts doppelt gezählt?"],
                ["P3", "Jede SELECT-Spalte aggregiert oder in GROUP BY? Nach Schlüssel gruppiert?"],
                ["P4", "„Anzahl Käufe“ = COUNT(*), „Anzahl Stück“ = SUM(Menge)?"],
                ["P5", "`IS NULL`, Text in `'…'`, Datum `'JJJJ-MM-TT'`?"],
                ["P6", "Klammern bei AND/OR, Grenzen (`>` oder `>=`), LIKE-Platzhalter?"],
                ["P7", "Präfix bei gleichnamigen Spalten, Alias richtig?"],
                ["P8", "Schlüssel für alle möglichen Daten eindeutig, nie NULL, minimal?"],
                ["P9", "Kardinalitäten, Schlüssel unterstrichen, FK auf der n-Seite, n:m eigene Relation?"],
                ["P10", "Operator erfüllt: im Sachzusammenhang, mit Datensatz, mit Begründung?"]],
                breiten=[1.55, 8.1], groesse=8.8),
        ]),
        (0.85, [
            "**Zeitbudget**",
            Liste(["Faustregel: **gut 2 Minuten je BE** (Generalprobe: 30 BE ≈ 70 Minuten).",
                   "Erst alle Teilaufgaben lesen, Schema/ER markieren (5 Minuten).",
                   "Mit sicheren Teilen beginnen; an einer Teilaufgabe nie länger als BE × 3 Minuten.",
                   "Am Ende 5 Minuten für den Abi-Check links."]),
            "**SQL-Muster im Abitur:** 2–3 Abfragen entwerfen + 1 gegebene Abfrage analysieren.",
        ]),
    ]))
    return k


def handbuch():
    d = AktenDokument("handbuch", 0, "Ermittler-Handbuch", fassung="schueler", material="Ermittler-Handbuch",
                      stempel=None)
    st = d.doc.styles
    st["DD Tabellentext"].font.size = Pt(9.6)
    st["SQL-Code"].font.size = Pt(9)
    st["SQL-Code"].paragraph_format.space_before = Pt(2)
    st["SQL-Code"].paragraph_format.space_after = Pt(3)
    s = d.doc.sections[0]
    s.top_margin = s.bottom_margin = Cm(1.0)
    s.left_margin = s.right_margin = Cm(1.2)
    s.header_distance = s.footer_distance = Cm(0.4)
    d.breite = 21.0 - 2.4
    hoehe = (29.7 - 2.0 - 1.3) / 2
    karten = karten_inhalte()
    for i in range(0, len(karten), 2):
        if i:
            d.seitenumbruch()
        paar = karten[i:i + 2]
        tab = d.doc.add_table(rows=len(paar), cols=1)
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        A._tabelle_breiten(tab, [d.breite])
        A._tabelle_ohne_rahmen(tab)
        for zi, (nr, titel, rechts, spalten, *rest) in enumerate(paar):
            zeile = tab.rows[zi]
            trpr = zeile._tr.get_or_add_trPr()
            trpr.append(A._el("w:trHeight", val=int(Cm(hoehe - 0.35) / 635), hRule="exact"))
            A._zeile_nicht_teilen(zeile)
            z = zeile.cells[0]
            A._zelle_innenabstand(z, 0.45, 0.7, 0.3, 0.7)
            A._zelle_rahmen(z, A.RAND_STARK, 0.5, art="dashed", seiten=("bottom",) if zi == 0 else ("top",))
            _karte(d, z, nr, titel, rechts, spalten, d.breite, unten=rest[0] if rest else ())
    return d.speichern(ZIEL / "DD_Ermittler-Handbuch.docx")


def main():
    ZIEL.mkdir(parents=True, exist_ok=True)
    for p in (gesamtuebersicht(), abi_kompass(), handbuch()):
        print(f"✔ {p.relative_to(REPO)}")


if __name__ == "__main__":
    main()
