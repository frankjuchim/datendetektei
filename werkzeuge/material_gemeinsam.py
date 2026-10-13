#!/usr/bin/env python3
"""Gemeinsamer Kern der Material-Pipeline (Spec §12.2, §12.4; docs/MATERIALIEN.md).

Zuordnung: seed/arbeitsblaetter.json (Gruppen = Akten, Gesamt, Kür, Videos; je Material art, titel, quelle, loesung,
datei …). Daraus entsteht der Plan aller erzeugten PDFs:

  Arbeitsblaetter_PDF/  Schülerfassungen (angemeldet): Arbeitsblätter, Unplugged, Handbuch, Kür, Präsentations-Handouts
                        (ohne [NUR LEHRKRAFT]-Folien, ohne Notizen), Warm-up-Karten ohne Lösungsteil
  Loesungen_Portal/     nur Lehrkraft: Lösungen (_Loesung), Erwartungshorizonte (_Erwartungshorizont),
                        Warm-up-Karten mit Lösungen (_Loesungen)
  Lehrkraft_PDF/        nur Lehrkraft: Präsentationen mit Notizenseiten, Lehrerhinweise, Gesamtmaterial, Generalprobe
                        (Standard), Klausuren, Video-Folien
  Originale/            nur verzeichnis.json: Whitelist der docx/pptx (Superadmin)

Je Zielordner: verzeichnis.json (Whitelist für den Server, aus dem Plan erzeugt) und quellen.json (SHA-256 der Quelle
je PDF → Aktualitätsprüfung ohne LibreOffice). Die Skripte arbeitsblaetter_pdf.py, loesungen_pdf.py, lehrkraft_pdf.py,
praesentationen_pdf.py und originale_verzeichnis.py sind dünne Aufrufer von main(erzeuger).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

WERKZEUGE = Path(__file__).resolve().parent
REPO = WERKZEUGE.parent

ZIELORDNER = ("Arbeitsblaetter_PDF", "Loesungen_Portal", "Lehrkraft_PDF")
ORIGINALE = "Originale"
PDF_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,120}\.pdf$")
ARTEN = ("ab", "praesentation", "warmup", "unplugged", "handbuch", "lehrerhinweise", "gesamt", "generalprobe", "klausur",
         "kuer", "video")
LOESUNG_PFLICHT = {"ab", "generalprobe", "klausur"}
PPTX_ARTEN = {"praesentation", "video"}
NOTIZ_MARKER = "[NUR LEHRKRAFT]"
LOESUNGSTEIL_WARMUP = "Lösungen der Warm-up-Karten"
ERZEUGER = ("arbeitsblaetter", "loesungen", "lehrkraft", "praesentationen", "originale")
#: Zielordner, deren Vollständigkeit/Verzeichnis ein Erzeuger prüft (praesentationen schreibt in fremde Ordner).
EIGENE_ORDNER = {"arbeitsblaetter": ("Arbeitsblaetter_PDF",), "loesungen": ("Loesungen_Portal",),
                 "lehrkraft": ("Lehrkraft_PDF",), "praesentationen": (), "originale": ()}
HINWEIS = {
    "Arbeitsblaetter_PDF": "Schülerfassungen (nur angemeldet, Spec §12.4/E13) – nie Lösungen.",
    "Loesungen_Portal": "Lösungen und Erwartungshorizonte – nur Lehrkraft.",
    "Lehrkraft_PDF": "Lehrkraft-Materialien – nur Lehrkraft.",
    "Originale": "docx/pptx-Originale (Pfade relativ zu materialien/) – nur Superadmin.",
}


@dataclass(frozen=True)
class Ziel:
    ordner: str      # Zielordner unter materialien/
    datei: str       # PDF-Name (flach, ohne Umlaute)
    quelle: str      # docx/pptx relativ zu materialien/
    verfahren: str   # pdf | ohne-loesungsteil | handout | notizen
    erzeuger: str    # arbeitsblaetter | loesungen | lehrkraft | praesentationen
    fassung: str     # schueler | handout | loesung | lehrkraft
    eintrag: tuple   # (gruppe-index, material-index) für das Verzeichnis
    zusatz: str      # Titelzusatz

    @property
    def pfad(self):
        return f"{self.ordner}/{self.datei}"


def mit_zusatz(datei, zusatz):
    return datei[:-4] + zusatz + ".pdf"


def lade_zuordnung(wurzel=REPO):
    datei = Path(wurzel) / "seed" / "arbeitsblaetter.json"
    return json.loads(datei.read_text(encoding="utf-8"))


def pruefe_zuordnung(z, wurzel=REPO, quellen_pruefen=True):
    """Formfehler der Zuordnung → Liste von Meldungen (leer = in Ordnung)."""
    m = []
    if not isinstance(z, dict) or not isinstance(z.get("gruppen"), list):
        return ["seed/arbeitsblaetter.json: Feld „gruppen“ (Liste) fehlt"]
    gesehen = {}
    mat = Path(wurzel) / "materialien"
    for gi, g in enumerate(z["gruppen"]):
        ort = f"Gruppe {g.get('gruppe', gi)}"
        for f in ("gruppe", "ordner", "titel", "materialien"):
            if f not in g:
                m.append(f"{ort}: Feld „{f}“ fehlt")
        for mi, x in enumerate(g.get("materialien", [])):
            o = f"{ort}, Material {mi + 1} ({x.get('titel', '?')})"
            art = x.get("art")
            if art not in ARTEN:
                m.append(f"{o}: unbekannte Art „{art}“ (erlaubt: {', '.join(ARTEN)})")
                continue
            for f in ("titel", "quelle", "datei"):
                if not x.get(f):
                    m.append(f"{o}: Feld „{f}“ fehlt")
            if art in LOESUNG_PFLICHT and not x.get("loesung"):
                m.append(f"{o}: loesung fehlt (Pflicht für {art})")
            datei = x.get("datei", "")
            if datei and not PDF_NAME.match(datei):
                m.append(f"{o}: PDF-Name „{datei}“ ungültig (nur A–Z, a–z, 0–9, . _ -, keine Umlaute, Endung .pdf)")
            if datei and datei in gesehen:
                m.append(f"{o}: PDF-Name „{datei}“ doppelt (auch {gesehen[datei]})")
            gesehen.setdefault(datei, o)
            quelle = x.get("quelle", "")
            if quelle:
                endung = Path(quelle).suffix
                if art in PPTX_ARTEN and endung != ".pptx":
                    m.append(f"{o}: Quelle muss eine .pptx sein")
                if art not in PPTX_ARTEN and endung != ".docx":
                    m.append(f"{o}: Quelle muss eine .docx sein")
            if x.get("freigabe") not in (None, "lehrkraft", "schueler"):
                m.append(f"{o}: freigabe muss „lehrkraft“ oder „schueler“ sein")
            if quellen_pruefen:
                for f in ("quelle", "loesung"):
                    if x.get(f) and not (mat / x[f]).exists():
                        m.append(f"{o}: {f} materialien/{x[f]} fehlt")
    return m


def plan(z):
    """Alle zu erzeugenden PDFs (Liste von Ziel) in fester Reihenfolge."""
    aus = []
    for gi, g in enumerate(z.get("gruppen", [])):
        for mi, x in enumerate(g.get("materialien", [])):
            art, d, q, lsg = x.get("art"), x.get("datei"), x.get("quelle"), x.get("loesung")
            if art not in ARTEN or not d or not q:
                continue
            e = (gi, mi)

            def ziel(ordner, datei, quelle, verfahren, erzeuger, fassung, zusatz=""):
                aus.append(Ziel(ordner, datei, quelle, verfahren, erzeuger, fassung, e, zusatz))

            if art in ("ab", "kuer", "unplugged", "handbuch"):
                ziel("Arbeitsblaetter_PDF", d, q, "pdf", "arbeitsblaetter", "schueler")
                if lsg:
                    ziel("Loesungen_Portal", mit_zusatz(d, "_Loesung"), lsg, "pdf", "loesungen", "loesung", " – Lösung")
            elif art == "warmup":
                ziel("Arbeitsblaetter_PDF", d, q, "ohne-loesungsteil", "arbeitsblaetter", "schueler")
                ziel("Loesungen_Portal", mit_zusatz(d, "_Loesungen"), q, "pdf", "loesungen", "loesung", " – mit Lösungen")
            elif art == "praesentation":
                ziel("Arbeitsblaetter_PDF", mit_zusatz(d, "_Handout"), q, "handout", "praesentationen", "handout", " (Handout)")
                ziel("Lehrkraft_PDF", d, q, "notizen", "praesentationen", "lehrkraft", " (mit Notizen)")
            elif art == "video":
                ziel("Lehrkraft_PDF", d, q, "pdf", "praesentationen", "lehrkraft")
            elif art in ("lehrerhinweise", "gesamt"):
                ziel("Lehrkraft_PDF", d, q, "pdf", "lehrkraft", "lehrkraft")
            elif art in ("generalprobe", "klausur"):
                if art == "generalprobe" and x.get("freigabe") == "schueler":
                    ziel("Arbeitsblaetter_PDF", d, q, "pdf", "arbeitsblaetter", "schueler")
                else:
                    ziel("Lehrkraft_PDF", d, q, "pdf", "lehrkraft", "lehrkraft")
                ziel("Loesungen_Portal", mit_zusatz(d, "_Erwartungshorizont"), lsg, "pdf", "loesungen", "loesung",
                     " – Erwartungshorizont")
    return aus


def _titel(x):
    return f"AB {x['nr']} · {x['titel']}" if x.get("art") == "ab" and x.get("nr") else x["titel"]


def _eintrag(g, x, datei, fassung, zusatz=""):
    e = {"datei": datei, "titel": _titel(x) + zusatz, "akte": g.get("akte"), "gruppe": g["gruppe"], "art": x["art"],
         "fassung": fassung}
    for f in ("niveau", "ermittlung", "ds"):
        if x.get(f) is not None:
            e[f] = x[f]
    return e


def verzeichnisse(z):
    """{ordner: {"_hinweis", "dateien": [...]}} für die drei PDF-Ordner und Originale (Whitelist-Format §12.4)."""
    aus = {o: [] for o in ZIELORDNER}
    gruppen = z.get("gruppen", [])
    for ziel in plan(z):
        g = gruppen[ziel.eintrag[0]]
        x = g["materialien"][ziel.eintrag[1]]
        aus[ziel.ordner].append(_eintrag(g, x, ziel.datei, ziel.fassung, ziel.zusatz))
    orig = []
    for g in gruppen:
        for x in g.get("materialien", []):
            if x.get("art") not in ARTEN:
                continue
            for feld, fassung, zusatz in (("quelle", "schueler" if x["art"] not in ("lehrerhinweise", "gesamt", "klausur", "video",
                                                                                     "generalprobe") else "lehrkraft", ""),
                                          ("loesung", "loesung", " – Lösung")):
                if x.get(feld):
                    orig.append(_eintrag(g, x, x[feld], fassung, zusatz))
    aus[ORIGINALE] = orig
    return {o: {"_hinweis": f"ERZEUGT aus seed/arbeitsblaetter.json (werkzeuge/*_pdf.py) – nicht von Hand ändern. {HINWEIS[o]}",
                "dateien": dateien} for o, dateien in aus.items()}


def json_text(daten):
    return json.dumps(daten, ensure_ascii=False, indent=1) + "\n"


def sha256(pfad):
    h = hashlib.sha256()
    with open(pfad, "rb") as f:
        for block in iter(lambda: f.read(1 << 16), b""):
            h.update(block)
    return h.hexdigest()


def lese_quellen(ordner):
    datei = ordner / "quellen.json"
    if not datei.exists():
        return {}
    try:
        return json.loads(datei.read_text(encoding="utf-8")).get("dateien", {})
    except json.JSONDecodeError:
        return {}


def schreibe_quellen(ordner, eintraege):
    ordner.mkdir(parents=True, exist_ok=True)
    (ordner / "quellen.json").write_text(json_text({
        "_hinweis": "ERZEUGT von werkzeuge/*_pdf.py – SHA-256 der Quelle je PDF; die CI prüft damit, ob die PDFs aktuell sind.",
        "dateien": dict(sorted(eintraege.items()))}), encoding="utf-8")


# ---------------------------------------------------------------- Prüfen (ohne LibreOffice)
def pruefe(wurzel, erzeuger):
    """Probleme eines Erzeugers: Quelle fehlt, PDF fehlt, PDF veraltet (Hash), überzählige PDFs, Verzeichnis veraltet."""
    wurzel = Path(wurzel)
    mat = wurzel / "materialien"
    z = lade_zuordnung(wurzel)
    probleme = pruefe_zuordnung(z, wurzel)
    ziele = plan(z)
    v = verzeichnisse(z)
    if erzeuger == "originale":
        datei = mat / ORIGINALE / "verzeichnis.json"
        if not datei.exists() or datei.read_text(encoding="utf-8") != json_text(v[ORIGINALE]):
            probleme.append("materialien/Originale/verzeichnis.json veraltet (python3 werkzeuge/originale_verzeichnis.py)")
        return probleme
    quellen_cache = {}
    for ziel in ziele:
        if ziel.erzeuger != erzeuger:
            continue
        pdf = mat / ziel.ordner / ziel.datei
        quelle = mat / ziel.quelle
        if not pdf.exists():
            probleme.append(f"fehlt: materialien/{ziel.pfad}")
            continue
        if ziel.ordner not in quellen_cache:
            quellen_cache[ziel.ordner] = lese_quellen(mat / ziel.ordner)
        q = quellen_cache[ziel.ordner].get(ziel.datei)
        if not quelle.exists():
            continue  # schon als Zuordnungsfehler gemeldet
        if not q or q.get("quelle") != ziel.quelle or q.get("verfahren") != ziel.verfahren or q.get("sha256") != sha256(quelle):
            probleme.append(f"veraltet: materialien/{ziel.pfad} (Quelle materialien/{ziel.quelle} geändert – neu erzeugen)")
    for ordner in EIGENE_ORDNER[erzeuger]:
        soll = {zz.datei for zz in ziele if zz.ordner == ordner}
        da = {p.name for p in (mat / ordner).glob("*.pdf")} if (mat / ordner).exists() else set()
        probleme += [f"überzählig: materialien/{ordner}/{d} (nicht in seed/arbeitsblaetter.json)" for d in sorted(da - soll)]
        datei = mat / ordner / "verzeichnis.json"
        if not datei.exists() or datei.read_text(encoding="utf-8") != json_text(v[ordner]):
            probleme.append(f"materialien/{ordner}/verzeichnis.json veraltet (npm run {erzeuger})")
    return probleme


# ---------------------------------------------------------------- Erzeugen (mit LibreOffice)
def fonts_installieren():
    """Schriften aus werkzeuge/fonts für LibreOffice bereitstellen (~/.local/share/fonts/datendetektei)."""
    ziel = Path.home() / ".local" / "share" / "fonts" / "datendetektei"
    neu = False
    for ttf in sorted((WERKZEUGE / "fonts").glob("*.ttf")):
        z = ziel / ttf.name
        if not z.exists() or z.stat().st_size != ttf.stat().st_size:
            ziel.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(ttf, z)
            neu = True
    if neu and shutil.which("fc-cache"):
        subprocess.run(["fc-cache", "-f", str(ziel)], capture_output=True)


def soffice():
    s = shutil.which("soffice") or shutil.which("libreoffice")
    if not s:
        sys.exit("soffice (LibreOffice) nicht gefunden – zum Erzeugen nötig (zum Prüfen: --pruefen).")
    return s


NOTIZEN_FILTER = ('pdf:impress_pdf_Export:{"ExportNotesPages":{"type":"boolean","value":"true"},'
                  '"ExportOnlyNotesPages":{"type":"boolean","value":"true"}}')


def nach_pdf(quelle, ausordner, profil, filter_="pdf"):
    """docx/pptx → PDF mit LibreOffice; gibt den Pfad der erzeugten PDF zurück."""
    r = subprocess.run([soffice(), f"-env:UserInstallation=file://{profil}", "--headless", "--convert-to", filter_,
                        "--outdir", str(ausordner), str(quelle)], capture_output=True, text=True, timeout=600)
    pdf = Path(ausordner) / (Path(quelle).stem + ".pdf")
    if r.returncode or not pdf.exists():
        raise RuntimeError(f"Umwandlung fehlgeschlagen: {quelle}\n{r.stderr.strip()}")
    return pdf


def handout_pptx(quelle, ziel):
    """Kopie der Präsentation ohne Folien mit „[NUR LEHRKRAFT]“ in den Notizen und ohne Sprechernotizen.
    → Anzahl entfernter Folien."""
    from pptx import Presentation
    prs = Presentation(str(quelle))
    sld_ids = prs.slides._sldIdLst
    entfernt = 0
    for folie, sld in list(zip(prs.slides, list(sld_ids))):
        notiz = folie.notes_slide.notes_text_frame.text if folie.has_notes_slide else ""
        if NOTIZ_MARKER in notiz:
            prs.part.drop_rel(sld.rId)
            sld_ids.remove(sld)
            entfernt += 1
        elif folie.has_notes_slide:
            folie.notes_slide.notes_text_frame.text = ""
    prs.save(str(ziel))
    return entfernt


def ohne_loesungsteil(pdf):
    """Schneidet die PDF vor der ersten Seite mit „Lösungen der Warm-up-Karten“ ab → (behalten, gesamt)."""
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(str(pdf))
    ab = next((i for i, s in enumerate(r.pages) if LOESUNGSTEIL_WARMUP in (s.extract_text() or "")), None)
    if not ab:
        raise RuntimeError(f"Lösungsteil „{LOESUNGSTEIL_WARMUP}“ nicht gefunden in {pdf.name} – nichts veröffentlichen.")
    w = PdfWriter()
    for s in r.pages[:ab]:
        w.add_page(s)
    with open(pdf, "wb") as f:
        w.write(f)
    return ab, len(r.pages)


def libreoffice_renderer(profil):
    def rendern(ziel, quelle, aus):
        tmp = Path(tempfile.mkdtemp(prefix="dd-pdf-", dir=profil.parent))
        try:
            if ziel.verfahren == "handout":
                kopie = tmp / quelle.name
                handout_pptx(quelle, kopie)
                pdf = nach_pdf(kopie, tmp, profil)
            elif ziel.verfahren == "notizen":
                pdf = nach_pdf(quelle, tmp, profil, NOTIZEN_FILTER)
            else:
                pdf = nach_pdf(quelle, tmp, profil)
                if ziel.verfahren == "ohne-loesungsteil":
                    ohne_loesungsteil(pdf)
            shutil.move(str(pdf), aus)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return rendern


def erzeuge(wurzel, erzeuger, renderer=None, alle=False, ausgabe=print):
    """Erzeugt die PDFs eines Erzeugers (nur fehlende/veraltete, mit alle=True alle), aktualisiert quellen.json,
    verzeichnis.json und entfernt überzählige PDFs in den eigenen Ordnern."""
    wurzel = Path(wurzel)
    mat = wurzel / "materialien"
    z = lade_zuordnung(wurzel)
    fehler = pruefe_zuordnung(z, wurzel)
    if fehler:
        raise SystemExit("seed/arbeitsblaetter.json fehlerhaft:\n  " + "\n  ".join(fehler))
    ziele = plan(z)
    v = verzeichnisse(z)
    if erzeuger == "originale":
        (mat / ORIGINALE).mkdir(parents=True, exist_ok=True)
        (mat / ORIGINALE / "verzeichnis.json").write_text(json_text(v[ORIGINALE]), encoding="utf-8")
        ausgabe(f"✔ materialien/Originale/verzeichnis.json ({len(v[ORIGINALE]['dateien'])} Originale)")
        return
    eigene = [zz for zz in ziele if zz.erzeuger == erzeuger]
    profil_dir = Path(tempfile.mkdtemp(prefix="dd-lo-"))
    profil = profil_dir / "profil"
    if renderer is None and eigene:
        fonts_installieren()
        renderer = libreoffice_renderer(profil)
    try:
        quellen = {o: lese_quellen(mat / o) for o in ZIELORDNER}
        erzeugt = 0
        for ziel in eigene:
            quelle = mat / ziel.quelle
            pdf = mat / ziel.ordner / ziel.datei
            hash_ = sha256(quelle)
            alt = quellen[ziel.ordner].get(ziel.datei, {})
            if not alle and pdf.exists() and alt.get("sha256") == hash_ and alt.get("verfahren") == ziel.verfahren \
                    and alt.get("quelle") == ziel.quelle:
                continue
            pdf.parent.mkdir(parents=True, exist_ok=True)
            renderer(ziel, quelle, pdf)
            quellen[ziel.ordner][ziel.datei] = {"quelle": ziel.quelle, "sha256": hash_, "verfahren": ziel.verfahren}
            erzeugt += 1
            ausgabe(f"✔ materialien/{ziel.quelle} → materialien/{ziel.pfad}" + (f" ({ziel.verfahren})" if ziel.verfahren != "pdf" else ""))
        betroffen = set(EIGENE_ORDNER[erzeuger]) | {zz.ordner for zz in eigene}
        for ordner in sorted(betroffen):
            soll = {zz.datei for zz in ziele if zz.ordner == ordner}
            if ordner in EIGENE_ORDNER[erzeuger] and (mat / ordner).exists():
                for p in sorted((mat / ordner).glob("*.pdf")):
                    if p.name not in soll:
                        p.unlink()
                        ausgabe(f"− entfernt: materialien/{ordner}/{p.name}")
            schreibe_quellen(mat / ordner, {d: e for d, e in quellen[ordner].items() if d in soll})
            (mat / ordner / "verzeichnis.json").write_text(json_text(v[ordner]), encoding="utf-8")
        ausgabe(f"{erzeugt} PDF(s) erzeugt, {len(eigene) - erzeugt} aktuell ({erzeuger}).")
    finally:
        shutil.rmtree(profil_dir, ignore_errors=True)


def main(erzeuger, beschreibung=""):
    ap = argparse.ArgumentParser(description=beschreibung)
    ap.add_argument("--pruefen", action="store_true", help="nichts erzeugen: Vollständigkeit + Aktualität prüfen (CI)")
    ap.add_argument("--alle", action="store_true", help="alle PDFs neu erzeugen (sonst nur fehlende/veraltete)")
    ap.add_argument("--wurzel", default=str(REPO), help=argparse.SUPPRESS)
    a = ap.parse_args()
    if a.pruefen:
        probleme = pruefe(a.wurzel, erzeuger)
        for p in probleme:
            print("✘ " + p)
        anzahl = len([z for z in plan(lade_zuordnung(a.wurzel)) if z.erzeuger == erzeuger]) if erzeuger != "originale" else None
        print(f"{erzeuger}: " + (f"{anzahl} PDF(s) geplant, " if anzahl is not None else "") + f"{len(probleme)} Problem(e).")
        sys.exit(1 if probleme else 0)
    erzeuge(a.wurzel, erzeuger, alle=a.alle)
    rest = pruefe(a.wurzel, erzeuger)
    if rest:
        for p in rest:
            print("✘ " + p)
        sys.exit(1)


if __name__ == "__main__":
    sys.exit("Bitte eines der Skripte arbeitsblaetter_pdf.py, loesungen_pdf.py, lehrkraft_pdf.py, praesentationen_pdf.py "
             "oder originale_verzeichnis.py aufrufen.")
