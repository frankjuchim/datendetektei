#!/usr/bin/env python3
"""Erklärvideos der Datendetektei (Spec §13): Folien (pptx) und Drehbuch (md) aus einer Video-Beschreibung (JSON).

  materialien/Videos/<NN>_<Kurz>/video.json            Quelle (Folien, Sprechertext, Metadaten) – hier wird bearbeitet
  → <NN>_<Kurz>/DD_Video<NN>_<Kurz>_Folien.pptx         Folien zum Abfilmen im Akten-Design; Sprechertext in den Notizen
  → <NN>_<Kurz>/DD_Video<NN>_<Kurz>_Skript.md           Drehbuch: Lernziele, YouTube-Text, Kapitel, Folie für Folie
  → <NN>_<Kurz>/bilder/*.png                            Bilder aus der Plattform (scripts/video_bilder.js, eigene Daten)

Aufbau nach Spec §13: Einstieg ≤ 30 s → Beispiel (eigene Daten) → ⏸ Pause-Aufgabe → Fallen (P-Nummern) → Merksätze
→ „Weiter im Lab“. Folien mit dem Design-Modul (AktenPraesentation(..., kamera_zone=True)): Unten rechts (ab x = 9,35 in,
y = 4,55 in) bleibt jede Folie frei – dort liegt das Kamerabild; geprüft mit pruefe_kamera_zone. SQL-Felder heißen
„SQL [db=…]“ und Ergebnistabellen „Ergebnistabelle“ – material_sql_pruefen.py prüft sie wie alle Materialien.

Aufruf: python3 werkzeuge/video_folien.py [Ordner …] [--bilder]   (ohne Ordner: alle Videos; --bilder rendert vorher die
                                                                  Plattform-Bilder neu: node scripts/video_bilder.js)
        python3 werkzeuge/video_folien.py --pruefen      (Vollständigkeit, Aktualität, Aufbau, Kamera-Zone – ohne Erzeugen)
Format der video.json und Folientypen: materialien/Videos/LIESMICH.md.
Benötigt python-pptx, Pillow, Node (SQL-Ergebnisse über scripts/sql_pruefen_cli.js); für --bilder playwright-core + Chromium.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import sys
from pathlib import Path

WERKZEUGE = Path(__file__).resolve().parent
REPO = WERKZEUGE.parent
if str(WERKZEUGE) not in sys.path:
    sys.path.insert(0, str(WERKZEUGE))

import akten_design as A  # noqa: E402
from akten_design import (AKTEN, BLAU, BLAU_FLAECHE, D_GRUEN, D_MESSING, D_SIEGELROT, D_TINTE, D_TINTE_2,  # noqa: E402
                          GRUEN, GRUEN_FLAECHE, KAMERA_X, KAMERA_Y, MESSING_FLAECHE, MESSING_TEXT, PAPIER, ROT_FLAECHE,
                          SCHRIFT, SCHRIFT_KOPF, SCHRIFT_MONO, SCHRIFT_SYMBOL, SIEGELROT, TINTE, TINTE_3, MSO_ANCHOR,
                          MSO_SHAPE, PP_ALIGN, Inches, PPt, nur_text, laeufe)

VIDEOS = REPO / "materialien" / "Videos"
BILDER_JS = REPO / "scripts" / "video_bilder.js"

# Geometrie (Zoll): linke Spalte über die volle Höhe neben der Kamera-Zone, rechte Spalte nur oberhalb
L, O, R = A.RAND_L, 1.85, A.FOLIE_B - A.RAND_L
LINKS_MAX = KAMERA_X - 0.25          # rechter Rand der linken Spalte
RECHTS_UNTEN = KAMERA_Y - 0.15       # unterer Rand der rechten Spalte bzw. eines breiten Bereichs („oben“)
UNTEN = 6.95                         # unterer Rand der linken Spalte (Fußzeile darunter)
TEILUNG = 6.6                        # Standard-Beginn der rechten Spalte

WORTE_PRO_MINUTE = 130
EINSTIEG_WORTE = 65                  # ≤ 30 s Sprechtext auf der Titelfolie
FOLIEN_MIN, FOLIEN_MAX = 8, 14
P_NUMMER = re.compile(r"\bP(?:10|[1-9])\b")

KASTEN_FARBEN = {  # Füllung, Etikett
    "blau": (BLAU_FLAECHE, BLAU), "gold": (MESSING_FLAECHE, MESSING_TEXT), "rot": (ROT_FLAECHE, SIEGELROT),
    "gruen": (GRUEN_FLAECHE, GRUEN), "papier": (PAPIER, TINTE_3)}


class Fehler(Exception):
    """Inhalt passt nicht auf die Folie bzw. die video.json ist unvollständig."""


# ================================================================== Höhen schätzen (Barlow ≈ 0,5 em je Zeichen)
def zeilen_text(text, w, g, faktor=0.5):
    """Zeilen nach Wortumbruch (gierig), Zeichenbreite faktor · g."""
    zpz = max(8, int((w - 0.12) * 72 / (g * faktor)))
    n = 0
    for absatz in str(text).split("\n"):
        zeile = 0
        n += 1
        for wort in nur_text(absatz).split(" "):
            lw = len(wort)
            if zeile and zeile + 1 + lw > zpz:
                n += 1
                zeile = lw
            else:
                zeile += (1 if zeile else 0) + lw
            while zeile > zpz:
                n += 1
                zeile -= zpz
    return n


def hoehe_text(text, w, g, zeilenabstand=1.2, faktor=0.47):
    return zeilen_text(text, w, g, faktor) * g * zeilenabstand / 72 + 0.08


def sql_groesse(text, w, g_max):
    laengste = max(len(z) for z in text.strip("\n").split("\n"))
    return max(10, min(g_max, int((w - 0.5) * 72 / (0.6 * max(laengste, 18)))))


def sql_hoehe(text, g):
    return len(text.strip("\n").split("\n")) * g * 1.25 / 72 + 0.35


def tabelle_hoehe(zeilen, g):
    return (len(zeilen) + 1) * g * 1.55 / 72


# ================================================================== Präsentation
class VideoFolien(A.AktenPraesentation):
    """AktenPraesentation mit Kamera-Zone und den Folientypen der Videos (inhalt, pause, merke, ende, titel)."""

    def __init__(self, spec, ordner, ergebnisse):
        akte = spec.get("akte") or 0
        super().__init__(akte, spec["titel"], kamera_zone=True, datenbank=spec.get("datenbank"))
        self.spec, self.ordner, self.ergebnisse = spec, ordner, ergebnisse
        a = AKTEN.get(akte, AKTEN[0])
        self.etikett = f"ERKLÄRVIDEO {spec['nr']} · " + (f"AKTE {akte} · {a['titel'].upper()}" if 1 <= akte <= 7 else "DATENDETEKTEI WESERSTEDT")

    def _folie(self, titel, etikett=None, notizen="", nur_lehrkraft=False, faden=True):
        return super()._folie(titel, etikett or self.etikett, notizen=notizen, nur_lehrkraft=nur_lehrkraft, faden=faden)

    # -------------------------------------------------------------- Bausteine
    def _absatz(self, tf, erster, text, g, farbe, schrift=SCHRIFT, fett=False, kursiv=False, abstand=4, zeichen=None,
                zeichen_farbe=None):
        p = tf.paragraphs[0] if erster else tf.add_paragraph()
        p.line_spacing = 1.05
        p.space_after = PPt(abstand)
        if zeichen:
            r = p.add_run()
            r.text = zeichen
            r.font.size, r.font.name, r.font.bold = PPt(g), SCHRIFT, True
            r.font.color.rgb = A._p_rgb(zeichen_farbe or farbe)
        self._laeufe(p, text, g, farbe, schrift, fett=fett, kursiv=kursiv)
        return p

    def _textblock(self, folie, x, y, w, h, absaetze, g, farbe, name=None):
        """absaetze: [(text, {zeichen, fett, kursiv, schrift, farbe})]"""
        tb = folie.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.03))
        for i, (t, o) in enumerate(absaetze):
            self._absatz(tf, i == 0, t, o.get("groesse", g), o.get("farbe", farbe), schrift=o.get("schrift", SCHRIFT),
                         fett=o.get("fett", False), kursiv=o.get("kursiv", False), abstand=o.get("abstand", 4),
                         zeichen=o.get("zeichen"), zeichen_farbe=o.get("zeichen_farbe"))
        return tb

    # -------------------------------------------------------------- Blöcke: Höhe und Setzen
    def _hoehe(self, b, w, s, dunkel=True):
        """Geschätzte Höhe eines Blocks bei Skalierung s (1.0 = Standardgrößen)."""
        g = 20 * s
        if isinstance(b, str):
            return hoehe_text(b, w, g) + 0.04
        if "text" in b and len(b) <= 2 and "falle" not in b:
            return hoehe_text(b["text"], w, b.get("groesse", 20) * s) + 0.04
        if "punkte" in b:
            gg = b.get("groesse", 20) * s
            return sum(hoehe_text(p, w - 0.35, gg - (3 if p.startswith("  - ") else 0)) + 0.02 for p in b["punkte"]) + 0.06
        if "sql" in b:
            gs = sql_groesse(b["sql"], w, min(b.get("groesse", 20), 20 * s))
            h = (0.34 if b.get("label") else 0) + sql_hoehe(b["sql"], gs)
            erg = self._ergebnis_tabelle(b)
            if erg is not None:
                h += 0.4 + tabelle_hoehe(erg.zeilen, b.get("tabellengroesse", 15) * s)
            return h + 0.12
        if "tabelle" in b:
            t = b["tabelle"]
            return (0.32 if t.get("titel") else 0) + tabelle_hoehe(t["zeilen"], t.get("groesse", 15) * s) + 0.12
        if "schema" in b:
            gg = b.get("groesse", 19) * s
            return sum(hoehe_text(z, w - 0.4, gg * 0.92, faktor=0.62) for z in b["schema"]) * 1.05 + 0.3 + (0.32 if b.get("titel") else 0) + 0.12
        if "bild" in b:
            return b.get("hoehe", 1.2) + 0.1
        if "kasten" in b:
            k = b["kasten"]
            gg = k.get("groesse", 17) * s
            h = 0.5 if k.get("titel") else 0.2
            h += sum(hoehe_text(t.lstrip("- "), w - 0.65, gg) + 0.03 for t in k.get("text", []))
            return h + 0.15 + 0.12
        if "falle" in b:
            gg = b.get("groesse", 16) * s
            return max(0.75, 0.42 + hoehe_text(b.get("text", ""), w - 1.25, gg)) + 0.12
        if "schritte" in b:
            gg = b.get("groesse", 16) * s
            return sum(max(0.6, hoehe_text(t, w - 0.85, gg) + 0.16) + 0.12 for t in b["schritte"])
        if "badge" in b:
            return 0.62
        if "abstand" in b:
            return b["abstand"]
        raise Fehler(f"unbekannter Block: {b!r}")

    def _ergebnis_tabelle(self, b):
        e = b.get("ergebnis")
        if not e:
            return None
        if e is True:
            kopf, zeilen = self.ergebnisse[b["_id"]]
        else:
            kopf, zeilen = e["kopf"], e["zeilen"]
        n = b.get("max_zeilen")
        if n and len(zeilen) > n:
            return A.Ergebnis(kopf, zeilen[:n], abgeschnitten=True)
        return A.Ergebnis(kopf, zeilen)

    def _setze(self, folie, b, x, y, w, s, h_bild=None):
        g = 20 * s
        if isinstance(b, str) or ("text" in b and len(b) <= 2 and "falle" not in b):
            text = b if isinstance(b, str) else b["text"]
            gg = g if isinstance(b, str) else b.get("groesse", 20) * s
            h = hoehe_text(text, w, gg)
            self._textblock(folie, x, y, w, h, [(t, {}) for t in text.split("\n")], gg, D_TINTE)
            return y + h + 0.04
        if "punkte" in b:
            gg = b.get("groesse", 20) * s
            absaetze = []
            for p in b["punkte"]:
                unter = p.startswith("  - ")
                absaetze.append((p[4:] if unter else p, {"zeichen": "      – " if unter else "▸  ", "groesse": gg - 3 if unter else gg,
                                                         "zeichen_farbe": D_TINTE_2 if unter else D_SIEGELROT,
                                                         "farbe": D_TINTE_2 if unter else D_TINTE, "abstand": 6}))
            h = self._hoehe(b, w, s)
            self._textblock(folie, x, y, w, h, absaetze, gg, D_TINTE)
            return y + h
        if "sql" in b:
            if b.get("label"):
                farbe = {"falsch": D_SIEGELROT, "richtig": D_GRUEN}.get(b.get("art"), D_MESSING)
                self._textblock(folie, x, y, w, 0.3, [(b["label"], {"fett": True, "schrift": SCHRIFT_MONO})], 11.5, farbe)
                y += 0.34
            gs = sql_groesse(b["sql"], w, min(b.get("groesse", 20), 20 * s))
            sql = A.SQL(b["sql"], db=b.get("db"), pruefung=b.get("pruefung", "ja"))
            y = self._sql_box(folie, x, y, w, sql, groesse=gs)
            if b.get("art") in ("falsch", "richtig"):  # farbiger Randstreifen
                self._form(folie, x, y - sql_hoehe(b["sql"], gs), 0.07, sql_hoehe(b["sql"], gs),
                           D_SIEGELROT if b["art"] == "falsch" else D_GRUEN)
            erg = self._ergebnis_tabelle(b)
            if erg is not None:
                n = len(erg.zeilen) - (1 if erg.abgeschnitten else 0)
                zusatz = " (Ausschnitt)" if erg.abgeschnitten else ""
                self._textblock(folie, x, y + 0.08, w, 0.3, [(f"ERGEBNIS · {n} ZEILE{'N' if n != 1 else ''}{zusatz}" if n else
                                                              "ERGEBNIS · KEINE ZEILE", {"fett": True, "schrift": SCHRIFT_MONO})],
                                10.5, D_MESSING)
                y = self._tabelle(folie, x, y + 0.42, w, erg, groesse=b.get("tabellengroesse", 15) * s)
            return y + 0.12
        if "tabelle" in b:
            t = b["tabelle"]
            if t.get("titel"):
                self._textblock(folie, x, y, w, 0.3, [(t["titel"], {"fett": True, "schrift": SCHRIFT_MONO})], 11, D_MESSING)
                y += 0.32
            tab = A.Tabelle(t["kopf"], t["zeilen"], pk=tuple(t.get("pk", ())), fk=tuple(t.get("fk", ())))
            return self._tabelle(folie, x, y, w, tab, groesse=t.get("groesse", 15) * s) + 0.12
        if "schema" in b:
            gg = b.get("groesse", 19) * s
            h = self._hoehe(b, w, s) - 0.12
            self._form(folie, x, y, w, h, PAPIER, form=MSO_SHAPE.ROUNDED_RECTANGLE)
            yy = y + 0.12
            absaetze = []
            if b.get("titel"):
                absaetze.append((b["titel"], {"fett": True, "schrift": SCHRIFT_MONO, "groesse": 11, "farbe": TINTE_3}))
            absaetze += [(z, {"schrift": SCHRIFT_MONO, "abstand": 5}) for z in b["schema"]]
            self._textblock(folie, x + 0.2, yy, w - 0.4, h - 0.2, absaetze, gg * 0.92, TINTE)
            return y + h + 0.12
        if "bild" in b:
            from PIL import Image
            pfad = self.ordner / "bilder" / f"{b['bild']}.png"
            if not pfad.exists():
                raise Fehler(f"Bild {self.ordner.name}/bilder/{pfad.name} fehlt (python3 werkzeuge/video_folien.py --bilder)")
            with Image.open(pfad) as im:
                bw, bh = im.size
            h_max = h_bild if h_bild is not None else b.get("hoehe", 3.0)
            f = min(w / bw, h_max / bh)
            bb, hh = bw * f, bh * f
            xx = x + (w - bb) / 2 if b.get("mitte", True) else x
            bild = folie.shapes.add_picture(str(pfad), Inches(xx), Inches(y), Inches(bb), Inches(hh))
            bild.name = f"Bild {b['bild']}"
            bild.line.color.rgb = A._p_rgb(D_TINTE_2)
            bild.line.width = PPt(0.75)
            return y + hh + 0.1
        if "kasten" in b:
            k = b["kasten"]
            fuellung, etikett = KASTEN_FARBEN.get(k.get("farbe", "papier"), KASTEN_FARBEN["papier"])
            gg = k.get("groesse", 17) * s
            h = self._hoehe(b, w, s) - 0.12
            self._form(folie, x, y, w, h, fuellung, form=MSO_SHAPE.ROUNDED_RECTANGLE)
            self._form(folie, x, y, 0.08, h, etikett)
            yy = y + 0.1
            if k.get("titel"):
                self._textblock(folie, x + 0.3, yy, w - 0.45, 0.32, [(k["titel"].upper(), {"fett": True, "schrift": SCHRIFT_MONO})],
                                11, etikett)
                yy += 0.4
            absaetze = [(t[2:] if t.startswith("- ") else t, {"zeichen": "▸  " if t.startswith("- ") else None,
                                                               "zeichen_farbe": etikett, "abstand": 3}) for t in k.get("text", [])]
            self._textblock(folie, x + 0.3, yy, w - 0.45, h - (yy - y) - 0.05, absaetze, gg, TINTE)
            return y + h + 0.12
        if "falle" in b:
            gg = b.get("groesse", 16) * s
            h = self._hoehe(b, w, s) - 0.12
            self._form(folie, x, y, w, h, PAPIER, rand=D_SIEGELROT, form=MSO_SHAPE.ROUNDED_RECTANGLE, staerke=1.5)
            chip = self._form(folie, x + 0.15, y + 0.15, 0.75, 0.42, SIEGELROT, form=MSO_SHAPE.ROUNDED_RECTANGLE)
            chip.name = f"Falle {b['falle']}"
            self._textblock(folie, x + 0.15, y + 0.17, 0.75, 0.38, [(b["falle"], {"fett": True})], 15, "FFFDF8").text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            absaetze = []
            if b.get("titel"):
                absaetze.append((b["titel"], {"fett": True, "groesse": gg + 1, "abstand": 2}))
            if b.get("text"):
                absaetze.append((b["text"], {}))
            self._textblock(folie, x + 1.05, y + 0.1, w - 1.2, h - 0.15, absaetze, gg, TINTE)
            return y + h + 0.12
        if "schritte" in b:
            gg = b.get("groesse", 16) * s
            for i, t in enumerate(b["schritte"], 1):
                h = max(0.6, hoehe_text(t, w - 0.85, gg) + 0.16)
                self._form(folie, x, y, w, h, PAPIER, form=MSO_SHAPE.ROUNDED_RECTANGLE)
                self._form(folie, x, y, 0.07, h, D_SIEGELROT)
                self._textblock(folie, x + 0.15, y + 0.04, 0.5, h - 0.08, [(str(i), {"schrift": SCHRIFT_KOPF})], 22, SIEGELROT)
                tb = self._textblock(folie, x + 0.6, y + 0.06, w - 0.7, h - 0.1, [(t, {})], gg, TINTE)
                tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                y += h + 0.12
            return y
        if "badge" in b:
            self._form(folie, x, y, 3.4, 0.5, D_MESSING, form=MSO_SHAPE.ROUNDED_RECTANGLE, name="Pause")
            tb = folie.shapes.add_textbox(Inches(x), Inches(y), Inches(3.4), Inches(0.5))
            tf = tb.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            for t, schrift in (("⏸", SCHRIFT_SYMBOL), ("  VIDEO ANHALTEN", SCHRIFT_MONO)):
                r = p.add_run()
                r.text = t
                r.font.name, r.font.size, r.font.bold = schrift, PPt(15), True
                r.font.color.rgb = A._p_rgb(TINTE)
            return y + 0.62
        if "abstand" in b:
            return y + b["abstand"]
        raise Fehler(f"unbekannter Block: {b!r}")

    def _bereich(self, folie, bloecke, x, y0, w, y_max, wo):
        """Blöcke untereinander in [y0, y_max] setzen; Schriften schrumpfen, bis alles passt. Bilder füllen den Rest."""
        if not bloecke:
            return y0
        bilder = [b for b in bloecke if isinstance(b, dict) and "bild" in b and "hoehe" not in b]
        for schritt in range(0, 9):
            s = 1.0 - 0.05 * schritt
            fest = sum(self._hoehe(b, w, s) for b in bloecke if not (isinstance(b, dict) and b in bilder))
            rest = y_max - y0 - fest
            if rest >= (1.3 + 0.1) * len(bilder) - 1e-6:
                break
        else:
            raise Fehler(f"{wo}: Inhalt passt nicht (≈ {fest:.2f} in für {y_max - y0:.2f} in) – kürzen oder aufteilen")
        h_bild = (rest / len(bilder) - 0.1) if bilder else None
        y = y0
        for b in bloecke:
            y = self._setze(folie, b, x, y, w, s, h_bild=h_bild if (isinstance(b, dict) and b in bilder) else None)
        if y > y_max + 0.05:
            raise Fehler(f"{wo}: Inhalt ragt über den Bereich ({y:.2f} > {y_max:.2f} in)")
        return y

    # -------------------------------------------------------------- Folientypen
    def folie_titel(self, f, notizen):
        folie = self._folie(None, notizen=notizen, faden=False)
        akte = self.spec.get("akte") or 0
        a = AKTEN.get(akte, AKTEN[0])
        breite = LINKS_MAX - L
        self._textblock(folie, L, 1.05, breite, 0.4, [(f"ERKLÄRVIDEO {self.spec['nr']}", {"fett": True, "schrift": SCHRIFT_MONO})],
                        14, D_MESSING)
        g = 46 if len(self.spec["titel"]) <= 40 else 40
        h_titel = hoehe_text(self.spec["titel"], breite, g, 1.05) + 0.1
        tb = self._textblock(folie, L, 1.5, breite, h_titel, [(self.spec["titel"], {"schrift": SCHRIFT_KOPF})], g, D_TINTE)
        tb.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        y = 1.5 + h_titel + 0.15
        self._roter_faden(folie, y)
        if self.spec.get("untertitel"):
            h = hoehe_text(self.spec["untertitel"], breite, 21)
            self._textblock(folie, L, y + 0.2, breite, h, [(self.spec["untertitel"], {})], 21, D_TINTE_2)
            y += 0.2 + h
        info = " · ".join(x for x in (self.spec.get("einsatz_kurz"), self.spec.get("laenge")) if x)
        self._textblock(folie, L, UNTEN - 0.45, breite, 0.4, [(info, {"schrift": SCHRIFT_MONO})], 12.5, D_TINTE_2)
        if y > UNTEN - 0.5:
            raise Fehler("Titelfolie: Titel und Untertitel zu lang")
        png = A.stempel_png(a["stempel"], self._tmp / "stempel.png", farbe=D_SIEGELROT)
        bild = folie.shapes.add_picture(str(png), Inches(LINKS_MAX + 0.6), Inches(1.15), Inches(R - LINKS_MAX - 0.6))
        if (bild.top + bild.height) / 914400 > RECHTS_UNTEN:
            raise Fehler("Titelfolie: Stempel ragt in die Kamera-Zone")
        return folie

    @staticmethod
    def _auto_teilung(links):
        """Rechte Spalte so weit nach rechts, dass SQL links mindestens in 15 pt passt (höchstens bis LINKS_MAX)."""
        breite = 0
        for b in links:
            if isinstance(b, dict) and "sql" in b:
                laengste = max(len(z) for z in b["sql"].split("\n"))
                breite = max(breite, 0.5 + 0.6 * laengste * 15 / 72)
        return min(LINKS_MAX, max(TEILUNG, L + breite + 0.3))

    def folie_inhalt(self, f, notizen, nr, vorne=(), titel=None):
        folie = self._folie(titel or f.get("titel", ""), notizen=notizen)
        wo = f"Folie {nr}"
        links = list(vorne) + list(f.get("links", []))
        rechts = list(f.get("rechts", []))
        oben = list(f.get("oben", []))
        y_links = O
        if oben:
            if rechts:
                raise Fehler(f"{wo}: „oben“ und „rechts“ zugleich geht nicht")
            y_links = self._bereich(folie, oben, L, O, R - L, RECHTS_UNTEN, wo + " (oben)") + 0.05
        teilung = f.get("teilung") or self._auto_teilung(links)
        if rechts:
            self._bereich(folie, rechts, teilung, O, R - teilung, RECHTS_UNTEN, wo + " (rechts)")
            breite_links = f.get("breite_links", teilung - L - 0.3)
        else:
            breite_links = LINKS_MAX - L
        self._bereich(folie, links, L, y_links, breite_links, UNTEN, wo + " (links)")
        return folie

    def folie(self, f, nr):
        notizen = ((f"[Regie: {f['regie']}]\n\n" if f.get("regie") else "") + f.get("sprechertext", "").strip())
        typ = f.get("typ", "inhalt")
        if typ == "titel":
            return self.folie_titel(f, notizen)
        if typ == "pause":
            return self.folie_inhalt(f, notizen, nr, vorne=[{"badge": True}], titel=f.get("titel", "Kurz anhalten: Jetzt du!"))
        if typ == "merke":
            punkte = {"punkte": f["punkte"], "groesse": f.get("groesse", 20)}
            g = dict(f)
            g["links"] = [{"kasten": {"farbe": "blau", "titel": "Ermittler-Notiz", "text": ["- " + p for p in f["punkte"]],
                                      "groesse": f.get("groesse", 21)}}] + list(f.get("links", []))
            del punkte
            return self.folie_inhalt(g, notizen, nr, titel=f.get("titel", "Das nimmst du mit"))
        if typ == "ende":
            g = dict(f)
            texte = []
            for e in f.get("lab", []):
                if e.get("text"):
                    texte.append("- " + e["text"])
                    continue
                texte.append(f"- **Ermittlung {e['ermittlung']} „{e['titel']}“** · {e['aufgaben']}"
                             + (f" · Werkzeug: {e['werkzeug']}" if e.get("werkzeug") else ""))
            bloecke = [{"kasten": {"farbe": "blau", "titel": "Spur im Lab", "text": texte, "groesse": 18}}]
            if f.get("material"):
                bloecke.append({"kasten": {"farbe": "papier", "titel": "Arbeitsblatt und Training",
                                           "text": ["- " + m for m in f["material"]], "groesse": 17}})
            if f.get("naechstes"):
                bloecke.append({"kasten": {"farbe": "gold", "titel": "Nächstes Video", "text": [f["naechstes"]], "groesse": 17}})
            g["links"] = bloecke + list(f.get("links", []))
            return self.folie_inhalt(g, notizen, nr, titel=f.get("titel", "Weiter im Lab"))
        if typ == "inhalt":
            return self.folie_inhalt(f, notizen, nr)
        raise Fehler(f"Folie {nr}: unbekannter Typ „{typ}“")


# ================================================================== SQL-Ergebnisse (ein Node-Aufruf je Video)
def sql_bloecke(spec):
    """Alle SQL-Blöcke (mit Ort) einer video.json; vergibt _id."""
    aus, nr = [], 0
    for i, f in enumerate(spec["folien"], 1):
        for bereich in ("oben", "links", "rechts"):
            for b in f.get(bereich, []):
                if isinstance(b, dict) and "sql" in b:
                    nr += 1
                    b["_id"] = str(nr)
                    aus.append((i, b))
    return aus


def zelle(v):
    if v is None:
        return "NULL"
    if isinstance(v, float):
        if abs(v - round(v)) < 1e-9:
            return str(int(round(v)))
        return f"{v:.2f}".rstrip("0").replace(".", ",")
    return str(v)


def ergebnisse_berechnen(spec):
    standard = spec.get("datenbank")
    auftraege = [{"id": b["_id"], "sql": b["sql"], "datenbank": b.get("db") or standard, "maxZeilen": 200}
                 for _, b in sql_bloecke(spec) if b.get("ergebnis") is True]
    if not auftraege:
        return {}
    aus = {}
    for antwort in A.sql_cli(auftraege, modus="ausfuehren"):
        if antwort.get("fehler"):
            raise Fehler(f"SQL {antwort['id']} nicht ausführbar: {antwort['fehler']}")
        aus[antwort["id"]] = (antwort["spalten"], [[zelle(v) for v in z] for z in antwort["zeilen"]])
    return aus


# ================================================================== Drehbuch (Markdown)
def worte(text):
    return len(re.findall(r"\S+", text or ""))


def sekunden(f):
    return worte(f.get("sprechertext", "")) / WORTE_PRO_MINUTE * 60 + f.get("pause_s", 0)


def mmss(s):
    s = int(round(s))
    return f"{s // 60}:{s % 60:02d}"


def _inhalt_md(f, ergebnisse):
    """Was auf der Folie steht (kurz), für das Drehbuch."""
    z = []
    for bereich in ("oben", "links", "rechts"):
        for b in f.get(bereich, []):
            if isinstance(b, str):
                z.append(f"> {b}")
            elif "sql" in b:
                kopf = b.get("label") or ("SQL (absichtlich fehlerhaft)" if b.get("pruefung") == "fehlerhaft" else "SQL")
                z += [f"*{kopf}:*", "", "```sql", b["sql"].strip(), "```"]
                if b.get("ergebnis") is True and b.get("_id") in ergebnisse:
                    k, zeilen = ergebnisse[b["_id"]]
                    z.append(f"*Ergebnis: {len(zeilen)} Zeile(n) – {', '.join(k)}*")
                z.append("")
            elif "tabelle" in b:
                z.append(f"*Tabelle{(' ' + b['tabelle']['titel']) if b['tabelle'].get('titel') else ''}: "
                         f"{', '.join(b['tabelle']['kopf'])} ({len(b['tabelle']['zeilen'])} Zeilen)*")
            elif "schema" in b:
                z += [f"- `{nur_text(x)}`" for x in b["schema"]]
            elif "bild" in b:
                z.append(f"*Bild: `bilder/{b['bild']}.png`*")
            elif "falle" in b:
                z.append(f"- **{b['falle']}** {b.get('titel', '')}: {b.get('text', '')}".rstrip(": "))
            elif "kasten" in b:
                k = b["kasten"]
                z.append(f"- *{k.get('titel', 'Kasten')}:* " + " · ".join(t.lstrip("- ") for t in k.get("text", [])))
            elif "punkte" in b:
                z += [("  " if p.startswith("  - ") else "") + "- " + p.replace("  - ", "", 1) for p in b["punkte"]]
            elif "schritte" in b:
                z += [f"{i}. {t}" for i, t in enumerate(b["schritte"], 1)]
            elif "text" in b:
                z.append(f"> {b['text']}")
    if f.get("typ") == "pause" and f.get("aufgabe"):
        z.insert(0, f"> ⏸ {f['aufgabe']}")
    if f.get("typ") == "merke":
        z += [f"- {p}" for p in f.get("punkte", [])]
    if f.get("typ") == "ende":
        for e in f.get("lab", []):
            z.append(f"- {e['text']}" if e.get("text") else f"- Ermittlung {e['ermittlung']} „{e['titel']}“ · {e['aufgaben']}")
        z += [f"- {m}" for m in f.get("material", [])]
        if f.get("naechstes"):
            z.append(f"- Nächstes Video: {f['naechstes']}")
    return z


def skript_md(spec, pptx_name, ergebnisse):
    folien = spec["folien"]
    dauer = sum(sekunden(f) for f in folien)
    yt = spec.get("youtube", {})
    z = [f"# Erklärvideo {spec['nr']}: {spec['titel']}", "",
         "> ERZEUGT aus `video.json` mit `python3 werkzeuge/video_folien.py` – Änderungen in der video.json vornehmen.", "",
         f"*{spec.get('untertitel', '')}*", "",
         "| | |", "|---|---|",
         f"| **Akte** | {spec.get('aktenname', '')} |",
         f"| **Einsatz** | {spec.get('einsatz', '')} |",
         f"| **Zielgruppe** | {spec.get('zielgruppe', '')} |",
         f"| **Länge** | Ziel {spec.get('laenge', '')} · Sprechtext ca. {mmss(dauer)} min (bei {WORTE_PRO_MINUTE} Wörtern/min, inkl. Pausen) |",
         f"| **Beispieldaten** | {spec.get('beispieldaten', '')} |",
         f"| **Folien** | `{pptx_name}` ({len(folien)} Folien, Kamera-Zone unten rechts frei, Sprechertext in den Notizen) |", ""]
    z += ["## Lernziele", ""] + [f"- {x}" for x in spec.get("lernziele", [])] + [""]
    if spec.get("hinweise"):
        z += ["## Hinweise für die Aufnahme", ""] + [f"- {x}" for x in spec["hinweise"]] + [""]
    z += ["## YouTube", "", f"**Titel:** {yt.get('titel', spec['titel'])}", "",
          "**Sichtbarkeit:** zuerst „nicht gelistet“, nach Freigabe öffentlich (Playlist „Datendetektei“).", "",
          "**Beschreibung:**", "", "```", yt.get("beschreibung", "").rstrip(), "", "Kapitel:"]
    t = 0.0
    for f in folien:
        if f.get("kapitel"):
            z.append(f"{mmss(t)} {f['kapitel']}")
        t += sekunden(f)
    z += ["```", "", "_Zeitmarken aus der Textlänge geschätzt – nach dem Schnitt anpassen._", ""]
    if yt.get("tags"):
        z += [f"**Tags:** {', '.join(yt['tags'])}", ""]
    z += ["## Drehbuch", ""]
    t = 0.0
    for i, f in enumerate(folien, 1):
        titel = spec["titel"] if f.get("typ") == "titel" else f.get("titel") or {
            "pause": "Kurz anhalten: Jetzt du!", "merke": "Das nimmst du mit", "ende": "Weiter im Lab"}.get(f.get("typ"), "")
        kapitel = f" · Kapitel „{f['kapitel']}“" if f.get("kapitel") else ""
        z += [f"### Folie {i} · {mmss(t)} – {titel}{kapitel}", ""]
        if f.get("regie"):
            z += [f"*Regie: {f['regie']}*", ""]
        inhalt = _inhalt_md(f, ergebnisse)
        if inhalt:
            z += ["**Auf der Folie:**", ""] + inhalt + [""]
        if f.get("sprechertext"):
            z += ["**Sprechertext:**", "", f["sprechertext"].strip(), ""]
        if f.get("pause_s"):
            z += [f"*(ca. {f['pause_s']} s Pause zum Mitdenken – im Schnitt den Hinweis „Video anhalten“ einblenden)*", ""]
        t += sekunden(f)
    return "\n".join(z).rstrip() + "\n", dauer


# ================================================================== Aufbau prüfen (Spec §13)
def pruefe_spec(spec):
    """Inhaltliche Vollständigkeit → Liste von Meldungen (leer = in Ordnung)."""
    m = []
    for feld in ("nr", "kurz", "titel", "untertitel", "laenge", "einsatz", "zielgruppe", "lernziele", "youtube", "folien",
                 "beispieldaten"):
        if not spec.get(feld) and spec.get(feld) != 0:
            m.append(f"Feld „{feld}“ fehlt")
    yt = spec.get("youtube", {})
    for feld in ("titel", "beschreibung", "tags"):
        if not yt.get(feld):
            m.append(f"youtube.{feld} fehlt")
    folien = spec.get("folien", [])
    if not FOLIEN_MIN <= len(folien) <= FOLIEN_MAX:
        m.append(f"{len(folien)} Folien (erwartet {FOLIEN_MIN}–{FOLIEN_MAX})")
    if not folien:
        return m
    typen = [f.get("typ", "inhalt") for f in folien]
    if typen[0] != "titel":
        m.append("erste Folie ist keine Titelfolie")
    if typen[-1] != "ende":
        m.append("letzte Folie ist nicht „Weiter im Lab“ (typ ende)")
    if "pause" not in typen:
        m.append("keine Pause-Aufgabe (typ pause)")
    if "merke" not in typen:
        m.append("keine Merksätze (typ merke)")
    elif "pause" in typen and typen.index("pause") > len(typen) - 1 - typen[::-1].index("merke"):
        m.append("Pause-Aufgabe nach den Merksätzen")
    fallen = json.dumps([f for f in folien if f.get("typ") not in ("titel", "ende")], ensure_ascii=False)
    if not P_NUMMER.search(fallen):
        m.append("keine Falle mit P-Nummer (P1–P10)")
    if worte(folien[0].get("sprechertext")) > EINSTIEG_WORTE:
        m.append(f"Einstieg zu lang: {worte(folien[0].get('sprechertext'))} Wörter (≤ {EINSTIEG_WORTE} ≈ 30 s)")
    if not folien[0].get("kapitel"):
        m.append("erste Folie ohne Kapitel (YouTube-Kapitel beginnen bei 0:00)")
    if sum(1 for f in folien if f.get("kapitel")) < 3:
        m.append("weniger als drei Kapitel")
    for i, f in enumerate(folien, 1):
        if not f.get("sprechertext", "").strip():
            m.append(f"Folie {i}: Sprechertext fehlt")
    return m


# ================================================================== Ablauf
def namen(ordner, spec):
    basis = f"DD_Video{spec['nr']:02d}_{spec['kurz']}"
    return ordner / f"{basis}_Folien.pptx", ordner / f"{basis}_Skript.md"


def quellen_hash(ordner, spec):
    """SHA-256 über video.json und alle verwendeten Bilder (Aktualität der Folien)."""
    h = hashlib.sha256((ordner / "video.json").read_bytes())
    for name in sorted(verwendete_bilder(spec)):
        p = ordner / "bilder" / f"{name}.png"
        h.update(name.encode("utf-8"))
        h.update(p.read_bytes() if p.exists() else b"FEHLT")
    return h.hexdigest()


def verwendete_bilder(spec):
    aus = set()
    for f in spec.get("folien", []):
        for bereich in ("oben", "links", "rechts"):
            for b in f.get(bereich, []):
                if isinstance(b, dict) and "bild" in b:
                    aus.add(b["bild"])
    return aus


def lade(ordner):
    return json.loads((ordner / "video.json").read_text(encoding="utf-8"))


def bilder_erzeugen(ordner):
    """Plattform-Bilder neu rendern (scripts/video_bilder.js); Bauplan-Diagramme auf den Inhalt zuschneiden."""
    subprocess.run(["node", str(BILDER_JS), str(ordner)], check=True, cwd=REPO)
    spec = lade(ordner)
    for name, b in spec.get("bilder", {}).items():
        if b.get("zuschneiden", b.get("werkzeug") == "bauplan"):
            zuschneiden(ordner / "bilder" / f"{name}.png")


def zuschneiden(pfad, rand=24):
    """Einfarbigen Rand (Farbe der linken oberen Ecke) bis auf `rand` Pixel abschneiden."""
    from PIL import Image, ImageChops
    with Image.open(pfad) as im:
        im = im.convert("RGB")
        im = im.crop((6, 6, im.width - 6, im.height - 6))  # Rahmen des Elements
        grund = Image.new("RGB", im.size, im.getpixel((2, 2)))
        diff = ImageChops.difference(im, grund).convert("L").point(lambda v: 255 if v > 18 else 0)
        box = diff.getbbox()
        if box:
            box = (max(0, box[0] - rand), max(0, box[1] - rand), min(im.width, box[2] + rand), min(im.height, box[3] + rand))
            im.crop(box).save(pfad, optimize=True)


def bauen(ordner, mit_bildern=False):
    spec = lade(ordner)
    m = pruefe_spec(spec)
    if m:
        raise Fehler(f"{ordner.name}: " + "; ".join(m))
    fehlend = [b for b in verwendete_bilder(spec) if b not in spec.get("bilder", {})]
    if fehlend:
        raise Fehler(f"{ordner.name}: Bilder ohne Beschreibung in „bilder“: {', '.join(sorted(fehlend))}")
    if mit_bildern:
        bilder_erzeugen(ordner)
    sql_bloecke(spec)
    ergebnisse = ergebnisse_berechnen(spec)
    p = VideoFolien(spec, ordner, ergebnisse)
    for i, f in enumerate(spec["folien"], 1):
        p.folie(f, i)
    pptx, md = namen(ordner, spec)
    p.speichern(pptx)
    A.setze_eigenschaften(pptx, {"DD-Art": "video", "DD-Akte": str(spec.get("akte") or 0), "DD-Kuer": "nein",
                                 "DD-Kamera-Zone": "ja", "DD-Video": str(spec["nr"]),
                                 "DD-Video-Quelle": quellen_hash(ordner, spec),
                                 **({"DD-Datenbank": spec["datenbank"]} if spec.get("datenbank") else {})})
    zone = A.pruefe_kamera_zone(pptx)
    if zone:
        raise Fehler(f"{ordner.name}: Kamera-Zone belegt: " + ", ".join(f"Folie {i} „{n}“" for i, n in zone))
    text, dauer = skript_md(spec, pptx.name, ergebnisse)
    md.write_text(text, encoding="utf-8")
    print(f"✔ {ordner.name}: {len(spec['folien'])} Folien, Sprechtext ca. {mmss(dauer)} min → {pptx.name}, {md.name}")
    return pptx, md


def alle_ordner(wurzel=VIDEOS):
    return sorted(p for p in Path(wurzel).iterdir() if (p / "video.json").exists()) if Path(wurzel).exists() else []


def pruefen(wurzel=VIDEOS, ausgabe=print):
    """Vollständigkeit und Aktualität ohne Erzeugen → Liste der Probleme."""
    from pptx import Presentation
    probleme = []
    for ordner in alle_ordner(wurzel):
        try:
            spec = lade(ordner)
        except (OSError, json.JSONDecodeError) as e:
            probleme.append(f"{ordner.name}: video.json unlesbar ({e})")
            continue
        probleme += [f"{ordner.name}: {x}" for x in pruefe_spec(spec)]
        for b in sorted(verwendete_bilder(spec)):
            if not (ordner / "bilder" / f"{b}.png").exists():
                probleme.append(f"{ordner.name}: Bild bilder/{b}.png fehlt (python3 werkzeuge/video_folien.py {ordner.name} --bilder)")
        pptx, md = namen(ordner, spec)
        if not pptx.exists() or not md.exists():
            probleme.append(f"{ordner.name}: Folien/Skript fehlen (python3 werkzeuge/video_folien.py {ordner.name})")
            continue
        eig = A.lese_eigenschaften(pptx)
        if eig.get("DD-Video-Quelle") != quellen_hash(ordner, spec):
            probleme.append(f"{ordner.name}: Folien veraltet – video.json oder Bilder geändert (python3 werkzeuge/video_folien.py {ordner.name})")
        if eig.get("DD-Kamera-Zone") != "ja":
            probleme.append(f"{ordner.name}: Eigenschaft DD-Kamera-Zone fehlt")
        sql_bloecke(spec)
        try:
            ergebnisse = ergebnisse_berechnen(spec)
        except Fehler as e:
            probleme.append(f"{ordner.name}: {e}")
            ergebnisse = {}
        erwartet, _ = skript_md(spec, pptx.name, ergebnisse)
        if md.read_text(encoding="utf-8") != erwartet:
            probleme.append(f"{ordner.name}: Skript veraltet (python3 werkzeuge/video_folien.py {ordner.name})")
        prs = Presentation(str(pptx))
        if len(prs.slides) != len(spec["folien"]):
            probleme.append(f"{ordner.name}: Folienzahl ({len(prs.slides)}) passt nicht zur video.json ({len(spec['folien'])})")
        for i, folie in enumerate(prs.slides, 1):
            notiz = folie.notes_slide.notes_text_frame.text if folie.has_notes_slide else ""
            if not notiz.strip():
                probleme.append(f"{ordner.name}: Folie {i} ohne Sprechertext in den Notizen")
        probleme += [f"{ordner.name}: Kamera-Zone belegt auf Folie {i} („{n}“)" for i, n in A.pruefe_kamera_zone(pptx)]
    for p in probleme:
        ausgabe("✘ " + p)
    ausgabe(f"Videos geprüft: {len(alle_ordner(wurzel))} Video(s), {len(probleme)} Problem(e).")
    return probleme


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if "--pruefen" in argv:
        return 1 if pruefen() else 0
    args = [a for a in argv if not a.startswith("--")]
    ordner = [(VIDEOS / a if not Path(a).is_absolute() else Path(a)).resolve() for a in args] or alle_ordner()
    code = 0
    for o in ordner:
        try:
            bauen(o, mit_bildern="--bilder" in argv)
        except Fehler as e:
            print(f"✘ {e}")
            code = 1
    return code


if __name__ == "__main__":
    sys.exit(main())
