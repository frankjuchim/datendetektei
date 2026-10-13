#!/usr/bin/env python3
"""Prüft die Unterrichtsmaterialien auf Konventionsverstöße (Spec §6, §12.3, §12.4; docs/MATERIALIEN.md).

Fehler (Exit-Code 1):
  - Altbegriffe „Entitätsmenge“, „Array“, „XP“ (außer als Erklärung: „statt/nie/kein … Entitätsmenge“)
  - „<>“ (Ungleich nach EH §8: „!=“) außer als Erklärung („!= statt <>“)
  - „JOIN“ außerhalb der Kür (Ordner Kuer/ bzw. Dokumenteigenschaft DD-Kuer = ja), außer verneint („ohne JOIN“)
  - SQL außerhalb des Absatzformats „SQL-Code“ (Absatz beginnt mit SELECT … FROM) bzw. außerhalb eines Textfelds „SQL“
  - Klarnamen aus werkzeuge/klarnamen.txt (auch in den Dokumenteigenschaften Autor/zuletzt geändert)
  - PDF-Dateinamen mit Umlauten/Sonderzeichen (materialien/**/*.pdf)
  - Schülerfassungen mit Lösungsmarker: Quellen der Schülerfassungen aus seed/arbeitsblaetter.json (docx-Text,
    Dokumenteigenschaft DD-Fassung) und alle PDFs in materialien/Arbeitsblaetter_PDF (PDF-Text)
Lösungsmarker: „NUR LEHRKRAFT“ (Fußzeile jeder Lösungs-/Lehrkraftfassung, Lösungsfolien), „[NUR LEHRKRAFT]“,
  „Lösungen der Warm-up-Karten“, „LÖSUNG ·“, „ERWARTUNGSHORIZONT ·“.

Aufruf: python3 werkzeuge/materialien_pruefen.py [--hinweise]
Benötigt python-docx und python-pptx; für den PDF-Text pypdf oder pdftotext (sonst nur Hinweis).
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

WERKZEUGE = Path(__file__).resolve().parent
REPO = WERKZEUGE.parent
sys.path.insert(0, str(WERKZEUGE))

from akten_design import lese_eigenschaften, LOESUNGSTEIL_WARMUP, MARKER_LEHRKRAFT, NOTIZ_NUR_LEHRKRAFT  # noqa: E402
import material_gemeinsam as M  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

ALTBEGRIFFE = re.compile(r"\b(Entitätsmengen?|Arrays?|XP)\b")
UNGLEICH = re.compile(r"<>")
JOIN = re.compile(r"(?<!Self-)\bJOIN\b", re.I)  # „Self-Join“ ist ein Fachbegriff (Verbund einer Tabelle mit sich selbst)
VERNEINUNG = re.compile(r"\b(statt|nicht|nie|kein|keine|keinen|ohne|verboten|anstelle)\b", re.I)
SQL_ROH = re.compile(r"^\s*SELECT\b.+\bFROM\b", re.S | re.I)
LOESUNGSMARKER = (MARKER_LEHRKRAFT, NOTIZ_NUR_LEHRKRAFT, LOESUNGSTEIL_WARMUP, "LÖSUNG ·", "ERWARTUNGSHORIZONT ·")
SCHUELER_ARTEN = {"ab", "kuer", "unplugged", "handbuch"}


def lade_klarnamen(datei=WERKZEUGE / "klarnamen.txt"):
    if not Path(datei).exists():
        return []
    return [z.strip() for z in Path(datei).read_text(encoding="utf-8").splitlines() if z.strip() and not z.startswith("#")]


def _verneint(text, m, fenster=30):
    return bool(VERNEINUNG.search(text[max(0, m.start() - fenster):m.end() + fenster]))


def _absatz(p):
    sichtbar = []
    for r in p.iter(W + "r"):
        rpr = r.find(W + "rPr")
        if rpr is not None and rpr.find(W + "vanish") is not None:
            continue
        for k in r:
            if k.tag == W + "t":
                sichtbar.append(k.text or "")
            elif k.tag in (W + "br", W + "cr"):
                sichtbar.append("\n")
    return "".join(sichtbar)


def texte_docx(pfad):
    """[(ort, text, ist_sql)] aller Absätze (Text, Kopf-/Fußzeilen); ist_sql = Absatzformat „SQL-Code“."""
    from docx import Document
    d = Document(str(pfad))
    namen = {s.style_id: s.name for s in d.styles}
    teile = [("", d.element.body)] + [("Kopfzeile", s.header._element) for s in d.sections] + \
            [("Fußzeile", s.footer._element) for s in d.sections]
    aus = []
    for bereich, teil in teile:
        for i, p in enumerate(teil.iter(W + "p"), 1):
            t = _absatz(p)
            if not t.strip():
                continue
            ppr = p.find(W + "pPr")
            stil = ppr.find(W + "pStyle") if ppr is not None else None
            ist_sql = stil is not None and namen.get(stil.get(W + "val")) == "SQL-Code"
            aus.append((f"{bereich} Absatz {i}".strip(), t, ist_sql))
    cp = d.core_properties
    aus.append(("Dokumenteigenschaften", f"{cp.author or ''} · {cp.last_modified_by or ''}", False))
    return aus


def texte_pptx(pfad):
    from pptx import Presentation
    prs = Presentation(str(pfad))
    aus = []
    for i, folie in enumerate(prs.slides, 1):
        for s in folie.shapes:
            ist_sql = (s.name or "").startswith("SQL") and s.name != "SQL-Hintergrund"
            if s.has_text_frame:
                t = "\n".join(p.text for p in s.text_frame.paragraphs)
                if t.strip():
                    aus.append((f"Folie {i}", t, ist_sql))
            if getattr(s, "has_table", False) and s.has_table:
                for r in s.table.rows:
                    t = " | ".join(c.text for c in r.cells)
                    if t.strip():
                        aus.append((f"Folie {i} (Tabelle)", t, False))
        if folie.has_notes_slide:
            t = folie.notes_slide.notes_text_frame.text
            if t.strip():
                aus.append((f"Folie {i} (Notiz)", t, False))
    cp = prs.core_properties
    aus.append(("Dokumenteigenschaften", f"{cp.author or ''} · {cp.last_modified_by or ''}", False))
    return aus


def ist_kuer(pfad, eig=None):
    eig = eig if eig is not None else lese_eigenschaften(pfad)
    return eig.get("DD-Kuer") == "ja" or "Kuer" in Path(pfad).parts


def pruefe_datei(pfad, wurzel=REPO, klarnamen=None):
    """→ (fehler, hinweise) einer docx/pptx."""
    pfad = Path(pfad)
    klarnamen = lade_klarnamen() if klarnamen is None else klarnamen
    kuer = ist_kuer(pfad)
    texte = texte_docx(pfad) if pfad.suffix == ".docx" else texte_pptx(pfad)
    fehler, hinweise = [], []
    for ort, t, ist_sql in texte:
        kurz = " ".join(t.split())[:110]
        for m in ALTBEGRIFFE.finditer(t):
            if not _verneint(t, m):
                fehler.append(f"{ort}: Altbegriff „{m.group(0)}“: {kurz}")
        for m in UNGLEICH.finditer(t):
            if not _verneint(t, m):
                fehler.append(f"{ort}: „<>“ – nach EH §8 „!=“ verwenden: {kurz}")
        if not kuer:
            for m in JOIN.finditer(t):
                if ist_sql or not VERNEINUNG.search(t):
                    fehler.append(f"{ort}: „JOIN“ außerhalb der Kür (Verbund nur implizit): {kurz}")
                    break
        if not ist_sql and SQL_ROH.match(t) and ort != "Dokumenteigenschaften":
            fehler.append(f"{ort}: SQL außerhalb des Absatzformats „SQL-Code“ (pptx: Textfeld „SQL“) – wird nicht geprüft: {kurz}")
        for name in klarnamen:
            if re.search(r"\b" + re.escape(name) + r"\b", t, re.I):
                fehler.append(f"{ort}: Klarname „{name}“ (werkzeuge/klarnamen.txt)")
    return fehler, hinweise


def hat_loesungsmarker(text):
    """Lösungsmarker im (PDF-)Text? Leerzeichen werden ignoriert (gesperrte Etiketten: „N U R  L E H R K R A F T“)."""
    kompakt = re.sub(r"\s+", "", text or "")
    return any(re.sub(r"\s+", "", m) in kompakt for m in LOESUNGSMARKER)


def pruefe_schuelerfassungen(wurzel=REPO):
    """Quellen der Schülerfassungen (laut seed/arbeitsblaetter.json) ohne Lösungsmarker und nicht als Lösungsfassung."""
    wurzel = Path(wurzel)
    try:
        z = M.lade_zuordnung(wurzel)
    except FileNotFoundError:
        return []
    fehler = []
    for ziel in M.plan(z):
        if ziel.ordner != "Arbeitsblaetter_PDF" or ziel.verfahren != "pdf":
            continue
        quelle = wurzel / "materialien" / ziel.quelle
        if not quelle.exists() or quelle.suffix != ".docx":
            continue
        eig = lese_eigenschaften(quelle)
        if eig.get("DD-Fassung") in ("loesung", "lehrkraft"):
            fehler.append(f"materialien/{ziel.quelle}: Schülerfassung ist als „{eig['DD-Fassung']}“ erzeugt (Lösungsmarker)")
            continue
        for ort, t, _ in texte_docx(quelle):
            if hat_loesungsmarker(t):
                fehler.append(f"materialien/{ziel.quelle}: {ort}: Lösungsmarker in einer Schülerfassung")
                break
    return fehler


def pdf_text(pdf):
    try:
        from pypdf import PdfReader
        return "".join(s.extract_text() or "" for s in PdfReader(str(pdf)).pages)
    except ImportError:
        pass
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True)
        return r.stdout
    return None


def pruefe_schueler_pdfs(wurzel=REPO):
    ordner = Path(wurzel) / "materialien" / "Arbeitsblaetter_PDF"
    fehler = []
    for pdf in sorted(ordner.glob("*.pdf")) if ordner.exists() else []:
        text = pdf_text(pdf)
        if text is None:
            return [f"Hinweis: weder pypdf noch pdftotext vorhanden – PDF-Texte nicht geprüft"]
        if hat_loesungsmarker(text):
            fehler.append(f"materialien/Arbeitsblaetter_PDF/{pdf.name}: Lösungsmarker in einer Schüler-PDF")
    return fehler


def pruefe_pdf_namen(wurzel=REPO):
    mat = Path(wurzel) / "materialien"
    return [f"materialien/{p.relative_to(mat)}: PDF-Name mit Umlaut/Sonderzeichen (nur A–Z, a–z, 0–9, . _ -)"
            for p in sorted(mat.rglob("*.pdf")) if not M.PDF_NAME.match(p.name)]


def main():
    mit_hinweisen = "--hinweise" in sys.argv
    mat = REPO / "materialien"
    dateien = sorted(p for p in mat.rglob("*") if p.suffix in (".docx", ".pptx") and not p.name.startswith("~$"))
    anzahl_fehler = anzahl_hinweise = 0
    klarnamen = lade_klarnamen()
    for p in dateien:
        fehler, hinweise = pruefe_datei(p, klarnamen=klarnamen)
        anzahl_fehler += len(fehler)
        anzahl_hinweise += len(hinweise)
        if fehler or (mit_hinweisen and hinweise):
            print(f"{'✘' if fehler else '·'} {p.relative_to(REPO)}")
            for f in fehler:
                print("    ! " + f)
            if mit_hinweisen:
                for h in hinweise:
                    print("    · " + h)
    allgemein = pruefe_pdf_namen() + pruefe_schuelerfassungen() + pruefe_schueler_pdfs()
    for f in allgemein:
        print(("· " if f.startswith("Hinweis") else "✘ ") + f)
        if f.startswith("Hinweis"):
            anzahl_hinweise += 1
        else:
            anzahl_fehler += 1
    print(f"{len(dateien)} Office-Dateien und {len(list(mat.rglob('*.pdf')))} PDFs geprüft: {anzahl_fehler} Fehler, "
          f"{anzahl_hinweise} Hinweis(e).")
    sys.exit(1 if anzahl_fehler else 0)


if __name__ == "__main__":
    main()
