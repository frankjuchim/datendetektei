#!/usr/bin/env python3
"""Design-Modul „Akten-Noir light“ für die Unterrichtsmaterialien der Datendetektei (Spec §11, §12; DESIGN.md §2, §4).

Bausteine für Word (python-docx) und PowerPoint (python-pptx) – die Erstellungsskripte unter quellen/material/
bauen damit Arbeitsblätter, Lösungen, Lehrerhinweise, Warm-up-Karten, Klausuren und Präsentationen. Die erzeugten
docx/pptx sind danach maßgeblich (Spec §12) und werden direkt gepflegt.

Kurzüberblick (Details in den Docstrings, Beispiele in quellen/material/muster.py):

  Inhalte (für docx und pptx gleich):
    "Text mit **fett**, *kursiv*, `Code` und <u>unterstrichen</u>"   – Absatz mit Mini-Auszeichnung
    SQL(text, db=None, pruefung="ja"|"fehlerhaft"|"aus", kuer=False)   – Absatzformat „SQL-Code“ / Textfeld „SQL“
    Tabelle(kopf, zeilen, ergebnis=False, pk=(), fk=(), breiten=None)   – Daten- bzw. „Ergebnistabelle“
    Ergebnis(kopf, zeilen, abgeschnitten=False)                          – Kurzform für Tabelle(..., ergebnis=True)
    Schema("Fundsache(<u>FundNr</u>, ↑StationID, …)")                    – Kurzschreibweise (PK unterstrichen, ↑FK)
    Liste(punkte, nummeriert=False) · Bild(pfad, breite_cm) · ER(graph, layout, breite_cm)
    Linien(n)                         – Schreiblinien (nur Schülerfassung)
    Loesung(*inhalt, fehler=(), varianten=())   – nur in der Lösungsfassung (grüner Kasten „LÖSUNG“)
    Notiz(titel, *inhalt)             – Merkkasten „Ermittler-Notiz“
    Auftrag(nr, stufe, ort, text, kompetenzen=(), inhalt=(), linien=0, loesung=None)

  Word: AktenDokument(art, akte, titel, ...)  .kopfband() .fall_einstieg() .beweisstueck() .ermittlungsauftraege()
        .abi_ecke() .spur_im_lab() .cold_case() .notiz() .ueberschrift() .absatz() .inhalt() .tabelle()
        .klausurkopf() .materialblock() .teilaufgabe() .bewertungsbogen() .seitenumbruch() .speichern(pfad)
        warmup_karten_dokument(karten, akte, titel, ...) – 2 × A5 je Seite, Lösungsteil ab eigener Seite
  PowerPoint: AktenPraesentation(akte, titel, kamera_zone=False)  .titelfolie() .briefing() .begriffe()
        .beispiel_schritte() .unplugged() .lab_auftrag() .merkfolie() .abi_format() .inhaltsfolie()
        .loesungsfolie() .speichern(pfad);  pruefe_kamera_zone(pfad)
  Grafik: stempel_png(text, pfad, farbe), er_png(graph, layout, pfad)   (Pillow, Schriften aus werkzeuge/fonts)
  Eigenschaften: setze_eigenschaften(pfad, {...}), lese_eigenschaften(pfad)  (docProps/custom.xml)
  SQL-Hilfe: abfrage_ergebnis(sql, db) → (kopf, zeilen)  (über scripts/sql_pruefen_cli.js)

Schriften wie die Plattform (DESIGN §2.7, SIL OFL 1.1, TTF in werkzeuge/fonts/): Barlow (Fließtext),
Barlow Condensed (Überschriften, Stempel), IBM Plex Mono (SQL, Etiketten). ★ kommt aus einer Symbolschrift
(„Segoe UI Symbol“, unter Linux ersetzt LibreOffice sie automatisch).
"""
from __future__ import annotations

import json
import math
import random
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

WERKZEUGE = Path(__file__).resolve().parent
REPO = WERKZEUGE.parent
FONTS = WERKZEUGE / "fonts"

# ------------------------------------------------------------------ Tokens (DESIGN.md §2.1/2.2, hell bzw. dunkel)
TINTE, TINTE_2, TINTE_3 = "1C2B3A", "46525E", "58626D"
PAPIER, KARTE, FLAECHE, FLAECHE_2 = "F4EDE1", "FFFDF8", "EBE2D2", "E3D8C4"
RAND, RAND_STARK = "D8CCB6", "8F8270"
SIEGELROT, MESSING, MESSING_TEXT, MESSING_GRAFIK = "B3261E", "B08D3C", "7A5E1F", "94742B"
GRUEN, BLAU = "2F6B4F", "2C5A85"
ROT_FLAECHE, GRUEN_FLAECHE, BLAU_FLAECHE, MESSING_FLAECHE = "F8E3DF", "E0EEE5", "E2EAF3", "F3E8CD"
SQL_GRAU = "ECEAE5"  # Absatzformat „SQL-Code“: neutral grau hinterlegt (druckt hell)
# dunkel (Fallbesprechung)
D_TINTE, D_PAPIER, D_KARTE, D_SIEGELROT, D_MESSING = "E8E1D3", "141B22", "1C252E", "E0574F", "D4B062"
D_TINTE_2, D_TINTE_3, D_BLAU, D_GRUEN = "BDB6A9", "9C978E", "7FA9D6", "5FAF86"
KONSOLE, KONSOLE_TEXT = "11161C", "D8E3D0"
SQL_FARBEN = {"schluessel": "F2A65A", "text": "9BD3A6", "zahl": "8CC4FF", "kommentar": "7F8F99", "aggregat": "E6C77A"}

SCHRIFT = "Barlow"
SCHRIFT_KOPF = "Barlow Condensed SemiBold"   # Überschriften (Condensed 600)
SCHRIFT_KOPF_FETT = "Barlow Condensed"       # mit fett = Condensed 700 (Stempel)
SCHRIFT_MONO = "IBM Plex Mono"
SCHRIFT_SYMBOL = "Segoe UI Symbol"           # ★ (Barlow hat keinen Stern)

# Akten (Spec §3.3, §12.1; Farben aus seed/akten; Stempel aus DS_PLANUNG „Kopfband“)
AKTEN = {
    1: {"titel": "Das Fundbüro", "ordner": "Akte1_Fundbuero", "stempel": "FUNDBÜRO", "farbe": "8C2F39", "db": "fundbuero"},
    2: {"titel": "Spur der Räder", "ordner": "Akte2_Raeder", "stempel": "STADTRAD", "farbe": "2C5A85", "db": "stadtrad"},
    3: {"titel": "Die Stadtfest-Kasse", "ordner": "Akte3_Stadtfest", "stempel": "STADTFEST", "farbe": "9C5B12", "db": "stadtfest"},
    4: {"titel": "Chaos im Vereinsarchiv", "ordner": "Akte4_Vereinsarchiv", "stempel": "VEREINSHEIM", "farbe": "5E4A8C", "db": "tus"},
    5: {"titel": "Der Bauplan", "ordner": "Akte5_Bauplan", "stempel": "DETEKTEI", "farbe": "17585C", "db": "detektei"},
    6: {"titel": "Datenspuren", "ordner": "Akte6_Datenspuren", "stempel": "RATHAUS", "farbe": SIEGELROT, "db": "datenspuren"},
    7: {"titel": "Die Ermittlerprüfung", "ordner": "Akte7_Generalprobe", "stempel": "GENERALPROBE", "farbe": SIEGELROT, "db": None},
    0: {"titel": "Gesamtmaterial", "ordner": "00_Gesamt", "stempel": "DETEKTEI", "farbe": SIEGELROT, "db": None},
    "K": {"titel": "Kür", "ordner": "Kuer", "stempel": "KÜR", "farbe": MESSING, "db": None},
}

#: Text, an dem Prüfungen eine Lehrkraft-Seite erkennen (Fußzeile jeder Lehrkraft-/Lösungsfassung, Lösungsfolien).
MARKER_LEHRKRAFT = "NUR LEHRKRAFT"
#: Überschrift des Lösungsteils der Warm-up-Karten (die Schülerfassung endet vor dieser Seite).
LOESUNGSTEIL_WARMUP = "Lösungen der Warm-up-Karten"
#: Notiz-Marker für Folien, die im Handout fehlen.
NOTIZ_NUR_LEHRKRAFT = "[NUR LEHRKRAFT]"

ARTEN_DOCX = ("ab", "loesung", "lehrerhinweise", "unplugged", "handbuch", "warmup", "klausur", "erwartungshorizont",
              "generalprobe", "gesamt", "kuer")
LEHRKRAFT_ARTEN = {"loesung", "lehrerhinweise", "erwartungshorizont", "gesamt"}


# ================================================================== Inhalte (gemeinsam für docx und pptx)
@dataclass
class SQL:
    """Eine SQL-Anweisung im Absatzformat „SQL-Code“ (docx) bzw. Textfeld „SQL“ (pptx).

    db: Datenbank (z. B. "fundbuero", "verhoer/kiosk", "keine"); ohne Angabe gilt die Dokumenteigenschaft
        DD-Datenbank. pruefung: "ja" (Syntax, EH §8, Ergebnis), "fehlerhaft" (absichtlich falsch – die Prüfung
        erwartet mindestens einen Befund), "aus" (nicht prüfen, z. B. Ausschnitt). kuer: JOIN … ON usw. erlaubt.
    Die Steuerangaben stehen als versteckter Text [[DD: …]] im Absatz (druckt nicht).
    """
    text: str
    db: str | None = None
    pruefung: str = "ja"
    kuer: bool = False

    def direktiven(self):
        d = []
        if self.db:
            d.append(f"db={self.db}")
        if self.pruefung != "ja":
            d.append(self.pruefung)
        if self.kuer:
            d.append("kuer")
        return d


@dataclass
class Tabelle:
    kopf: list
    zeilen: list
    ergebnis: bool = False        # True → Tabellenformat „Ergebnistabelle“ (wird mit dem SQL davor verglichen)
    pk: tuple = ()                # Spaltennamen, die im Kopf unterstrichen werden
    fk: tuple = ()                # Spaltennamen mit ↑ im Kopf
    breiten: list | None = None   # cm je Spalte
    abgeschnitten: bool = False   # Ergebnistabelle zeigt nur die ersten Zeilen (letzte Zeile „…“)
    groesse: float = 9.0


def Ergebnis(kopf, zeilen, abgeschnitten=False, breiten=None):
    """Ergebnistabelle zum unmittelbar davor stehenden SQL (Spec §12.3)."""
    zeilen = [list(z) for z in zeilen]
    if abgeschnitten:
        zeilen.append(["…"] * len(kopf))
    return Tabelle(list(kopf), zeilen, ergebnis=True, breiten=breiten, abgeschnitten=abgeschnitten)


@dataclass
class Schema:
    """Kurzschreibweise nach EH §7: Primärschlüssel mit <u>…</u>, Fremdschlüssel mit ↑ direkt vor dem Namen."""
    text: str


@dataclass
class Liste:
    punkte: list
    nummeriert: bool = False


@dataclass
class Bild:
    pfad: str
    breite_cm: float = 8.0


@dataclass
class ER:
    """ER-Diagramm in Chen-Notation, Format wie Block `er` (AUFGABENFORMAT §5.7): graph + layout (+ markiert)."""
    graph: dict
    layout: dict
    breite_cm: float | None = None
    markiert: tuple = ()


@dataclass
class Linien:
    """Schreiblinien für Antworten (nur in der Schülerfassung)."""
    anzahl: int = 3


@dataclass
class Loesung:
    """Nur in der Lösungsfassung: grüner Kasten „LÖSUNG“; fehler = [("P6", "Klammern vergessen …")],
    varianten = weitere richtige Lösungen (Inhalte, z. B. SQL)."""
    inhalt: tuple = ()
    fehler: tuple = ()
    varianten: tuple = ()

    def __init__(self, *inhalt, fehler=(), varianten=()):
        self.inhalt, self.fehler, self.varianten = tuple(inhalt), tuple(fehler), tuple(varianten)


@dataclass
class Notiz:
    titel: str
    inhalt: tuple = ()

    def __init__(self, titel, *inhalt):
        self.titel, self.inhalt = titel, tuple(inhalt)


@dataclass
class Auftrag:
    """Ein Ermittlungsauftrag: stufe 1–3 (★–★★★), ort "PAPIER" oder "LAB"."""
    nr: int | str
    stufe: int
    ort: str
    text: str
    kompetenzen: tuple = ()
    inhalt: tuple = ()
    linien: int = 0
    loesung: Loesung | None = None


@dataclass
class Bewertung:
    """Zeile des Bewertungsbogens: teil ("a"), inhalt (Erwartungen, Kernpunkte …), afb = (I, II, III)."""
    teil: str
    inhalt: tuple
    afb: tuple = (0, 0, 0)

    @property
    def be(self):
        return sum(self.afb)


# ================================================================== Mini-Auszeichnung
_MUSTER = re.compile(r"(`[^`]+`|\*\*.+?\*\*|(?<![\w(*\\])\*(?![\s*)]).+?(?<![\s(*\\])\*(?![\w)*])|<u>.+?</u>)", re.S)


def laeufe(text, basis=None):
    """Zerlegt Mini-Auszeichnung in Läufe [(text, {fett, kursiv, code, unter})]; ↑ und ★ werden eigene Läufe."""
    basis = dict(basis or {})
    aus = []
    for teil in _MUSTER.split(text):
        if not teil:
            continue
        if teil.startswith("`") and teil.endswith("`") and len(teil) > 1:
            aus.append((teil[1:-1], {**basis, "code": True}))
        elif teil.startswith("**") and teil.endswith("**") and len(teil) > 4:
            aus += laeufe(teil[2:-2], {**basis, "fett": True})
        elif teil.startswith("<u>") and teil.endswith("</u>"):
            aus += laeufe(teil[3:-4], {**basis, "unter": True})
        elif teil.startswith("*") and teil.endswith("*") and len(teil) > 2 and not teil.startswith("**"):
            aus += laeufe(teil[1:-1], {**basis, "kursiv": True})
        else:
            for stueck in re.split(r"(↑|★+|☆+)", teil.replace("\\*", "*")):
                if not stueck:
                    continue
                if stueck == "↑":
                    aus.append((stueck, {**basis, "pfeil": True}))
                elif stueck[0] in "★☆":
                    aus.append((stueck, {**basis, "stern": True}))
                else:
                    aus.append((stueck, basis))
    return aus


def nur_text(text):
    """Text ohne Auszeichnung (für Prüfungen und Titel)."""
    return "".join(t for t, _ in laeufe(text))


# ================================================================== Grafik (Pillow): Stempel und ER-Diagramm
def _font(name, groesse):
    from PIL import ImageFont
    datei = {"kopf": "BarlowCondensed-SemiBold.ttf", "kopf_fett": "BarlowCondensed-Bold.ttf", "text": "Barlow-Regular.ttf",
             "text_halb": "Barlow-SemiBold.ttf", "kursiv": "Barlow-Italic.ttf", "mono": "IBMPlexMono-SemiBold.ttf",
             "mono_normal": "IBMPlexMono-Regular.ttf"}[name]
    return ImageFont.truetype(str(FONTS / datei), int(round(groesse)))


def stempel_png(text, pfad, farbe=SIEGELROT, groesse=64, drehung=-4, rauschen=True):
    """Stempel (Barlow Condensed 700, Versalien, doppelter Rahmen, −4°) als PNG mit transparentem Grund."""
    from PIL import Image, ImageDraw
    text = text.upper()
    f = _font("kopf_fett", groesse)
    sperrung = groesse * 0.1
    breiten = [f.getlength(z) for z in text]
    tb = sum(breiten) + sperrung * (len(text) - 1)
    asc, desc = f.getmetrics()
    pad_x, pad_y, rahmen = groesse * 0.45, groesse * 0.22, max(3, int(groesse * 0.075))
    w, h = int(tb + 2 * pad_x), int(asc * 0.78 + 2 * pad_y + groesse * 0.25)
    rand = int(groesse * 0.35)
    bild = Image.new("RGBA", (w + 2 * rand, h + 2 * rand), (0, 0, 0, 0))
    d = ImageDraw.Draw(bild)
    rgb = tuple(int(farbe[i:i + 2], 16) for i in (0, 2, 4))
    x0, y0 = rand, rand
    d.rounded_rectangle([x0, y0, x0 + w, y0 + h], radius=int(groesse * 0.08), outline=rgb + (255,), width=rahmen)
    innen = rahmen + max(2, int(groesse * 0.05))
    d.rounded_rectangle([x0 + innen, y0 + innen, x0 + w - innen, y0 + h - innen], radius=int(groesse * 0.05),
                        outline=rgb + (255,), width=max(1, rahmen // 3))
    x = x0 + pad_x
    y = y0 + (h - asc * 0.78) / 2 - asc * 0.2
    for z, b in zip(text, breiten):
        d.text((x, y), z, font=f, fill=rgb + (255,))
        x += b + sperrung
    if rauschen:  # leichte Stempel-Optik: einzelne Pixel heller (deterministisch)
        zufall = random.Random(sum(map(ord, text)))
        px = bild.load()
        for _ in range(int(bild.width * bild.height * 0.05)):
            i, j = zufall.randrange(bild.width), zufall.randrange(bild.height)
            r, g, b, a = px[i, j]
            if a:
                px[i, j] = (r, g, b, int(a * 0.35))
    bild = bild.rotate(-drehung, resample=Image.BICUBIC, expand=True)
    bild = bild.crop(bild.getbbox())
    bild.save(pfad)
    return pfad


# -- ER-Diagramm: Geometrie wie public/static/js/bauplan/chen.js (gleiche Layout-Koordinaten wie die Akten)
ER_MASSE = {"entB": 160, "entH": 60, "rauteH": 74, "rauteMin": 120, "rauteMax": 160, "attrRy": 22, "plaettchen": 28, "rand": 36}


def _raute_breite(name):
    return max(ER_MASSE["rauteMin"], min(ER_MASSE["rauteMax"], 44 + len(name or "") * 9.5))


def _attr_rx(name):
    return max(46, min(104, 18 + len(name or "") * 4.7))


def _er_elemente(graph):
    aus = []
    for e in graph.get("entitaeten", []):
        aus.append({"id": e["id"], "art": "entitaet", "el": e})
    for b in graph.get("beziehungen", []):
        aus.append({"id": b["id"], "art": "beziehung", "el": b})
    for besitzer in graph.get("entitaeten", []) + graph.get("beziehungen", []):
        for i, a in enumerate(besitzer.get("attribute", [])):
            aus.append({"id": f"{besitzer['id']}.{i}", "art": "attribut", "el": a, "besitzer": besitzer})
    return aus


def _ausdehnung(e):
    if e["art"] == "entitaet":
        return ER_MASSE["entB"] / 2, ER_MASSE["entH"] / 2
    if e["art"] == "beziehung":
        return _raute_breite(e["el"].get("name")) / 2, ER_MASSE["rauteH"] / 2
    return _attr_rx(e["el"].get("name")), ER_MASSE["attrRy"]


def _rechteck_rand(c, hb, hh, z):
    dx, dy = z[0] - c[0], z[1] - c[1]
    if not dx and not dy:
        return c
    t = min(hb / abs(dx) if dx else math.inf, hh / abs(dy) if dy else math.inf)
    t = min(1, t)
    return (c[0] + dx * t, c[1] + dy * t)


def _rauten_rand(c, hb, hh, z):
    dx, dy = z[0] - c[0], z[1] - c[1]
    if not dx and not dy:
        return c
    t = min(1, 1 / (abs(dx) / hb + abs(dy) / hh))
    return (c[0] + dx * t, c[1] + dy * t)


def er_kanten(bez, layout, ent_ids):
    """Kanten eines Beziehungstyps wie chen.js: [{ende, start, von, bis, mitte, normal}]."""
    p = layout.get(bez["id"])
    if not p:
        return []
    p = (p["x"], p["y"])
    hb, hh = _raute_breite(bez.get("name")) / 2, ER_MASSE["rauteH"] / 2
    enden = bez.get("enden", [])
    zaehler = {}
    for e in enden:
        zaehler[e["entitaet"]] = zaehler.get(e["entitaet"], 0) + 1
    gesehen, aus = {}, []
    for ende in enden:
        q = layout.get(ende["entitaet"])
        if not q or ende["entitaet"] not in ent_ids:
            continue
        q = (q["x"], q["y"])
        if zaehler[ende["entitaet"]] > 1:
            nr = gesehen.get(ende["entitaet"], 0)
            gesehen[ende["entitaet"]] = nr + 1
            bis = (p[0] - hb, p[1]) if nr == 0 else (p[0] + hb, p[1]) if nr == 1 else (p[0], p[1] + hh)
        else:
            bis = _rauten_rand(p, hb, hh, q)
        von = _rechteck_rand(q, ER_MASSE["entB"] / 2, ER_MASSE["entH"] / 2, bis)
        mitte = ((von[0] + bis[0]) / 2, (von[1] + bis[1]) / 2)
        laenge = math.hypot(bis[0] - von[0], bis[1] - von[1]) or 1
        aus.append({"ende": ende, "start": q, "von": von, "bis": bis, "mitte": mitte,
                    "normal": (-(bis[1] - von[1]) / laenge, (bis[0] - von[0]) / laenge)})
    return aus


def er_grenzen(graph, layout):
    x0 = y0 = math.inf
    x1 = y1 = -math.inf
    for e in _er_elemente(graph):
        p = layout.get(e["id"])
        if not p:
            continue
        hb, hh = _ausdehnung(e)
        x0, y0 = min(x0, p["x"] - hb), min(y0, p["y"] - hh)
        x1, y1 = max(x1, p["x"] + hb), max(y1, p["y"] + hh)
    if x0 == math.inf:
        raise ValueError("ER-Diagramm ohne Positionen")
    return x0, y0, x1, y1


def er_png(graph, layout, pfad, massstab=3.0, markiert=(), hintergrund="FFFFFF"):
    """Zeichnet ein ER-Diagramm in Chen-Notation (EH §7, DESIGN §4.14) als PNG. Jedes Element braucht eine Position
    im layout (Mittelpunkt, Einheiten wie im Bauplan-Tisch ≈ px bei 96 dpi). → (pfad, breite_px, hoehe_px)"""
    from PIL import Image, ImageDraw
    fehlend = [e["id"] for e in _er_elemente(graph) if e["id"] not in layout]
    if fehlend:
        raise ValueError(f"ER-Diagramm: Position fehlt für {', '.join(fehlend)}")
    s = massstab
    x0, y0, x1, y1 = er_grenzen(graph, layout)
    rand = ER_MASSE["rand"]
    ox, oy = x0 - rand, y0 - rand
    w, h = int((x1 - x0 + 2 * rand) * s), int((y1 - y0 + 2 * rand) * s)
    bild = Image.new("RGB", (w, h), "#" + hintergrund)
    d = ImageDraw.Draw(bild)
    P = lambda x, y: ((x - ox) * s, (y - oy) * s)  # noqa: E731
    tinte, rot = "#" + TINTE, "#" + SIEGELROT
    lw = lambda px: max(1, int(round(px * s)))  # noqa: E731
    elemente = _er_elemente(graph)
    ent_ids = {e["id"] for e in elemente if e["art"] == "entitaet"}

    def text_mitte(cx, cy, t, font, farbe=tinte, unter=False, max_b=None, art="text"):
        g = font.size
        while max_b and font.getlength(t) > max_b and g > 8:
            g -= 1
            font = _font(art, g)
        b = font.getlength(t)
        asc, desc = font.getmetrics()
        d.text((cx - b / 2, cy - (asc + desc) / 2 + desc * 0.15), t, font=font, fill=farbe)
        if unter:
            uy = cy + asc * 0.42
            d.line([(cx - b / 2, uy), (cx + b / 2, uy)], fill=farbe, width=lw(1.3))

    for e in elemente:  # Attributlinien
        if e["art"] == "attribut":
            a, b = layout[e["id"]], layout[e["besitzer"]["id"]]
            d.line([P(a["x"], a["y"]), P(b["x"], b["y"])], fill=tinte, width=lw(1.2))
    plaettchen = []
    for e in elemente:
        if e["art"] != "beziehung":
            continue
        for k in er_kanten(e["el"], layout, ent_ids):
            d.line([P(*k["start"]), P(*k["bis"])], fill=tinte, width=lw(1.6))
            plaettchen.append((e, k))
    for e in elemente:
        p = layout[e["id"]]
        cx, cy = P(p["x"], p["y"])
        hb, hh = _ausdehnung(e)
        hb, hh = hb * s, hh * s
        rahmen = rot if e["id"] in markiert else tinte
        name = e["el"].get("name") or "…"
        if e["art"] == "entitaet":
            d.rectangle([cx - hb, cy - hh, cx + hb, cy + hh], fill="#" + hintergrund, outline=rahmen, width=lw(2))
            text_mitte(cx, cy, name, _font("kopf_fett", 23 * s), max_b=2 * hb - 14 * s, art="kopf_fett")
        elif e["art"] == "beziehung":
            d.polygon([(cx - hb, cy), (cx, cy - hh), (cx + hb, cy), (cx, cy + hh)], fill="#" + hintergrund, outline=rahmen,
                      width=lw(2))
            text_mitte(cx, cy, name, _font("kursiv", 18 * s), max_b=2 * hb - 40 * s, art="kursiv")
        else:
            d.ellipse([cx - hb, cy - hh, cx + hb, cy + hh], fill="#" + hintergrund, outline=rahmen, width=lw(1.5))
            schluessel = bool(e["el"].get("schluessel"))
            art = "text_halb" if schluessel else "text"
            text_mitte(cx, cy, name, _font(art, 16.5 * s), unter=schluessel, max_b=2 * hb - 14 * s, art=art)
    for e, k in plaettchen:  # Kardinalitäten zuletzt (liegen oben)
        kard = str(k["ende"].get("kard") or "").strip()
        mx, my = P(*k["mitte"])
        if kard:
            g = ER_MASSE["plaettchen"] * s / 2
            d.rounded_rectangle([mx - g, my - g, mx + g, my + g], radius=4 * s, fill="#" + KARTE, outline="#" + RAND, width=lw(1))
            text_mitte(mx, my, kard, _font("mono", 19 * s), farbe=rot, art="mono")
        rolle = str(k["ende"].get("rolle") or "").strip()
        if rolle:
            p = layout[e["id"]]
            n = k["normal"]
            # Abstand so, dass die Rolle das Kardinalitätsplättchen nicht verdeckt (auch bei steilen Kanten)
            rb = _font("kursiv", 14).getlength(rolle)
            ab = max(26, ER_MASSE["plaettchen"] / 2 + 6 + abs(n[0]) * rb / 2 + abs(n[1]) * 9)
            a1 = (k["mitte"][0] + n[0] * ab, k["mitte"][1] + n[1] * ab)
            a2 = (k["mitte"][0] - n[0] * ab, k["mitte"][1] - n[1] * ab)
            q = a1 if math.hypot(a1[0] - p["x"], a1[1] - p["y"]) >= math.hypot(a2[0] - p["x"], a2[1] - p["y"]) else a2
            text_mitte(*P(*q), rolle, _font("kursiv", 14 * s), farbe="#" + TINTE_2, art="kursiv")
    bild.save(pfad, dpi=(int(96 * s), int(96 * s)))
    return pfad, w, h


# ================================================================== Dokumenteigenschaften (docProps/custom.xml)
_CUSTOM_CT = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
_CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
_FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"


def setze_eigenschaften(pfad, eigenschaften):
    """Schreibt benutzerdefinierte Dokumenteigenschaften (Text) in eine docx/pptx, z. B. {"DD-Datenbank": "fundbuero"}.
    Vorhandene benutzerdefinierte Eigenschaften werden ersetzt (Word: Datei → Eigenschaften → Anpassen)."""
    pfad = Path(pfad)
    props = "".join(
        f'<property fmtid="{_FMTID}" pid="{i}" name="{escape(str(k), {chr(34): "&quot;"})}">'
        f"<vt:lpwstr>{escape(str(v))}</vt:lpwstr></property>"
        for i, (k, v) in enumerate(sorted(eigenschaften.items()), start=2))
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">' + props + "</Properties>")
    tmp = pfad.with_suffix(pfad.suffix + ".tmp")
    with zipfile.ZipFile(pfad) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            daten = zin.read(info.filename)
            if info.filename == "docProps/custom.xml":
                continue
            if info.filename == "[Content_Types].xml":
                t = daten.decode("utf-8")
                if "/docProps/custom.xml" not in t:
                    t = t.replace("</Types>", f'<Override PartName="/docProps/custom.xml" ContentType="{_CUSTOM_CT}"/></Types>')
                daten = t.encode("utf-8")
            elif info.filename == "_rels/.rels":
                t = daten.decode("utf-8")
                if "docProps/custom.xml" not in t:
                    ids = [int(x) for x in re.findall(r'Id="rId(\d+)"', t)] or [0]
                    t = t.replace("</Relationships>",
                                  f'<Relationship Id="rId{max(ids) + 1}" Type="{_CUSTOM_REL}" Target="docProps/custom.xml"/></Relationships>')
                daten = t.encode("utf-8")
            zout.writestr(info, daten)
        zout.writestr("docProps/custom.xml", xml)
    tmp.replace(pfad)


def lese_eigenschaften(pfad):
    """Benutzerdefinierte Dokumenteigenschaften einer docx/pptx als dict (leer, wenn keine)."""
    try:
        with zipfile.ZipFile(pfad) as z:
            if "docProps/custom.xml" not in z.namelist():
                return {}
            xml = z.read("docProps/custom.xml").decode("utf-8")
    except (zipfile.BadZipFile, FileNotFoundError):
        return {}
    from xml.etree import ElementTree as ET
    wurzel = ET.fromstring(xml)
    aus = {}
    for p in wurzel:
        name = p.get("name")
        wert = "".join(x.text or "" for x in p)
        aus[name] = wert
    return aus


# ================================================================== SQL-Hilfe (Node-CLI)
def sql_cli(eintraege, modus="pruefen"):
    """Ruft scripts/sql_pruefen_cli.js auf (JSON über stdin) und gibt die Antwortliste zurück."""
    node = shutil.which("node")
    if not node:
        raise RuntimeError("node nicht gefunden – die SQL-Prüfung braucht Node (npm ci).")
    r = subprocess.run([node, str(REPO / "scripts" / "sql_pruefen_cli.js")] + (["--ausfuehren"] if modus == "ausfuehren" else []),
                       input=json.dumps(eintraege, ensure_ascii=False), capture_output=True, text=True, cwd=REPO, timeout=600)
    if r.returncode not in (0, 1):
        raise RuntimeError(f"sql_pruefen_cli.js: {r.stderr.strip()[:500]}")
    return json.loads(r.stdout)


def abfrage_ergebnis(sql, datenbank, max_zeilen=200):
    """Führt eine Abfrage auf der Haupt-DB aus → (kopf, zeilen) mit Zellen als Text (NULL → "NULL", Kommazahlen mit
    Komma). Für Ergebnistabellen in Lösungen: Ergebnis(*abfrage_ergebnis(sql, "fundbuero"))."""
    antwort = sql_cli([{"id": "1", "sql": sql, "datenbank": datenbank, "maxZeilen": max_zeilen}], modus="ausfuehren")[0]
    if antwort.get("fehler"):
        raise ValueError(f"Abfrage nicht ausführbar: {antwort['fehler']}")
    return antwort["spalten"], [[zelle_text(v) for v in z] for z in antwort["zeilen"]]


def zelle_text(v):
    if v is None:
        return "NULL"
    if isinstance(v, float):
        t = f"{v:.2f}".rstrip("0").rstrip(".") if abs(v - round(v, 2)) < 1e-9 else repr(v)
        return t.replace(".", ",")
    return str(v)


# ================================================================== docx: Grundlagen
def _rgb(h):
    return RGBColor.from_string(h.upper())


def _setze_schrift(rpr_besitzer, name):
    """rFonts für alle Schriftbereiche setzen und Theme-Schriften entfernen (sonst gewinnt Calibri)."""
    rpr = rpr_besitzer.get_or_add_rPr() if hasattr(rpr_besitzer, "get_or_add_rPr") else rpr_besitzer
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k if ":" in k else "w:" + k), str(v))
    return e


def _rahmen_xml(behaelter_tag, seiten, farbe, staerke_pt=0.75, art="single", abstand=0):
    b = OxmlElement(behaelter_tag)
    for s in seiten:
        b.append(_el("w:" + s, val=art if staerke_pt else "nil", sz=int(staerke_pt * 8), space=abstand, color=farbe))
    return b


def _zelle_rahmen(zelle, farbe=RAND, staerke_pt=0.75, art="single", seiten=("top", "left", "bottom", "right")):
    tcpr = zelle._tc.get_or_add_tcPr()
    alt = tcpr.find(qn("w:tcBorders"))
    if alt is None:
        alt = OxmlElement("w:tcBorders")
        tcpr.append(alt)
    for s in seiten:
        vorher = alt.find(qn("w:" + s))
        if vorher is not None:
            alt.remove(vorher)
        alt.append(_el("w:" + s, val=art if staerke_pt else "nil", sz=int(staerke_pt * 8), space=0, color=farbe))
    # Reihenfolge laut Schema: top, start/left, bottom, end/right
    reihenfolge = {"top": 0, "left": 1, "start": 1, "bottom": 2, "right": 3, "end": 3, "insideH": 4, "insideV": 5}
    kinder = sorted(list(alt), key=lambda e: reihenfolge.get(e.tag.split("}")[1], 9))
    for k in kinder:
        alt.remove(k)
        alt.append(k)


def _zelle_flaeche(zelle, farbe):
    tcpr = zelle._tc.get_or_add_tcPr()
    for alt in tcpr.findall(qn("w:shd")):
        tcpr.remove(alt)
    tcpr.append(_el("w:shd", val="clear", color="auto", fill=farbe))


def _zelle_innenabstand(zelle, oben=0.12, links=0.2, unten=0.12, rechts=0.2):
    tcpr = zelle._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for s, v in (("top", oben), ("left", links), ("bottom", unten), ("right", rechts)):
        m.append(_el("w:" + s, w=int(Cm(v).twips), type="dxa"))
    tcpr.append(m)


def _tabelle_breiten(tabelle, breiten_cm):
    tabelle.autofit = False
    tbl = tabelle._tbl
    tblpr = tbl.tblPr
    for alt in tblpr.findall(qn("w:tblW")):
        tblpr.remove(alt)
    gesamt = sum(breiten_cm)
    tblpr.append(_el("w:tblW", w=int(gesamt * 567), type="dxa"))
    if tblpr.find(qn("w:tblLayout")) is None:
        tblpr.append(_el("w:tblLayout", type="fixed"))
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(breiten_cm):
            gc.set(qn("w:w"), str(int(breiten_cm[i] * 567)))
    for zeile in tabelle.rows:
        for i, z in enumerate(zeile.cells):
            if i < len(breiten_cm):
                z.width = Cm(breiten_cm[i])


def _tabelle_ohne_rahmen(tabelle):
    tblpr = tabelle._tbl.tblPr
    for alt in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(alt)
    tblpr.append(_rahmen_xml("w:tblBorders", ("top", "left", "bottom", "right", "insideH", "insideV"), "auto", 0))


def _zeile_nicht_teilen(zeile):
    trpr = zeile._tr.get_or_add_trPr()
    trpr.append(_el("w:cantSplit"))


def _zeile_kopf(zeile):
    trpr = zeile._tr.get_or_add_trPr()
    trpr.append(_el("w:tblHeader"))


def _absatz_rahmen(absatz, farbe, staerke_pt, seiten=("left",), abstand=4):
    ppr = absatz._p.get_or_add_pPr()
    ppr.append(_rahmen_xml("w:pBdr", seiten, farbe, staerke_pt, abstand=abstand))


def _feld(absatz, befehl, groesse=7.5, farbe=TINTE_3):
    """Feld (z. B. PAGE, NUMPAGES) als einfache Feldfunktion."""
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), befehl)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = _el("w:rFonts", ascii=SCHRIFT_MONO, hAnsi=SCHRIFT_MONO, cs=SCHRIFT_MONO, eastAsia=SCHRIFT_MONO)
    rpr.append(rf)
    rpr.append(_el("w:color", val=farbe))
    rpr.append(_el("w:sz", val=int(groesse * 2)))
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    f.append(r)
    absatz._p.append(f)


def _sperrung(lauf, pt):
    rpr = lauf._r.get_or_add_rPr()
    rpr.append(_el("w:spacing", val=int(pt * 20)))


def _lauf_flaeche(lauf, farbe):
    rpr = lauf._r.get_or_add_rPr()
    rpr.append(_el("w:shd", val="clear", color="auto", fill=farbe))


# ================================================================== docx: AktenDokument
class AktenDokument:
    """Word-Dokument im Akten-Design.

    art: "ab" | "unplugged" | "handbuch" | "lehrerhinweise" | "warmup" | "klausur" | "generalprobe" | "gesamt" | "kuer"
    fassung: "schueler" (Aufgaben, Schreiblinien) | "loesung" (Lösungskästen, Fußzeile „NUR LEHRKRAFT“) |
             "lehrkraft" (Lehrkraft-Dokument ohne Lösungsbezug, ebenfalls „NUR LEHRKRAFT“)
    material: Kennung für die Kompetenzmatrix, z. B. "AB 1.1", "GP-eA", "KL-A" (Bezüge "AB 1.1/3" bzw. "GP-eA-a").
    datenbank: Standard-DB für alle SQL-Absätze (Dokumenteigenschaft DD-Datenbank).

    Typischer Ablauf (eine Funktion baut beide Fassungen):
        def baue(fassung):
            d = AktenDokument("ab", akte=1, ermittlung="1.1", ds=1, titel="Der erste Auftrag", fassung=fassung,
                              material="AB 1.1", datenbank="fundbuero")
            d.kopfband(); d.fall_einstieg("…"); d.beweisstueck(1, "Aufbau einer Tabelle", "…", SQL("SELECT …"))
            d.ermittlungsauftraege([Auftrag(1, 1, "PAPIER", "…", kompetenzen=["K01"], linien=3, loesung=Loesung("…"))])
            d.abi_ecke("**Entwerfen** Sie …", be=2, ...); d.spur_im_lab(...); d.cold_case(...)
            d.speichern(pfad)
    """

    def __init__(self, art, akte, titel, *, fassung="schueler", ermittlung=None, ds=None, material=None,
                 datenbank=None, kuer=False, untertitel=None, querformat=False, stempel=None):
        if fassung not in ("schueler", "loesung", "lehrkraft"):
            raise ValueError(f"unbekannte Fassung {fassung}")
        self.art, self.akte, self.titel, self.fassung = art, akte, titel, fassung
        self.ermittlung, self.ds, self.material, self.datenbank = ermittlung, ds, material, datenbank
        self.kuer = kuer or art == "kuer"
        self.untertitel = untertitel
        self.stempel = stempel if stempel is not None else AKTEN.get(akte, AKTEN[0])["stempel"]
        self.kompetenzen = {}
        self._tmp = Path(tempfile.mkdtemp(prefix="dd-design-"))
        self.doc = Document()
        self._seite(querformat)
        self._formate()
        self._kopf_fuss()

    # -------------------------------------------------------------- Grundaufbau
    @property
    def ist_loesung(self):
        return self.fassung == "loesung"

    @property
    def lehrkraft(self):
        return self.fassung in ("loesung", "lehrkraft")

    def _seite(self, quer):
        s = self.doc.sections[0]
        s.page_width, s.page_height = (Cm(29.7), Cm(21.0)) if quer else (Cm(21.0), Cm(29.7))
        if quer:
            s.orientation = WD_ORIENT.LANDSCAPE
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(1.6), Cm(1.6)
        s.header_distance, s.footer_distance = Cm(0.8), Cm(0.8)
        self.breite = (s.page_width - s.left_margin - s.right_margin) / 360000  # cm

    def _formate(self):
        st = self.doc.styles
        # Standardschrift (docDefaults ohne Theme-Schriften)
        rpr_def = self.doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        _setze_schrift(rpr_def, SCHRIFT)
        n = st["Normal"]
        n.font.name, n.font.size, n.font.color.rgb = SCHRIFT, Pt(10.5), _rgb(TINTE)
        _setze_schrift(n.element.get_or_add_rPr(), SCHRIFT)
        n.paragraph_format.space_after = Pt(4)
        n.paragraph_format.line_spacing = 1.12
        for name, groesse, vor, nach in (("Title", 22, 0, 4), ("Heading 1", 17, 10, 4), ("Heading 2", 13.5, 10, 3),
                                         ("Heading 3", 11.5, 8, 2)):
            h = st[name]
            h.font.name, h.font.size, h.font.bold, h.font.italic = SCHRIFT_KOPF, Pt(groesse), False, False
            h.font.color.rgb = _rgb(TINTE)
            _setze_schrift(h.element.get_or_add_rPr(), SCHRIFT_KOPF)
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(vor), Pt(nach)
            h.paragraph_format.keep_with_next = True
            ppr = h.element.get_or_add_pPr()
            for alt in ppr.findall(qn("w:pBdr")):
                ppr.remove(alt)
        # Überschrift 2 mit rotem Faden links
        _p = st["Heading 2"].element.get_or_add_pPr()
        _p.append(_rahmen_xml("w:pBdr", ("left",), SIEGELROT, 3, abstand=6))
        # SQL-Code: Plex Mono, grau hinterlegt, Ränder in Hintergrundfarbe als Innenabstand
        sql = st.add_style("SQL-Code", WD_STYLE_TYPE.PARAGRAPH)
        sql.base_style = n
        sql.font.name, sql.font.size, sql.font.color.rgb = SCHRIFT_MONO, Pt(9), _rgb(TINTE)
        _setze_schrift(sql.element.get_or_add_rPr(), SCHRIFT_MONO)
        pf = sql.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(3), Pt(5), 1.0
        pf.left_indent, pf.right_indent = Cm(0.15), Cm(0.15)
        pf.keep_together = True
        ppr = sql.element.get_or_add_pPr()
        ppr.append(_rahmen_xml("w:pBdr", ("top", "left", "bottom", "right"), SQL_GRAU, 1, abstand=4))
        ppr.append(_el("w:shd", val="clear", color="auto", fill=SQL_GRAU))
        # Tabellenformate „Ergebnistabelle“ und „Datentabelle“ (die Gestaltung kommt zusätzlich direkt an die Zellen)
        for name in ("Ergebnistabelle", "Datentabelle"):
            t = st.add_style(name, WD_STYLE_TYPE.TABLE)
            t.base_style = st["Table Grid"]
            t.font.size = Pt(9)
        # Zellenabsatz, Etikett, Schreiblinie, Fall-Einstieg
        z = st.add_style("DD Tabellentext", WD_STYLE_TYPE.PARAGRAPH)
        z.base_style = n
        z.font.size = Pt(9.5)
        z.paragraph_format.space_after, z.paragraph_format.space_before = Pt(1), Pt(1)
        z.paragraph_format.line_spacing = 1.0
        lin = st.add_style("DD Schreiblinie", WD_STYLE_TYPE.PARAGRAPH)
        lin.base_style = n
        lin.paragraph_format.space_before, lin.paragraph_format.space_after = Pt(11), Pt(0)
        lin.font.color.rgb = _rgb(RAND_STARK)
        # Abstandhalter: winziger Absatz nach Tabellen/Kästen (Word verlangt nach einer Tabelle einen Absatz)
        ah = st.add_style("DD Abstand", WD_STYLE_TYPE.PARAGRAPH)
        ah.base_style = n
        ah.font.size = Pt(2)
        ah.paragraph_format.space_before = ah.paragraph_format.space_after = Pt(0)
        ah.paragraph_format.line_spacing = Pt(4)
        fe = st.add_style("DD Fall-Einstieg", WD_STYLE_TYPE.PARAGRAPH)
        fe.base_style = n
        fe.font.italic, fe.font.size = True, Pt(11)
        fe.paragraph_format.left_indent = Cm(0.35)
        fe.paragraph_format.space_before, fe.paragraph_format.space_after = Pt(6), Pt(8)
        fe.element.get_or_add_pPr().append(_rahmen_xml("w:pBdr", ("left",), SIEGELROT, 3, abstand=8))

    def _kopf_fuss(self):
        s = self.doc.sections[0]
        akte = AKTEN.get(self.akte, AKTEN[0])
        tab = Cm(self.breite)
        kopf = s.header.paragraphs[0]
        kopf.style = self.doc.styles["Normal"]
        kopf.paragraph_format.space_after = Pt(0)
        kopf.paragraph_format.tab_stops.add_tab_stop(tab, WD_TAB_ALIGNMENT.RIGHT)
        links = "DATENDETEKTEI WESERSTEDT"
        if isinstance(self.akte, int) and 1 <= self.akte <= 7:
            rechts = f"AKTE {self.akte} · {akte['titel'].upper()}"
        else:
            rechts = akte["titel"].upper()
        self._etikett(kopf, links + "\t" + rechts, groesse=7)
        fuss = s.footer.paragraphs[0]
        fuss.style = self.doc.styles["Normal"]
        fuss.paragraph_format.space_after = Pt(0)
        fuss.paragraph_format.tab_stops.add_tab_stop(Cm(self.breite / 2), WD_TAB_ALIGNMENT.CENTER)
        fuss.paragraph_format.tab_stops.add_tab_stop(tab, WD_TAB_ALIGNMENT.RIGHT)
        kennung = self.material or self.titel
        self._etikett(fuss, nur_text(kennung) + "\t", groesse=7)
        if self.lehrkraft:
            r = fuss.add_run(MARKER_LEHRKRAFT + (" · LÖSUNG" if self.ist_loesung else ""))
            self._lauf_stil(r, code=True, groesse=7, farbe=SIEGELROT, fett=True)
            _sperrung(r, 0.6)
        self._etikett(fuss, "\tSeite ", groesse=7)
        _feld(fuss, "PAGE")
        self._etikett(fuss, " von ", groesse=7)
        _feld(fuss, "NUMPAGES")

    # -------------------------------------------------------------- Läufe und Absätze
    def _lauf_stil(self, r, fett=False, kursiv=False, code=False, unter=False, farbe=None, groesse=None, schrift=None,
                   pfeil=False, stern=False):
        f = r.font
        if code or pfeil:
            f.name = SCHRIFT_MONO
            _setze_schrift(r._r.get_or_add_rPr(), SCHRIFT_MONO)
        if stern:
            f.name = SCHRIFT_SYMBOL
            _setze_schrift(r._r.get_or_add_rPr(), SCHRIFT_SYMBOL)
            farbe = farbe or MESSING_GRAFIK
        if schrift:
            f.name = schrift
            _setze_schrift(r._r.get_or_add_rPr(), schrift)
        if fett:
            f.bold = True
        if kursiv:
            f.italic = True
        if unter:
            f.underline = True
        if pfeil:
            farbe = farbe or SIEGELROT
        if farbe:
            f.color.rgb = _rgb(farbe)
        if groesse:
            f.size = Pt(groesse * (0.92 if code and not schrift else 1))
        elif code:
            f.size = Pt(9.5)
        return r

    def _text(self, absatz, text, **stil):
        for t, o in laeufe(text):
            r = absatz.add_run(t)
            self._lauf_stil(r, **{**stil, **o, "fett": o.get("fett") or stil.get("fett", False)})
        return absatz

    def _etikett(self, absatz, text, groesse=7.5, farbe=TINTE_3, fett=False):
        r = absatz.add_run(text)
        self._lauf_stil(r, code=True, groesse=groesse / 0.92, farbe=farbe, fett=fett)
        _sperrung(r, 0.8)
        return r

    def _absatz(self, behaelter, text="", style=None):
        p = behaelter.add_paragraph(style=style)
        if text:
            self._text(p, text)
        return p

    # -------------------------------------------------------------- Inhalte
    def _fuelle(self, behaelter, inhalt, breite=None):
        """Fügt eine Liste von Inhalten (str, SQL, Tabelle, …) in das Dokument oder eine Zelle ein."""
        breite = breite or self.breite
        for x in inhalt:
            if x is None:
                continue
            if isinstance(x, str):
                p = self._absatz(behaelter, x, style=None if behaelter is self.doc else "DD Tabellentext")
                if behaelter is not self.doc:
                    p.paragraph_format.space_after = Pt(3)
            elif isinstance(x, SQL):
                self._sql(behaelter, x)
            elif isinstance(x, Tabelle):
                self._datentabelle(behaelter, x, breite)
            elif isinstance(x, Schema):
                p = self._absatz(behaelter)
                self._text(p, x.text.replace("<u>", "<u>").replace("`", ""), code=True, groesse=10)
                p.paragraph_format.space_after = Pt(5)
            elif isinstance(x, Liste):
                for i, punkt in enumerate(x.punkte, 1):
                    p = self._absatz(behaelter, style=None if behaelter is self.doc else "DD Tabellentext")
                    p.paragraph_format.left_indent = Cm(0.55)
                    p.paragraph_format.first_line_indent = Cm(-0.45)
                    p.paragraph_format.space_after = Pt(2)
                    marke = f"{i}.\t" if x.nummeriert else "▸\t"
                    r = p.add_run(marke)
                    self._lauf_stil(r, farbe=SIEGELROT if not x.nummeriert else TINTE_2, fett=x.nummeriert)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.55))
                    self._text(p, punkt)
            elif isinstance(x, Bild):
                p = self._absatz(behaelter)
                p.add_run().add_picture(str(x.pfad), width=Cm(min(x.breite_cm, breite - 0.4)))
            elif isinstance(x, ER):
                pfad = self._tmp / f"er_{id(x)}.png"
                _, w, _h = er_png(x.graph, x.layout, pfad, markiert=x.markiert)
                natuerlich = w / 3 / 96 * 2.54
                p = self._absatz(behaelter)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(pfad), width=Cm(min(x.breite_cm or natuerlich, breite - 0.4)))
            elif isinstance(x, Linien):
                if self.fassung == "schueler":
                    for _ in range(x.anzahl):
                        self._schreiblinie(behaelter, breite - (0 if behaelter is self.doc else 0.45))
            elif isinstance(x, Loesung):
                if self.ist_loesung:
                    self._loesungskasten(behaelter, x, breite)
            elif isinstance(x, Notiz):
                self.notiz(x.titel, *x.inhalt, behaelter=behaelter, breite=breite)
            else:
                raise TypeError(f"unbekannter Inhalt: {x!r}")

    def _schreiblinie(self, behaelter, breite):
        p = self._absatz(behaelter, style="DD Schreiblinie")
        tabs = p.paragraph_format.tab_stops
        # linksbündiger Tabulator + Nullbreite-Leerzeichen: LibreOffice zeichnet Füllpunkte in Tabellenzellen sonst nicht
        tabs.add_tab_stop(Cm(max(1.0, breite)), WD_TAB_ALIGNMENT.LEFT, leader=WD_TAB_LEADER.DOTS)
        p.add_run("\t\u200b")
        return p

    def inhalt(self, *inhalt):
        """Inhalte direkt ins Dokument (Text, SQL, Tabelle, Schema, Liste, Bild, ER, Linien, Loesung, Notiz)."""
        self._fuelle(self.doc, inhalt)

    def absatz(self, text):
        return self._absatz(self.doc, text)

    def ueberschrift(self, text, ebene=2):
        p = self._absatz(self.doc, style=f"Heading {ebene}")
        self._text(p, text)
        return p

    def seitenumbruch(self):
        self._absatz(self.doc).add_run().add_break(WD_BREAK.PAGE)

    def tabelle(self, kopf, zeilen, breiten=None, ergebnis=False, pk=(), fk=(), groesse=9.0):
        self._datentabelle(self.doc, Tabelle(kopf, zeilen, ergebnis=ergebnis, pk=pk, fk=fk, breiten=breiten, groesse=groesse),
                           self.breite)

    def _sql(self, behaelter, s):
        p = self._absatz(behaelter, style="SQL-Code")
        direktiven = s.direktiven()
        if self.kuer and "kuer" not in direktiven:
            direktiven.append("kuer")
        if direktiven:
            r = p.add_run(f"[[DD: {', '.join(direktiven)}]]")
            r.font.hidden = True
        r = p.add_run(s.text.strip("\n"))
        return p

    def _datentabelle(self, behaelter, t, breite):
        n = len(t.kopf)
        tab = behaelter.add_table(rows=1 + len(t.zeilen), cols=n)
        tab.style = self.doc.styles["Ergebnistabelle" if t.ergebnis else "Datentabelle"]
        tab.alignment = WD_TABLE_ALIGNMENT.LEFT
        breiten = t.breiten
        if not breiten:
            laengen = [max([len(nur_text(str(t.kopf[i])))] + [len(str(z[i])) for z in t.zeilen if i < len(z)]) for i in range(n)]
            roh = [max(1.4, min(6.0, 0.2 * l + 0.6)) for l in laengen]
            faktor = min(1.0, (breite - 0.3) / sum(roh))
            breiten = [b * faktor for b in roh]
        _tabelle_breiten(tab, breiten)
        _zeile_kopf(tab.rows[0])
        for i, k in enumerate(t.kopf):
            z = tab.rows[0].cells[i]
            _zelle_flaeche(z, FLAECHE)
            p = z.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            name = str(k)
            if name in t.fk and not name.startswith("↑"):
                name = "↑" + name
            if name.strip("↑") in t.pk:
                name = f"<u>{name}</u>"
            self._text(p, name, code=True, fett=True, groesse=t.groesse * 0.95 / 0.92, farbe=TINTE)
        for zi, zeile in enumerate(t.zeilen, start=1):
            for i in range(n):
                wert = zeile[i] if i < len(zeile) else ""
                z = tab.rows[zi].cells[i]
                p = z.paragraphs[0]
                p.style = self.doc.styles["DD Tabellentext"]
                text = "" if wert is None else str(wert)
                if text == "NULL":
                    self._lauf_stil(p.add_run("NULL"), code=True, kursiv=True, farbe=TINTE_3, groesse=t.groesse * 0.9 / 0.92)
                else:
                    self._text(p, text, groesse=t.groesse)
                if re.fullmatch(r"-?\d+([.,]\d+)?( ?€)?", text):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for zeile in tab.rows:
            for z in zeile.cells:
                _zelle_rahmen(z, RAND, 0.5)
        self._nach_kasten(behaelter, klein=True)
        return tab

    def _kasten(self, behaelter, breite, rand=RAND, staerke=0.75, art="single", flaeche=None, nicht_teilen=False):
        tab = behaelter.add_table(rows=1, cols=1)
        _tabelle_breiten(tab, [breite])
        z = tab.rows[0].cells[0]
        _zelle_rahmen(z, rand, staerke, art)
        if flaeche:
            _zelle_flaeche(z, flaeche)
        _zelle_innenabstand(z, 0.15, 0.3, 0.15, 0.3)
        if nicht_teilen:
            _zeile_nicht_teilen(tab.rows[0])
        z.paragraphs[0].style = self.doc.styles["DD Tabellentext"]
        return tab, z

    def _kastenkopf(self, zelle, etikett, titel=None, farbe=TINTE_3, rechts=None, breite=None):
        p = zelle.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        self._etikett(p, etikett, groesse=7.5, farbe=farbe, fett=True)
        if titel:
            r = p.add_run("  " + nur_text(titel))
            self._lauf_stil(r, schrift=SCHRIFT_KOPF, groesse=12, farbe=TINTE)
        if rechts:
            p.paragraph_format.tab_stops.add_tab_stop(Cm((breite or self.breite) - 0.65), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run("\t" + rechts)
            self._lauf_stil(r, code=True, fett=True, groesse=9, farbe=farbe)
        return p

    def _nach_kasten(self, behaelter, klein=False):
        """Abstand nach einer Tabelle/einem Kasten. In Zellen hängt python-docx nach add_table schon einen leeren Absatz
        an (Word verlangt ihn) – der wird zum winzigen Abstandhalter."""
        if behaelter is self.doc:
            p = self._absatz(self.doc)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(4 if klein else 8)
            return
        letzter = behaelter.paragraphs[-1] if behaelter.paragraphs else None
        if letzter is not None and not letzter.text and letzter._p.getprevious() is not None \
                and letzter._p.getprevious().tag == qn("w:tbl"):
            letzter.style = self.doc.styles["DD Abstand"]
        else:
            behaelter.add_paragraph(style="DD Abstand")

    def _loesungskasten(self, behaelter, l, breite):
        tab, z = self._kasten(behaelter, breite - 0.8, rand=GRUEN, staerke=1.25, flaeche=GRUEN_FLAECHE)
        self._kastenkopf(z, "LÖSUNG", farbe=GRUEN)
        self._fuelle(z, l.inhalt, breite - 1.2)
        if l.varianten:
            self._text(z.add_paragraph(style="DD Tabellentext"), "**Auch richtig:**")
            self._fuelle(z, l.varianten, breite - 1.2)
        if l.fehler:
            p = z.add_paragraph(style="DD Tabellentext")
            self._text(p, "**Typische Fehler:**")
            for falle, text in l.fehler:
                p = z.add_paragraph(style="DD Tabellentext")
                r = p.add_run(f"◆ {falle}")
                self._lauf_stil(r, code=True, fett=True, farbe=SIEGELROT, groesse=8.5)
                _lauf_flaeche(r, ROT_FLAECHE)
                self._text(p, "  " + text)
        self._nach_kasten(behaelter)

    def _kompetenz(self, schluessel, kompetenzen):
        if kompetenzen and self.material:
            norm = [k if k.startswith("DB-") else "DB-" + k for k in kompetenzen]
            self.kompetenzen[schluessel] = norm

    # -------------------------------------------------------------- Bausteine Arbeitsblatt
    def kopfband(self, name_datum=True):
        """Kopfband „AKTE n · ERMITTLUNG n.k · DS x“ + Titel + Name/Datum, rechts der Stempel der Akte."""
        teile = []
        if isinstance(self.akte, int) and 1 <= self.akte <= 7:
            teile.append(f"AKTE {self.akte}")
        elif self.akte == "K":
            teile.append("KÜR · NICHT IM ABITUR")
        if self.ermittlung:
            teile.append(f"ERMITTLUNG {self.ermittlung}")
        if self.ds:
            teile.append(f"DS {self.ds}")
        if not teile and self.material:
            teile.append(nur_text(self.material).upper())
        tab = self.doc.add_table(rows=1, cols=2)
        _tabelle_breiten(tab, [self.breite - 4.2, 4.2])
        _tabelle_ohne_rahmen(tab)
        links, rechts = tab.rows[0].cells
        for z in (links, rechts):
            _zelle_flaeche(z, PAPIER)
            _zelle_rahmen(z, SIEGELROT, 2.25, seiten=("bottom",))
            _zelle_innenabstand(z, 0.2, 0.35, 0.2, 0.25)
        p = links.paragraphs[0]
        p.style = self.doc.styles["DD Tabellentext"]
        self._etikett(p, " · ".join(teile), groesse=8, farbe=TINTE_2, fett=True)
        p = links.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(2)
        self._lauf_stil(p.add_run(nur_text(self.titel)), schrift=SCHRIFT_KOPF, groesse=21, farbe=TINTE)
        if self.untertitel:
            p = links.add_paragraph(style="DD Tabellentext")
            self._text(p, self.untertitel, farbe=TINTE_2, groesse=10.5)
        p = links.add_paragraph(style="DD Tabellentext")
        p.paragraph_format.space_before = Pt(5)
        if self.lehrkraft:
            r = p.add_run(("LÖSUNG · " if self.ist_loesung else "") + MARKER_LEHRKRAFT)
            self._lauf_stil(r, code=True, fett=True, groesse=8.5 / 0.92, farbe=SIEGELROT)
            _sperrung(r, 0.8)
        elif name_datum:
            self._etikett(p, "NAME ", groesse=7.5)
            self._lauf_stil(p.add_run("_" * 30), farbe=RAND_STARK)
            self._etikett(p, "   DATUM ", groesse=7.5)
            self._lauf_stil(p.add_run("_" * 13), farbe=RAND_STARK)
        rechts.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr = rechts.paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.stempel:
            farbe = MESSING_TEXT if self.akte == "K" else SIEGELROT
            png = stempel_png(self.stempel, self._tmp / "stempel.png", farbe=farbe)
            pr.add_run().add_picture(str(png), width=Cm(3.3))
        self._nach_kasten(self.doc)
        return tab

    def fall_einstieg(self, text):
        p = self._absatz(self.doc, style="DD Fall-Einstieg")
        self._etikett(p, "FALL-EINSTIEG  ", groesse=7.5, farbe=SIEGELROT, fett=True)
        self._text(p, text)
        return p

    def beweisstueck(self, nr, titel, *inhalt):
        """Beweisstück-Kasten (Infoteil): Messing-Rahmen, Kopf „BEWEISSTÜCK n“ + Titel."""
        tab, z = self._kasten(self.doc, self.breite, rand=MESSING, staerke=1.0, flaeche=KARTE)
        self._kastenkopf(z, f"BEWEISSTÜCK {nr}", titel, farbe=MESSING_TEXT)
        self._fuelle(z, inhalt, self.breite - 0.6)
        self._nach_kasten(self.doc)
        return tab

    def notiz(self, titel, *inhalt, behaelter=None, breite=None):
        """Merkkasten „Ermittler-Notiz“ (blau hinterlegt)."""
        behaelter = behaelter if behaelter is not None else self.doc
        breite = (breite or self.breite) - (0 if behaelter is self.doc else 0.6)
        tab, z = self._kasten(behaelter, breite, rand=BLAU_FLAECHE, staerke=0.5, flaeche=BLAU_FLAECHE)
        self._kastenkopf(z, "ERMITTLER-NOTIZ", titel, farbe=BLAU)
        self._fuelle(z, inhalt, breite - 0.6)
        self._nach_kasten(behaelter)
        return tab

    def ermittlungsauftraege(self, auftraege, ueberschrift="Ermittlungsaufträge"):
        """Tabelle der Ermittlungsaufträge: Nr. | Stufe (★–★★★) | Ort (PAPIER/LAB) | Auftrag (+ Inhalt, Linien, Lösung)."""
        if ueberschrift:
            self.ueberschrift(ueberschrift, 2)
        breiten = [1.0, 1.7, 1.7, self.breite - 4.4]
        tab = self.doc.add_table(rows=1 + len(auftraege), cols=4)
        _tabelle_breiten(tab, breiten)
        _tabelle_ohne_rahmen(tab)
        _zeile_kopf(tab.rows[0])
        for i, k in enumerate(("NR.", "STUFE", "ORT", "AUFTRAG")):
            z = tab.rows[0].cells[i]
            _zelle_flaeche(z, FLAECHE)
            p = z.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            self._etikett(p, k, groesse=7.5, farbe=TINTE_2, fett=True)
        for zi, a in enumerate(auftraege, start=1):
            if isinstance(a, dict):
                a = Auftrag(**a)
            zeile = tab.rows[zi]
            nr, stufe, ort, text = zeile.cells
            for z in zeile.cells:
                _zelle_rahmen(z, RAND, 0.5, seiten=("bottom",))
                _zelle_innenabstand(z, 0.12, 0.15, 0.12, 0.15)
            p = nr.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            self._lauf_stil(p.add_run(str(a.nr)), schrift=SCHRIFT_KOPF, groesse=15)
            p = stufe.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            self._lauf_stil(p.add_run("★" * a.stufe), stern=True, groesse=10)
            self._lauf_stil(p.add_run("☆" * (3 - a.stufe)), stern=True, groesse=10, farbe=RAND_STARK)
            p = ort.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            ort_text = a.ort.upper()
            r = p.add_run(f" {ort_text} ")
            if ort_text == "LAB":
                self._lauf_stil(r, code=True, fett=True, groesse=7.5 / 0.92, farbe="FFFFFF")
                _lauf_flaeche(r, BLAU)
            else:
                self._lauf_stil(r, code=True, fett=True, groesse=7.5 / 0.92, farbe=TINTE)
                _lauf_flaeche(r, FLAECHE_2)
            p = text.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            self._text(p, a.text)
            self._fuelle(text, a.inhalt, breiten[3] - 0.3)
            if a.linien and self.fassung == "schueler":
                self._fuelle(text, [Linien(a.linien)], breiten[3])
            if a.loesung is not None:
                self._fuelle(text, [a.loesung], breiten[3] - 0.3)
            self._kompetenz(f"{self.material}/{a.nr}", a.kompetenzen)
        self._nach_kasten(self.doc)
        return tab

    def abi_ecke(self, text, be, *, operator=None, afb=None, kompetenzen=(), inhalt=(), linien=4, loesung=None):
        """Abi-Ecke (rot gerahmt): eine Teilaufgabe im Abitur-Stil, gesiezt, Operator fett, „[x BE]“ wird ergänzt."""
        tab, z = self._kasten(self.doc, self.breite, rand=SIEGELROT, staerke=1.5, flaeche=KARTE)
        info = []
        if self.ist_loesung and operator:
            info.append(f"Operator {operator}")
        if self.ist_loesung and afb:
            info.append(f"AFB {afb}")
        self._kastenkopf(z, "ABI-ECKE", "Im Stil des Abiturs" + (f" ({', '.join(info)})" if info else ""), farbe=SIEGELROT)
        p = z.add_paragraph(style="DD Tabellentext")
        if not re.search(r"\[\s*\d+\s*BE\s*\]", text):
            text = f"{text} [{be} BE]"
        self._text(p, text)
        self._fuelle(z, inhalt, self.breite - 0.6)
        if self.fassung == "schueler" and linien:
            self._fuelle(z, [Linien(linien)])
        if loesung is not None:
            self._fuelle(z, [loesung], self.breite - 0.6)
        self._kompetenz(f"{self.material}/Abi-Ecke", kompetenzen)
        self._nach_kasten(self.doc)
        return tab

    def spur_im_lab(self, ermittlung, aufgaben, werkzeug=None, text=None):
        """Kasten „Spur im Lab“: welche Ermittlung, welche Aufgaben, welches Werkzeug."""
        tab, z = self._kasten(self.doc, self.breite, rand=BLAU, staerke=0.75, flaeche=BLAU_FLAECHE, nicht_teilen=True)
        self._kastenkopf(z, "SPUR IM LAB", f"Ermittlung {ermittlung}", farbe=BLAU)
        p = z.add_paragraph(style="DD Tabellentext")
        self._text(p, f"**Aufgaben:** {aufgaben}" + (f"  ·  **Werkzeug:** {werkzeug}" if werkzeug else ""))
        if text:
            self._text(z.add_paragraph(style="DD Tabellentext"), text)
        self._nach_kasten(self.doc)
        return tab

    def cold_case(self, titel, text, *inhalt, kompetenzen=(), linien=3, loesung=None):
        """Cold Case für Schnelle (ohne Punkte), gestrichelter Rahmen."""
        tab, z = self._kasten(self.doc, self.breite, rand=TINTE_3, staerke=1.0, art="dashed")
        self._kastenkopf(z, "COLD CASE · OHNE PUNKTE", titel, farbe=TINTE_2)
        self._text(z.add_paragraph(style="DD Tabellentext"), text)
        self._fuelle(z, inhalt, self.breite - 0.6)
        if self.fassung == "schueler" and linien:
            self._fuelle(z, [Linien(linien)])
        if loesung is not None:
            self._fuelle(z, [loesung], self.breite - 0.6)
        self._kompetenz(f"{self.material}/Cold Case", kompetenzen)
        self._nach_kasten(self.doc)
        return tab

    # -------------------------------------------------------------- Prüfungen (Generalprobe, Klausur)
    def klausurkopf(self, titel, untertitel=None, *, kurs="Informatik · Leistungskurs (eA)", dauer=None, hilfsmittel=None,
                    be=None):
        """Nüchterner Kopf für Generalprobe/Klausur (ohne Story-Optik): Titel, Kurs, Dauer, Hilfsmittel, Name/Datum."""
        tab = self.doc.add_table(rows=1, cols=2)
        _tabelle_breiten(tab, [self.breite - 5.0, 5.0])
        _tabelle_ohne_rahmen(tab)
        l, r = tab.rows[0].cells
        for z in (l, r):
            _zelle_rahmen(z, TINTE, 1.5, seiten=("top", "bottom"))
            _zelle_innenabstand(z, 0.2, 0.1, 0.2, 0.1)
        p = l.paragraphs[0]
        p.style = self.doc.styles["DD Tabellentext"]
        self._etikett(p, kurs.upper(), groesse=7.5, farbe=TINTE_2, fett=True)
        p = l.add_paragraph(style="DD Tabellentext")
        self._lauf_stil(p.add_run(titel), schrift=SCHRIFT_KOPF, groesse=19)
        if untertitel:
            self._text(l.add_paragraph(style="DD Tabellentext"), untertitel, farbe=TINTE_2)
        angaben = [x for x in (f"Bearbeitungszeit: {dauer}" if dauer else None, f"Hilfsmittel: {hilfsmittel}" if hilfsmittel else None,
                               f"Gesamt: {be} BE" if be else None) if x]
        if angaben:
            self._text(l.add_paragraph(style="DD Tabellentext"), " · ".join(angaben), groesse=9, farbe=TINTE_2)
        p = r.paragraphs[0]
        p.style = self.doc.styles["DD Tabellentext"]
        if self.lehrkraft:
            rr = p.add_run(("ERWARTUNGSHORIZONT · " if self.ist_loesung else "") + MARKER_LEHRKRAFT)
            self._lauf_stil(rr, code=True, fett=True, groesse=8 / 0.92, farbe=SIEGELROT)
        else:
            for etikett in ("NAME", "DATUM"):
                p = r.add_paragraph(style="DD Tabellentext") if etikett == "DATUM" else p
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(4.7), WD_TAB_ALIGNMENT.LEFT, leader=WD_TAB_LEADER.DOTS)
                self._etikett(p, etikett + "\t\u200b", groesse=7)
        self._nach_kasten(self.doc)

    def materialblock(self, titel, *inhalt):
        """Materialblock einer Prüfung („Material 1: …“), dünn gerahmt."""
        tab, z = self._kasten(self.doc, self.breite, rand=RAND_STARK, staerke=0.75)
        self._kastenkopf(z, "MATERIAL", titel, farbe=TINTE_2)
        self._fuelle(z, inhalt, self.breite - 0.6)
        self._nach_kasten(self.doc)
        return tab

    def teilaufgabe(self, teil, text, be, *, kennung=None, kompetenzen=(), inhalt=(), linien=0, loesung=None):
        """Teilaufgabe im Abitur-Stil: „a) **Erläutern** Sie … [4 BE]“; kennung z. B. "GP-eA-a" (Kompetenzmatrix)."""
        p = self._absatz(self.doc)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = bool(inhalt)
        self._lauf_stil(p.add_run(f"{teil})\t"), fett=True)
        if not re.search(r"\[\s*\d+\s*BE\s*\]", text):
            text = f"{text} [{be} BE]"
        self._text(p, text)
        if inhalt:
            self._fuelle(self.doc, inhalt)
        if linien and self.fassung == "schueler":
            self._fuelle(self.doc, [Linien(linien)])
        if loesung is not None:
            self._fuelle(self.doc, [loesung])
        self._kompetenz(kennung or f"{self.material}-{teil}", kompetenzen)
        return p

    def bewertungsbogen(self, zeilen, titel="Bewertungsbogen", kennungen=None, kompetenzen=None):
        """Erwartungshorizont: „Teil | Erwartete Prüfungsleistungen | I | II | III | BE“ mit Summenzeile.
        zeilen: [Bewertung(teil, inhalt, afb=(I, II, III))]; kompetenzen: {teil: [K…]} (optional)."""
        if titel:
            self.ueberschrift(titel, 2)
        breiten = [1.2, self.breite - 5.2, 1.0, 1.0, 1.0, 1.0]
        tab = self.doc.add_table(rows=2 + len(zeilen), cols=6)
        tab.style = self.doc.styles["Datentabelle"]
        _tabelle_breiten(tab, breiten)
        _zeile_kopf(tab.rows[0])
        for i, k in enumerate(("Teil", "Erwartete Prüfungsleistungen", "I", "II", "III", "BE")):
            z = tab.rows[0].cells[i]
            _zelle_flaeche(z, FLAECHE)
            p = z.paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            if i >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._lauf_stil(p.add_run(k), code=True, fett=True, groesse=8.5 / 0.92)
        summen = [0, 0, 0]
        for zi, b in enumerate(zeilen, start=1):
            zeile = tab.rows[zi]
            p = zeile.cells[0].paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            self._lauf_stil(p.add_run(f"{b.teil})"), fett=True)
            zelle = zeile.cells[1]
            zelle.paragraphs[0].style = self.doc.styles["DD Tabellentext"]
            erste = True
            for x in b.inhalt:
                if erste and isinstance(x, str):
                    self._text(zelle.paragraphs[0], x)
                    erste = False
                else:
                    self._fuelle(zelle, [x], breiten[1] - 0.3)
                    erste = False
            for i, wert in enumerate(b.afb):
                summen[i] += wert
                p = zeile.cells[2 + i].paragraphs[0]
                p.style = self.doc.styles["DD Tabellentext"]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._lauf_stil(p.add_run(str(wert) if wert else ""), code=True)
            p = zeile.cells[5].paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._lauf_stil(p.add_run(str(b.be)), code=True, fett=True)
            if kompetenzen and b.teil in kompetenzen:
                self._kompetenz((kennungen or {}).get(b.teil, f"{self.material}-{b.teil}"), kompetenzen[b.teil])
        letzte = tab.rows[-1]
        p = letzte.cells[1].paragraphs[0]
        p.style = self.doc.styles["DD Tabellentext"]
        gesamt = sum(summen)
        anteile = " · ".join(f"{n} {round(100 * s / gesamt) if gesamt else 0} %" for n, s in zip(("I", "II", "III"), summen))
        self._text(p, f"**Summe** (AFB {anteile})")
        for i, wert in enumerate(summen + [gesamt]):
            p = letzte.cells[2 + i].paragraphs[0]
            p.style = self.doc.styles["DD Tabellentext"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._lauf_stil(p.add_run(str(wert)), code=True, fett=True)
        for z in letzte.cells:
            _zelle_flaeche(z, FLAECHE)
        for zeile in tab.rows:
            for z in zeile.cells:
                _zelle_rahmen(z, RAND_STARK, 0.5)
        self._nach_kasten(self.doc)
        return tab

    # -------------------------------------------------------------- Speichern
    def eigenschaften(self):
        e = {"DD-Art": self.art, "DD-Fassung": self.fassung, "DD-Akte": str(self.akte), "DD-Kuer": "ja" if self.kuer else "nein"}
        if self.datenbank:
            e["DD-Datenbank"] = self.datenbank
        if self.material:
            e["DD-Material"] = self.material
        if self.kompetenzen:
            e["DD-Kompetenzen"] = json.dumps(self.kompetenzen, ensure_ascii=False, sort_keys=True)
        return e

    def speichern(self, pfad):
        pfad = Path(pfad)
        pfad.parent.mkdir(parents=True, exist_ok=True)
        cp = self.doc.core_properties
        cp.author = cp.last_modified_by = "Datendetektei Weserstedt"
        cp.title = nur_text(self.titel)
        cp.subject = AKTEN.get(self.akte, AKTEN[0])["titel"]
        cp.language = "de-DE"
        cp.keywords = "Datendetektei; Datenbanken; Informatik"
        cp.comments = ""
        self.doc.save(pfad)
        setze_eigenschaften(pfad, self.eigenschaften())
        shutil.rmtree(self._tmp, ignore_errors=True)
        return pfad


# ================================================================== docx: Warm-up-Karten (2 × A5 je Seite)
def _warmup_absaetze(d, zelle, text, breite):
    """Text einer Warm-up-Karte (Markdown-light aus seed/warmup) in eine Zelle: Absätze, Listen („- …“) und
    Tabellenzeilen („`a | b | c`“) – zeilenweise gruppiert."""
    gruppen = []  # [(art, [zeilen])]
    for zeile in text.strip().split("\n"):
        if not zeile.strip():
            gruppen.append(("leer", []))
            continue
        art = ("tabelle" if re.fullmatch(r"\s*`[^`]*\|[^`]*`\s*", zeile) else
               "liste" if re.match(r"\s*[-*] ", zeile) else "text")
        if gruppen and gruppen[-1][0] == art:
            gruppen[-1][1].append(zeile)
        else:
            gruppen.append((art, [zeile]))
    for art, zeilen in gruppen:
        if art == "tabelle":
            zellen = [[c.strip() for c in z.strip().strip("`").split("|")] for z in zeilen]
            d._datentabelle(zelle, Tabelle(zellen[0], zellen[1:], groesse=9), breite - 0.4)
        elif art == "liste":
            d._fuelle(zelle, [Liste([re.sub(r"^\s*[-*] ", "", z) for z in zeilen])], breite)
        elif art == "text":
            d._fuelle(zelle, ["\n".join(zeilen)], breite)


def warmup_karten_dokument(karten, akte, titel=None, fassung="loesung"):
    """Warm-up-Karten als docx: je Seite zwei Karten im A5-Querformat (Schnittlinie in der Mitte); ab eigener Seite der
    Lösungsteil mit der Überschrift „Lösungen der Warm-up-Karten“ (die Schüler-PDF endet davor). karten = Einträge aus
    seed/warmup/warmup_a<n>.json. Gibt das AktenDokument zurück (→ .speichern(pfad))."""
    d = AktenDokument("warmup", akte, titel or f"Warm-up-Karten Akte {akte}", fassung="schueler",
                      material=f"Warm-up A{akte}", stempel=None)
    s = d.doc.sections[0]
    s.top_margin = s.bottom_margin = Cm(1.0)
    s.left_margin = s.right_margin = Cm(1.2)
    s.header_distance = s.footer_distance = Cm(0.4)
    d.breite = 21.0 - 2.4
    hoehe_karte = (29.7 - 2.0 - 1.3) / 2  # cm, zwei Karten je Seite (inkl. Schnittlinie)
    for i in range(0, len(karten), 2):
        if i:
            d.seitenumbruch()
        paar = karten[i:i + 2]
        tab = d.doc.add_table(rows=len(paar), cols=1)
        _tabelle_breiten(tab, [d.breite])
        _tabelle_ohne_rahmen(tab)
        for zi, k in enumerate(paar):
            zeile = tab.rows[zi]
            zeile.height = Cm(hoehe_karte - 0.35)
            trpr = zeile._tr.get_or_add_trPr()
            for alt in trpr.findall(qn("w:trHeight")):
                trpr.remove(alt)
            trpr.append(_el("w:trHeight", val=int(Cm(hoehe_karte - 0.35) / 635), hRule="exact"))
            _zeile_nicht_teilen(zeile)
            z = zeile.cells[0]
            _zelle_innenabstand(z, 0.5, 0.7, 0.4, 0.7)
            _zelle_rahmen(z, RAND_STARK, 0.5, art="dashed", seiten=("bottom",) if zi == 0 else ("top",))
            p = z.paragraphs[0]
            p.style = d.doc.styles["DD Tabellentext"]
            p.paragraph_format.tab_stops.add_tab_stop(Cm(d.breite - 1.5), WD_TAB_ALIGNMENT.RIGHT)
            d._etikett(p, f"WARM-UP · {k['id']} · {k.get('kategorie', '').upper()}", groesse=7.5, farbe=SIEGELROT, fett=True)
            d._etikett(p, f"\t{k.get('minuten', '')} MIN", groesse=7.5, farbe=TINTE_2, fett=True)
            p = z.add_paragraph(style="DD Tabellentext")
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(6)
            d._lauf_stil(p.add_run(k["titel"]), schrift=SCHRIFT_KOPF, groesse=17)
            _warmup_absaetze(d, z, k["aufgabe"], d.breite - 1.4)
            if k.get("sql"):
                d._sql(z, SQL(k["sql"], db=k.get("datenbank"), pruefung="aus" if k.get("loesungSql") else "ja"))
            p = z.add_paragraph(style="DD Tabellentext")
            p.paragraph_format.space_before = Pt(8)
            fallen = ", ".join(k.get("falle") or [])
            d._etikett(p, "Kompetenz " + ", ".join(x.replace("DB-", "") for x in k.get("kompetenz", [])) +
                       (f" · Falle {fallen}" if fallen else ""), groesse=6.5)
    # Lösungsteil ab eigener Seite
    d.seitenumbruch()
    d.ueberschrift(LOESUNGSTEIL_WARMUP, 1)
    p = d.absatz("")
    r = p.add_run(MARKER_LEHRKRAFT + " – dieser Teil fehlt in der Schülerfassung.")
    d._lauf_stil(r, code=True, fett=True, groesse=8, farbe=SIEGELROT)
    for k in karten:
        p = d.ueberschrift(f"{k['id']} · {k['titel']}", 3)
        d._fuelle(d.doc, [k["loesung"]])
        if k.get("loesungSql"):
            d._sql(d.doc, SQL(k["loesungSql"], db=k.get("datenbank")))
    d.datenbank = None
    return d


# ================================================================== pptx: Fallbesprechung (dunkel, 16:9)
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor as PRGB  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt as PPt  # noqa: E402

FOLIE_B, FOLIE_H = 13.333, 7.5
KAMERA_X, KAMERA_Y = 9.35, 4.55  # Kamera-Zone unten rechts (Videos) bleibt frei
RAND_L = 0.6

_SQL_SCHLUESSEL = set("SELECT DISTINCT FROM WHERE GROUP BY HAVING ORDER ASC DESC LIMIT OFFSET AS AND OR NOT IS NULL IN "
                      "BETWEEN LIKE JOIN ON LEFT INNER CREATE TABLE INSERT INTO VALUES PRIMARY KEY FOREIGN REFERENCES "
                      "INTEGER TEXT REAL".split())
_SQL_AGGREGATE = {"COUNT", "SUM", "AVG", "MIN", "MAX"}


def _p_rgb(h):
    return PRGB.from_string(h.upper())


def sql_tokens(text):
    """Grobe Syntax-Färbung für Folien: [(text, art)] mit art in schluessel/aggregat/text/zahl/kommentar/sonst."""
    aus = []
    for m in re.finditer(r"(--[^\n]*)|('(?:[^']|'')*')|(\b\d+(?:\.\d+)?\b)|([A-Za-zÄÖÜäöüß_][\wÄÖÜäöüß]*)|(\s+)|(.)", text):
        kom, txt, zahl, wort, ws, sonst = m.groups()
        if kom:
            aus.append((kom, "kommentar"))
        elif txt:
            aus.append((txt, "text"))
        elif zahl:
            aus.append((zahl, "zahl"))
        elif wort:
            art = "schluessel" if wort.upper() in _SQL_SCHLUESSEL else "aggregat" if wort.upper() in _SQL_AGGREGATE else "sonst"
            aus.append((wort, art))
        else:
            aus.append((ws or sonst, "sonst"))
    return aus


class AktenPraesentation:
    """Präsentation im Stil „Fallbesprechung“: Tinte-Hintergrund, Papierkarten, roter Faden (Spec §11.2).

    Jede Folienmethode nimmt notizen (Sprechernotizen mit Moderationshinweisen) und nur_lehrkraft (setzt den Marker
    „[NUR LEHRKRAFT]“ an den Anfang der Notizen – die Folie fehlt im Handout; Lösungsfolien immer so markieren).
    kamera_zone=True (Videos): unten rechts ab x = 9,35 in, y = 4,55 in bleibt frei.
    SQL-Textfelder heißen „SQL“ (Steuerangaben im Namen: „SQL [db=stadtfest, fehlerhaft]“), Ergebnistabellen
    „Ergebnistabelle“ – beides prüft material_sql_pruefen.py.
    """

    def __init__(self, akte, titel, *, kamera_zone=False, datenbank=None, kuer=False):
        self.akte, self.titel, self.kamera_zone, self.datenbank, self.kuer = akte, titel, kamera_zone, datenbank, kuer
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(FOLIE_B), Inches(FOLIE_H)
        self._leer = self.prs.slide_layouts[6]
        self._tmp = Path(tempfile.mkdtemp(prefix="dd-folien-"))
        self._nr = 0

    # -------------------------------------------------------------- Grundbausteine
    @property
    def inhalt_rechts(self):
        return KAMERA_X - 0.25 if self.kamera_zone else FOLIE_B - RAND_L

    def _textfeld(self, folie, x, y, w, h, absaetze, groesse=20, farbe=D_TINTE, schrift=SCHRIFT, anker=MSO_ANCHOR.TOP,
                  ausrichtung=None, name=None, zeilenabstand=1.05):
        tb = folie.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anker
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.04))
        for i, a in enumerate(absaetze if isinstance(absaetze, list) else [absaetze]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = zeilenabstand
            if ausrichtung:
                p.alignment = ausrichtung
            if isinstance(a, tuple):  # (text, optionen)
                text, o = a
            else:
                text, o = a, {}
            self._laeufe(p, text, groesse=o.get("groesse", groesse), farbe=o.get("farbe", farbe),
                         schrift=o.get("schrift", schrift), fett=o.get("fett", False), kursiv=o.get("kursiv", False),
                         sperrung=o.get("sperrung"))
            if o.get("abstand_nach") is not None:
                p.space_after = PPt(o["abstand_nach"])
        return tb

    def _laeufe(self, p, text, groesse, farbe, schrift, fett=False, kursiv=False, sperrung=None):
        for t, o in laeufe(text):
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = PPt(groesse * (0.9 if o.get("code") else 1))
            f.name = SCHRIFT_MONO if (o.get("code") or o.get("pfeil")) else SCHRIFT_SYMBOL if o.get("stern") else schrift
            f.bold = fett or o.get("fett", False)
            f.italic = kursiv or o.get("kursiv", False)
            f.underline = o.get("unter", False)
            c = farbe
            if o.get("pfeil"):
                c = D_SIEGELROT if farbe in (D_TINTE, D_TINTE_2) else SIEGELROT
            elif o.get("stern"):
                c = D_MESSING
            elif o.get("code") and farbe == D_TINTE:
                c = D_MESSING
            f.color.rgb = _p_rgb(c)
            if sperrung:
                r._r.get_or_add_rPr().set("spc", str(int(sperrung * 100)))

    def _form(self, folie, x, y, w, h, fuellung, rand=None, form=MSO_SHAPE.RECTANGLE, staerke=1.0, name=None):
        s = folie.shapes.add_shape(form, Inches(x), Inches(y), Inches(w), Inches(h))
        if fuellung:
            s.fill.solid()
            s.fill.fore_color.rgb = _p_rgb(fuellung)
        else:
            s.fill.background()
        if rand:
            s.line.color.rgb = _p_rgb(rand)
            s.line.width = PPt(staerke)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        if form == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.04
        if name:
            s.name = name
        s.text_frame.text = ""
        return s

    def _folie(self, titel, etikett=None, notizen="", nur_lehrkraft=False, faden=True):
        folie = self.prs.slides.add_slide(self._leer)
        self._nr += 1
        bg = folie.background.fill
        bg.solid()
        bg.fore_color.rgb = _p_rgb(TINTE)
        akte = AKTEN.get(self.akte, AKTEN[0])
        if etikett is None:
            etikett = f"AKTE {self.akte} · {akte['titel'].upper()}" if isinstance(self.akte, int) and 1 <= self.akte <= 7 else akte["titel"].upper()
        self._textfeld(folie, RAND_L, 0.32, 9.0, 0.35, [(etikett, {"sperrung": 1.5})], groesse=11, farbe=D_MESSING,
                       schrift=SCHRIFT_MONO)
        if titel:
            self._textfeld(folie, RAND_L, 0.62, FOLIE_B - 2 * RAND_L - (1.9 if nur_lehrkraft else 0), 0.75,
                           [titel], groesse=32, farbe=D_TINTE, schrift=SCHRIFT_KOPF, anker=MSO_ANCHOR.BOTTOM)
        if faden:
            self._roter_faden(folie, 1.48)
        if nur_lehrkraft:
            chip = self._form(folie, FOLIE_B - RAND_L - 1.75, 0.34, 1.75, 0.34, D_SIEGELROT, form=MSO_SHAPE.ROUNDED_RECTANGLE)
            self._textfeld(folie, FOLIE_B - RAND_L - 1.75, 0.34, 1.75, 0.34, [(MARKER_LEHRKRAFT, {"fett": True})], groesse=10,
                           farbe=D_PAPIER, schrift=SCHRIFT_MONO, anker=MSO_ANCHOR.MIDDLE, ausrichtung=PP_ALIGN.CENTER)
            chip.name = "Marker Lehrkraft"
        # Fußzeile
        fuss = f"DATENDETEKTEI WESERSTEDT · {self._nr}"
        self._textfeld(folie, RAND_L, FOLIE_H - 0.42, 5.0, 0.3, [(fuss, {"sperrung": 1})], groesse=9, farbe=D_TINTE_3,
                       schrift=SCHRIFT_MONO)
        text = (NOTIZ_NUR_LEHRKRAFT + " " if nur_lehrkraft else "") + (notizen or "")
        if text.strip():
            folie.notes_slide.notes_text_frame.text = text.strip()
        return folie

    def _roter_faden(self, folie, y):
        rechts = FOLIE_B - RAND_L
        linie = folie.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(RAND_L), Inches(y), Inches(rechts), Inches(y))
        linie.line.color.rgb = _p_rgb(D_SIEGELROT)
        linie.line.width = PPt(2.25)
        linie.name = "Roter Faden"
        for x in (RAND_L, rechts):
            self._form(folie, x - 0.07, y - 0.07, 0.14, 0.14, D_SIEGELROT, form=MSO_SHAPE.OVAL)

    def _karte(self, folie, x, y, w, h, etikett=None, farbe_etikett=SIEGELROT, fuellung=PAPIER, rand=None):
        self._form(folie, x, y, w, h, fuellung, rand=rand, form=MSO_SHAPE.ROUNDED_RECTANGLE, staerke=2)
        if etikett:
            self._textfeld(folie, x + 0.25, y + 0.15, w - 0.5, 0.3, [(etikett, {"sperrung": 1.2, "fett": True})], groesse=10.5,
                           farbe=farbe_etikett, schrift=SCHRIFT_MONO)
            return y + 0.5
        return y + 0.2

    def _inhalt(self, folie, x, y, w, h, inhalt, groesse=18, farbe=TINTE):
        """Inhalte (str, SQL, Tabelle, Schema, Liste, Bild, ER, Notiz) untereinander in einen Bereich setzen."""
        texte = []

        def flush(y):
            if not texte:
                return y
            zeilen = sum(max(1, math.ceil(len(nur_text(t)) * groesse * 0.0075 / max(w, 1))) for t in texte)
            hh = min(h - (y - y0), zeilen * groesse * 1.25 / 72 + 0.1 * len(texte))
            self._textfeld(folie, x, y, w, hh, [(t, {"abstand_nach": 6}) for t in texte], groesse=groesse, farbe=farbe)
            texte.clear()
            return y + hh + 0.08

        y0 = y
        for el in inhalt:
            if isinstance(el, str):
                texte.append(el)
                continue
            y = flush(y)
            if isinstance(el, Liste):
                punkte = [(("" if el.nummeriert else "▸ ") + (f"{i}. " if el.nummeriert else "") + p, {"abstand_nach": 6})
                          for i, p in enumerate(el.punkte, 1)]
                hh = min(h - (y - y0), len(el.punkte) * groesse * 1.45 / 72 + 0.1)
                self._textfeld(folie, x, y, w, hh, punkte, groesse=groesse, farbe=farbe)
                y += hh + 0.08
            elif isinstance(el, SQL):
                y = self._sql_box(folie, x, y, w, el) + 0.12
            elif isinstance(el, Schema):
                hh = groesse * 1.6 / 72 + 0.1
                self._textfeld(folie, x, y, w, hh, [el.text], groesse=groesse * 0.9, farbe=farbe, schrift=SCHRIFT_MONO)
                y += hh + 0.08
            elif isinstance(el, Tabelle):
                y = self._tabelle(folie, x, y, w, el, groesse=min(16, groesse - 2)) + 0.12
            elif isinstance(el, (Bild, ER)):
                if isinstance(el, ER):
                    pfad = self._tmp / f"er_{id(el)}.png"
                    _, bw, bh = er_png(el.graph, el.layout, pfad, markiert=el.markiert, hintergrund=PAPIER if farbe == TINTE else "FFFFFF")
                    breite = min(w, 1.6 * bw / 3 / 96)  # auf Folien so groß wie möglich (Höhe begrenzt unten)
                else:
                    pfad, breite = el.pfad, min(w, el.breite_cm / 2.54)
                    from PIL import Image
                    with Image.open(pfad) as im:
                        bw, bh = im.size
                hoehe = breite * bh / bw
                rest = h - (y - y0)
                if hoehe > rest:
                    breite, hoehe = breite * rest / hoehe, rest
                folie.shapes.add_picture(str(pfad), Inches(x + (w - breite) / 2), Inches(y), Inches(breite), Inches(hoehe))
                y += hoehe + 0.1
            elif isinstance(el, Notiz):
                texte_n = [(f"ERMITTLER-NOTIZ · {el.titel.upper()}", {"groesse": 11, "fett": True, "farbe": BLAU, "abstand_nach": 4})]
                texte_n += [(t, {"abstand_nach": 4}) for t in el.inhalt if isinstance(t, str)]
                hh = 0.35 + sum(max(1, math.ceil(len(nur_text(t)) * groesse * 0.0075 / w)) for t in el.inhalt if isinstance(t, str)) * groesse * 1.3 / 72
                self._form(folie, x, y, w, hh + 0.2, BLAU_FLAECHE, form=MSO_SHAPE.ROUNDED_RECTANGLE)
                self._textfeld(folie, x + 0.2, y + 0.1, w - 0.4, hh, texte_n, groesse=groesse - 2, farbe=TINTE)
                y += hh + 0.3
            elif isinstance(el, (Linien, Loesung)):
                continue  # auf Folien ohne Bedeutung (Lösungen gehören auf eigene Lösungsfolien)
            else:
                raise TypeError(f"unbekannter Inhalt: {el!r}")
        return flush(y)

    def _sql_box(self, folie, x, y, w, s, groesse=None):
        zeilen = s.text.strip("\n").split("\n")
        laengste = max(len(z) for z in zeilen)
        groesse = groesse or max(11, min(20, int((w - 0.5) * 72 / (0.6 * max(laengste, 20)))))
        h = len(zeilen) * groesse * 1.25 / 72 + 0.35
        self._form(folie, x, y, w, h, KONSOLE, form=MSO_SHAPE.ROUNDED_RECTANGLE, name="SQL-Hintergrund")
        direktiven = s.direktiven() + (["kuer"] if self.kuer and not s.kuer else [])
        name = "SQL" + (f" [{', '.join(direktiven)}]" if direktiven else "")
        tb = folie.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(h - 0.3))
        tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True  # ohne Umbruch zentriert LibreOffice den Text; die Schriftgröße passt ohnehin in die Breite
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.02))
        for i, zeile in enumerate(zeilen):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.0
            for t, art in sql_tokens(zeile) or [("", "sonst")]:
                r = p.add_run()
                r.text = t
                r.font.name, r.font.size = SCHRIFT_MONO, PPt(groesse)
                r.font.color.rgb = _p_rgb(SQL_FARBEN.get(art, KONSOLE_TEXT))
                r.font.italic = art == "kommentar"
        return y + h

    def _tabelle(self, folie, x, y, w, t, groesse=14):
        n, m = len(t.kopf), len(t.zeilen) + 1
        zeile_h = groesse * 1.55 / 72
        breiten = t.breiten
        if breiten:
            breiten = [b / 2.54 for b in breiten]
        else:
            laengen = [max([len(nur_text(str(t.kopf[i])))] + [len(str(z[i])) for z in t.zeilen if i < len(z)]) for i in range(n)]
            roh = [max(0.8, l * groesse * 0.0085 + 0.3) for l in laengen]
            faktor = min(1.0, w / sum(roh))
            breiten = [b * faktor for b in roh]
        form = folie.shapes.add_table(m, n, Inches(x), Inches(y), Inches(sum(breiten)), Inches(zeile_h * m))
        form.name = "Ergebnistabelle" if t.ergebnis else "Datentabelle"
        tab = form.table
        tab.first_row = True
        tab.horz_banding = False
        for i, b in enumerate(breiten):
            tab.columns[i].width = Inches(b)
        for zi in range(m):
            tab.rows[zi].height = Inches(zeile_h)
            for si in range(n):
                z = tab.cell(zi, si)
                z.fill.solid()
                z.fill.fore_color.rgb = _p_rgb(FLAECHE if zi == 0 else KARTE)
                z.margin_left = z.margin_right = Inches(0.06)
                z.margin_top = z.margin_bottom = Inches(0.02)
                tf = z.text_frame
                tf.paragraphs[0].text = ""
                if zi == 0:
                    name = str(t.kopf[si])
                    if name in t.fk and not name.startswith("↑"):
                        name = "↑" + name
                    if name.strip("↑") in t.pk:
                        name = f"<u>{name}</u>"
                    self._laeufe(tf.paragraphs[0], name, groesse * 0.92, TINTE, SCHRIFT_MONO, fett=True)
                else:
                    zeile = t.zeilen[zi - 1]
                    wert = "" if si >= len(zeile) or zeile[si] is None else str(zeile[si])
                    if wert == "NULL":
                        self._laeufe(tf.paragraphs[0], "NULL", groesse * 0.85, TINTE_3, SCHRIFT_MONO, kursiv=True)
                    else:
                        self._laeufe(tf.paragraphs[0], wert, groesse, TINTE, SCHRIFT)
                    if re.fullmatch(r"-?\d+([.,]\d+)?( ?€)?", wert):
                        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        return y + zeile_h * m

    # -------------------------------------------------------------- Folientypen
    def titelfolie(self, titel, untertitel=None, notizen=""):
        folie = self._folie(None, notizen=notizen, faden=False)
        akte = AKTEN.get(self.akte, AKTEN[0])
        self._textfeld(folie, RAND_L, 2.0, 9.0, 1.4, [titel], groesse=54, farbe=D_TINTE, schrift=SCHRIFT_KOPF,
                       anker=MSO_ANCHOR.BOTTOM)
        self._roter_faden(folie, 3.6)
        if untertitel:
            self._textfeld(folie, RAND_L, 3.8, 9.0, 1.2, [untertitel], groesse=22, farbe=D_TINTE_2)
        png = stempel_png(akte["stempel"], self._tmp / "stempel.png", farbe=D_SIEGELROT if self.akte != "K" else D_MESSING)
        if not self.kamera_zone:
            folie.shapes.add_picture(str(png), Inches(9.9), Inches(1.6), Inches(2.8))
        else:
            folie.shapes.add_picture(str(png), Inches(6.6), Inches(0.9), Inches(2.4))
        return folie

    def briefing(self, titel, text, *, von="Edda Lorenz", bild=None, notizen="", nur_lehrkraft=False):
        """Briefing-Folie (Story): Papierkarte mit dem Auftrag, max. 3 Sätze."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        breite = (rechts - RAND_L) * (0.62 if bild else 1)
        y = self._karte(folie, RAND_L, 1.85, breite, 4.6 if not self.kamera_zone else 2.55, etikett=f"BRIEFING · {von.upper()}")
        self._textfeld(folie, RAND_L + 0.3, y + 0.1, breite - 0.6, 3.8 if not self.kamera_zone else 1.8,
                       [(text, {"kursiv": True})], groesse=22, farbe=TINTE)
        if bild:
            self._inhalt(folie, RAND_L + breite + 0.3, 1.85, rechts - RAND_L - breite - 0.3, 4.6, [bild])
        return folie

    def begriffe(self, titel, paare=None, *, inhalt=(), notizen="", nur_lehrkraft=False):
        """Begriffe/Notation: zweispaltige Begriffstabelle (paare = [(Begriff, Erklärung)]) und/oder freie Inhalte."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        y = 1.85
        if paare:
            y = self._tabelle(folie, RAND_L, y, rechts - RAND_L,
                              Tabelle(["Fachbegriff", "Bedeutung"], [list(p) for p in paare],
                                      breiten=[7.0, (rechts - RAND_L) * 2.54 - 7.0]), groesse=17) + 0.2
        if inhalt:
            self._inhalt(folie, RAND_L, y, rechts - RAND_L, FOLIE_H - 0.6 - y, list(inhalt), farbe=D_TINTE)
        return folie

    def beispiel_schritte(self, titel, sql, schritte, *, ergebnis=None, notizen="", nur_lehrkraft=False):
        """Beispiel Schritt für Schritt: links das SQL (Textfeld „SQL“), rechts nummerierte Schritte; darunter optional
        die Ergebnistabelle (Ergebnis(...))."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        sql = sql if isinstance(sql, SQL) else SQL(sql)
        rechts = self.inhalt_rechts
        breite_l = (rechts - RAND_L) * 0.55
        y = self._sql_box(folie, RAND_L, 1.85, breite_l, sql)
        if ergebnis is not None:
            self._textfeld(folie, RAND_L, y + 0.1, breite_l, 0.3, [("ERGEBNIS", {"sperrung": 1.2, "fett": True})], groesse=10,
                           farbe=D_MESSING, schrift=SCHRIFT_MONO)
            self._tabelle(folie, RAND_L, y + 0.45, breite_l, ergebnis, groesse=13)
        x = RAND_L + breite_l + 0.3
        w = rechts - x
        yy = 1.85
        for i, s in enumerate(schritte, 1):
            hoehe = max(0.75, math.ceil(len(nur_text(s)) * 16 * 0.0075 / (w - 0.75)) * 16 * 1.3 / 72 + 0.25)
            self._form(folie, x, yy, w, hoehe, PAPIER, form=MSO_SHAPE.ROUNDED_RECTANGLE)
            self._form(folie, x, yy, 0.08, hoehe, D_SIEGELROT)
            self._textfeld(folie, x + 0.15, yy + 0.08, 0.5, 0.5, [(str(i), {"fett": False})], groesse=24, farbe=SIEGELROT,
                           schrift=SCHRIFT_KOPF)
            self._textfeld(folie, x + 0.65, yy + 0.08, w - 0.8, hoehe - 0.12, [s], groesse=16, farbe=TINTE,
                           anker=MSO_ANCHOR.MIDDLE)
            yy += hoehe + 0.15
        return folie

    def unplugged(self, titel, schritte, *, material=None, dauer=None, notizen="", nur_lehrkraft=False):
        """Unplugged-Anleitung: nummerierte Schritte auf einer Papierkarte, Material/Dauer als Etikett."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        etikett = "UNPLUGGED" + (f" · {dauer}" if dauer else "") + (f" · MATERIAL: {material.upper()}" if material else "")
        hoehe = (4.75 if not self.kamera_zone else 2.5)
        y = self._karte(folie, RAND_L, 1.85, rechts - RAND_L, hoehe, etikett=etikett, farbe_etikett=GRUEN)
        self._inhalt(folie, RAND_L + 0.3, y + 0.05, rechts - RAND_L - 0.6, hoehe - 0.7, [Liste(list(schritte), nummeriert=True)],
                     groesse=19)
        return folie

    def lab_auftrag(self, ermittlung, titel, aufgaben, *, werkzeug=None, hinweise=(), notizen="", nur_lehrkraft=False):
        """Lab-Auftrag: Ermittlung, Aufgaben, Werkzeug – blaue Karte wie „Spur im Lab“."""
        folie = self._folie(f"Ab ins Lab: Ermittlung {ermittlung}", notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        y = self._karte(folie, RAND_L, 1.85, rechts - RAND_L, 3.2 if not self.kamera_zone else 2.5, etikett="SPUR IM LAB",
                        farbe_etikett=BLAU, fuellung=BLAU_FLAECHE)
        zeilen = [(f"„{titel}“", {"groesse": 26, "schrift": SCHRIFT_KOPF, "abstand_nach": 6}),
                  (f"**Aufgaben:** {aufgaben}", {"abstand_nach": 4})]
        if werkzeug:
            zeilen.append((f"**Werkzeug:** {werkzeug}", {"abstand_nach": 4}))
        self._textfeld(folie, RAND_L + 0.3, y + 0.05, rechts - RAND_L - 0.6, 2.4, zeilen, groesse=19, farbe=TINTE)
        if hinweise:
            self._inhalt(folie, RAND_L, 5.25 if not self.kamera_zone else 4.5, rechts - RAND_L, 1.6, [Liste(list(hinweise))],
                         groesse=17, farbe=D_TINTE)
        return folie

    def merkfolie(self, titel, punkte, *, inhalt=(), notizen="", nur_lehrkraft=False):
        """Merkfolie „Ermittler-Notiz“: Merksätze auf blauer Karte."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        hoehe = 4.75 if not self.kamera_zone else 2.55
        y = self._karte(folie, RAND_L, 1.85, rechts - RAND_L, hoehe, etikett="ERMITTLER-NOTIZ", farbe_etikett=BLAU,
                        fuellung=BLAU_FLAECHE)
        self._inhalt(folie, RAND_L + 0.3, y + 0.05, rechts - RAND_L - 0.6, hoehe - 0.7, [Liste(list(punkte))] + list(inhalt),
                     groesse=20)
        return folie

    def abi_format(self, text, be, *, operator=None, inhalt=(), notizen="", nur_lehrkraft=False):
        """Abi-Format-Beispiel: rot gerahmte Papierkarte, gesiezt, Operator fett, [x BE]."""
        folie = self._folie("Abi-Format", notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        if not re.search(r"\[\s*\d+\s*BE\s*\]", text):
            text = f"{text} [{be} BE]"
        hoehe = 4.75 if not self.kamera_zone else 2.55
        y = self._karte(folie, RAND_L, 1.85, rechts - RAND_L, hoehe, etikett="ABI-ECKE" + (f" · OPERATOR {operator.upper()}" if operator else ""),
                        rand=D_SIEGELROT)
        self._inhalt(folie, RAND_L + 0.3, y + 0.05, rechts - RAND_L - 0.6, hoehe - 0.7, [text] + list(inhalt), groesse=20)
        return folie

    def inhaltsfolie(self, titel, *inhalt, karte=True, notizen="", nur_lehrkraft=False):
        """Allgemeine Folie: Inhalte auf einer Papierkarte (karte=True) oder direkt auf dem dunklen Grund."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=nur_lehrkraft)
        rechts = self.inhalt_rechts
        hoehe = 4.75 if not self.kamera_zone else 2.55
        if karte:
            y = self._karte(folie, RAND_L, 1.85, rechts - RAND_L, hoehe)
            self._inhalt(folie, RAND_L + 0.3, y + 0.05, rechts - RAND_L - 0.6, hoehe - 0.4, list(inhalt), groesse=19)
        else:
            self._inhalt(folie, RAND_L, 1.85, rechts - RAND_L, hoehe, list(inhalt), groesse=19, farbe=D_TINTE)
        return folie

    def loesungsfolie(self, titel, *inhalt, notizen=""):
        """Lösungsfolie: immer „[NUR LEHRKRAFT]“ – fehlt im Handout; sichtbarer Marker oben rechts."""
        folie = self._folie(titel, notizen=notizen, nur_lehrkraft=True)
        rechts = self.inhalt_rechts
        hoehe = 4.75 if not self.kamera_zone else 2.55
        y = self._karte(folie, RAND_L, 1.85, rechts - RAND_L, hoehe, etikett="LÖSUNG", farbe_etikett=GRUEN, rand=D_GRUEN)
        self._inhalt(folie, RAND_L + 0.3, y + 0.05, rechts - RAND_L - 0.6, hoehe - 0.6, list(inhalt), groesse=18)
        return folie

    def speichern(self, pfad):
        pfad = Path(pfad)
        pfad.parent.mkdir(parents=True, exist_ok=True)
        cp = self.prs.core_properties
        cp.author = cp.last_modified_by = "Datendetektei Weserstedt"
        cp.title = nur_text(self.titel)
        cp.language = "de-DE"
        self.prs.save(pfad)
        e = {"DD-Art": "praesentation", "DD-Akte": str(self.akte), "DD-Kuer": "ja" if self.kuer else "nein"}
        if self.datenbank:
            e["DD-Datenbank"] = self.datenbank
        if self.kamera_zone:
            e["DD-Kamera-Zone"] = "ja"
        setze_eigenschaften(pfad, e)
        shutil.rmtree(self._tmp, ignore_errors=True)
        return pfad


def pruefe_kamera_zone(pfad):
    """Formen, die in die Kamera-Zone (unten rechts) ragen → [(folie, name)] (Hintergrund-/Fußzeilenformen zählen mit)."""
    prs = Presentation(pfad)
    aus = []
    for i, folie in enumerate(prs.slides, 1):
        for s in folie.shapes:
            if s.left is None:
                continue
            x1, y1 = (s.left + s.width) / 914400, (s.top + s.height) / 914400
            if x1 > KAMERA_X + 0.01 and y1 > KAMERA_Y + 0.01 and s.left / 914400 < FOLIE_B and s.name != "Roter Faden":
                # nur Formen, die wirklich in die Zone ragen (rechts von KAMERA_X und unterhalb KAMERA_Y)
                if (s.left + s.width) / 914400 > KAMERA_X and (s.top + s.height) / 914400 > KAMERA_Y:
                    aus.append((i, s.name))
    return aus
