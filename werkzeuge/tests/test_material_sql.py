"""material_sql_pruefen.py (Spec §12.3): SQL aus docx („SQL-Code“) und pptx (Textfeld „SQL“) sammeln und prüfen."""
import shutil
import unittest

from hilfen import A, mini_repo

import material_sql_pruefen as S

SQL_OK = "SELECT FundNr, Bezeichnung\nFROM Fundsache\nWHERE FundNr <= 2\nORDER BY FundNr;"
ERG_OK = (["FundNr", "Bezeichnung"], [["1", "E-Reader"], ["2", "Regenjacke"]])


class SammelnTest(unittest.TestCase):
    def setUp(self):
        self.wurzel = mini_repo()

    def tearDown(self):
        shutil.rmtree(self.wurzel, ignore_errors=True)

    def doc(self, *inhalt, datenbank="fundbuero", kuer=False, name="t.docx"):
        d = A.AktenDokument("ab", 1, "Test", fassung="loesung", material="AB 1.1", datenbank=datenbank, kuer=kuer)
        d.inhalt(*inhalt)
        d.beweisstueck(1, "Im Kasten", A.SQL("SELECT Kategorie\nFROM Fundsache\nWHERE Wert > 100", db="fundbuero"))
        return d.speichern(self.wurzel / name)

    def test_docx_absaetze_direktiven_und_ergebnistabelle(self):
        pfad = self.doc(A.SQL(SQL_OK), A.Ergebnis(*ERG_OK), A.SQL("-- DB: kiosk\nSELECT Bezeichnung FROM Produkt"),
                        A.SQL("SELECT x FROM y", pruefung="fehlerhaft", db="keine"),
                        A.SQL("SELECT 1; SELECT 2", pruefung="aus"))
        anw = S.sammle(pfad)
        self.assertEqual(len(anw), 4, [a.sql for a in anw])
        a = anw[0]
        self.assertEqual(a.db, "fundbuero")
        self.assertEqual(a.erwartet["kopf"], ["FundNr", "Bezeichnung"])
        self.assertEqual(a.erwartet["zeilen"], ERG_OK[1])
        self.assertNotIn("[[DD", a.sql)
        self.assertEqual(anw[1].db, "kiosk")
        self.assertEqual(anw[2].pruefung, "fehlerhaft")
        self.assertEqual(anw[2].db, "keine")
        self.assertIn("Wert > 100", anw[3].sql, "SQL in Kästen (Tabellenzellen) wird gefunden")

    def test_mehrere_anweisungen_und_zusammenhaengende_absaetze(self):
        from docx import Document
        d = Document()
        d.styles.add_style("SQL-Code", 1)
        d.add_paragraph("SELECT FundNr", style="SQL-Code")
        d.add_paragraph("FROM Fundsache;", style="SQL-Code")
        d.add_paragraph("Zwischentext")
        d.add_paragraph("SELECT 1;\nSELECT 'a;b';", style="SQL-Code")
        pfad = self.wurzel / "roh.docx"
        d.save(pfad)
        anw = S.sammle(pfad)
        self.assertEqual([a.sql.strip() for a in anw], ["SELECT FundNr\nFROM Fundsache;", "SELECT 1;", "SELECT 'a;b';"])
        self.assertIsNone(anw[0].db)

    def test_pptx_textfeld_sql_und_ergebnistabelle(self):
        p = A.AktenPraesentation(1, "T", datenbank="fundbuero")
        p.beispiel_schritte("Beispiel", SQL_OK, ["Erst FROM", "dann WHERE"], ergebnis=A.Ergebnis(*ERG_OK))
        p.inhaltsfolie("Fehler", A.SQL("SELECT Kategorie, COUNT(*) FROM Fundsache", pruefung="fehlerhaft"))
        pfad = p.speichern(self.wurzel / "f.pptx")
        anw = S.sammle(pfad)
        self.assertEqual(len(anw), 2)
        self.assertEqual(anw[0].erwartet["zeilen"], ERG_OK[1])
        self.assertEqual(anw[0].db, "fundbuero")
        self.assertEqual(anw[1].pruefung, "fehlerhaft")


class PruefenTest(unittest.TestCase):
    def setUp(self):
        self.wurzel = mini_repo()

    def tearDown(self):
        shutil.rmtree(self.wurzel, ignore_errors=True)

    def speichere(self, name, *inhalt, kuer=False, datenbank="fundbuero"):
        d = A.AktenDokument("ab", 1, "Test", fassung="loesung", material="AB 1.1", datenbank=datenbank, kuer=kuer)
        d.inhalt(*inhalt)
        return d.speichern(self.wurzel / name)

    def test_gruen_und_absichtliche_fehler(self):
        gut = self.speichere("gut.docx", A.SQL(SQL_OK), A.Ergebnis(*ERG_OK),
                             A.SQL("SELECT Kategorie, Farbe, COUNT(*)\nFROM Fundsache\nGROUP BY Kategorie", pruefung="fehlerhaft"))
        falsches_ergebnis = self.speichere("erg.docx", A.SQL(SQL_OK), A.Ergebnis(["FundNr", "Bezeichnung"], [["1", "E-Reader"], ["2", "Schirm"]]))
        join = self.speichere("join.docx", A.SQL("SELECT f.FundNr\nFROM Fundsache f JOIN Fundsache g ON f.FundNr = g.FundNr"))
        kuer = self.speichere("kuer.docx", A.SQL("SELECT f.FundNr\nFROM Fundsache f JOIN Fundsache g ON f.FundNr = g.FundNr"), kuer=True)
        ungleich = self.speichere("ungl.docx", A.SQL("SELECT FundNr FROM Fundsache WHERE Wert <> 3"))
        ohne_db = self.speichere("ohne.docx", A.SQL("SELECT FundNr FROM Fundsache"), datenbank=None)
        ergebnis = S.pruefe_dateien([gut, falsches_ergebnis, join, kuer, ungleich, ohne_db])
        fehler = {p.name: [b for b in bs if b["art"] != "hinweis"] for p, bs in ergebnis.items()}
        self.assertEqual(fehler["gut.docx"], [])
        self.assertEqual(fehler["kuer.docx"], [])
        self.assertTrue(any(b["art"] == "ergebnis" for b in fehler["erg.docx"]))
        self.assertTrue(any(b["art"] == "eh" for b in fehler["join.docx"]))
        self.assertTrue(fehler["ungl.docx"])
        self.assertTrue(any("Datenbank" in b["text"] for b in fehler["ohne.docx"]))


if __name__ == "__main__":
    unittest.main()
