"""akten_design.py: Auszeichnung, Formate, Fassungen, Eigenschaften, Grafiken, Folien."""
import json
import shutil
import unittest

from docx import Document
from pptx import Presentation

from hilfen import A, mini_repo


class AuszeichnungTest(unittest.TestCase):
    def test_laeufe(self):
        l = A.laeufe("**fett** und *kursiv*, `Code`, <u>PK</u>, ↑FK, ★★ und COUNT(*)")
        self.assertIn(("fett", {"fett": True}), l)
        self.assertIn(("kursiv", {"kursiv": True}), l)
        self.assertIn(("Code", {"code": True}), l)
        self.assertIn(("PK", {"unter": True}), l)
        self.assertIn(("↑", {"pfeil": True}), l)
        self.assertIn(("★★", {"stern": True}), l)
        self.assertEqual(A.nur_text("COUNT(*) und **x**"), "COUNT(*) und x")
        self.assertEqual(A.nur_text("<u>↑SanNr, ↑EinsatzNr</u>"), "↑SanNr, ↑EinsatzNr")


class DokumentTest(unittest.TestCase):
    def setUp(self):
        self.wurzel = mini_repo()

    def tearDown(self):
        shutil.rmtree(self.wurzel, ignore_errors=True)

    def baue(self, fassung):
        d = A.AktenDokument("ab", 2, "Der rote Faden", fassung=fassung, ermittlung="2.1", ds=3, material="AB 2.1",
                            datenbank="stadtrad")
        d.kopfband()
        d.fall_einstieg("Ein Rad fehlt.")
        d.beweisstueck(1, "Schema", A.Schema("Rad(<u>RadID</u>, Typ, ↑StationID)"), A.SQL("SELECT RadID\nFROM Rad"))
        d.ermittlungsauftraege([A.Auftrag(1, 2, "PAPIER", "Erläutere ↑StationID.", kompetenzen=["K03"], linien=3,
                                          loesung=A.Loesung("GEHEIMLÖSUNG", fehler=[("P7", "Präfix fehlt")]))])
        d.abi_ecke("**Erläutern** Sie die Kennzeichnung.", be=3, kompetenzen=["K05"], loesung=A.Loesung("Abi-Antwort"))
        d.spur_im_lab("2.1", "A2-E1-01 … 08", werkzeug="Spurenbrett")
        d.cold_case("Nachts", "Wer fuhr?", kompetenzen=["K23"])
        return d.speichern(self.wurzel / f"{fassung}.docx")

    def text(self, pfad):
        d = Document(str(pfad))
        teile = [p.text for p in d.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")]
        return "\n".join(teile)

    def test_fassungen(self):
        s = self.baue("schueler")
        l = self.baue("loesung")
        ts, tl = self.text(s), self.text(l)
        self.assertNotIn("GEHEIMLÖSUNG", ts)
        self.assertNotIn("Abi-Antwort", ts)
        self.assertNotIn(A.MARKER_LEHRKRAFT, ts)
        self.assertIn("GEHEIMLÖSUNG", tl)
        self.assertIn("P7", tl)
        self.assertIn("[3 BE]", ts)
        fuss = Document(str(l)).sections[0].footer.paragraphs[0].text
        self.assertIn(A.MARKER_LEHRKRAFT, fuss)
        self.assertIn("NAME", ts)

    def test_formate_und_eigenschaften(self):
        pfad = self.baue("loesung")
        d = Document(str(pfad))
        namen = {s.name for s in d.styles}
        self.assertTrue({"SQL-Code", "Ergebnistabelle", "Datentabelle"} <= namen)
        sql = [p for p in d.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
               if "SELECT RadID" in "".join(t.text or "" for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))]
        self.assertEqual(len(sql), 1)
        e = A.lese_eigenschaften(pfad)
        self.assertEqual(e["DD-Datenbank"], "stadtrad")
        self.assertEqual(e["DD-Fassung"], "loesung")
        self.assertEqual(json.loads(e["DD-Kompetenzen"]), {"AB 2.1/1": ["DB-K03"], "AB 2.1/Abi-Ecke": ["DB-K05"],
                                                            "AB 2.1/Cold Case": ["DB-K23"]})
        self.assertEqual(d.core_properties.author, "Datendetektei Weserstedt")
        # Unterstreichung des Primärschlüssels in der Kurzschreibweise
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        unter = ["".join(t.text or "" for t in r.iter(W + "t")) for r in d.element.body.iter(W + "r")
                 if r.find(W + "rPr") is not None and r.find(W + "rPr").find(W + "u") is not None]
        self.assertIn("RadID", unter)

    def test_eigenschaften_ueberschreiben(self):
        pfad = self.baue("schueler")
        A.setze_eigenschaften(pfad, {"DD-Datenbank": "fundbuero"})
        self.assertEqual(A.lese_eigenschaften(pfad), {"DD-Datenbank": "fundbuero"})
        Document(str(pfad))  # bleibt lesbar

    def test_bewertungsbogen_summen(self):
        d = A.AktenDokument("erwartungshorizont", 7, "EH", fassung="loesung", material="GP-eA")
        d.bewertungsbogen([A.Bewertung("a", ("x",), afb=(2, 2, 0)), A.Bewertung("b", ("y",), afb=(1, 3, 1))],
                          kompetenzen={"a": ["K02"]})
        pfad = d.speichern(self.wurzel / "eh.docx")
        tab = Document(str(pfad)).tables[-1]
        self.assertEqual([c.text for c in tab.rows[-1].cells][2:], ["3", "5", "1", "9"])
        self.assertEqual(json.loads(A.lese_eigenschaften(pfad)["DD-Kompetenzen"]), {"GP-eA-a": ["DB-K02"]})

    def test_warmup_loesungsteil_auf_eigener_seite(self):
        from hilfen import warmup_karten
        pfad = A.warmup_karten_dokument(warmup_karten(), 1).speichern(self.wurzel / "w.docx")
        d = Document(str(pfad))
        idx = next(i for i, p in enumerate(d.paragraphs) if p.text == A.LOESUNGSTEIL_WARMUP)
        vorher = d.paragraphs[idx - 1]
        self.assertTrue(any(br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page"
                            for br in vorher._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br")))


class GrafikUndFolienTest(unittest.TestCase):
    def setUp(self):
        self.wurzel = mini_repo()

    def tearDown(self):
        shutil.rmtree(self.wurzel, ignore_errors=True)

    def test_er_png_und_stempel(self):
        from PIL import Image
        graph = {"entitaeten": [{"id": "e1", "name": "Person", "attribute": [{"name": "PersonNr", "schluessel": True}]}],
                 "beziehungen": [{"id": "b1", "name": "kennt", "enden": [{"entitaet": "e1", "kard": "n", "rolle": "kennt"},
                                                                        {"entitaet": "e1", "kard": "m", "rolle": "wirdGekannt"}],
                                  "attribute": []}]}
        layout = {"e1": {"x": 100, "y": 200}, "e1.0": {"x": 100, "y": 100}, "b1": {"x": 330, "y": 200}}
        pfad, w, h = A.er_png(graph, layout, self.wurzel / "er.png")
        self.assertGreater(w, 300)
        with Image.open(pfad) as im:
            self.assertEqual(im.size, (w, h))
        with self.assertRaises(ValueError):
            A.er_png(graph, {"e1": {"x": 1, "y": 1}}, self.wurzel / "x.png")
        A.stempel_png("FUNDBÜRO", self.wurzel / "s.png")
        with Image.open(self.wurzel / "s.png") as im:
            self.assertEqual(im.mode, "RGBA")

    def test_folien(self):
        p = A.AktenPraesentation(3, "Kasse", datenbank="stadtfest")
        p.titelfolie("Kasse")
        p.beispiel_schritte("B", A.SQL("SELECT 1", db="keine"), ["eins"], ergebnis=A.Ergebnis(["1"], [["1"]]))
        p.loesungsfolie("L", "Text", notizen="Nur zeigen, wenn …")
        pfad = p.speichern(self.wurzel / "p.pptx")
        prs = Presentation(str(pfad))
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)
        namen = [s.name for s in prs.slides[1].shapes]
        self.assertIn("SQL [db=keine]", namen)
        self.assertIn("Ergebnistabelle", namen)
        self.assertTrue(prs.slides[2].notes_slide.notes_text_frame.text.startswith(A.NOTIZ_NUR_LEHRKRAFT))
        self.assertEqual(A.lese_eigenschaften(pfad)["DD-Datenbank"], "stadtfest")

    def test_kamera_zone(self):
        p = A.AktenPraesentation(1, "Video", kamera_zone=True)
        p.titelfolie("Video")
        p.merkfolie("M", ["a", "b"])
        p.beispiel_schritte("B", "SELECT 1", ["x", "y", "z", "w"])
        frei = p.speichern(self.wurzel / "v.pptx")
        self.assertEqual(A.pruefe_kamera_zone(frei), [])
        q = A.AktenPraesentation(1, "Ohne")
        q.merkfolie("M", ["a"])
        self.assertTrue(A.pruefe_kamera_zone(q.speichern(self.wurzel / "o.pptx")))


if __name__ == "__main__":
    unittest.main()
