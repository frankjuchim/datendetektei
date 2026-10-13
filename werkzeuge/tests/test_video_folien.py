"""Tests für werkzeuge/video_folien.py: Aufbau-Prüfung, Folien mit Kamera-Zone, Drehbuch, Aktualität (--pruefen)."""
import copy
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from hilfen import A  # noqa: F401  (setzt sys.path)
import video_folien as V

SQL = "SELECT Name, Preis\nFROM Sorte\nWHERE Preis >= 1.8"


def mini_spec():
    folien = [{"typ": "titel", "kapitel": "Einstieg", "sprechertext": "Hallo und willkommen."},
              {"titel": "Beispiel", "kapitel": "Beispiel", "links": [{"sql": SQL, "ergebnis": True}],
               "rechts": [{"kasten": {"farbe": "blau", "titel": "Merke", "text": ["- WHERE filtert Zeilen"]}}],
               "sprechertext": "Ein Beispiel."},
              {"typ": "pause", "kapitel": "Jetzt du", "links": [{"text": "Schreib die Abfrage."}], "pause_s": 10,
               "sprechertext": "Halte an."}]
    folien += [{"titel": f"Folie {i}", "links": [{"punkte": ["Punkt eins", "  - Unterpunkt"]}], "sprechertext": "Text."}
               for i in range(4)]
    folien += [{"titel": "Fallen", "links": [{"falle": "P6", "titel": "Bedingungslogik", "text": "Klammern setzen."}],
                "sprechertext": "Die Falle P6."},
               {"typ": "merke", "punkte": ["Eins.", "Zwei."], "sprechertext": "Das nimmst du mit."},
               {"typ": "ende", "lab": [{"ermittlung": "1.1", "titel": "Test", "aufgaben": "A1-E1-01 … 08"}],
                "naechstes": "Video 2", "sprechertext": "Ab ins Lab."}]
    return {"nr": 9, "kurz": "Test", "akte": 1, "titel": "Testvideo", "untertitel": "Nur für Tests", "laenge": "1 min",
            "einsatz": "Test", "zielgruppe": "Tests", "datenbank": "eisdiele", "beispieldaten": "Eisdiele",
            "lernziele": ["Testen"], "youtube": {"titel": "T", "beschreibung": "B", "tags": ["x"]}, "bilder": {},
            "folien": folien}


class AufbauPruefen(unittest.TestCase):
    def test_vollstaendig(self):
        self.assertEqual(V.pruefe_spec(mini_spec()), [])

    def test_fehlende_teile(self):
        s = mini_spec()
        s["folien"] = [f for f in s["folien"] if f.get("typ") not in ("pause", "merke")]
        s["folien"][0]["sprechertext"] = "wort " * 80
        m = " | ".join(V.pruefe_spec(s))
        for teil in ("Pause-Aufgabe", "Merksätze", "Einstieg zu lang"):
            self.assertIn(teil, m)

    def test_ohne_p_nummer_und_folienzahl(self):
        s = mini_spec()
        s["folien"] = [f for f in s["folien"] if "falle" not in json.dumps(f)][:5] + [s["folien"][-1]]
        m = " | ".join(V.pruefe_spec(s))
        self.assertIn("P-Nummer", m)
        self.assertIn("Folien (erwartet 8–14)", m)


class Bauen(unittest.TestCase):
    def setUp(self):
        self.wurzel = Path(tempfile.mkdtemp(prefix="dd-video-test-"))
        self.ordner = self.wurzel / "09_Test"
        self.ordner.mkdir()
        (self.ordner / "video.json").write_text(json.dumps(mini_spec(), ensure_ascii=False), encoding="utf-8")

    def tearDown(self):
        shutil.rmtree(self.wurzel, ignore_errors=True)

    def test_folien_skript_und_aktualitaet(self):
        from pptx import Presentation
        pptx, md = V.bauen(self.ordner)
        self.assertEqual(A.pruefe_kamera_zone(pptx), [])
        eig = A.lese_eigenschaften(pptx)
        self.assertEqual(eig["DD-Kamera-Zone"], "ja")
        self.assertEqual(eig["DD-Art"], "video")
        prs = Presentation(str(pptx))
        self.assertEqual(len(prs.slides), len(mini_spec()["folien"]))
        namen = [s.name for s in prs.slides[1].shapes]
        self.assertIn("SQL", namen)
        self.assertIn("Ergebnistabelle", namen)
        self.assertIn("Ein Beispiel.", prs.slides[1].notes_slide.notes_text_frame.text)
        text = md.read_text(encoding="utf-8")
        self.assertIn("0:00 Einstieg", text)
        self.assertIn("## Lernziele", text)
        self.assertIn("### Folie 3", text)
        # Ergebnis aus der DB: Sorten ab 1,80 € (Sanddorn-Mango, Karamell, Pistazie)
        tab = next(s for s in prs.slides[1].shapes if s.name == "Ergebnistabelle").table
        self.assertEqual(len(tab.rows), 4)
        self.assertEqual(V.pruefen(self.wurzel, ausgabe=lambda *_: None), [])
        # video.json geändert → Folien und Skript veraltet
        s = mini_spec()
        s["folien"][1]["sprechertext"] = "Geändert."
        (self.ordner / "video.json").write_text(json.dumps(s, ensure_ascii=False), encoding="utf-8")
        probleme = " | ".join(V.pruefen(self.wurzel, ausgabe=lambda *_: None))
        self.assertIn("veraltet", probleme)

    def test_zu_viel_inhalt_bricht_ab(self):
        s = mini_spec()
        s["folien"][3]["links"] = [{"punkte": ["Ein sehr langer Punkt " * 12] * 8}]
        (self.ordner / "video.json").write_text(json.dumps(s, ensure_ascii=False), encoding="utf-8")
        with self.assertRaises(V.Fehler):
            V.bauen(self.ordner)

    def test_fehlendes_bild(self):
        s = copy.deepcopy(mini_spec())
        s["folien"][3]["links"] = [{"bild": "gibtsnicht"}]
        s["bilder"] = {"gibtsnicht": {"werkzeug": "bauplan"}}
        (self.ordner / "video.json").write_text(json.dumps(s, ensure_ascii=False), encoding="utf-8")
        with self.assertRaises(V.Fehler):
            V.bauen(self.ordner)


class EchteVideos(unittest.TestCase):
    def test_alle_videos_vollstaendig_und_aktuell(self):
        ordner = V.alle_ordner()
        self.assertEqual(len(ordner), 8, [o.name for o in ordner])
        self.assertEqual(V.pruefen(ausgabe=lambda *_: None), [])
        for o in ordner:
            spec = V.lade(o)
            self.assertTrue(8 <= len(spec["folien"]) <= 14, o.name)


if __name__ == "__main__":
    unittest.main()
